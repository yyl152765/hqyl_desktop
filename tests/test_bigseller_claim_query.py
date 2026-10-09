from __future__ import annotations

import hashlib
import json
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from contextlib import ExitStack, contextmanager
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from backend.services import bigseller_claim_query as service
from backend.services.bigseller_claim_checkpoint import (
    ClaimCheckpoint,
    ClaimCheckpointError,
    create_claim_checkpoint,
    read_claim_checkpoint,
)


def listing(sku="sku", item="2", created="2026-09-01 09:00:00", listed="2026-08-30 08:00:00", site="MY", shop="店铺A"):
    return {"itemSku": sku, "id": "internal-" + item, "itemId": item, "shopId": "shop-" + shop,
            "shopName": shop, "site": site, "createTimeStr": created, "platformCreateTime": listed}


def result(sku="sku", status="matched"):
    if status != "matched":
        return service.empty_row(sku, "测试结果", status)
    return {"sku": sku, "shop_name": "店铺A", "item_id": "2", "created_time": "2026-09-01 09:00:00",
            "listed_time": "2026-08-30 08:00:00", "status": "matched", "message": ""}


class ClaimQueryTests(unittest.TestCase):
    def query(self, pages, site="all"):
        with patch.object(service, "request_benchmark_page", side_effect=pages) as request:
            value = service.query_claim_sku(Mock(), Mock(), {}, "sku", site, Mock())
            return value, request

    def test_exact_parent_query_is_ascending_read_only_and_has_no_invented_site_filter(self):
        query = service.build_claim_query("SKU-001", 3)
        self.assertEqual(query["searchType"], "parentSku")
        self.assertEqual(query["inquireType"], 2)
        self.assertEqual(query["orderBy"], "create_time")
        self.assertIs(query["desc"], False)
        self.assertEqual(query["searchContent"], "SKU-001")
        self.assertEqual(query["pageNo"], 3)
        self.assertEqual(query["shopeeStatus"], "live")
        self.assertNotIn("site", query)

    def test_shared_http_reader_receives_claim_query_without_changing_existing_defaults(self):
        response = Mock(status_code=200, headers={})
        response.json.return_value = {"code": 0, "data": {"page": {"rows": [listing()], "totalSize": 1}}}
        with patch("backend.services.bigseller_sku_benchmark._post_listing", return_value=response) as post:
            value = service.query_claim_sku(Mock(), Mock(), {}, "sku", "all", Mock())
        self.assertEqual(value["status"], "matched")
        self.assertEqual(post.call_args.args[3], service.build_claim_query("sku"))
        response.close.assert_called_once()

    def test_earliest_is_chosen_across_pages_and_all_fields_from_that_listing(self):
        later = listing(item="1", created="2026-09-02 09:00:00", listed="2026-01-01 00:00:00", shop="后来店铺")
        earlier = listing(item="99", created="2026-08-01 09:00:00", listed="2026-07-31 12:34:56", shop="最早店铺")
        value, request = self.query([([later], 2), ([earlier], 2)])
        self.assertEqual(value, {"sku": "sku", "shop_name": "最早店铺", "item_id": "99",
                                "created_time": "2026-08-01 09:00:00", "listed_time": "2026-07-31 12:34:56",
                                "status": "matched", "message": ""})
        self.assertEqual(request.call_count, 2)
        self.assertEqual(request.call_args.kwargs["query"]["pageNo"], 2)

    def test_exact_parent_and_site_scope_exclude_other_candidates(self):
        rows = [listing(sku="sku-more", item="1", created="2020-01-01 00:00:00"),
                listing(item="2", site="PH", created="2020-01-01 00:00:00"), listing(item="3")]
        value, _ = self.query([(rows, 3)], site="MY")
        self.assertEqual(value["item_id"], "3")

    def test_required_creation_parent_site_and_shop_missing_never_choose_other_row(self):
        for key, site in (("createTimeStr", "all"), ("itemSku", "all"), ("shopName", "all"), ("itemId", "all"), ("site", "MY")):
            bad = listing(item="1")
            bad.pop(key)
            with self.subTest(key=key), self.assertRaises(service.BigSellerQueryError):
                self.query([([bad, listing(item="2")], 2)], site=site)

    def test_platform_time_missing_is_blank_and_warns_without_selecting_other_listing(self):
        value, _ = self.query([([listing(item="1", listed=None), listing(item="2", created="2026-09-02 09:00:00")], 2)])
        self.assertEqual(value["item_id"], "1")
        self.assertEqual(value["listed_time"], "")
        self.assertIn("上架时间留空", value["message"])

    def test_creation_tie_has_stable_numeric_item_id_tiebreak_and_message(self):
        a, b = listing(item="10"), listing(item="2")
        for rows in ([a, b], [b, a]):
            value, _ = self.query([(rows, 2)])
            self.assertEqual(value["item_id"], "2")
            self.assertIn("2 条", value["message"])
            self.assertIn("Item ID", value["message"])

    def test_display_time_tie_uses_raw_milliseconds_before_item_id_without_changing_output(self):
        early = {**listing(item="99", created="2025-05-29 14:59"), "createTime": 1748530757000}
        late = {**listing(item="1", created="2025-05-29 14:59"), "createTime": "1748530777000"}
        for rows in ([early, late], [late, early]):
            value, _ = self.query([(rows, 2)])
            self.assertEqual(value["item_id"], "99")
            self.assertEqual(value["created_time"], "2025-05-29 14:59:00")
            self.assertEqual(value["message"], "")

    def test_true_raw_timestamp_tie_counts_only_final_tied_candidates(self):
        rows = [{**listing(item=item, created="2025-05-29 14:59"), "createTime": raw}
                for item, raw in (("1", 1748530799000), ("3", 1748530757000), ("2", "1748530757000"))]
        value, _ = self.query([(rows, 3)])
        self.assertEqual(value["item_id"], "2")
        self.assertIn("2 条", value["message"])
        self.assertIn("原始创建时间", value["message"])
        self.assertNotIn("3 条", value["message"])

    def test_missing_or_invalid_raw_timestamp_falls_back_to_display_precision_with_explanation(self):
        early = {**listing(item="99", created="2025-05-29 14:59"), "createTime": 1748530757000}
        for invalid in (None, "", True, 1748530757, -1748530757000, 1748530757000.0,
                        "1.748530757e12", " 1748530757000", "1748530757000x"):
            late = {**listing(item="1", created="2025-05-29 14:59"), "createTime": invalid}
            with self.subTest(invalid=invalid):
                value, _ = self.query([([early, late], 2)])
                self.assertEqual(value["item_id"], "1")
                self.assertIn("按 BS 显示时间精度并列", value["message"])
                self.assertIn("原始时间缺失或无效", value["message"])

    def test_raw_timestamp_never_reorders_different_display_times_or_changes_their_timezone(self):
        earlier_display = {**listing(item="99", created="2025-05-29 14:59"), "createTime": 1748530799000}
        later_display = {**listing(item="1", created="2025-05-29 15:00"), "createTime": 1748530757000}
        value, _ = self.query([([earlier_display, later_display], 2)])
        self.assertEqual(value["item_id"], "99")
        self.assertEqual(value["created_time"], "2025-05-29 14:59:00")

    def test_empty_and_out_of_scope_are_confirmed_not_found(self):
        for rows in ([], [listing(site="PH")], [listing(sku="different")]):
            value, _ = self.query([(rows, len(rows))], site="MY")
            self.assertEqual(value["status"], "not_found")
            self.assertEqual(value["created_time"], "")

    def test_incomplete_repeated_changed_or_duplicate_pages_fail(self):
        a, b = listing(item="1"), listing(item="2")
        cases = [([([a], 2), ([], 2)]), ([([a], 2), ([a], 2)]),
                 ([([a], 2), ([b], 3)]), ([([a, a], 2)]), ([([a, b], 1)])]
        for pages in cases:
            with self.subTest(pages=pages), self.assertRaises(service.BigSellerQueryError):
                self.query(pages)

    def test_unknown_total_full_page_continues_and_page_limit_does_not_return_partial_winner(self):
        with patch.object(service, "PAGE_SIZE", 1):
            value, request = self.query([([listing()], None), ([], None)])
            self.assertEqual(value["status"], "matched")
            self.assertEqual(request.call_count, 2)
            with patch.object(service, "MAX_PAGES", 1), self.assertRaises(service.BigSellerQueryError):
                self.query([([listing()], None)])

    def test_dates_are_parsed_not_compared_lexically_and_invalid_values_rejected(self):
        self.assertEqual(service.parse_listing_time("2026/09/01 01:23").strftime("%Y-%m-%d %H:%M:%S"), "2026-09-01 01:23:00")
        self.assertEqual(service.parse_listing_time("2026-09-01T01:23:45Z").hour, 9)
        for value in (None, "", "--", 1720000000, True, "2026-02-30 00:00:00", "2026-09-01"):
            with self.subTest(value=value), self.assertRaises(service.BigSellerQueryError):
                service.parse_listing_time(value)

    def source(self, directory, skus=("sku", "other")):
        path = Path(directory) / "source.xlsx"
        path.write_bytes(b"source-workbook-identity")
        return SimpleNamespace(source_path=path, source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                               sheet_name="新品认领", skus=skus, total_rows=len(skus) + 1, warnings=[])

    def job(self, source, directory):
        return service.BigSellerClaimJob(username="account", password="password-secret", source_file=source.source_path,
                                        sheet_name=source.sheet_name, output_dir=Path(directory))

    @contextmanager
    def runtime(self, source, outcomes):
        with ExitStack() as stack:
            stack.enter_context(patch.object(service, "load_claim_input", return_value=source))
            stack.enter_context(patch.object(service, "query_process_config", return_value={}))
            stack.enter_context(patch.object(service, "load_bigseller_request_util", return_value=Mock()))
            login = stack.enter_context(patch.object(service, "login_for_query", return_value=Mock()))
            query = stack.enter_context(patch.object(service, "query_claim_sku", side_effect=outcomes))
            export = stack.enter_context(patch.object(service, "export_claim_workbook", return_value=source.source_path.parent / "output.xlsx"))
            yield query, login, export

    def test_run_commits_each_sku_and_resume_only_queries_failed(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            job = self.job(source, directory)
            with self.runtime(source, [result(), service.BigSellerQueryError("失败")]):
                first = service.run_bigseller_claim_query(job, Mock())
            self.assertFalse(first["is_complete"])
            self.assertEqual(first["confirmed_count"], 1)
            data, _ = read_claim_checkpoint(Path(first["checkpoint_file"]), service.checkpoint_identity(job, source))
            self.assertEqual(data["sku"], result())
            with self.runtime(source, [result("other")]) as (query, _, export):
                second = service.run_bigseller_claim_query(replace(job, resume_checkpoint=Path(first["checkpoint_file"])), Mock())
            self.assertTrue(second["is_complete"])
            self.assertEqual(query.call_count, 1)
            self.assertEqual(query.call_args.args[3], "other")
            self.assertNotEqual(first["checkpoint_file"], second["checkpoint_file"])
            self.assertEqual(len(export.call_args.args[1]), 2)

    def test_completed_resume_reexports_without_login(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory, ("sku",))
            job = self.job(source, directory)
            with self.runtime(source, [result()]):
                first = service.run_bigseller_claim_query(job, Mock())
            with self.runtime(source, []) as (query, login, _):
                second = service.run_bigseller_claim_query(replace(job, resume_checkpoint=Path(first["checkpoint_file"])), Mock())
            self.assertTrue(second["is_complete"])
            query.assert_not_called()
            login.assert_not_called()

    def test_cooldown_stops_next_sku_and_persists_deadline(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            job = self.job(source, directory)
            with self.runtime(source, [service.BenchmarkCooldownError("限流", 120)]) as (query, _, _):
                first = service.run_bigseller_claim_query(job, Mock())
            self.assertFalse(first["is_complete"])
            self.assertEqual(query.call_count, 1)
            _, deadline = read_claim_checkpoint(Path(first["checkpoint_file"]), service.checkpoint_identity(job, source))
            self.assertGreater(deadline, service.time.time() + 110)
            with self.runtime(source, [result(), result("other")]), patch.object(service, "wait_before_retry") as wait:
                second = service.run_bigseller_claim_query(replace(job, resume_checkpoint=Path(first["checkpoint_file"])), Mock())
            self.assertTrue(second["is_complete"])
            wait.assert_called_once()

    def test_auth_refresh_once_permission_stops_and_unexpected_exception_is_sanitized(self):
        cases = [([service.BigSellerAuthenticationError("expired"), result(), result("other")], 2, True),
                 ([service.BigSellerPermissionError("无权限")], 1, False),
                 ([RuntimeError("password-secret"), result("other")], 1, False)]
        for outcomes, login_count, complete in cases:
            with self.subTest(outcomes=outcomes), tempfile.TemporaryDirectory() as directory:
                source = self.source(directory)
                with self.runtime(source, outcomes) as (_, login, _):
                    value = service.run_bigseller_claim_query(self.job(source, directory), Mock())
                self.assertEqual(login.call_count, login_count)
                self.assertEqual(value["is_complete"], complete)
                self.assertNotIn("password-secret", json.dumps(value))

    def test_storage_failure_stops_before_next_request_and_retains_prior_commit(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            original = ClaimCheckpoint.save_result
            def save(checkpoint, row):
                if row["sku"] == "other":
                    raise ClaimCheckpointError("磁盘失败")
                return original(checkpoint, row)
            with self.runtime(source, [result(), result("other")]), patch.object(ClaimCheckpoint, "save_result", save):
                value = service.run_bigseller_claim_query(self.job(source, directory), Mock())
            self.assertFalse(value["is_complete"])
            self.assertEqual(value["confirmed_count"], 1)
            self.assertIn("磁盘失败", value["completion_message"])

    def test_export_failure_keeps_complete_query_checkpoint_for_reexport(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory, ("sku",))
            with self.runtime(source, [result()]) as (_, _, export):
                export.side_effect = RuntimeError("credential-secret")
                value = service.run_bigseller_claim_query(self.job(source, directory), Mock())
            self.assertFalse(value["is_complete"])
            self.assertEqual(value["confirmed_count"], 1)
            self.assertTrue(Path(value["checkpoint_file"]).is_file())
            self.assertNotIn("credential-secret", value["completion_message"])

    def test_checkpoint_identity_rejects_account_source_and_scope_mismatch(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            job = self.job(source, directory)
            identity = service.checkpoint_identity(job, source)
            checkpoint = create_claim_checkpoint(Path(directory), identity)
            checkpoint.save_result(result())
            checkpoint.close()
            for key in ("source_sha256", "account_key", "site", "listing_scope", "sheet_name", "query_contract"):
                wrong = {**identity, key: "different"}
                with self.subTest(key=key), self.assertRaises(ClaimCheckpointError):
                    read_claim_checkpoint(checkpoint.path, wrong)
            self.assertTrue(identity["query_contract"].endswith("-v2"))
            self.assertNotIn("password-secret", checkpoint.path.read_bytes().decode("utf-8", errors="ignore"))

    def test_checkpoint_rejects_corrupted_confirmed_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            identity = service.checkpoint_identity(self.job(source, directory), source)
            checkpoint = create_claim_checkpoint(Path(directory), identity)
            checkpoint.save_result(result())
            path = checkpoint.path
            checkpoint.close()
            db = sqlite3.connect(path)
            try:
                bad = {**result(), "created_time": "invalid"}
                with db:
                    db.execute("UPDATE results SET row_json=?", (json.dumps(bad),))
            finally:
                db.close()
            with self.assertRaises(ClaimCheckpointError):
                read_claim_checkpoint(path, identity)

    def test_payload_defaults_all_sites_live_and_rejects_invalid_scope_before_login(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            payload = {"username": "a", "password": "b", "source_file": str(source.source_path), "output_dir": directory}
            with patch.object(service, "load_claim_input", return_value=source):
                job = service.validate_bigseller_claim_payload(payload)
                self.assertEqual((job.site, job.listing_scope), ("all", "live"))
                for extra in ({"site": "XX"}, {"listing_scope": "all"}, {"output_dir": str(source.source_path)}):
                    with self.subTest(extra=extra), self.assertRaises(ValueError):
                        service.validate_bigseller_claim_payload({**payload, **extra})

    def test_progress_reports_confirmed_count_and_does_not_count_failed_attempt(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            progress = Mock()
            with self.runtime(source, [service.BigSellerQueryError("失败"), result("other", "not_found")]):
                value = service.run_bigseller_claim_query(self.job(source, directory), progress)
            updates = [call.args[0] for call in progress.call_args_list if "BS认领进度" in call.args[0]]
            self.assertTrue(updates[0].startswith("[BS认领进度 0/2]"))
            self.assertTrue(updates[1].startswith("[BS认领进度 1/2]"))
            self.assertEqual(value["confirmed_count"], 1)

    def test_initial_login_failure_explains_unexecuted_rows_and_persists_diagnostic(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            job = self.job(source, directory)
            with self.runtime(source, []) as (query, login, _):
                login.side_effect = service.BigSellerAuthenticationError("账号登录失败")
                value = service.run_bigseller_claim_query(job, Mock())
            query.assert_not_called()
            self.assertTrue(all("账号登录失败" in row["message"] for row in value["rows"]))
            rows, _ = read_claim_checkpoint(Path(value["checkpoint_file"]), service.checkpoint_identity(job, source))
            self.assertTrue(all("账号登录失败" in row["message"] for row in rows.values()))

    def test_shorter_cooldown_never_overwrites_existing_longer_deadline(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            identity = service.checkpoint_identity(self.job(source, directory), source)
            checkpoint = create_claim_checkpoint(Path(directory), identity)
            checkpoint.save_deadline(200.0)
            checkpoint.save_deadline(100.0)
            checkpoint.close()
            _, deadline = read_claim_checkpoint(checkpoint.path, identity)
            self.assertEqual(deadline, 200.0)

    def test_process_crash_during_next_save_recovers_commits_without_modifying_source(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            identity = service.checkpoint_identity(self.job(source, directory), source)
            checkpoint = create_claim_checkpoint(Path(directory), identity)
            checkpoint.save_result(result())
            path = checkpoint.path
            checkpoint.close()
            code = """import os, sqlite3, sys
connection = sqlite3.connect(sys.argv[1])
connection.execute('PRAGMA cache_size=1')
connection.execute('BEGIN IMMEDIATE')
connection.execute('UPDATE results SET row_json=?', ('x' * 100000,))
os._exit(0)
"""
            subprocess.run([sys.executable, "-c", code, str(path)], check=True, capture_output=True, timeout=20)
            journal = Path(str(path) + "-journal")
            before, journal_before = path.read_bytes(), journal.read_bytes()
            rows, _ = read_claim_checkpoint(path, identity)
            self.assertEqual(rows["sku"], result())
            self.assertEqual(path.read_bytes(), before)
            self.assertEqual(journal.read_bytes(), journal_before)

    def test_real_workbook_pipeline_preserves_source_duplicate_rows_and_existing_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            source_path = Path(directory) / "新品.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "新品"
            sheet.append(["SKU", "上架时间", "计算"])
            sheet.append(["SKU001", "手工时间", "=1+2"])
            sheet.append(["SKU001", "原始内容", "=3+4"])
            book.save(source_path)
            book.close()
            before = source_path.read_bytes()
            job = service.validate_bigseller_claim_payload({
                "username": "account", "password": "secret", "source_file": str(source_path), "output_dir": directory,
            })
            with patch.object(service, "query_process_config", return_value={}), \
                    patch.object(service, "load_bigseller_request_util", return_value=Mock()), \
                    patch.object(service, "login_for_query", return_value=Mock()), \
                    patch.object(service, "query_claim_sku", return_value=result("SKU001")) as query:
                value = service.run_bigseller_claim_query(job, Mock())
            self.assertTrue(value["is_complete"])
            self.assertEqual(value["sku_count"], 1)
            self.assertEqual(value["total_rows"], 2)
            query.assert_called_once()
            self.assertEqual(source_path.read_bytes(), before)
            output = load_workbook(value["output_file"])
            try:
                sheet = output["新品"]
                headers = {cell.value: cell.column for cell in sheet[1]}
                self.assertEqual(sheet.cell(2, headers["BS最早创建店铺"]).value, "店铺A")
                self.assertEqual(sheet.cell(3, headers["BS最早创建店铺"]).value, "店铺A")
                self.assertEqual(sheet["B2"].value, "手工时间")
                self.assertEqual(sheet["C3"].value, "=3+4")
            finally:
                output.close()


if __name__ == "__main__":
    unittest.main()
