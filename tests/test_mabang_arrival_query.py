from __future__ import annotations

import tempfile
import unittest
from collections import Counter
from io import BytesIO
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import parse_qs, urlparse
from unittest.mock import patch

import httpx
from openpyxl import Workbook, load_workbook

from backend.core.mabang_client import MabangApiError
from backend.services.mabang_arrival_query import (
    ARRIVAL_OUTPUT_DIR_NAME,
    EXPORT_ACTION_PATH,
    EXPORT_ENTRY_PATH,
    ExportTemplate,
    OLE2_MAGIC,
    SEARCH_PATH,
    SHIPMENTS_PAGE_PATH,
    build_arrival_output_filename,
    build_arrival_search_payload,
    build_export_entry_payload,
    build_export_step1_payload,
    parse_arrival_batches,
    parse_arrival_total_pages,
    parse_export_template,
    parse_remark_lines,
    resolve_arrival_output_dir,
    run_mabang_arrival_query,
    _run_export_steps,
    validate_arrival_query_payload,
    validate_arrival_export_bytes,
    validate_excel_bytes,
    validate_xlsx_bytes,
)


def _xlsx_bytes(
    batch_codes: Sequence[str] = ("BATCH-1",),
    *,
    remark: str = "备注一",
) -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "到货查询"
    worksheet.append(["批次编号", "备注"])
    for batch_code in batch_codes:
        worksheet.append([batch_code, remark])
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def _xls_bytes() -> bytes:
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    header[24:26] = (0x003E).to_bytes(2, "little")
    header[26:28] = (3).to_bytes(2, "little")
    header[28:30] = b"\xfe\xff"
    header[30:32] = (9).to_bytes(2, "little")
    return bytes(header)


def _batch_html(allocation_id: str, batch_code: str, remark: str) -> str:
    return f"""
    <li class="allocation-row">
      <ul id="allowlist-ul">
        <li><input class="allow" name="allot[]" value="{allocation_id}"
                   data-code="{batch_code}"></li>
        <li><textarea class="form-control fhRemark">{remark}</textarea></li>
      </ul>
    </li>
    """


def _template_html(field_count: int = 49) -> str:
    fields = [
        '<label><input name="fieldlabel" value="all">全选</label>',
        *[
            f'<label class="checkbox-inline"><input name="fieldlabel" '
            f'value="uq{index:03d}"><span>字段{index:02d}</span></label>'
            for index in range(1, field_count + 1)
        ],
    ]
    return f"""
    <html><body>
      <form id="theform">
        <input type="hidden" name="csrf" value="token-1">
        <input type="hidden" name="templateId" value="0">
        <input type="hidden" name="orderIds" value="stale-id">
        <input type="hidden" name="isMerage" value="1">
        <input type="hidden" name="sn" value="stale-sn">
        <input type="hidden" name="taskId" value="stale-task">
        <input type="hidden" name="sub_no" value="99">
        {''.join(fields)}
      </form>
    </body></html>
    """


def _response(
    method: str,
    url: str,
    *,
    status: int = 200,
    html: str | None = None,
    json_data: dict[str, Any] | None = None,
    content: bytes | None = None,
    headers: dict[str, str] | None = None,
) -> httpx.Response:
    request = httpx.Request(method, url)
    if json_data is not None:
        return httpx.Response(status, json=json_data, headers=headers, request=request)
    if content is not None:
        return httpx.Response(status, content=content, headers=headers, request=request)
    return httpx.Response(status, text=html or "", headers=headers, request=request)


class FakeArrivalHttpClient:
    BATCH_CODES = {
        "101": "BATCH-101",
        "102": "BATCH-102",
        "201": "BATCH-201",
    }

    def __init__(
        self,
        *,
        async_export: bool,
        legacy_excel: bool = False,
        header_only_export_numbers: set[int] | None = None,
        step2_no_data_export_numbers: set[int] | None = None,
    ) -> None:
        self.async_export = async_export
        self.legacy_excel = legacy_excel
        self.header_only_export_numbers = set(header_only_export_numbers or set())
        self.step2_no_data_export_numbers = set(step2_no_data_export_numbers or set())
        self.calls: list[dict[str, Any]] = []
        self.search_calls: list[dict[str, str]] = []
        self.entry_calls: list[list[tuple[str, str]]] = []
        self.action_calls: list[list[tuple[str, str]]] = []
        self.events: list[str] = []
        self.step4_count = 0
        self.export_count = 0
        self.active_search_ids: set[str] = set()
        self.active_remark = ""
        self.downloads: dict[str, bytes] = {}
        self.exports_by_sn: dict[str, dict[str, str]] = {}
        self.exports_by_task: dict[str, dict[str, str]] = {}
        self.step4_counts_by_task: Counter[str] = Counter()

    def get(self, url: str, **kwargs: Any) -> httpx.Response:
        self.calls.append({"method": "GET", "url": url, "kwargs": kwargs})
        if SHIPMENTS_PAGE_PATH in url:
            return _response("GET", url, html="<html><body>shipment page</body></html>")
        download_path = urlparse(url).path
        if download_path in self.downloads:
            export_number = Path(download_path).stem.rsplit("-", 1)[-1]
            self.events.append(f"download:{export_number}")
            suffix = ".xls" if self.legacy_excel else ".xlsx"
            return _response(
                "GET",
                url,
                content=self.downloads[download_path],
                headers={
                    "Content-Type": (
                        "application/vnd.ms-excel"
                        if self.legacy_excel
                        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    ),
                    "Content-Disposition": (
                        "attachment; filename*=UTF-8''%E5%88%B0%E8%B4%A7%E6%9F%A5%E8%AF%A2"
                        f"{suffix}"
                    ),
                },
            )
        raise AssertionError(f"unexpected GET: {url}")

    def post(
        self,
        url: str,
        *,
        content: str | bytes = "",
        headers: dict[str, str] | None = None,
        **kwargs: Any,
    ) -> httpx.Response:
        body = content.decode() if isinstance(content, bytes) else str(content or "")
        pairs = [
            (key, value)
            for key, values in parse_qs(body, keep_blank_values=True).items()
            for value in values
        ]
        form = {
            key: values[-1]
            for key, values in parse_qs(body, keep_blank_values=True).items()
        }
        self.calls.append(
            {"method": "POST", "url": url, "content": body, "headers": headers or {}}
        )

        if SEARCH_PATH in url:
            self.search_calls.append(form)
            remark = form["search-content-text1"]
            page = form["page"]
            self.active_remark = remark
            self.events.append(f"search:{remark}:{page}")
            if remark == "备注一" and page == "1":
                self.active_search_ids = {"101", "102"}
                payload = {
                    "success": True,
                    "message": _batch_html("101", "BATCH-101", "备注一"),
                    "pageHtml": '<span class="semibold">1/2</span>页 共 3 条',
                }
            elif remark == "备注一" and page == "2":
                self.active_search_ids = {"101", "102"}
                payload = {
                    "success": True,
                    "message": (
                        _batch_html("101", "BATCH-101", "备注一")
                        + _batch_html("102", "BATCH-102", "备注一-补充")
                    ),
                    "pageHtml": '<span class="semibold">2/2</span>页 共 3 条',
                }
            elif remark == "未匹配":
                self.active_search_ids = set()
                payload = {
                    "success": True,
                    "message": '<div class="group-nodata">暂无数据</div>',
                    "pageHtml": False,
                }
            elif remark == "备注二" and page == "1":
                self.active_search_ids = {"102", "201"}
                payload = {
                    "success": True,
                    "message": (
                        _batch_html("102", "BATCH-102", "备注一-补充")
                        + _batch_html("201", "BATCH-201", "备注二")
                    ),
                    "pageHtml": '<span class="semibold">1/1</span>页 共 2 条',
                }
            else:
                raise AssertionError(f"unexpected search: {remark}, page={page}")
            return _response("POST", url, json_data=payload)

        if EXPORT_ENTRY_PATH in url:
            self.entry_pairs = pairs
            self.entry_calls.append(pairs)
            self.events.append("entry")
            return _response("POST", url, html=_template_html())

        if EXPORT_ACTION_PATH in url:
            self.action_calls.append(pairs)
            step = form.get("step")
            self.events.append(f"action:{step}")
            if step == "1":
                self.export_count += 1
                export_number = self.export_count
                allocation_ids = [
                    value for value in form.get("orderIds", "").split(",") if value
                ]
                has_active_context = set(allocation_ids).issubset(self.active_search_ids)
                header_only = (
                    export_number in self.header_only_export_numbers
                    or not has_active_context
                )
                batch_codes = (
                    []
                    if header_only
                    else [self.BATCH_CODES[value] for value in allocation_ids]
                )
                suffix = "xls" if self.legacy_excel else "xlsx"
                download_path = f"/download/arrival-{export_number}.{suffix}"
                self.downloads[download_path] = (
                    _xls_bytes()
                    if self.legacy_excel
                    else _xlsx_bytes(batch_codes, remark=self.active_remark)
                )
                sn = f"SN-ARRIVAL-{export_number}"
                task_id = f"TASK-{export_number}"
                export = {
                    "sn": sn,
                    "taskId": task_id,
                    "file_url": download_path,
                }
                self.exports_by_sn[sn] = export
                self.exports_by_task[task_id] = export
                if self.async_export:
                    payload = {
                        "success": True,
                        "success_type": 2,
                        "sn": sn,
                        "subtask_num": 2,
                    }
                else:
                    payload = {
                        "success": True,
                        "success_type": 1,
                        "gourl": download_path,
                    }
            elif step == "2":
                if form.get("sn") not in self.exports_by_sn:
                    raise AssertionError(f"unknown export sn: {form}")
                export_number = int(form["sn"].rsplit("-", 1)[-1])
                payload = {
                    "success": True,
                    "message": (
                        "请选择填写要导出的数据"
                        if export_number in self.step2_no_data_export_numbers
                        else "完成"
                    ),
                }
            elif step == "3":
                export = self.exports_by_sn[form.get("sn", "")]
                payload = {
                    "success": True,
                    "async": True,
                    "taskId": export["taskId"],
                }
            elif step == "4":
                export = self.exports_by_task[form.get("taskId", "")]
                task_id = export["taskId"]
                self.step4_count += 1
                self.step4_counts_by_task[task_id] += 1
                # Mabang can expose an allocated URL before the file is ready.
                # The service must wait for state=true instead of downloading it.
                payload = {
                    "success": True,
                    "state": self.step4_counts_by_task[task_id] > 1,
                    "file_url": export["file_url"],
                }
            else:
                raise AssertionError(f"unexpected export step: {form}")
            return _response("POST", url, json_data=payload)

        raise AssertionError(f"unexpected POST: {url}")


class FakeMabangClient:
    instances: list["FakeMabangClient"] = []
    async_export = True
    legacy_excel = False
    header_only_export_numbers: set[int] = set()
    step2_no_data_export_numbers: set[int] = set()

    def __init__(self, base_url: str, **_kwargs: Any) -> None:
        self.base_url = base_url.rstrip("/")
        self.client = FakeArrivalHttpClient(
            async_export=self.async_export,
            legacy_excel=self.legacy_excel,
            header_only_export_numbers=self.header_only_export_numbers,
            step2_no_data_export_numbers=self.step2_no_data_export_numbers,
        )
        self.login_args: tuple[str, str] | None = None
        self.__class__.instances.append(self)

    def __enter__(self) -> "FakeMabangClient":
        return self

    def __exit__(self, *_args: Any) -> None:
        return None

    def login(self, username: str, password: str) -> None:
        self.login_args = (username, password)

    def _url(self, path: str) -> str:
        if path.startswith(("http://", "https://")):
            return path
        return self.base_url + path


class MabangArrivalQueryTests(unittest.TestCase):
    def setUp(self) -> None:
        FakeMabangClient.instances.clear()
        FakeMabangClient.async_export = True
        FakeMabangClient.legacy_excel = False
        FakeMabangClient.header_only_export_numbers = set()
        FakeMabangClient.step2_no_data_export_numbers = set()

    def test_parse_and_validate_multiline_remarks_stably_deduplicates(self) -> None:
        self.assertEqual(
            parse_remark_lines(" 备注一\r\n\n备注二\n备注一 "),
            ("备注一", "备注二"),
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            job = validate_arrival_query_payload(
                {
                    "username": " user ",
                    "password": "secret",
                    "remarks": ["备注一\n备注二", "备注一"],
                    "output_dir": temp_dir,
                    "rows_per_page": 999,
                    "poll_interval_seconds": 0,
                }
            )
        self.assertEqual(job.username, "user")
        self.assertEqual(job.remarks, ("备注一", "备注二"))
        self.assertEqual(job.rows_per_page, 500)
        self.assertEqual(job.poll_interval_seconds, 0)

    def test_validate_rejects_missing_credentials_remarks_and_output(self) -> None:
        with self.assertRaisesRegex(ValueError, "马帮账号"):
            validate_arrival_query_payload({})
        with self.assertRaisesRegex(ValueError, "马帮密码"):
            validate_arrival_query_payload({"username": "user"})
        with self.assertRaisesRegex(ValueError, "至少输入一个备注"):
            validate_arrival_query_payload({"username": "user", "password": "pw"})
        with self.assertRaisesRegex(ValueError, "输出目录"):
            validate_arrival_query_payload(
                {"username": "user", "password": "pw", "remarks": "备注"}
            )

    def test_search_payload_uses_all_status_and_remark_mode(self) -> None:
        payload = build_arrival_search_payload("测试备注", page=3, rows_per_page=200)
        self.assertEqual(payload["allocationstatus"], "")
        self.assertEqual(payload["search-content1"], "remark")
        self.assertEqual(payload["search-content-text1"], "测试备注")
        self.assertEqual(payload["page"], "3")
        self.assertEqual(payload["rowsPerPage"], "200")

    def test_batch_and_pagination_parsers_read_required_source_fields(self) -> None:
        html = _batch_html("2505787", "2026072016520080", "7.20菲律宾2仓敏感3件")
        batches = parse_arrival_batches(html + html)
        self.assertEqual(len(batches), 1)
        self.assertEqual(batches[0].allocation_id, "2505787")
        self.assertEqual(batches[0].batch_code, "2026072016520080")
        self.assertEqual(batches[0].data_code, "2026072016520080")
        self.assertEqual(batches[0].remark, "7.20菲律宾2仓敏感3件")
        self.assertEqual(
            parse_arrival_total_pages(
                '共<span class="semibold">12,345</span>条 '
                '<span class="semibold">1/124</span>页'
            ),
            124,
        )
        self.assertEqual(parse_arrival_total_pages(False), 0)

    def test_template_discovers_all_49_fields_and_builds_parallel_arrays(self) -> None:
        template = parse_export_template(
            _template_html(),
            page_url="https://example.test/export-template",
        )
        self.assertEqual(len(template.fields), 49)
        self.assertNotIn("all", [field.code for field in template.fields])
        payload = build_export_step1_payload(template, ["101", "102", "101"])
        counts = Counter(key for key, _ in payload)
        self.assertEqual(counts["fieldlabel"], 49)
        self.assertEqual(counts["map-uq[]"], 49)
        self.assertEqual(counts["map-name[]"], 49)
        self.assertEqual(counts["map-text[]"], 49)
        payload_map = {key: value for key, value in payload}
        self.assertEqual(payload_map["orderIds"], "101,102,")
        self.assertEqual(payload_map["isMerage"], "2")
        self.assertEqual(payload_map["version"], "v2")
        self.assertEqual(payload_map["step"], "1")
        self.assertEqual(payload_map["csrf"], "token-1")
        self.assertNotIn("sn", payload_map)
        self.assertNotIn("taskId", payload_map)
        self.assertNotIn("sub_no", payload_map)

    def test_export_entry_has_exact_allocation_warehouse_contract(self) -> None:
        client = FakeMabangClient("https://example.test")
        payload = dict(build_export_entry_payload(client, ["101", "102", "101"]))
        self.assertEqual(payload["mod"], "export.exportTemplate")
        self.assertEqual(payload["datasOpen"], "2")
        self.assertEqual(payload["data"], "101,102,")
        self.assertEqual(payload["type"], "1")
        self.assertEqual(payload["menu"], "allocationWarehouse")
        self.assertEqual(
            payload["exportUrl"],
            "https://example.test/index.php?mod=export.doAllocationWarehouseExportFile",
        )
        self.assertEqual(payload["mainMenu"], "")
        self.assertEqual(payload["showRmbColumn"], "0")

    def test_module_output_directory_is_fixed_and_not_nested_twice(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output_root = Path(temp_dir)
            module_dir = output_root / ARRIVAL_OUTPUT_DIR_NAME
            self.assertEqual(resolve_arrival_output_dir(output_root), module_dir)
            self.assertEqual(resolve_arrival_output_dir(module_dir), module_dir)

    def test_output_filename_sanitizes_search_name_and_keeps_real_suffix(self) -> None:
        self.assertEqual(
            build_arrival_output_filename(
                '  报表/一:*?"<>|   名称.  ',
                timestamp="2026/08/25 10:20:30",
                suffix=".XLS",
            ),
            "到货查询_20260825102030_报表_一_ 名称.xls",
        )
        self.assertEqual(
            build_arrival_output_filename(
                "CON",
                timestamp="20260825_102030",
                suffix=".xlsx",
            ),
            "到货查询_20260825_102030__CON.xlsx",
        )
        self.assertEqual(
            build_arrival_output_filename(
                ':*?"<>|',
                timestamp="20260825_102030",
                suffix=".csv",
            ),
            "到货查询_20260825_102030_未命名.xlsx",
        )
        long_name = build_arrival_output_filename(
            "长" * 120,
            timestamp="20260825_102030",
            suffix=".xlsx",
        )
        component = Path(long_name).stem.removeprefix("到货查询_20260825_102030_")
        self.assertEqual(len(component), 80)

    def test_full_async_flow_is_serial_deduplicated_and_saves_one_matched_remark_file(self) -> None:
        progress: list[str] = []
        with tempfile.TemporaryDirectory() as temp_dir:
            job = validate_arrival_query_payload(
                {
                    "username": "user",
                    "password": "password",
                    "remarks": "备注一\n未匹配\n备注一",
                    "output_dir": temp_dir,
                    "base_url": "https://example.test",
                    "poll_interval_seconds": 0,
                    "poll_timeout_seconds": 1,
                }
            )
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                result = run_mabang_arrival_query(job, progress.append)

            self.assertEqual(result.remark_count, 2)
            self.assertEqual(result.matched_remark_count, 1)
            self.assertEqual(result.matched_batch_count, 2)
            self.assertEqual(result.exported_batch_count, 2)
            self.assertEqual(result.unmatched_remarks, ("未匹配",))
            self.assertEqual([batch.allocation_id for batch in result.batches], ["101", "102"])
            self.assertEqual(result.batches[0].matched_remarks, ("备注一",))
            self.assertEqual(len(result.exports), 1)
            self.assertEqual(result.exports[0].remark, "备注一")
            self.assertEqual(result.exports[0].matched_batch_count, 2)
            self.assertEqual(result.output_files, (result.output_file,))
            self.assertTrue(result.output_file.is_file())
            self.assertEqual(result.output_root, Path(temp_dir))
            self.assertEqual(
                result.output_dir,
                Path(temp_dir) / ARRIVAL_OUTPUT_DIR_NAME,
            )
            self.assertEqual(list(Path(temp_dir).glob("*.xlsx")), [])
            self.assertEqual(list(result.output_dir.glob("*.xlsx")), [result.output_file])
            self.assertIn("备注一", result.output_file.name)
            workbook = load_workbook(result.output_file, read_only=True)
            self.assertEqual(workbook.sheetnames, ["到货查询"])
            workbook.close()

        fake_http = FakeMabangClient.instances[-1].client
        self.assertEqual(
            [(call["search-content-text1"], call["page"]) for call in fake_http.search_calls],
            [("备注一", "1"), ("备注一", "2"), ("未匹配", "1")],
        )
        self.assertTrue(all(call["allocationstatus"] == "" for call in fake_http.search_calls))
        entry = {key: value for key, value in fake_http.entry_pairs}
        self.assertEqual(entry["data"], "101,102,")
        step1_pairs = fake_http.action_calls[0]
        self.assertEqual(sum(key == "fieldlabel" for key, _ in step1_pairs), 49)
        steps = [dict(call).get("step") for call in fake_http.action_calls]
        self.assertEqual(steps, ["1", "2", "2", "3", "4", "4"])
        self.assertEqual(fake_http.step4_count, 2)
        self.assertEqual(
            [event for event in fake_http.events if event.startswith("download:")],
            ["download:1"],
        )
        self.assertTrue(any("匹配 2 个唯一调拨批次" in message for message in progress))

    def test_each_matched_remark_exports_independently_and_returns_multiple_files(self) -> None:
        progress: list[str] = []
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                result = run_mabang_arrival_query(
                    {
                        "username": "user",
                        "password": "password",
                        "remarks": "备注一\n备注二\n未匹配",
                        "output_dir": temp_dir,
                        "base_url": "https://example.test",
                        "poll_interval_seconds": 0,
                    },
                    progress.append,
                )

            self.assertEqual(result.remark_count, 3)
            self.assertEqual(result.matched_remark_count, 2)
            self.assertEqual(result.matched_batch_count, 3)
            self.assertEqual(result.exported_batch_count, 4)
            self.assertEqual(result.unmatched_remarks, ("未匹配",))
            self.assertEqual(
                [batch.allocation_id for batch in result.batches],
                ["101", "102", "201"],
            )
            self.assertEqual(result.batches[1].matched_remarks, ("备注一", "备注二"))
            self.assertEqual([item.remark for item in result.exports], ["备注一", "备注二"])
            self.assertEqual(
                [item.matched_batch_count for item in result.exports],
                [2, 2],
            )
            self.assertEqual(
                [[batch.allocation_id for batch in item.batches] for item in result.exports],
                [["101", "102"], ["102", "201"]],
            )
            self.assertEqual(len(result.output_files), 2)
            self.assertEqual(result.output_file, result.output_files[0])
            self.assertEqual(result.output_dir, Path(temp_dir) / ARRIVAL_OUTPUT_DIR_NAME)
            self.assertTrue(all(path.is_file() for path in result.output_files))
            self.assertTrue(all(path.parent == result.output_dir for path in result.output_files))
            self.assertTrue(all(path.suffix == ".xlsx" for path in result.output_files))
            self.assertIn("备注一", result.output_files[0].name)
            self.assertIn("备注二", result.output_files[1].name)
            exported_batch_codes: list[list[str]] = []
            for output_file in result.output_files:
                workbook = load_workbook(output_file, read_only=True, data_only=True)
                try:
                    rows = list(workbook.active.iter_rows(values_only=True))
                finally:
                    workbook.close()
                self.assertGreater(len(rows), 1)
                exported_batch_codes.append(
                    [str(row[0]) for row in rows[1:] if row and row[0]]
                )
            self.assertEqual(
                exported_batch_codes,
                [["BATCH-101", "BATCH-102"], ["BATCH-102", "BATCH-201"]],
            )

        fake_http = FakeMabangClient.instances[-1].client
        self.assertEqual(
            [(call["search-content-text1"], call["page"]) for call in fake_http.search_calls],
            [
                ("备注一", "1"),
                ("备注一", "2"),
                ("备注二", "1"),
                ("未匹配", "1"),
            ],
        )
        self.assertEqual(
            [dict(pairs)["data"] for pairs in fake_http.entry_calls],
            ["101,102,", "102,201,"],
        )
        step1_calls = [
            dict(pairs)
            for pairs in fake_http.action_calls
            if dict(pairs).get("step") == "1"
        ]
        self.assertEqual(
            [payload["orderIds"] for payload in step1_calls],
            ["101,102,", "102,201,"],
        )
        self.assertEqual(
            [dict(call).get("sn") for call in fake_http.action_calls if dict(call).get("step") == "3"],
            ["SN-ARRIVAL-1", "SN-ARRIVAL-2"],
        )
        self.assertEqual(
            fake_http.events,
            [
                "search:备注一:1",
                "search:备注一:2",
                "entry",
                "action:1",
                "action:2",
                "action:2",
                "action:3",
                "action:4",
                "action:4",
                "download:1",
                "search:备注二:1",
                "entry",
                "action:1",
                "action:2",
                "action:2",
                "action:3",
                "action:4",
                "action:4",
                "download:2",
                "search:未匹配:1",
            ],
        )
        self.assertTrue(any("共生成 2 个 Excel 文件" in message for message in progress))

    def test_success_type_one_downloads_without_async_steps(self) -> None:
        FakeMabangClient.async_export = False
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                result = run_mabang_arrival_query(
                    {
                        "username": "user",
                        "password": "password",
                        "remarks": "备注一",
                        "output_dir": temp_dir,
                        "base_url": "https://example.test",
                        "poll_interval_seconds": 0,
                    }
                )
            self.assertEqual(result.output_root, Path(temp_dir))
            self.assertEqual(result.output_dir, Path(temp_dir) / ARRIVAL_OUTPUT_DIR_NAME)
            self.assertEqual(len(result.output_files), 1)
            self.assertTrue(result.output_file.is_file())
        fake_http = FakeMabangClient.instances[-1].client
        self.assertTrue(result.output_file.name.endswith(".xlsx"))
        self.assertEqual([dict(call).get("step") for call in fake_http.action_calls], ["1"])

    def test_xlsx_validation_rejects_html_and_invalid_archives(self) -> None:
        with self.assertRaisesRegex(MabangApiError, "HTML"):
            validate_xlsx_bytes(
                b"<!doctype html><html><body>login</body></html>",
                content_type="text/html",
            )
        with self.assertRaisesRegex(MabangApiError, "XLSX"):
            validate_xlsx_bytes(b"not a zip file", content_type="application/octet-stream")
        validate_xlsx_bytes(_xlsx_bytes())

    def test_excel_validation_accepts_mabang_legacy_xls(self) -> None:
        self.assertEqual(
            validate_excel_bytes(
                _xls_bytes(),
                content_type="application/vnd.ms-excel",
            ),
            ".xls",
        )
        validate_xlsx_bytes(_xls_bytes(), content_type="application/vnd.ms-excel")

    def test_arrival_content_validation_rejects_header_only_and_wrong_batch(self) -> None:
        with self.assertRaisesRegex(MabangApiError, "无数据"):
            validate_arrival_export_bytes(
                _xlsx_bytes([]),
                expected_batch_codes=["BATCH-101"],
            )
        with self.assertRaisesRegex(MabangApiError, "不一致"):
            validate_arrival_export_bytes(
                _xlsx_bytes(["BATCH-999"]),
                expected_batch_codes=["BATCH-101"],
            )
        with self.assertRaisesRegex(MabangApiError, "额外批次"):
            validate_arrival_export_bytes(
                _xlsx_bytes(["BATCH-101", "BATCH-999"]),
                expected_batch_codes=["BATCH-101"],
            )
        with self.assertRaisesRegex(MabangApiError, "缺少批次编号"):
            validate_arrival_export_bytes(
                _xlsx_bytes(["BATCH-101"]),
                expected_batch_codes=[""],
            )
        self.assertEqual(
            validate_arrival_export_bytes(
                _xlsx_bytes(["BATCH-101", "BATCH-102"]),
                expected_batch_codes=["BATCH-101", "BATCH-102"],
            ),
            ".xlsx",
        )

    def test_header_only_export_is_requeried_and_retried_once(self) -> None:
        FakeMabangClient.async_export = False
        FakeMabangClient.header_only_export_numbers = {1}
        progress: list[str] = []
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                result = run_mabang_arrival_query(
                    {
                        "username": "user",
                        "password": "password",
                        "remarks": "备注一",
                        "output_dir": temp_dir,
                        "base_url": "https://example.test",
                        "poll_interval_seconds": 0,
                    },
                    progress.append,
                )
            workbook = load_workbook(result.output_file, read_only=True, data_only=True)
            try:
                rows = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
            self.assertEqual([row[0] for row in rows[1:]], ["BATCH-101", "BATCH-102"])

        fake_http = FakeMabangClient.instances[-1].client
        self.assertEqual(len(fake_http.entry_calls), 2)
        self.assertEqual(
            [(call["search-content-text1"], call["page"]) for call in fake_http.search_calls],
            [("备注一", "1"), ("备注一", "2"), ("备注一", "1"), ("备注一", "2")],
        )
        self.assertTrue(any("重新查询并重试" in message for message in progress))

    def test_step2_no_data_requeries_before_retrying_the_whole_export(self) -> None:
        FakeMabangClient.step2_no_data_export_numbers = {1}
        progress: list[str] = []
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                result = run_mabang_arrival_query(
                    {
                        "username": "user",
                        "password": "password",
                        "remarks": "备注一",
                        "output_dir": temp_dir,
                        "base_url": "https://example.test",
                        "poll_interval_seconds": 0,
                        "poll_timeout_seconds": 1,
                    },
                    progress.append,
                )
            self.assertTrue(result.output_file.is_file())

        fake_http = FakeMabangClient.instances[-1].client
        self.assertEqual(len(fake_http.entry_calls), 2)
        self.assertEqual(
            [(call["search-content-text1"], call["page"]) for call in fake_http.search_calls],
            [("备注一", "1"), ("备注一", "2"), ("备注一", "1"), ("备注一", "2")],
        )
        self.assertTrue(any("重新查询并重试" in message for message in progress))

    def test_step3_failure_with_stale_task_id_never_polls_step4(self) -> None:
        client = FakeMabangClient("https://example.test")
        template = ExportTemplate(
            page_url="https://example.test/export-template",
            hidden_defaults=(),
            fields=(),
        )
        responses = [
            {"success": True, "success_type": 2, "sn": "SN-1", "subtask_num": 1},
            {"success": True, "message": "完成"},
            {"success": False, "taskId": "STALE-TASK"},
        ]
        with patch(
            "backend.services.mabang_arrival_query._post_export_action",
            side_effect=responses,
        ) as post_action:
            with self.assertRaisesRegex(MabangApiError, "step3"):
                _run_export_steps(
                    client,
                    template=template,
                    allocation_ids=["101"],
                    step_retries=1,
                    poll_interval_seconds=0,
                    poll_timeout_seconds=1,
                )
        self.assertEqual(post_action.call_count, 3)

    def test_step4_terminal_url_without_state_uses_content_validation_fallback(self) -> None:
        client = FakeMabangClient("https://example.test")
        template = ExportTemplate(
            page_url="https://example.test/export-template",
            hidden_defaults=(),
            fields=(),
        )
        responses = [
            {"success": True, "success_type": 2, "sn": "SN-1", "subtask_num": 1},
            {"success": True, "message": "完成"},
            {"success": True, "taskId": "TASK-1"},
            {"success": True, "file_url": "/download/arrival-1.xlsx"},
        ]
        with patch(
            "backend.services.mabang_arrival_query._post_export_action",
            side_effect=responses,
        ):
            payload = _run_export_steps(
                client,
                template=template,
                allocation_ids=["101"],
                step_retries=1,
                poll_interval_seconds=0,
                poll_timeout_seconds=1,
            )
        self.assertEqual(payload["file_url"], "/download/arrival-1.xlsx")

    def test_legacy_xls_business_rows_are_checked_with_corruption_tolerance(self) -> None:
        class FakeSheet:
            nrows = 2

            @staticmethod
            def row_values(index: int) -> list[str]:
                return ["批次编号", "备注"] if index == 0 else ["BATCH-101", "备注一"]

        class FakeBook:
            nsheets = 1

            def __init__(self) -> None:
                self.released = False

            @staticmethod
            def sheet_by_index(_index: int) -> FakeSheet:
                return FakeSheet()

            def release_resources(self) -> None:
                self.released = True

        book = FakeBook()
        with patch(
            "backend.services.mabang_arrival_query.xlrd.open_workbook",
            return_value=book,
        ) as open_workbook:
            suffix = validate_arrival_export_bytes(
                _xls_bytes(),
                expected_batch_codes=["BATCH-101"],
                content_type="application/vnd.ms-excel",
            )
        self.assertEqual(suffix, ".xls")
        self.assertTrue(book.released)
        self.assertTrue(open_workbook.call_args.kwargs["ignore_workbook_corruption"])
        self.assertEqual(open_workbook.call_args.kwargs["file_contents"], _xls_bytes())

    def test_persistent_header_only_export_fails_without_saving_a_file(self) -> None:
        FakeMabangClient.async_export = False
        FakeMabangClient.header_only_export_numbers = {1, 2}
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ):
                with self.assertRaisesRegex(MabangApiError, "无数据"):
                    run_mabang_arrival_query(
                        {
                            "username": "user",
                            "password": "password",
                            "remarks": "备注一",
                            "output_dir": temp_dir,
                            "base_url": "https://example.test",
                            "poll_interval_seconds": 0,
                        }
                    )
            self.assertEqual(
                list((Path(temp_dir) / ARRIVAL_OUTPUT_DIR_NAME).glob("*.xlsx")),
                [],
            )

    def test_full_flow_preserves_mabang_legacy_xls_extension(self) -> None:
        FakeMabangClient.async_export = False
        FakeMabangClient.legacy_excel = True
        with tempfile.TemporaryDirectory() as temp_dir:
            with patch(
                "backend.services.mabang_arrival_query.MabangClient",
                FakeMabangClient,
            ), patch(
                "backend.services.mabang_arrival_query.validate_arrival_export_bytes",
                return_value=".xls",
            ):
                result = run_mabang_arrival_query(
                    {
                        "username": "user",
                        "password": "password",
                        "remarks": "备注一",
                        "output_dir": temp_dir,
                        "base_url": "https://example.test",
                        "poll_interval_seconds": 0,
                    }
                )
            self.assertEqual(result.output_root, Path(temp_dir))
            self.assertEqual(result.output_dir, Path(temp_dir) / ARRIVAL_OUTPUT_DIR_NAME)
            self.assertEqual(result.output_files, (result.output_file,))
            self.assertEqual(result.output_file.suffix, ".xls")
            self.assertIn("备注一", result.output_file.name)
            self.assertTrue(result.output_file.read_bytes().startswith(OLE2_MAGIC))

    def test_service_source_has_no_browser_automation_dependency(self) -> None:
        source = (
            Path(__file__).resolve().parents[1]
            / "backend"
            / "services"
            / "mabang_arrival_query.py"
        ).read_text(encoding="utf-8").lower()
        for forbidden in ("selenium", "playwright", "webdriver", "pyautogui"):
            self.assertNotIn(forbidden, source)


if __name__ == "__main__":
    unittest.main()
