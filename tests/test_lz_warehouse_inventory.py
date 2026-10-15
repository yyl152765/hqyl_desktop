from __future__ import annotations

import struct
import tempfile
import traceback
import unittest
from io import BytesIO
from pathlib import Path
from typing import Any, Sequence
from urllib.parse import parse_qs
from unittest.mock import patch

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._writer import ALL_TEMP_FILES

from backend.core.mabang_client import MabangApiError, MabangClient
from backend.services import lz_warehouse_inventory as inventory


WAREHOUSES = (
    ("1093125", "泰国TH04仓"),
    ("1101002", "泰国TH05仓"),
    ("1102469", "泰国TH06仓"),
)
HEADERS = ("库存SKU编号", "商品中文名称", "仓库", "仓位库存", "原生额外字段")
SEARCH_TOKEN = "snapshot-token+with&special=characters"


def _inventory_html(warehouses: Sequence[tuple[str, str]] = WAREHOUSES) -> str:
    choices = [
        ("all", "全部仓库"),
        ("1093126", "泰国TH04中转仓"),
        ("1094491", "泰国TH04陆运仓（卡运）"),
        ("1101004", "泰国TH05中转仓"),
        ("1094492", "泰国TH06陆运仓（卡运）"),
        ("fake-name", "泰国TH06仓-测试"),
        *warehouses,
    ]
    checkboxes = "".join(
        f'<label><input type="checkbox" name="warehouseIds[]" value="{identity}">'
        f'<span class="text">{name}</span></label>'
        for identity, name in choices
    )
    return (
        '<html><body><form name="AdvanceSearch-form">'
        '<input type="hidden" name="stockOrderby" value="">'
        '<input name="stockNameValue" value="">'
        '<select name="stockName"><option value="nameCN" selected>商品中文名</option></select>'
        f'{checkboxes}</form></body></html>'
    )


def _xlsx_bytes(rows: Sequence[Sequence[Any]], headers: Sequence[str] = HEADERS) -> bytes:
    buffer = BytesIO()
    book = Workbook()
    sheet = book.active
    sheet.title = "库存统计"
    sheet.append(list(headers))
    for values in rows:
        sheet.append(list(values))
    book.save(buffer)
    book.close()
    return buffer.getvalue()


def _xls_bytes(rows: Sequence[Sequence[Any]], headers: Sequence[str] = HEADERS) -> bytes:
    """Build a small, genuine BIFF8/OLE workbook without an xlwt dependency."""

    def record(code: int, value: bytes = b"") -> bytes:
        return struct.pack("<HH", code, len(value)) + value

    def bof(kind: int) -> bytes:
        return record(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0x41, 6))

    def label(row: int, col: int, value: Any) -> bytes:
        text = str(value)
        encoded = text.encode("utf-16le")
        return record(0x0204, struct.pack("<HHHHB", row, col, 0, len(text), 1) + encoded)

    sheet_name = "库存统计".encode("utf-16le")
    bound_sheet_size = 4 + 8 + len(sheet_name)
    globals_prefix = bof(5) + record(0x0042, struct.pack("<H", 1200))
    sheet_offset = len(globals_prefix) + bound_sheet_size + 4
    globals_stream = globals_prefix + record(
        0x0085, struct.pack("<IBBBB", sheet_offset, 0, 0, len(sheet_name) // 2, 1) + sheet_name
    ) + record(0x000A)
    values = [tuple(headers), *rows]
    sheet_stream = bof(0x0010) + record(
        0x0200, struct.pack("<IIHHH", 0, len(values), 0, len(headers), 0)
    )
    for row_index, row in enumerate(values):
        for col_index, value in enumerate(row):
            if isinstance(value, (int, float)):
                sheet_stream += record(0x0203, struct.pack("<HHHd", row_index, col_index, 0, value))
            elif value is not None:
                sheet_stream += label(row_index, col_index, value)
    workbook_stream = globals_stream + sheet_stream + record(0x000A)
    if len(workbook_stream) > 4096:
        raise ValueError("The XLS test fixture must fit in eight OLE sectors")
    workbook_stream = workbook_stream.ljust(4096, b"\x00")

    end = 0xFFFFFFFE
    free = 0xFFFFFFFF
    header = bytearray(512)
    header[:8] = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
    struct.pack_into("<HHHHH", header, 24, 0x003E, 3, 0xFFFE, 9, 6)
    struct.pack_into("<IIIIIIIII", header, 40, 0, 1, 8, 0, 4096, end, 0, end, 0)
    struct.pack_into("<109I", header, 76, 9, *([free] * 108))

    def directory_entry(name: str, kind: int, start: int, size: int, child: int = free) -> bytes:
        entry = bytearray(128)
        encoded = (name + "\x00").encode("utf-16le")
        entry[:len(encoded)] = encoded
        struct.pack_into("<HBBIII", entry, 64, len(encoded), kind, 1, free, free, child)
        struct.pack_into("<IQ", entry, 116, start, size)
        return bytes(entry)

    directory = (
        directory_entry("Root Entry", 5, end, 0, child=1)
        + directory_entry("Workbook", 2, 0, 4096)
    ).ljust(512, b"\x00")
    fat_values = [*range(1, 8), end, end, 0xFFFFFFFD] + [free] * 118
    return bytes(header) + workbook_stream + directory + struct.pack("<128I", *fat_values)


class InventoryHttpFixture:
    def __init__(
        self,
        total: int,
        exports: dict[int, bytes],
        *,
        html: str | None = None,
        search_token: str = SEARCH_TOKEN,
        search_success: bool = True,
        scopes: dict[str, tuple[int, dict[int, bytes]]] | None = None,
    ) -> None:
        self.scopes = scopes if scopes is not None else {WAREHOUSES[0][0]: (total, exports)}
        self.html = html if html is not None else _inventory_html()
        self.search_token = search_token
        self.search_success = search_success
        self.active_warehouse: str | None = None
        self.requests: list[tuple[str, str, dict[str, list[str]]]] = []

    def token_for(self, warehouse_id: str) -> str:
        return f"{self.search_token}:{warehouse_id}" if self.search_token else ""

    def handle(self, request: httpx.Request) -> httpx.Response:
        query = parse_qs(request.url.query.decode(), keep_blank_values=True)
        form = parse_qs(request.content.decode(), keep_blank_values=True) if request.content else {}
        mod = query.get("mod", [""])[0]
        self.requests.append((request.method, mod, form if request.method == "POST" else query))
        if mod == "main.doLogin":
            if form.get("username") != ["test-account"] or form.get("password") != ["test-password"]:
                raise AssertionError("Only fixture credentials may be used")
            return httpx.Response(200, json={"success": True})
        if mod == "main.plogin" or request.url.path == "/index.htm":
            return httpx.Response(200, text="<html><body>logged in</body></html>")
        if mod == "warehouse.inventorydetail":
            if request.url.host != "private-amz.mabangerp.com":
                return httpx.Response(
                    200,
                    text='<iframe id="iframeContent" src="https://private-amz.mabangerp.com/index.php?mod=warehouse.inventorydetail"></iframe>',
                )
            return httpx.Response(200, text=self.html)
        if mod == "warehouse.searchwarehousestock":
            selected = form.get("warehouseIds[]", [])
            if len(selected) != 1 or selected[0] not in self.scopes:
                raise AssertionError(f"Each search must query exactly one selected warehouse: {selected}")
            self.active_warehouse = selected[0]
            return httpx.Response(200, json={
                "success": self.search_success,
                "message": "<table><tr><td>库存</td></tr></table>" if self.search_success else "没有权限",
                "searchToken": self.token_for(self.active_warehouse),
            })
        if mod == "warehouse.getSearchWarehouseStockPage":
            if self.active_warehouse is None:
                raise AssertionError("Pagination must follow a warehouse search")
            total, _exports = self.scopes[self.active_warehouse]
            page_html = (
                f'<div class="static-text">共<span class="semibold">{total:,}</span>条 '
                f'<span class="semibold">1/{max(1, (total + 49) // 50)}</span>页</div>'
            )
            return httpx.Response(200, json={"success": True, "pageHtml": page_html})
        if mod == "warehouse.doexportwarehousestock":
            if self.active_warehouse is None:
                raise AssertionError("Export must follow a warehouse search")
            if query.get("searchToken") != [self.token_for(self.active_warehouse)]:
                raise AssertionError("Export token belongs to a different warehouse search")
            _total, exports = self.scopes[self.active_warehouse]
            page = int(query.get("page", ["1"])[0])
            if page not in exports:
                raise AssertionError(f"Unexpected export page: {page}")
            return httpx.Response(200, content=exports[page], headers={
                "Content-Type": "application/vnd.ms-excel",
                "Content-Disposition": f'attachment; filename="inventory-page-{page}.xls"',
            })
        raise AssertionError(f"Unexpected HTTP request: {request.method} {request.url}")


class LzWarehouseInventoryParsingTests(unittest.TestCase):
    def test_warehouse_selection_uses_exact_names_and_live_ids(self) -> None:
        options = inventory.parse_target_warehouses(_inventory_html())
        self.assertEqual(
            [(item["id"], item["name"]) for item in options], list(WAREHOUSES)
        )

    def test_missing_or_ambiguous_target_warehouse_is_rejected(self) -> None:
        with self.assertRaises(MabangApiError):
            inventory.parse_target_warehouses(_inventory_html(WAREHOUSES[:-1]))
        with self.assertRaises(MabangApiError):
            inventory.parse_target_warehouses(_inventory_html((*WAREHOUSES, ("other-id", "泰国TH05仓"))))

    def test_unselected_warehouses_need_no_permission_or_unique_name(self) -> None:
        selected = (WAREHOUSES[1][1],)
        html = _inventory_html((WAREHOUSES[1], ("unselected-id1", "泰国TH04仓"), ("unselected-id2", "泰国TH04仓")))
        options = inventory.parse_target_warehouses(html, selected)
        self.assertEqual([(item["id"], item["name"]) for item in options], [WAREHOUSES[1]])

    def test_record_count_uses_first_semibold_instead_of_page_fraction(self) -> None:
        html = '<span class="semibold">98,533</span>条 <span class="semibold">1/1971</span>页'
        self.assertEqual(inventory.parse_record_count(html), 98533)
        self.assertEqual(inventory.parse_record_count('<span class="semibold">0</span>条'), 0)
        with self.assertRaises(MabangApiError):
            inventory.parse_record_count('<span class="semibold">加载中</span>')

    def test_query_selects_all_three_warehouses_without_keyword_or_stock_threshold(self) -> None:
        payload = inventory.build_inventory_search_payload(tuple(identity for identity, _ in WAREHOUSES))
        self.assertEqual(
            [value for key, value in payload if key == "warehouseIds[]"],
            [identity for identity, _ in WAREHOUSES],
        )
        form = dict(payload)
        self.assertEqual(form.get("page"), "1")
        self.assertEqual(form.get("rowsPerPage"), "50")
        self.assertFalse(form.get("stockNameValue"))
        self.assertFalse(form.get("stockQuantitygt"))

    def test_xlsx_reader_preserves_all_native_columns_zero_stock_and_text_sku(self) -> None:
        values = [("000012345678901234", "原生商品", "泰国TH04仓", 0, "页面原生字段")]
        headers, rows = inventory.read_export_table(_xlsx_bytes(values))
        self.assertEqual(tuple(headers), HEADERS)
        self.assertEqual(rows, values)

    def test_legacy_xls_is_read_as_a_real_ole_workbook(self) -> None:
        values = [("000001", "原生商品", "泰国TH06仓", 0, "页面字段")]
        content = _xls_bytes(values)
        self.assertTrue(content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"))
        headers, rows = inventory.read_export_table(content)
        self.assertEqual(tuple(headers), HEADERS)
        self.assertEqual(rows, values)

    def test_error_html_and_invalid_files_cannot_become_successful_exports(self) -> None:
        for content in (b"", b"<html><body>login required</body></html>", b"not an Excel workbook", b"PK\x03\x04corrupt"):
            with self.subTest(content=content[:20]):
                with self.assertRaises(MabangApiError):
                    inventory.read_export_table(content)

    def test_missing_duplicate_or_blank_native_columns_are_rejected(self) -> None:
        for headers in (("库存SKU编号", "商品中文名称"), ("库存SKU编号", "仓库", "仓库"), ("库存SKU编号", "仓库", "")):
            with self.subTest(headers=headers):
                with self.assertRaises(MabangApiError):
                    inventory.read_export_table(_xlsx_bytes([], headers))

    def test_job_requires_credentials(self) -> None:
        for missing in ("username", "password"):
            payload = {"username": "test-account", "password": "test-password", "output_dir": "D:/test-output", "warehouse_names": [WAREHOUSES[0][1]]}
            del payload[missing]
            with self.subTest(missing=missing):
                with self.assertRaises(ValueError):
                    inventory.validate_lz_inventory_payload(payload)

    def test_job_rejects_empty_or_missing_warehouse_selection(self) -> None:
        for selected in (None, [], (), "", ["  "]):
            payload = {"username": "test-account", "password": "test-password", "output_dir": "D:/test-output"}
            if selected is not None:
                payload["warehouse_names"] = selected
            with self.subTest(selected=selected):
                with self.assertRaises(ValueError):
                    inventory.validate_lz_inventory_payload(payload)


class LzWarehouseInventoryHttpTests(unittest.TestCase):
    def _run(self, fixture: InventoryHttpFixture, output: Path, warehouse_names: tuple[str, ...] = (WAREHOUSES[0][1],)):
        with httpx.Client(transport=httpx.MockTransport(fixture.handle), follow_redirects=True) as http:
            real_client = MabangClient(inventory.BASE_URL, client=http)
            with patch.object(inventory, "MabangClient", return_value=real_client):
                job = inventory.LzInventoryJob(
                    username="test-account", password="test-password", output_dir=output,
                    warehouse_names=warehouse_names,
                )
                return inventory.run_lz_inventory_export(job)

    def _assert_failure_without_published_file(self, fixture: InventoryHttpFixture, output: Path) -> None:
        result = self._run(fixture, output)
        self.assertFalse(result.is_complete)
        self.assertEqual(len(result.failures), 1)
        self.assertEqual(result.failures[0]["warehouse_name"], WAREHOUSES[0][1])
        self.assertEqual(tuple(result.output_files), ())
        self.assertEqual(tuple(result.exports), ())
        self.assertEqual(list(output.rglob("*.xlsx")), [])
        self.assertEqual(list(output.rglob("*.tmp")), [])

    def test_http_export_merges_every_native_export_page_and_keeps_one_header(self) -> None:
        first_rows = [
            (f"{index:018d}", f"商品{index}", WAREHOUSES[0][1], index % 9, f"原生字段{index}")
            for index in range(10000)
        ]
        last_rows = [("000000000000010000", "最后一条", WAREHOUSES[0][1], 0, "最后字段")]
        fixture = InventoryHttpFixture(10001, {1: _xlsx_bytes(first_rows), 2: _xlsx_bytes(last_rows)})
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory))
            self.assertEqual(result.record_count, 10001)
            self.assertEqual(result.source_record_count, 10001)
            self.assertEqual(result.export_page_count, 2)
            self.assertEqual(sum(result.warehouse_counts.values()), 10001)
            self.assertEqual(len(result.output_files), 1)
            self.assertEqual(len(result.exports), 1)
            self.assertEqual(result.exports[0].warehouse_name, WAREHOUSES[0][1])
            self.assertEqual(result.exports[0].output_file, result.output_file)
            self.assertIn(WAREHOUSES[0][1], result.output_file.name)
            self.assertTrue(result.is_complete)
            self.assertEqual(tuple(result.failures), ())
            self.assertEqual(tuple(result.preview_columns), HEADERS)
            self.assertEqual(len(result.preview_rows), inventory.PREVIEW_ROW_LIMIT)
            book = load_workbook(result.output_file, read_only=True, data_only=False)
            try:
                self.assertEqual(len(book.worksheets), 1)
                values = list(book.worksheets[0].iter_rows(values_only=True))
                self.assertEqual(values[0], HEADERS)
                self.assertEqual(len(values), 10002)
                self.assertEqual(values[1], first_rows[0])
                self.assertEqual(values[-1], last_rows[0])
                self.assertEqual(sum(row == HEADERS for row in values), 1)
                self.assertEqual(book.worksheets[0]["A2"].data_type, "s")
            finally:
                book.close()
        searches = [form for method, mod, form in fixture.requests if mod == "warehouse.searchwarehousestock"]
        self.assertEqual(len(searches), 1)
        self.assertEqual(searches[0]["warehouseIds[]"], [WAREHOUSES[0][0]])
        exports = [(method, form) for method, mod, form in fixture.requests if mod == "warehouse.doexportwarehousestock"]
        self.assertEqual([form["page"] for _, form in exports], [["1"], ["2"]])
        for method, form in exports:
            self.assertEqual(method, "GET")
            self.assertEqual(form["flag"], ["1"])
            self.assertEqual(form["showRmbColumn"], ["0"])
            self.assertEqual(form["searchToken"], [fixture.token_for(WAREHOUSES[0][0])])
        self.assertTrue(any(mod == "warehouse.getSearchWarehouseStockPage" for _, mod, _ in fixture.requests))

    def test_multi_selection_queries_each_warehouse_and_writes_isolated_files(self) -> None:
        first_rows = [("04-1", "商品1", WAREHOUSES[0][1], 0, "字段1"), ("04-2", "商品2", WAREHOUSES[0][1], 2, "字段2")]
        first_last = [("04-3", "商品3", WAREHOUSES[0][1], 3, "字段3")]
        second_rows = [("06-1", "商品4", WAREHOUSES[2][1], 0, "字段4"), ("06-2", "商品5", WAREHOUSES[2][1], 5, "字段5")]
        fixture = InventoryHttpFixture(0, {}, scopes={
            WAREHOUSES[0][0]: (3, {1: _xlsx_bytes(first_rows), 2: _xlsx_bytes(first_last)}),
            WAREHOUSES[2][0]: (2, {1: _xlsx_bytes(second_rows)}),
        })
        selected = (WAREHOUSES[0][1], WAREHOUSES[2][1])
        with patch.object(inventory, "EXPORT_PAGE_SIZE", 2), tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory), selected)
            self.assertTrue(result.is_complete)
            self.assertEqual(result.record_count, 5)
            self.assertEqual(result.source_record_count, 5)
            self.assertEqual(result.export_page_count, 3)
            self.assertEqual(len(result.output_files), 2)
            self.assertEqual(len(set(result.output_files)), 2)
            self.assertEqual(tuple(item.output_file for item in result.exports), tuple(result.output_files))
            self.assertEqual(tuple(item.warehouse_name for item in result.exports), selected)
            self.assertEqual(result.output_file, result.output_files[0])
            expected = {WAREHOUSES[0][1]: first_rows + first_last, WAREHOUSES[2][1]: second_rows}
            for export in result.exports:
                self.assertIn(export.warehouse_name, export.output_file.name)
                self.assertEqual(export.record_count, len(expected[export.warehouse_name]))
                self.assertEqual(export.source_record_count, export.record_count)
                self.assertEqual(tuple(export.preview_columns), HEADERS)
                book = load_workbook(export.output_file, read_only=True)
                try:
                    self.assertEqual(len(book.worksheets), 1)
                    values = list(book.worksheets[0].iter_rows(values_only=True))
                    self.assertEqual(values, [HEADERS, *expected[export.warehouse_name]])
                    self.assertEqual({row[2] for row in values[1:]}, {export.warehouse_name})
                finally:
                    book.close()
            self.assertEqual(len(list(Path(directory).rglob("*.xlsx"))), 2)
        business_calls = [(mod, form) for _, mod, form in fixture.requests if mod in {
            "warehouse.searchwarehousestock", "warehouse.getSearchWarehouseStockPage", "warehouse.doexportwarehousestock"
        }]
        self.assertEqual([mod for mod, _ in business_calls], [
            "warehouse.searchwarehousestock", "warehouse.getSearchWarehouseStockPage",
            "warehouse.doexportwarehousestock", "warehouse.doexportwarehousestock",
            "warehouse.searchwarehousestock", "warehouse.getSearchWarehouseStockPage", "warehouse.doexportwarehousestock",
        ])
        searches = [form["warehouseIds[]"] for mod, form in business_calls if mod == "warehouse.searchwarehousestock"]
        self.assertEqual(searches, [[WAREHOUSES[0][0]], [WAREHOUSES[2][0]]])
        tokens = [form["searchToken"] for mod, form in business_calls if mod == "warehouse.doexportwarehousestock"]
        self.assertEqual(tokens, [
            [fixture.token_for(WAREHOUSES[0][0])], [fixture.token_for(WAREHOUSES[0][0])], [fixture.token_for(WAREHOUSES[2][0])]
        ])
        self.assertNotEqual(tokens[0], tokens[-1])

    def test_single_selection_works_without_permission_to_unselected_warehouses(self) -> None:
        warehouse_id, warehouse_name = WAREHOUSES[1]
        rows = [("05-1", "商品", warehouse_name, 0, "原生字段")]
        fixture = InventoryHttpFixture(0, {}, html=_inventory_html((WAREHOUSES[1],)), scopes={
            warehouse_id: (1, {1: _xlsx_bytes(rows)}),
        })
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory), (warehouse_name,))
            self.assertTrue(result.is_complete)
            self.assertEqual(len(result.output_files), 1)
            self.assertIn(warehouse_name, result.output_files[0].name)
        searches = [form["warehouseIds[]"] for _, mod, form in fixture.requests if mod == "warehouse.searchwarehousestock"]
        self.assertEqual(searches, [[warehouse_id]])

    def test_a_visible_warehouse_outside_default_three_can_be_selected(self) -> None:
        warehouse_id, warehouse_name = "1093126", "泰国TH04中转仓"
        rows = [("TRANSFER-1", "商品", warehouse_name, 1, "原生字段")]
        fixture = InventoryHttpFixture(0, {}, scopes={warehouse_id: (1, {1: _xlsx_bytes(rows)})})
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory), (warehouse_name,))
            self.assertTrue(result.is_complete)
            self.assertEqual(len(result.output_files), 1)
            self.assertEqual(result.exports[0].warehouse_name, warehouse_name)
        searches = [form["warehouseIds[]"] for _, mod, form in fixture.requests if mod == "warehouse.searchwarehousestock"]
        self.assertEqual(searches, [[warehouse_id]])

    def test_missing_selected_warehouse_fails_before_searching_any_warehouse(self) -> None:
        fixture = InventoryHttpFixture(1, {}, html=_inventory_html((WAREHOUSES[0],)))
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaises(MabangApiError):
                self._run(fixture, Path(directory), (WAREHOUSES[0][1], "不存在的仓库"))
            self.assertEqual(list(Path(directory).rglob("*.xlsx")), [])
        self.assertFalse(any(mod == "warehouse.searchwarehousestock" for _, mod, _ in fixture.requests))

    def test_one_failed_warehouse_does_not_stop_other_selected_warehouses(self) -> None:
        rows_04 = [("04-1", "商品1", WAREHOUSES[0][1], 1, "字段1")]
        rows_06 = [("06-1", "商品3", WAREHOUSES[2][1], 3, "字段3")]
        fixture = InventoryHttpFixture(0, {}, scopes={
            WAREHOUSES[0][0]: (1, {1: _xlsx_bytes(rows_04)}),
            WAREHOUSES[1][0]: (1, {1: b"<html>permission denied</html>"}),
            WAREHOUSES[2][0]: (1, {1: _xlsx_bytes(rows_06)}),
        })
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory), tuple(name for _, name in WAREHOUSES))
            self.assertFalse(result.is_complete)
            self.assertEqual(len(result.failures), 1)
            self.assertEqual(result.failures[0]["warehouse_name"], WAREHOUSES[1][1])
            self.assertEqual({item.warehouse_name for item in result.exports}, {WAREHOUSES[0][1], WAREHOUSES[2][1]})
            self.assertEqual(len(result.output_files), 2)
            self.assertEqual(result.record_count, 2)
            paths = list(Path(directory).rglob("*.xlsx"))
            self.assertEqual(set(paths), set(result.output_files))
            self.assertFalse(any(WAREHOUSES[1][1] in path.name for path in paths))
        searches = [form["warehouseIds[]"] for _, mod, form in fixture.requests if mod == "warehouse.searchwarehousestock"]
        self.assertEqual(searches, [[identity] for identity, _ in WAREHOUSES])

    def test_missing_search_token_stops_before_any_export(self) -> None:
        fixture = InventoryHttpFixture(1, {}, search_token="")
        with tempfile.TemporaryDirectory() as directory:
            self._assert_failure_without_published_file(fixture, Path(directory))
        self.assertFalse(any(mod == "warehouse.doexportwarehousestock" for _, mod, _ in fixture.requests))

    def test_permission_failure_stops_before_any_export(self) -> None:
        fixture = InventoryHttpFixture(1, {}, search_success=False)
        with tempfile.TemporaryDirectory() as directory:
            self._assert_failure_without_published_file(fixture, Path(directory))
        self.assertFalse(any(mod == "warehouse.doexportwarehousestock" for _, mod, _ in fixture.requests))

    def test_empty_or_short_native_exports_are_rejected_as_incomplete(self) -> None:
        one_row = [("SKU-1", "商品", "泰国TH04仓", 0, "完整字段")]
        for exported_rows in ([], one_row):
            with self.subTest(rows=len(exported_rows)), tempfile.TemporaryDirectory() as directory:
                fixture = InventoryHttpFixture(2, {1: _xlsx_bytes(exported_rows)})
                self._assert_failure_without_published_file(fixture, Path(directory))

    def test_unselected_warehouse_cannot_be_mixed_into_merged_result(self) -> None:
        for warehouse_name in ("泰国TH04中转仓", "泰国TH05仓"):
            rows = [("SKU-1", "商品", warehouse_name, 7, "原生字段")]
            fixture = InventoryHttpFixture(1, {1: _xlsx_bytes(rows)})
            with self.subTest(warehouse=warehouse_name), tempfile.TemporaryDirectory() as directory:
                self._assert_failure_without_published_file(fixture, Path(directory))

    def test_duplicate_export_page_is_rejected_even_when_row_count_matches(self) -> None:
        rows = [("SKU-1", "商品1", "泰国TH04仓", 0, "字段1"), ("SKU-2", "商品2", "泰国TH04仓", 1, "字段2")]
        content = _xlsx_bytes(rows)
        fixture = InventoryHttpFixture(4, {1: content, 2: content})
        with patch.object(inventory, "EXPORT_PAGE_SIZE", 2), tempfile.TemporaryDirectory() as directory:
            self._assert_failure_without_published_file(fixture, Path(directory))

    def test_changed_native_header_between_pages_is_rejected(self) -> None:
        row = [("SKU-1", "商品1", "泰国TH04仓", 0, "字段1")]
        fixture = InventoryHttpFixture(2, {
            1: _xlsx_bytes(row),
            2: _xlsx_bytes([("SKU-2", "商品2", "泰国TH04仓", 1, "字段2")], (*HEADERS[:-1], "另一字段")),
        })
        with patch.object(inventory, "EXPORT_PAGE_SIZE", 1), tempfile.TemporaryDirectory() as directory:
            self._assert_failure_without_published_file(fixture, Path(directory))

    def test_zero_result_exports_an_empty_workbook_with_native_headers(self) -> None:
        fixture = InventoryHttpFixture(0, {1: _xlsx_bytes([])})
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory))
            self.assertTrue(result.is_complete)
            self.assertEqual(result.record_count, 0)
            self.assertEqual(len(result.output_files), 1)
            self.assertIn(WAREHOUSES[0][1], result.output_file.name)
            book = load_workbook(result.output_file, read_only=True)
            try:
                self.assertEqual(list(book.worksheets[0].iter_rows(values_only=True)), [HEADERS])
            finally:
                book.close()
        self.assertEqual(sum(mod == "warehouse.doexportwarehousestock" for _, mod, _ in fixture.requests), 1)

    def test_native_business_duplicates_and_formula_like_text_are_preserved(self) -> None:
        row = ("000001", "商品", "泰国TH04仓", 0, '=HYPERLINK("https://example.test","literal")')
        fixture = InventoryHttpFixture(2, {1: _xlsx_bytes([row, row])})
        with tempfile.TemporaryDirectory() as directory:
            result = self._run(fixture, Path(directory))
            book = load_workbook(result.output_file, read_only=True, data_only=False)
            try:
                values = list(book.worksheets[0].iter_rows(values_only=True))
                self.assertEqual(values, [HEADERS, row, row])
                self.assertEqual(book.worksheets[0]["E2"].data_type, "s")
            finally:
                book.close()

    def test_failure_on_later_page_publishes_no_partial_workbook(self) -> None:
        fixture = InventoryHttpFixture(2, {
            1: _xlsx_bytes([("SKU-1", "商品1", "泰国TH04仓", 1, "字段1")]),
            2: b"<html>session expired</html>",
        })
        with patch.object(inventory, "EXPORT_PAGE_SIZE", 1), tempfile.TemporaryDirectory() as directory:
            self._assert_failure_without_published_file(fixture, Path(directory))

    def test_failed_jobs_release_write_only_xml_temporary_files(self) -> None:
        before_append = InventoryHttpFixture(1, {}, search_token="")
        after_append = InventoryHttpFixture(2, {
            1: _xlsx_bytes([("SKU-1", "商品1", "泰国TH04仓", 1, "字段1")]),
            2: b"<html>session expired</html>",
        })
        for stage, fixture in (("before rows", before_append), ("after rows", after_append)):
            with self.subTest(stage=stage), patch.object(inventory, "EXPORT_PAGE_SIZE", 1), tempfile.TemporaryDirectory() as directory:
                original_temp_files = set(ALL_TEMP_FILES)
                self._assert_failure_without_published_file(fixture, Path(directory))
                self.assertEqual(set(ALL_TEMP_FILES) - original_temp_files, set())

    def test_login_errors_do_not_expose_remote_payload_or_authentication_url(self) -> None:
        sensitive = "fixture-private-authentication-value"
        failures = (
            MabangApiError(f"登录失败: {{'private': '{sensitive}'}}"),
            httpx.ReadError(f"https://example.test/login?auth={sensitive}"),
        )
        for failure in failures:
            with self.subTest(error_type=type(failure).__name__), tempfile.TemporaryDirectory() as directory:
                fixture = InventoryHttpFixture(1, {})
                with patch.object(MabangClient, "login", side_effect=failure):
                    try:
                        self._run(fixture, Path(directory))
                    except MabangApiError:
                        displayed_error = traceback.format_exc()
                    else:
                        self.fail("An unsuccessful login must stop the export")
                self.assertNotIn(sensitive, displayed_error)
                self.assertNotIn("https://example.test", displayed_error)
                self.assertEqual(list(Path(directory).rglob("*.xlsx")), [])
