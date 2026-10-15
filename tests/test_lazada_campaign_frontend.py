"""Drive the local Excel UI through AppBridge and read back actual output files."""
from __future__ import annotations

import functools
import json
import tempfile
import threading
import unittest
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from backend.app_bridge import AppBridge
from backend.config_store import AppSettings
from tests.test_lazada_campaign_bridge import write_prices

try:
    from playwright.sync_api import Error, expect, sync_playwright
except ImportError:
    sync_playwright = None

ROOT = Path(__file__).resolve().parents[1]


class BridgeHandler(SimpleHTTPRequestHandler):
    bridge: AppBridge

    def log_message(self, *_args):
        pass

    def do_POST(self):
        request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        method = request["method"]
        allowed = {"get_lazada_campaign_workbook_info", "choose_lazada_campaign_files", "inspect_lazada_campaign_files",
                   "export_lazada_campaign_template", "start_lazada_campaign_workbook", "get_task_status", "get_latest_task_status", "open_path"}
        if method == "get_app_info":
            settings = self.bridge.config_store.load()
            result = {"ok": True, "app": {"name": "测试工作台", "version": "test"},
                      "settings": {"output_dir": settings.output_dir}, "account_state": {"accounts": [], "active_account_ids": {}}}
        elif method == "get_sidebar_preferences":
            result = {"ok": True, "preferences": {}}
        elif method in allowed:
            result = getattr(self.bridge, method)(*request.get("args", []))
        else:
            result = {"ok": False, "error": "测试接口不允许调用"}
        payload = json.dumps(result, ensure_ascii=False).encode("utf-8")
        self.send_response(200)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(payload)))
        self.end_headers()
        self.wfile.write(payload)


class CampaignFrontendTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        if sync_playwright is None:
            raise unittest.SkipTest("Playwright unavailable")
        cls.playwright = sync_playwright().start()
        cls.addClassCleanup(cls.playwright.stop)
        for options in ({"channel": "msedge"}, {}):
            try:
                cls.browser = cls.playwright.chromium.launch(headless=True, **options)
                break
            except Error:
                continue
        else:
            raise unittest.SkipTest("Chromium/Edge unavailable")
        cls.addClassCleanup(cls.browser.close)

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.sources = [write_prices(self.root / "第一.xlsx"), write_prices(self.root / "第二.xlsx")]
        self.selected = [str(path) for path in self.sources]
        self.bridge = AppBridge()
        self.bridge.config_store = Mock()
        self.bridge.config_store.load.return_value = AppSettings(output_dir=str(self.root / "输出"))
        self.bridge._payload_with_account = Mock(side_effect=AssertionError("不能读取业务账号"))
        self.window = Mock()
        self.window.create_file_dialog.side_effect = lambda *args, **kwargs: self.selected
        picker = patch.dict("sys.modules", {"webview": SimpleNamespace(OPEN_DIALOG=1, windows=[self.window])})
        picker.start()
        self.addCleanup(picker.stop)
        BridgeHandler.bridge = self.bridge
        self.server = ThreadingHTTPServer(("127.0.0.1", 0), functools.partial(BridgeHandler, directory=str(ROOT / "frontend")))
        self.server.daemon_threads = True
        threading.Thread(target=self.server.serve_forever, daemon=True).start()
        self.addCleanup(self.server.server_close)
        self.addCleanup(self.server.shutdown)
        self.context = self.browser.new_context(viewport={"width": 1440, "height": 1000})
        self.addCleanup(self.context.close)
        self.context.add_init_script("""window.pywebview = {api: new Proxy({}, {get: (_, method) => (...args) =>
            fetch('/api', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({method,args})}).then(r=>r.json())})};""")
        self.page = self.context.new_page()
        self.page.set_default_timeout(10000)
        self.errors = []
        self.page.on("pageerror", lambda error: self.errors.append(str(error)))
        self.opened = []
        opener = patch("backend.app_bridge.os.startfile", side_effect=lambda path: self.opened.append(path))
        opener.start()
        self.addCleanup(opener.stop)

    def open_page(self):
        self.page.goto(f"http://127.0.0.1:{self.server.server_port}/pages/lazada-campaign-workbook.html")
        expect(self.page.locator("#chooseFilesBtn")).to_be_enabled()

    def choose_files(self):
        self.page.locator("#chooseFilesBtn").click()
        expect(self.page.locator("#sourceFileList [data-remove-file]")).to_have_count(len(self.selected))

    def test_batch_without_accounts_then_open_outputs_and_restore(self):
        originals = [path.read_bytes() for path in self.sources]
        self.open_page()
        self.assertEqual(self.page.locator("#campaignAccount, #bindAccountBtn, #campaignStores, #campaignUrl, #clientPath, [data-active-account-name]").count(), 0)
        form = self.page.locator("#campaignForm").bounding_box()
        log = self.page.locator(".log-surface").bounding_box()
        self.assertGreater(log["x"], form["x"] + form["width"])
        self.choose_files()
        self.page.locator("#inspectBtn").click()
        expect(self.page.locator("#inspectionMessage")).to_contain_text("第一.xlsx")
        expect(self.page.locator("#inspectionMessage")).to_contain_text("3 行")
        self.page.locator("#startBtn").click()
        expect(self.page.locator("#completedCount")).to_have_text("完成 2 · 失败 0")
        expect(self.page.locator("#keptCount")).to_have_text("2")
        expect(self.page.locator("#totalCount")).to_have_text("6")
        expect(self.page.locator("#exceptionCount")).to_contain_text("异常行 2")
        task = self.bridge.get_latest_task_status("lazada_campaign_workbook")
        for file in task["result"]["files"]:
            book = load_workbook(file["processed_file"])
            try:
                self.assertEqual((book.active["A2"].value, book.active["D2"].value, book.active.max_row), ("0001", 95, 2))
            finally:
                book.close()
        self.assertEqual([path.read_bytes() for path in self.sources], originals)
        self.page.locator("#openArchiveBtn").click()
        self.page.wait_for_function("() => document.querySelector('#logBox').textContent.includes('任务完成')")
        self.page.reload()
        expect(self.page.locator("#completedCount")).to_have_text("完成 2 · 失败 0")
        expect(self.page.locator("#sourceFileList [data-remove-file]")).to_have_count(2)
        expect(self.page.locator("#discountThreshold")).to_have_value("95")
        self.assertFalse(self.errors)
        self.assertIn(task["result"]["archive_file"], self.opened)
        self.bridge._payload_with_account.assert_not_called()

    def test_corrupt_file_does_not_block_successful_file(self):
        self.sources[1].write_bytes(b"corrupt Excel")
        self.open_page()
        self.choose_files()
        self.page.locator("#startBtn").click()
        expect(self.page.locator("#completedCount")).to_have_text("完成 1 · 失败 1")
        expect(self.page.locator("#resultBadge")).to_have_text("部分完成")
        expect(self.page.get_by_role("button", name="筛选 Excel", exact=True)).to_be_enabled()
        expect(self.page.locator("#openArchiveBtn")).to_be_enabled()
        self.assertFalse(self.errors)

    def test_recommended_mode_and_decimal_threshold_restore(self):
        self.selected = [str(write_prices(self.root / "推荐.xlsx", recommended=True))]
        self.open_page()
        self.choose_files()
        self.page.locator("#processingMode").select_option("invited")
        self.page.locator("#discountThreshold").fill("95.05")
        self.page.locator("#startBtn").click()
        expect(self.page.locator("#completedCount")).to_have_text("完成 1 · 失败 0")
        expect(self.page.locator("#keptCount")).to_have_text("0")
        expect(self.page.locator("#updatedCount")).to_have_text("填充活动价 2 行")
        task = self.bridge.get_latest_task_status("lazada_campaign_workbook")
        self.assertEqual(task["context"]["threshold"], "0.9505")
        self.assertEqual(task["context"]["mode"], "invited")
        self.page.reload()
        expect(self.page.locator("#processingMode")).to_have_value("invited")
        expect(self.page.locator("#discountThreshold")).to_have_value("95.05")
        self.assertFalse(self.errors)

    def test_localized_invited_template_detected_and_filled_automatically(self):
        source = self.root / "新版特邀.xlsx"
        book = Workbook()
        book.active.append(["Sku ID", "卖家SKU", "产品名称", "提醒", "当前价格", "当前存货",
                            "市场活动价", "流量加权价", "推荐价格", "Is Hero Product"])
        book.active.append(["1", "0001", "商品", None, 100, 5, None, 1, "<=95", "No"])
        book.active.append(["2", "0002", "商品", None, 100, 5, None, 999, "<=94", "No"])
        book.save(source)
        book.close()
        original = source.read_bytes()
        self.selected = [str(source)]
        self.open_page()
        self.choose_files()
        self.page.locator("#inspectBtn").click()
        expect(self.page.locator("#inspectionMessage")).to_contain_text("按推荐价填充")
        self.page.locator("#startBtn").click()
        expect(self.page.locator("#completedCount")).to_have_text("完成 1 · 失败 0")
        expect(self.page.locator("#keptCount")).to_have_text("1")
        expect(self.page.locator("#updatedCount")).to_have_text("填充活动价 2 行")
        task = self.bridge.get_latest_task_status("lazada_campaign_workbook")
        output = load_workbook(task["result"]["files"][0]["processed_file"])
        try:
            self.assertEqual((output.active["G2"].value, output.active["H2"].value), (95, 1))
            self.assertEqual(output.active.max_row, 2)
        finally:
            output.close()
        self.assertEqual(source.read_bytes(), original)
        self.assertFalse(self.errors)

    def test_export_empty_template_before_importing_files(self):
        self.open_page()
        self.page.locator("#exportTemplateBtn").click()
        expect(self.page.locator("#templateNotice")).to_contain_text("模板已保存")
        self.page.wait_for_function("() => !document.querySelector('#exportTemplateBtn').disabled")
        self.assertEqual(len(self.opened), 1)
        book = load_workbook(self.opened[0])
        try:
            self.assertEqual(book.active.max_row, 1)
            self.assertEqual(book.active["C1"].value, "原价")
        finally:
            book.close()
        self.assertFalse(self.errors)

    def test_manual_paths_deduplicate_limit_and_picker_cancel(self):
        self.open_page()
        self.choose_files()
        self.selected = None
        self.page.locator("#chooseFilesBtn").click()
        expect(self.page.locator("#chooseFilesBtn")).to_be_enabled()
        expect(self.page.locator("#sourceFileList [data-remove-file]")).to_have_count(2)
        self.page.locator(".campaign-paths summary").click()
        self.page.locator("#sourcePaths").fill(f'"{self.sources[0]}"\n\n{self.sources[0]}')
        expect(self.page.locator("#fileCount")).to_have_text("1")
        self.page.locator("#sourcePaths").fill("\n".join(f"C:\\商品{i}.xlsx" for i in range(101)))
        expect(self.page.locator("#startBtn")).to_be_disabled()
        expect(self.page.locator("#formIssues")).to_contain_text("100")
        self.page.locator("#clearFilesBtn").click()
        expect(self.page.locator("#sourceFileList")).to_contain_text("尚未选择文件")
        expect(self.page.locator("#startBtn")).to_be_disabled()
        self.assertFalse(self.errors)

    def test_old_download_task_does_not_restore_or_request_account(self):
        self.bridge.tasks.start("旧紫鸟任务", lambda progress: {"file_count": 99, "kept_rows": 99},
                                tool="lazada_campaign_workbook", context={"workflow": "download", "account_id": "z1", "store_names": ["old"]})
        self.open_page()
        expect(self.page.locator("#keptCount")).to_have_text("0")
        expect(self.page.locator("#openArchiveBtn")).to_be_disabled()
        expect(self.page.locator("#sourcePaths")).to_have_value("")
        self.assertFalse(self.errors)

    def test_dashboard_general_tools_filter_links_to_local_excel(self):
        self.page.goto(f"http://127.0.0.1:{self.server.server_port}/pages/dashboard.html")
        card = self.page.locator(".dashboard-tool", has_text="Excel 折扣筛选")
        expect(card).to_have_count(1)
        self.page.locator("#vendorTabs [data-vendor=local]").click()
        expect(self.page.locator(".dashboard-group")).to_have_count(1)
        expect(self.page.locator(".dashboard-group[data-vendor=local]")).to_be_visible()
        expect(card).to_contain_text("本地处理 · 无需绑定账号")
        card.click()
        expect(self.page.locator("#chooseFilesBtn")).to_be_enabled()
        self.assertFalse(self.errors)


if __name__ == "__main__":
    unittest.main()
