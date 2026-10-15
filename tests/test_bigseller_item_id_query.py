from __future__ import annotations

import tempfile
import threading
import unittest
from dataclasses import replace
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import Mock, patch

import requests
from openpyxl import load_workbook

from backend.services import bigseller_item_id_query as service


class BigSellerItemIdQueryTests(unittest.TestCase):
    def setUp(self):
        self.clock = 100.0

        def advance(seconds):
            self.clock += seconds

        patches = (
            patch.object(service, "_REQUEST_LOCK", threading.Lock()),
            patch.object(service, "_LAST_REQUEST", None),
            patch.object(service, "_COOLDOWN_UNTIL", 0.0),
            patch.object(service.time, "monotonic", side_effect=lambda: self.clock),
            patch.object(service.time, "sleep", side_effect=advance),
        )
        for mocked in patches:
            mocked.start()
            self.addCleanup(mocked.stop)

    def make_job(self, output_dir: Path, skus: tuple[str, ...] = ("sku-a",)):
        return service.BigSellerItemIdQueryJob(
            username="private-account", password="private-password",
            skus=skus, output_dir=output_dir,
            captcha_username="private-captcha-user", captcha_password="private-captcha-password",
        )

    def make_util(self):
        util = Mock()
        util.build_api_url.side_effect = lambda path: "https://www.bigseller.pro/api" + path
        util.get_request_user_agent.return_value = "test-agent"
        util.get_listing_page_url.return_value = "https://www.bigseller.pro/web/listing/shopee/active/index.htm?bsStatus=4"
        util.build_request_headers.return_value = {"Content-Type": "application/json"}
        return util

    def response(self, rows=None, *, status=200, payload=None, headers=None):
        response = Mock(status_code=status, headers=headers or {})
        response.json.return_value = payload if payload is not None else {"code": 0, "data": {"page": {"rows": rows or []}}}
        return response

    def test_parse_keeps_order_case_leading_zeroes_and_internal_spaces(self):
        self.assertEqual(
            service.parse_skus("\ufeff 00123\nsku-A,sku-A，SKU-A;two words；x\ty\r\n"),
            ("00123", "sku-A", "SKU-A", "two words", "x", "y"),
        )

    def test_parse_rejects_empty_wrong_type_controls_and_oversize(self):
        for value in (None, [], " \n,;\t", "a\x00b", "a" * 32768):
            with self.subTest(value=repr(value)[:50]), self.assertRaises(ValueError):
                service.parse_skus(value)
        with self.assertRaisesRegex(ValueError, "5000"):
            service.parse_skus("\n".join(str(i) for i in range(5001)))
        self.assertEqual(service.parse_skus("a\n" * 6000), ("a",))

    def test_validation_requires_output_and_preserves_password(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            payload = {"username": " user ", "password": " pass ", "sku_text": "001", "output_dir": temp_dir}
            job = service.validate_bigseller_item_id_query_payload(payload)
            self.assertEqual(job.password, " pass ")
            self.assertEqual(job.username, "user")
            self.assertEqual(job.skus, ("001",))
            self.assertNotIn("pass", repr(job))
            for key in ("username", "password", "sku_text", "output_dir"):
                with self.subTest(key=key), self.assertRaises(ValueError):
                    service.validate_bigseller_item_id_query_payload({**payload, key: " "})
            file_path = Path(temp_dir) / "file.txt"
            file_path.write_text("test", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "不能是文件"):
                service.validate_bigseller_item_id_query_payload({**payload, "output_dir": str(file_path)})

    def test_listing_contract_fuzzy_sku_views_desc_first_page_all_shops(self):
        self.assertEqual(service.build_listing_query("sku / 001"), {
            "searchType": "sku", "searchContent": "sku / 001", "inquireType": 0,
            "shopeeStatus": "live", "status": "active", "orderBy": "views", "desc": True,
            "pageNo": 1, "pageSize": 50, "timeType": "create_time", "startDateStr": "", "endDateStr": "",
        })

    def test_first_result_is_authoritative_even_on_views_ties(self):
        rows = [
            {"itemId": "001234567890123456789", "id": "internal-1", "views": 30, "shopName": "Z"},
            {"itemId": "999", "views": 30, "shopName": "A"},
        ]
        self.assertEqual(service.first_item_id(rows), "001234567890123456789")
        self.assertEqual(service.first_item_id([{ "itemId": 123456789012345678901}]), "123456789012345678901")
        self.assertEqual(service.first_item_id([]), "")

    def test_internal_ids_and_next_result_are_never_substitutes(self):
        for invalid in (None, "", 0, "0", False, [], "=123", 123.0, "1e12"):
            with self.subTest(invalid=invalid), self.assertRaises(service.BigSellerQueryError):
                service.first_item_id([
                    {"itemId": invalid, "id": 123, "productId": 456, "listingId": 789},
                    {"itemId": "999"},
                ])

    def test_only_explicit_empty_rows_count_as_no_match(self):
        self.assertEqual(service.extract_listing_rows({"code": "0", "data": {"page": {"rows": []}}}), [])
        for payload in (None, [], {}, {"code": 0}, {"code": 0, "data": {"page": {}}},
                        {"code": 0, "data": {"page": {"rows": None}}},
                        {"code": 0, "data": {"page": {"rows": ["invalid"]}}},
                        {"code": 500, "data": {"page": {"rows": []}}}):
            with self.subTest(payload=payload), self.assertRaises(service.BigSellerQueryError):
                service.extract_listing_rows(payload)

    def test_api_login_expiration_and_permission_are_not_empty_results(self):
        for payload in ({"code": 2001}, {"code": "401006"}, {"code": 1, "msg": "please login"}):
            with self.subTest(payload=payload), self.assertRaises(service.BigSellerAuthenticationError):
                service.extract_listing_rows(payload)
        with self.assertRaises(service.BigSellerPermissionError):
            service.extract_listing_rows({"code": 403})
        with self.assertRaises(service.BigSellerRateLimitError):
            service.extract_listing_rows({"code": 429})

    def test_business_frequency_errors_are_classified_without_echoing_server_message(self):
        messages = ("The request is too frequent. Please try again later", "TOO MANY requests",
                    "rate limit exceeded", "请求过于频繁", "平台限流")
        for message in messages:
            with self.subTest(message=message), self.assertRaises(service.BigSellerRateLimitError) as caught:
                service.extract_listing_rows({"code": -1, "msg": message + " private-password"})
            self.assertEqual(str(caught.exception), "BigSeller 限制访问频率，请稍后重新查询未完成的 SKU")
        for code, error in ((401, service.BigSellerAuthenticationError), (403, service.BigSellerPermissionError)):
            with self.subTest(code=code), self.assertRaises(error):
                service.extract_listing_rows({"code": code, "msg": "login too frequent private-password"})
        self.assertEqual(service.extract_listing_rows({
            "code": 0, "msg": "too frequent", "data": {"page": {"rows": []}},
        }), [])

    def test_nonzero_business_errors_preserve_safe_numeric_codes_only(self):
        for code in (500, "500", -1, "-1", 9999999999, "-9999999999"):
            payload = {"code": code, "msg": "private-password", "data": {"page": {"rows": []}}}
            with self.subTest(code=code), self.assertRaises(service.BigSellerQueryError) as caught:
                service.extract_listing_rows(payload)
            self.assertEqual(
                str(caught.exception),
                f"BigSeller 商品查询接口返回失败，请检查账号权限或稍后重试（错误码：{code}）",
            )
            self.assertNotIn("private-password", str(caught.exception))

    def test_unknown_or_malicious_business_codes_do_not_leak_response_data(self):
        private_value = "private-password"
        invalid_codes = (None, True, False, 500.0, "", " 500", "500\n", "+500", "５０１",
                         "12345678901", "-12345678901", 10 ** 20, -10 ** 20,
                         {"password": private_value}, [private_value], f"500 {private_value}")
        for code in invalid_codes:
            payload = {"code": code, "msg": private_value,
                       "data": {"page": {"rows": [{"itemId": "123", "password": private_value}]}}}
            with self.subTest(code=code), self.assertRaises(service.BigSellerQueryError) as caught:
                service.extract_listing_rows(payload)
            self.assertEqual(
                str(caught.exception),
                "BigSeller 商品查询接口返回失败，请检查账号权限或稍后重试（错误码：未知）",
            )
            self.assertNotIn(private_value, str(caught.exception))

    def test_request_uses_readonly_listing_post_not_sync_endpoint(self):
        session, util = Mock(), self.make_util()
        response = self.response([{"itemId": "123"}])
        session.post.return_value = response
        self.assertEqual(service.request_first_item_id(session, util, {}, "hqsZ12027A"), "123")
        args, kwargs = session.post.call_args
        self.assertEqual(args[0], "https://www.bigseller.pro/api/v1/product/listing/shopee/pageList.json")
        self.assertEqual(kwargs["json"]["searchContent"], "hqsZ12027A")
        self.assertEqual(kwargs["json"]["searchType"], "sku")
        self.assertEqual(kwargs["json"]["inquireType"], 0)
        self.assertEqual(kwargs["json"]["orderBy"], "views")
        self.assertIs(kwargs["json"]["desc"], True)
        self.assertFalse(kwargs["allow_redirects"])
        self.assertEqual(kwargs["timeout"], (10, 45))
        response.close.assert_called_once()

    def test_transient_failures_retry_and_do_not_leak_exception_text(self):
        session, util = Mock(), self.make_util()
        session.post.side_effect = [requests.Timeout("password=secret"), self.response(status=503), self.response([{"itemId": "123"}])]
        with patch.object(service.time, "sleep"):
            self.assertEqual(service.request_first_item_id(session, util, {}, "sku-a"), "123")
        self.assertEqual(session.post.call_count, 3)
        session.post.side_effect = requests.ConnectionError("Authorization: private-password")
        with patch.object(service.time, "sleep"), self.assertRaises(service.BigSellerQueryError) as caught:
            service.request_first_item_id(session, util, {}, "sku-a")
        self.assertNotIn("private-password", str(caught.exception))
        self.assertTrue(caught.exception.__suppress_context__)

    def test_http_failures_and_bad_json_never_become_no_match(self):
        for status, error_type in ((401, service.BigSellerAuthenticationError), (302, service.BigSellerAuthenticationError),
                                   (403, service.BigSellerPermissionError), (429, service.BigSellerRateLimitError),
                                   (400, service.BigSellerQueryError)):
            session = Mock()
            session.post.return_value = self.response(status=status)
            with self.subTest(status=status), self.assertRaises(error_type):
                service.request_first_item_id(session, self.make_util(), {}, "sku-a")
            self.assertEqual(session.post.call_count, service.REQUEST_ATTEMPTS if status == 429 else 1)
        response = self.response()
        response.json.side_effect = ValueError("secret-response")
        session.post.return_value = response
        with self.assertRaises(service.BigSellerQueryError) as caught:
            service.request_first_item_id(session, self.make_util(), {}, "sku-a")
        self.assertNotIn("secret-response", str(caught.exception))

    def test_generic_http_failure_keeps_status_without_reading_response_payload(self):
        response = self.response(status=400, payload={"code": "private-password", "msg": "private-password"})
        session = Mock()
        session.post.return_value = response
        with self.assertRaises(service.BigSellerQueryError) as caught:
            service.request_first_item_id(session, self.make_util(), {}, "sku-a")
        self.assertEqual(
            str(caught.exception),
            "BigSeller 商品查询请求失败，请检查账号权限或稍后重试（HTTP 状态码：400）",
        )
        response.json.assert_not_called()
        response.close.assert_called_once()
        self.assertEqual(session.post.call_count, 1)

    def test_successful_requests_share_minimum_start_interval_across_sessions(self):
        starts = []

        def success(*args, **kwargs):
            starts.append(self.clock)
            return self.response([{"itemId": "123"}])

        first, second = Mock(), Mock()
        first.post.side_effect = success
        second.post.side_effect = success
        self.assertEqual(service.request_first_item_id(first, self.make_util(), {}, "a"), "123")
        self.clock += 0.2
        self.assertEqual(service.request_first_item_id(second, self.make_util(), {}, "b"), "123")
        self.assertEqual(starts, [100.0, 100.75])
        service.time.sleep.assert_called_once()
        self.assertAlmostEqual(service.time.sleep.call_args.args[0], 0.55)

    def test_concurrent_sessions_cannot_overlap_requests(self):
        entered, release, second_started = threading.Event(), threading.Event(), threading.Event()
        starts, results, errors = [], [], []
        first, second = Mock(), Mock()

        def first_response(*args, **kwargs):
            starts.append(self.clock)
            entered.set()
            if not release.wait(2):
                raise RuntimeError("test request was not released")
            return self.response([{"itemId": "123"}])

        def second_response(*args, **kwargs):
            starts.append(self.clock)
            return self.response([{"itemId": "456"}])

        def query(session, sku, ready=None):
            try:
                if ready is not None:
                    ready.set()
                results.append(service.request_first_item_id(session, self.make_util(), {}, sku))
            except Exception as exc:
                errors.append(exc)

        first.post.side_effect = first_response
        second.post.side_effect = second_response
        workers = [threading.Thread(target=query, args=(first, "a")),
                   threading.Thread(target=query, args=(second, "b", second_started))]
        workers[0].start()
        self.assertTrue(entered.wait(1))
        workers[1].start()
        try:
            self.assertTrue(second_started.wait(1))
            second.post.assert_not_called()
        finally:
            release.set()
            for worker in workers:
                worker.join(2)
        self.assertTrue(all(not worker.is_alive() for worker in workers))
        self.assertEqual(errors, [])
        self.assertEqual(results, ["123", "456"])
        self.assertEqual(starts, [100.0, 100.75])

    def test_business_frequency_retries_current_sku_after_ten_then_twenty_seconds(self):
        session, util, logs, starts = Mock(), self.make_util(), [], []
        responses = [self.response(payload={"code": -1, "msg": "too frequent private-password"}),
                     self.response(payload={"code": -1, "msg": "too frequent private-password"}),
                     self.response([{"itemId": "123"}])]

        def respond(*args, **kwargs):
            starts.append(self.clock)
            return responses[len(starts) - 1]

        session.post.side_effect = respond
        self.assertEqual(service.request_first_item_id(session, util, {}, "HQsy000864B", logs.append), "123")
        self.assertEqual(starts, [100.0, 110.0, 130.0])
        self.assertEqual([call.kwargs["json"]["searchContent"] for call in session.post.call_args_list],
                         ["HQsy000864B"] * 3)
        self.assertNotIn("private-password", str(logs))
        self.assertNotIn("too frequent", str(logs))
        for response in responses:
            response.close.assert_called_once()

    def test_http_frequency_retries_respect_retry_after_seconds(self):
        session, starts = Mock(), []
        responses = [self.response(status=429, headers={"Retry-After": "3"}),
                     self.response([{"itemId": "123"}])]

        def respond(*args, **kwargs):
            starts.append(self.clock)
            return responses[len(starts) - 1]

        session.post.side_effect = respond
        self.assertEqual(service.request_first_item_id(session, self.make_util(), {}, "a"), "123")
        self.assertEqual(starts, [100.0, 103.0])
        service.time.sleep.assert_called_once_with(3.0)

    def test_retry_after_accepts_http_date_and_rejects_unsafe_values(self):
        instant = datetime(2026, 10, 14, 6, 0, 0, tzinfo=timezone.utc)
        with patch.object(service, "datetime") as date:
            date.now.return_value = instant
            self.assertEqual(service._retry_after("Wed, 14 Oct 2026 06:00:07 GMT"), 7.0)
            self.assertEqual(service._retry_after("Wed, 14 Oct 2026 05:59:00 GMT"), 0.0)
        for value in (None, True, False, [], {}, "private-password", "nan", "inf", "-1",
                      float("nan"), float("inf"), 10 ** 5000, "1" * 1000):
            with self.subTest(type=type(value).__name__):
                self.assertEqual(service._retry_after(value), 0.0)

    def test_long_cooldown_is_registered_and_prevents_another_session_request(self):
        first, second, logs = Mock(), Mock(), []
        first.post.return_value = self.response(status=429, headers={"Retry-After": "120"})
        with self.assertRaises(service.BigSellerRateLimitError) as caught:
            service.request_first_item_id(first, self.make_util(), {}, "a", logs.append)
        self.assertIn("较长冷却时间", str(caught.exception))
        self.assertEqual(service._COOLDOWN_UNTIL, 220.0)
        with self.assertRaises(service.BigSellerRateLimitError):
            service.request_first_item_id(second, self.make_util(), {}, "b", logs.append)
        first.post.assert_called_once()
        second.post.assert_not_called()
        service.time.sleep.assert_not_called()
        self.clock = 221.0
        second.post.return_value = self.response([{"itemId": "123"}])
        self.assertEqual(service.request_first_item_id(second, self.make_util(), {}, "b"), "123")

    def test_unknown_negative_code_remains_an_error_without_rate_limit_retry(self):
        session = Mock()
        session.post.return_value = self.response(payload={"code": -1, "msg": "unrecognized private-password"})
        with self.assertRaises(service.BigSellerQueryError) as caught:
            service.request_first_item_id(session, self.make_util(), {}, "a")
        self.assertNotIsInstance(caught.exception, service.BigSellerRateLimitError)
        self.assertIn("错误码：-1", str(caught.exception))
        self.assertNotIn("private-password", str(caught.exception))
        session.post.assert_called_once()
        service.time.sleep.assert_not_called()

    def test_long_server_cooldown_stops_batch_without_querying_remaining_skus(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), ("a", "b", "c"))
            session, logs = Mock(), []
            session.post.return_value = self.response(status=429, headers={"Retry-After": "120"})
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=self.make_util()), \
                 patch.object(service, "login_for_query", return_value=session):
                result = service.run_bigseller_item_id_query(job, logs.append)
            session.post.assert_called_once()
            service.time.sleep.assert_not_called()
            self.assertFalse(result["is_complete"])
            self.assertEqual(result["failed_count"], 3)
            self.assertIn("较长冷却时间", result["rows"][0]["message"])
            self.assertTrue(all("未执行" in row["message"] for row in result["rows"][1:]))

    def test_frequency_retry_exhaustion_stops_batch_and_retains_unexecuted_skus(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), ("a", "b", "c", "d"))
            session, util, logs = Mock(), self.make_util(), []
            session.post.side_effect = [self.response([{"itemId": "123"}])] + [
                self.response(payload={"code": -1, "msg": "too frequent private-password"}) for _ in range(3)
            ]
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=util), \
                 patch.object(service, "login_for_query", return_value=session):
                result = service.run_bigseller_item_id_query(job, logs.append)
            self.assertEqual(session.post.call_count, 4)
            self.assertEqual([call.kwargs["json"]["searchContent"] for call in session.post.call_args_list],
                             ["a", "b", "b", "b"])
            self.assertEqual((result["matched_count"], result["not_found_count"], result["failed_count"]), (1, 0, 3))
            self.assertFalse(result["is_complete"])
            self.assertIn("结果不完整", result["completion_message"])
            self.assertEqual([row["sku"] for row in result["rows"]], list(job.skus))
            self.assertTrue(all("未执行" in row["message"] for row in result["rows"][2:]))
            self.assertNotIn(job.password, str(result) + str(logs))
            self.assertNotIn("too frequent", str(result) + str(logs))
            workbook = load_workbook(result["output_file"], read_only=True)
            try:
                self.assertEqual(workbook.active.max_row, 5)
                self.assertEqual(workbook.active.max_column, 2)
            finally:
                workbook.close()

    def test_confirmed_matches_and_empty_rows_mark_batch_complete(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), ("a", "b"))
            session = Mock()
            session.post.side_effect = [self.response([{"itemId": "123"}]), self.response([])]
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=self.make_util()), \
                 patch.object(service, "login_for_query", return_value=session):
                result = service.run_bigseller_item_id_query(job, lambda _: None)
            self.assertTrue(result["is_complete"])
            self.assertEqual(result["completion_message"], "查询完成")
            self.assertEqual((result["matched_count"], result["not_found_count"], result["failed_count"]), (1, 1, 0))

    def test_excel_has_only_two_text_columns_and_preserves_every_sku(self):
        rows = [
            {"sku": "00123", "item_id": "001234567890123456789"},
            {"sku": "=HYPERLINK(\"https://example.invalid\")", "item_id": ""},
            {"sku": "+123", "item_id": ""},
            {"sku": "@sku", "item_id": "123"},
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            path = service.export_item_ids(Path(temp_dir), rows)
            other_path = service.export_item_ids(Path(temp_dir), rows)
            self.assertNotEqual(path, other_path)
            workbook = load_workbook(path)
            try:
                sheet = workbook.active
                self.assertEqual(workbook.sheetnames, ["SKU-Item ID"])
                self.assertEqual(sheet.max_column, 2)
                self.assertEqual(sheet.max_row, 5)
                self.assertEqual(tuple(cell.value for cell in sheet[1]), ("SKU", "Item ID"))
                self.assertEqual(sheet["A2"].value, "00123")
                self.assertEqual(sheet["B2"].value, "001234567890123456789")
                self.assertEqual(sheet["A3"].value, rows[1]["sku"])
                self.assertEqual(sheet["A3"].data_type, "s")
                self.assertEqual(sheet["B3"].value, None)
                self.assertEqual(sheet.freeze_panes, "A2")
                for row in sheet.iter_rows(min_row=2):
                    for cell in row:
                        self.assertNotEqual(cell.data_type, "f")
                        self.assertEqual(cell.number_format, "@")
            finally:
                workbook.close()

    def test_batch_preserves_failures_and_sanitizes_logs_and_rows(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), ("sku-a", "sku-b", "sku-c", "sku-d"))
            session, util, logs = Mock(), self.make_util(), []
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=util), \
                 patch.object(service, "login_for_query", return_value=session), \
                 patch.object(service, "request_first_item_id", side_effect=["123", "", service.BigSellerQueryError("第一条缺少 Item ID"), RuntimeError(job.password)]):
                result = service.run_bigseller_item_id_query(job, logs.append)
            self.assertEqual((result["sku_count"], result["matched_count"], result["not_found_count"], result["failed_count"]), (4, 1, 1, 2))
            self.assertEqual([row["sku"] for row in result["rows"]], list(job.skus))
            self.assertEqual([row["item_id"] for row in result["rows"]], ["123", "", "", ""])
            self.assertNotIn(job.password, str(result) + "\n".join(logs))
            self.assertIn("[BS进度 4/4]", "\n".join(logs))
            self.assertTrue(Path(result["output_file"]).exists())
            session.close.assert_called_once()

    def test_login_refresh_once_then_stop_requests_keep_remaining_skus(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), ("a", "b", "c"))
            first, second, util = Mock(), Mock(), self.make_util()
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=util), \
                 patch.object(service, "login_for_query", side_effect=[first, second]) as login, \
                 patch.object(service, "request_first_item_id", side_effect=[service.BigSellerAuthenticationError("登录失效"), "123", service.BigSellerAuthenticationError("登录失效")]) as query:
                result = service.run_bigseller_item_id_query(job, lambda _: None)
            self.assertEqual(login.call_count, 2)
            self.assertEqual(query.call_count, 3)
            self.assertEqual(result["matched_count"], 1)
            self.assertEqual(result["failed_count"], 2)
            self.assertIn("未执行", result["rows"][2]["message"])
            first.close.assert_called_once()
            second.close.assert_called_once()

    def test_initial_login_error_is_sanitized_and_legacy_logger_disabled(self):
        util, logs = self.make_util(), []
        secrets = ("private-account", "private-password", "private-session", "private-captcha-1234")
        util.login_bigseller_by_request.side_effect = RuntimeError(
            "unknown upstream failure account=private-account password=private-password "
            "Cookie=private-session captcha=private-captcha-1234")
        with self.assertRaises(service.BigSellerAuthenticationError) as caught:
            service.login_for_query(util, {}, logs.append)
        util.login_bigseller_by_request.assert_called_once_with({}, logger=None)
        for secret in secrets:
            self.assertNotIn(secret, str(caught.exception) + str(logs))
        self.assertNotIn("unknown upstream failure", str(caught.exception))
        self.assertTrue(caught.exception.__suppress_context__)

    def test_known_login_rejections_keep_safe_reason_without_raw_upstream_data(self):
        secrets = ("private-account", "private-password", "private-session", "private-captcha-1234")
        raw_details = (" account=private-account password=private-password "
                       "Cookie=private-session captcha=private-captcha-1234")
        for reason in ("账号或密码错误", "账号已停用"):
            messages = []
            for suffix in ("", raw_details):
                with self.subTest(reason=reason, includes_raw_details=bool(suffix)):
                    util, logs = self.make_util(), []
                    util.login_bigseller_by_request.side_effect = RuntimeError(
                        f"BigSeller 登录失败: {{'code': -1, 'msg': '{reason}'}}{suffix}")
                    with self.assertRaises(service.BigSellerAuthenticationError) as caught:
                        service.login_for_query(util, {}, logs.append)
                    message = str(caught.exception)
                    messages.append(message)
                    self.assertIn(reason, message)
                    self.assertNotIn("'code'", message)
                    for secret in secrets:
                        self.assertNotIn(secret, message + str(logs))
                    self.assertTrue(caught.exception.__suppress_context__)
                    util.login_bigseller_by_request.assert_called_once_with({}, logger=None)
            self.assertEqual(messages[0], messages[1])

    def test_permission_or_rate_limit_stops_batch_without_dropping_rows(self):
        for error in (service.BigSellerPermissionError("无权访问"), service.BigSellerRateLimitError("平台限流")):
            with self.subTest(error=type(error).__name__), tempfile.TemporaryDirectory() as temp_dir:
                job = self.make_job(Path(temp_dir), ("a", "b", "c"))
                with patch.object(service, "query_process_config", return_value={}), \
                     patch.object(service, "load_bigseller_request_util", return_value=self.make_util()), \
                     patch.object(service, "login_for_query", return_value=Mock()) as login, \
                     patch.object(service, "request_first_item_id", side_effect=error) as query:
                    result = service.run_bigseller_item_id_query(job, lambda _: None)
                self.assertEqual(query.call_count, 1)
                self.assertEqual(login.call_count, 1)
                self.assertEqual(result["failed_count"], 3)
                self.assertEqual(result["not_found_count"], 0)
                self.assertEqual([row["sku"] for row in result["rows"]], ["a", "b", "c"])
                self.assertEqual([row["item_id"] for row in result["rows"]], ["", "", ""])
                self.assertIn("未执行", result["rows"][1]["message"])

    def test_preview_is_limited_but_export_is_complete(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            job = self.make_job(Path(temp_dir), tuple(f"sku-{i}" for i in range(501)))
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=self.make_util()), \
                 patch.object(service, "login_for_query", return_value=Mock()), \
                 patch.object(service, "request_first_item_id", return_value="123"):
                result = service.run_bigseller_item_id_query(job, lambda _: None)
            self.assertEqual(len(result["rows"]), 500)
            self.assertTrue(result["preview_limited"])
            workbook = load_workbook(result["output_file"], read_only=True)
            try:
                self.assertEqual(workbook.active.max_row, 502)
            finally:
                workbook.close()

    def test_query_config_uses_saved_account_and_captcha_without_environment(self):
        job = self.make_job(Path("."))
        reference_config = {
            "bigseller_account": {"username": "reference-bs-user", "password": "reference-bs-password"},
            "captcha_service": {
                "username": "reference-captcha-user", "password": "reference-captcha-password",
                "max_attempts": 8,
            },
        }
        with patch("backend.services.bigseller_sync.load_reference_bigseller_config", return_value=reference_config), \
             patch.dict("os.environ", {}, clear=True):
            config = service.query_process_config(job)

        self.assertEqual(config["bigseller_account"], {"username": job.username, "password": job.password})
        self.assertEqual(config["captcha_service"]["username"], job.captcha_username)
        self.assertEqual(config["captcha_service"]["password"], job.captcha_password)
        self.assertEqual(config["captcha_service"]["max_attempts"], 3)
        util = self.make_util()
        service.login_for_query(util, config, lambda _: None)
        util.login_bigseller_by_request.assert_called_once_with(config, logger=None)

    def test_query_config_environment_captcha_overrides_saved_settings(self):
        job = self.make_job(Path("."))
        environment = {
            "HQYL_TTSHITU_USERNAME": "environment-captcha-user",
            "HQYL_TTSHITU_PASSWORD": "environment-captcha-password",
        }
        with patch("backend.services.bigseller_sync.load_reference_bigseller_config", return_value={}), \
             patch.dict("os.environ", environment, clear=True):
            config = service.query_process_config(job)

        self.assertEqual(config["bigseller_account"], {"username": job.username, "password": job.password})
        self.assertEqual(config["captcha_service"]["username"], environment["HQYL_TTSHITU_USERNAME"])
        self.assertEqual(config["captcha_service"]["password"], environment["HQYL_TTSHITU_PASSWORD"])
        self.assertEqual(config["captcha_service"]["max_attempts"], 3)

    def test_query_config_trims_login_credentials_like_web_form_and_keeps_internal_spaces(self):
        job = replace(self.make_job(Path(".")),
                      username=" \tfixture-user@13CT0000000\r\n", password="\t pass word @ \r\n")
        with patch("backend.services.bigseller_sync.load_reference_bigseller_config", return_value={}), \
             patch.dict("os.environ", {}, clear=True):
            config = service.query_process_config(job)

        self.assertEqual(config["bigseller_account"], {
            "username": "fixture-user@13CT0000000", "password": "pass word @",
        })
        util = self.make_util()
        service.login_for_query(util, config, lambda _: None)
        util.login_bigseller_by_request.assert_called_once_with(config, logger=None)


if __name__ == "__main__":
    unittest.main()
