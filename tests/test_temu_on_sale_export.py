from __future__ import annotations

import copy
import json
import shutil
import tempfile
import unittest
from collections import Counter
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from backend.core.ziniao_cli import ZiniaoCli, ZiniaoCliError
from backend.services import temu_on_sale_export as service
from backend.services.temu_on_sale_gateway import TemuOnSaleGateway
from backend.services.temu_on_sale_workbook import COLUMNS, SHEET_NAME, file_sha256, inspect_source, merge_sources


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = ROOT / "docs/requirements/temu-on-sale-export-20260915/111商品基础信息 (1).xlsx"


def store(identity, name=None):
    return {"store_id": str(identity), "store_name": name or f"店铺 {identity}", "platform_name": "TEMU 东南亚"}


def source(path, rows=None, headers=COLUMNS):
    book = Workbook()
    sheet = book.active
    sheet.title = SHEET_NAME
    sheet.append(list(headers))
    if rows is None:
        rows = [{"店铺名": "错误旧店名", "商品标题": "=商品标题", "SPU ID": "000123", "SKC ID": 12345678901,
                 "SKU ID": "001234567890123456789", "SKU货号": "0007", "商品状态": "在售中",
                 "申报价格(USD)": 1.2345, "申报价格状态": "已作废", "库存": 0, "经营站点": "泰国站"}]
    for row in rows:
        sheet.append([row.get(name) for name in headers])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.row_dimensions[2].hidden = True
    sheet.auto_filter.ref = sheet.dimensions
    sheet.auto_filter.add_filter_column(14, ["已生效"])
    book.save(path)
    book.close()
    return path


class FakeGateway:
    profile = "测试授权账号"
    auth_failed = False

    def __init__(self, original, failures=(), empty=()):
        self.original, self.failures, self.empty = original, set(failures), set(empty)
        self.calls = []

    def preflight(self, stores):
        pass

    def export(self, store, directory):
        identity = store["store_id"]
        self.calls.append(identity)
        if identity in self.failures:
            raise ValueError("测试下载失败")
        if identity in self.empty:
            return {"no_data": True, "evidence": {"tab": "在售中0", "empty_result": True}}
        path = directory / "平台导出.xlsx"
        shutil.copyfile(self.original, path)
        return {"path": str(path)}


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_all_hidden_rows_ids_numeric_values_and_single_store_column(self):
        original = source(self.root / "input.xlsx")
        before = file_sha256(original)
        info = inspect_source(original)
        output = self.root / "summary.xlsx"
        merge_sources([("紫鸟真实店名", info), ("另一个店", info)], output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(tuple(cell.value for cell in sheet[1]), COLUMNS)
        self.assertEqual(sheet["A2"].value, "紫鸟真实店名")
        self.assertEqual(sheet["A3"].value, "另一个店")
        self.assertEqual(sheet["B2"].value, "=商品标题")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["C2"].value, "000123")
        self.assertEqual(sheet["E2"].value, "001234567890123456789")
        self.assertEqual(sheet["E2"].number_format, "@")
        self.assertEqual(sheet["N2"].value, 1.2345)
        self.assertEqual(sheet["P2"].value, 0)
        self.assertFalse(sheet.row_dimensions[2].hidden)
        self.assertEqual(sheet.auto_filter.filterColumn, [])
        self.assertEqual(file_sha256(original), before)

    def test_different_field_order_and_new_columns(self):
        columns = tuple(reversed(COLUMNS[1:])) + ("新字段",)
        data = {"商品标题": "标题", "SPU ID": "1", "SKC ID": "2", "SKU ID": "3", "商品状态": "在售中", "新字段": "新值"}
        first = inspect_source(source(self.root / "a.xlsx"))
        second = inspect_source(source(self.root / "b.xlsx", [data], columns))
        self.assertEqual(second.added_columns, ["新字段"])
        output = self.root / "out.xlsx"
        merge_sources([("A", first), ("B", second)], output)
        book = load_workbook(output)
        self.addCleanup(book.close)
        self.assertEqual(book.active["E3"].value, "3")
        self.assertEqual(book.active["R3"].value, "新值")

    def test_numeric_identifier_zero_format(self):
        original = source(self.root / "input.xlsx")
        book = load_workbook(original)
        book.active["C2"] = 123
        book.active["C2"].number_format = "000000"
        book.save(original)
        book.close()
        out = self.root / "out.xlsx"
        merge_sources([("店", inspect_source(original))], out)
        book = load_workbook(out)
        self.addCleanup(book.close)
        self.assertEqual(book.active["C2"].value, "000123")

    def test_invalid_download_or_non_sale_or_formulas_are_not_zero_success(self):
        invalid = self.root / "error.xlsx"
        invalid.write_text("<html>请登录</html>", encoding="utf-8")
        with self.assertRaises(ValueError):
            inspect_source(invalid)
        original = source(self.root / "input.xlsx")
        for cell, value in (("J2", "已下架"), ("E2", None), ("B2", "=1+1")):
            with self.subTest(cell=cell):
                source(original)
                book = load_workbook(original)
                book.active[cell] = value
                book.save(original)
                book.close()
                with self.assertRaises(ValueError):
                    inspect_source(original)

    def test_real_requirement_sample_14340_rows(self):
        if not SAMPLE.is_file():
            self.skipTest("需求样例未随仓库提供")
        info = inspect_source(SAMPLE)
        self.assertEqual(info.row_count, 14340)
        self.assertEqual(info.sha256.upper(), "6C29B9F2513561E1BAE72D88501FA2A7871175C5B0ED04F980BC9D6407466158")
        output = self.root / "sample_merged.xlsx"
        merge_sources([("样例验收店铺", info)], output)
        book = load_workbook(output, read_only=True)
        self.addCleanup(book.close)
        rows = list(book.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 14340)
        self.assertTrue(all(row[0] == "样例验收店铺" for row in rows))
        self.assertEqual(Counter(row[14] for row in rows), {"已生效": 9119, "已作废": 5221})


class BatchTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.original = source(self.root / "input.xlsx")

    def batch(self, ids=("a", "b", "c")):
        return service.create_batch(self.root, FakeGateway.profile, [store(identity) for identity in ids])

    def test_failure_continues_retry_only_failures_rebuilds_in_order(self):
        path = self.batch()
        gateway = FakeGateway(self.original, failures=["b"])
        result = service.run_batch(path, gateway, lambda _: None)
        self.assertEqual(gateway.calls, ["a", "b", "c"])
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["row_count"], 2)
        self.assertIn("部分结果", result["output_file"])
        old_partial = Path(result["output_file"])
        original_a = result["stores"][0]["raw_file"]
        gateway.failures.clear()
        gateway.calls.clear()
        result = service.run_batch(path, gateway, lambda _: None)
        self.assertEqual(gateway.calls, ["b"])
        self.assertTrue(result["is_complete"])
        self.assertEqual(result["row_count"], 3)
        self.assertEqual(result["stores"][0]["raw_file"], original_a)
        self.assertTrue(old_partial.is_file())
        book = load_workbook(result["output_file"], read_only=True)
        self.addCleanup(book.close)
        self.assertEqual([row[0] for row in book.active.iter_rows(min_row=2, values_only=True)], ["店铺 a", "店铺 b", "店铺 c"])

    def test_all_failed_has_no_summary_all_confirmed_empty_has_header(self):
        path = self.batch()
        result = service.run_batch(path, FakeGateway(self.original, failures=["a", "b", "c"]), lambda _: None)
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["output_file"], "")
        result = service.run_batch(path, FakeGateway(self.original, empty=["a", "b", "c"]), lambda _: None)
        self.assertTrue(result["is_complete"])
        self.assertEqual(result["row_count"], 0)
        book = load_workbook(result["output_file"], read_only=True)
        self.addCleanup(book.close)
        self.assertEqual(list(book.active.values), [COLUMNS])

    def test_empty_download_does_not_mean_no_data(self):
        empty = source(self.root / "empty.xlsx", [])
        result = service.run_batch(self.batch(["a"]), FakeGateway(empty), lambda _: None)
        self.assertEqual(result["failed_count"], 1)
        self.assertEqual(result["no_data_count"], 0)

    def test_previous_download_outside_attempt_is_rejected(self):
        gateway = FakeGateway(self.original)
        gateway.export = lambda *_: {"path": str(self.original)}
        result = service.run_batch(self.batch(["a"]), gateway, lambda _: None)
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("本店本次", result["stores"][0]["error"])

    def test_merge_failure_can_retry_without_exporting_successful_store(self):
        path = self.batch(["a"])
        gateway = FakeGateway(self.original)
        with patch.object(service, "merge_sources", side_effect=PermissionError("文件正在使用")):
            result = service.run_batch(path, gateway, lambda _: None)
        self.assertFalse(result["is_complete"])
        self.assertTrue(result["can_retry"])
        gateway.calls.clear()
        result = service.run_batch(path, gateway, lambda _: None)
        self.assertEqual(gateway.calls, [])
        self.assertTrue(result["is_complete"])

    def test_auth_error_stops_further_store_requests(self):
        gateway = FakeGateway(self.original)
        def denied(store, directory):
            gateway.calls.append(store["store_id"])
            gateway.auth_failed = True
            raise ValueError("紫鸟认证失败")
        gateway.export = denied
        result = service.run_batch(self.batch(), gateway, lambda _: None)
        self.assertEqual(gateway.calls, ["a"])
        self.assertEqual(result["failed_count"], 3)

    def test_interrupted_batch_preserves_successes_and_resumes_pending(self):
        path = self.batch(["a", "b"])
        gateway = FakeGateway(self.original, failures=["b"])
        service.run_batch(path, gateway, lambda _: None)
        record = service.load_batch(path)
        record["status"] = "running"
        record["stores"][1]["status"] = "running"
        service.write_json(path, record)
        gateway.failures.clear()
        gateway.calls.clear()
        result = service.run_batch(path, gateway, lambda _: None)
        self.assertEqual(gateway.calls, ["b"])
        self.assertTrue(result["is_complete"])

    def test_native_gateway_closes_every_store_before_next_even_after_failure(self):
        path = self.batch()
        browser = Mock(profile=FakeGateway.profile, auth_failed=False, cleanup_failed=False)
        browser.list_stores.return_value = [store(identity) for identity in ('a','b','c')]
        gateway = TemuOnSaleGateway(browser.profile, browser=browser)
        def export_one(item, raw_dir):
            browser.open_store(item)
            if item['store_id'] == 'b':
                raise ValueError('download failed')
            return {'no_data':True, 'evidence':{'empty_result':True}}
        gateway._export = export_one
        result = service.run_batch(path, gateway, lambda _:None)
        order = [(call[0], call.args[0]['store_id'] if call[0] == 'open_store' else call.args[0])
                 for call in browser.mock_calls if call[0] in ('open_store','close_store')]
        self.assertEqual(order, [('open_store','a'),('close_store','a'),('open_store','b'),
                                 ('close_store','b'),('open_store','c'),('close_store','c')])
        self.assertEqual(result['failed_count'], 1)

    def test_native_gateway_close_failure_halts_batch_before_next_store(self):
        browser = Mock(profile=FakeGateway.profile, auth_failed=False, cleanup_failed=False)
        browser.list_stores.return_value = [store(identity) for identity in ('a','b','c')]
        def close_failed(identity):
            browser.cleanup_failed = True
            raise ValueError('店铺未确认关闭')
        browser.close_store.side_effect = close_failed
        gateway = TemuOnSaleGateway(browser.profile, browser=browser)
        gateway._export = Mock(return_value={'no_data':True, 'evidence':{'empty_result':True}})
        result = service.run_batch(self.batch(), gateway, lambda _:None)
        gateway._export.assert_called_once()
        browser.close_store.assert_called_once_with('a')
        self.assertEqual(result['failed_count'], 3)
        self.assertIn('未确认关闭', result['stores'][1]['error'])

    def test_zero_data_page_executes_filter_clicks_then_closes_store(self):
        item = store('a','泰国-testSeller')
        browser = Mock(profile=FakeGateway.profile, auth_failed=False, cleanup_failed=False)
        gateway = TemuOnSaleGateway(browser.profile, browser=browser)
        gateway.current_stores = {'a':item}
        found = {label:[{'selector':'#'+str(index), 'text':label, 'selected':True}]
                 for index,label in enumerate(('商品列表','重置','在售中','查询','下载查询结果'))}
        found['在售中'][0]['text'] = '在售中0'
        found['下载查询结果'][0]['text'] = '下载查询结果0'
        browser.evaluate.return_value = {'url':'https://agentseller.temu.com/goods/list',
            'found':found, 'filters':[{'tag':'INPUT','value':''}], 'sellers':['testSeller'], 'body':'暂无数据'}
        result = gateway.export(item, self.root)
        self.assertTrue(result['no_data'])
        self.assertEqual([call.args for call in browser.click.call_args_list],
                         [('a','#1'),('a','#2'),('a','#3')])
        browser.close_store.assert_called_once_with('a')

    def test_known_todo_notice_closes_before_filter_click(self):
        browser=Mock(profile=FakeGateway.profile)
        gateway=TemuOnSaleGateway(browser.profile,browser=browser)
        browser.evaluate.side_effect=[{'notices':['#todo-close']},{'notices':[]},
            {'url':'https://agentseller.temu.com/goods/list','found':{'重置':[{'selector':'#reset'}]}}]
        with patch('backend.services.temu_on_sale_gateway.time.sleep'):
            gateway.click_label('a','重置')
        self.assertEqual([call.args for call in browser.click.call_args_list],[('a','#todo-close'),('a','#reset')])

    def test_ambiguous_todo_notice_does_not_click_page(self):
        browser=Mock(profile=FakeGateway.profile)
        gateway=TemuOnSaleGateway(browser.profile,browser=browser)
        browser.evaluate.return_value={'notices':['#first','#second']}
        with self.assertRaisesRegex(ValueError,'多个'):
            gateway.click_label('a','重置')
        browser.click.assert_not_called()

    def test_late_todo_notice_retries_only_confirmed_intercepted_click(self):
        from backend.core.ziniao_browser import ZiniaoClickInterceptedError
        browser=Mock(profile=FakeGateway.profile)
        gateway=TemuOnSaleGateway(browser.profile,browser=browser)
        state={'url':'https://agentseller.temu.com/goods/list','found':{'重置':[{'selector':'#reset'}]}}
        browser.evaluate.side_effect=[{'notices':[]},state,{'notices':['#todo-close']},
                                      {'notices':[]},{'notices':[]},state]
        browser.click.side_effect=[ZiniaoClickInterceptedError('blocked'),None,None]
        with patch('backend.services.temu_on_sale_gateway.time.sleep'):
            gateway.click_label('a','重置')
        self.assertEqual([call.args for call in browser.click.call_args_list],
                         [('a','#reset'),('a','#todo-close'),('a','#reset')])

    def test_unknown_overlay_and_uncertain_click_are_not_retried(self):
        from backend.core.ziniao_browser import ZiniaoClickInterceptedError, ZiniaoBrowserError
        for error in (ZiniaoClickInterceptedError('unknown overlay'),ZiniaoBrowserError('uncertain')):
            with self.subTest(error=type(error).__name__):
                browser=Mock(profile=FakeGateway.profile)
                gateway=TemuOnSaleGateway(browser.profile,browser=browser)
                browser.evaluate.side_effect=[{'notices':[]},
                    {'url':'https://agentseller.temu.com/goods/list','found':{'下载查询结果':[{'selector':'#download'}]}},
                    {'notices':[]}]
                browser.click.side_effect=error
                with self.assertRaises(type(error)):
                    gateway.click_label('a','下载查询结果')
                browser.click.assert_called_once_with('a','#download')

    def test_batch_directories_unique_and_account_path_lock_validation(self):
        first, second = self.batch(), self.batch()
        self.assertNotEqual(first, second)
        with self.assertRaisesRegex(ValueError, "账号"):
            service.load_batch(first, "其他账号")
        with service.batch_lock(first):
            with self.assertRaisesRegex(ValueError, "正在运行"):
                with service.batch_lock(first):
                    pass
        record = service.load_batch(first)
        record["stores"][0]["raw_file"] = "../input.xlsx"
        service.write_json(first, record)
        with self.assertRaisesRegex(ValueError, "超出"):
            service.load_batch(first)


class SelectionAndCliTests(unittest.TestCase):
    def test_name_dedup_ambiguity_and_platform_metadata(self):
        stores = [store("1", "泰国 A"), store("2", "重名"), store("3", "重名"), {**store("4", "菲律宾 TEMU 假店"), "platform_name": "Lazada"}]
        result = service.match_stores("泰国 A\n\n泰国 A\n重名\n没匹配\n菲律宾 TEMU 假店", stores)
        self.assertFalse(result["can_start"])
        self.assertEqual(len(result["stores"]), 1)
        self.assertEqual(len(result["issues"]), 3)
        self.assertEqual(service.select_stores(["2", "2"], stores), [stores[1]])

    def test_filters_fail_closed_for_country_keyword_and_unknown_controls(self):
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": None}))
        self.assertTrue(TemuOnSaleGateway.filters_clear({"filters": [{"tag": "INPUT", "value": ""}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{"tag": "INPUT", "value": "关键词"}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{"tag": "SELECT", "text": "泰国站", "value": "th"}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{"role": "combobox", "selectedText": "泰国站"}]}))

    def test_live_beast_controls_allow_modes_but_reject_actual_filters(self):
        control = {"tag": "INPUT", "readOnly": True, "control": "beast-core-select-htmlInput"}
        for mode in ("SKC", "SKU(精准查询)", "包含", "不包含"):
            self.assertTrue(TemuOnSaleGateway.filters_clear({"filters": [{**control, "value": mode}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{**control, "value": "泰国站"}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{"tag": "INPUT", "value": "SKC"}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"filters": [{**control, "value": "", "selectedText": "泰国站"}]}))
        self.assertFalse(TemuOnSaleGateway.filters_clear({"quick_filters_clear": False, "filters": [{"tag": "INPUT", "value": ""}]}))

    def test_query_waits_for_matching_download_count_and_seller(self):
        state = {"found": {"在售中": [{"text": "在售中650", "selected": True}],
                           "下载查询结果": [{"text": "下载查询结果2578"}]},
                 "filters": [{"tag": "INPUT", "value": ""}], "sellers": ["herefn"]}
        self.assertFalse(TemuOnSaleGateway.query_ready(state))
        state["found"]["下载查询结果"][0]["text"] = "下载查询结果650"
        self.assertTrue(TemuOnSaleGateway.query_ready(state))
        self.assertEqual(TemuOnSaleGateway.verify_seller(store("id", "袁晶-菲律宾021-和韌-herefn"), state), "herefn")
        with self.assertRaisesRegex(ValueError, "卖家名"):
            TemuOnSaleGateway.verify_seller(store("id", "泰国-homasrapkA"), state)

    def test_export_confirmation_rejects_truncation_and_count_mismatch(self):
        cli = Mock(profile="测试")
        gateway = TemuOnSaleGateway("测试", cli=cli)
        for expected, returned in ((650, 2578), (40001, 40001)):
            cli.evaluate.return_value = {"dialogs": [{"skcCount": returned}]}
            with self.assertRaisesRegex(ValueError, "数量|上限"):
                gateway.confirm_export("store", expected)
        cli.click.assert_not_called()

    def cli(self):
        cli = object.__new__(ZiniaoCli)
        cli.profile = "当前账号"
        cli.auth_failed = False
        cli.current_profile = Mock(return_value=cli.profile)
        return cli

    def test_cli_auth_stops_requests_and_redacts_raw_response(self):
        cli = self.cli()
        cli._process = Mock(return_value=SimpleNamespace(returncode=1, stdout=json.dumps({"ok": False, "error": {"type": "auth", "message": "SECRET-key"}}), stderr=""))
        for _ in range(2):
            with self.assertRaises(ZiniaoCliError) as error:
                cli.list_stores()
            self.assertNotIn("SECRET-key", str(error.exception))
        cli._process.assert_called_once()

    def test_cli_plain_mouse_ack_is_unverified_and_does_not_repeat_click(self):
        cli = self.cli()
        cli._process = Mock(return_value=SimpleNamespace(returncode=0, stdout="", stderr="✓ 鼠标指令：已发送\n! DOM 点击事件：Bridge 未返回回执\n- 导航观察：未请求\n"))
        self.assertEqual(cli.run(["page", "click", "--store-id", "id", "--selector", "button"]), {"sent": True, "verified": False})
        cli._process.assert_called_once()
        cli._process.return_value.stderr = "unrecognized response"
        with self.assertRaises(ZiniaoCliError):
            cli.run(["page", "click"])

    def test_cli_identity_and_exec_exception_checks(self):
        cli = self.cli()
        cli.run = Mock(return_value={"storeId": "wrong", "name": "A"})
        with self.assertRaisesRegex(ZiniaoCliError, "不一致"):
            cli.open_store(store("id", "A"))
        cli.run.return_value = {"data": {"exceptionDetails": {"text": "bad"}, "result": "null"}}
        with self.assertRaises(ZiniaoCliError):
            cli.evaluate("id", "script")
        cli.run.return_value = {"data": {"result": '{"found":true}'}}
        self.assertEqual(cli.evaluate("id", "script"), {"found": True})


if __name__ == "__main__":
    unittest.main()
