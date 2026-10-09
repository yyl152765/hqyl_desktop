from __future__ import annotations

import json
import tempfile
import unittest
from contextlib import ExitStack, contextmanager
from dataclasses import replace
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

import requests
from openpyxl import Workbook, load_workbook

from backend.services import bigseller_sku_benchmark as service
from backend.services.bigseller_benchmark_checkpoint import Checkpoint
from backend.app_bridge import AppBridge
from backend.config_store import AppSettings, BoundAccount


def listing(shop: str, views=10, sales=1, item="1"):
    return {"shopId": shop, "shopName": "店铺" + shop, "views": views, "sales": sales, "itemId": item}


class BenchmarkTests(unittest.TestCase):
    def setUp(self):
        for target, value in (("_LAST_REQUEST", 0.0), ("_COOLDOWN_UNTIL", 0.0)):
            replacement = patch.object(service, target, value)
            replacement.start()
            self.addCleanup(replacement.stop)
        sleep = patch.object(service.time, "sleep")
        sleep.start()
        self.addCleanup(sleep.stop)

    def source(self, directory, skus=("a", "b", "a", "c")):
        path = Path(directory) / "source.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "滞销"
        sheet.append(["月度滞销 SKU"])
        sheet.append(["库存SKU编号", "浏览量第一店铺", "其他计算"])
        for sku in skus:
            sheet.append([sku, "旧店铺", "=1+2"])
        book.save(path)
        book.close()
        return path

    def response(self, rows=(), *, status=200, total=None, headers=None, payload=None):
        page = {"rows": list(rows)}
        if total is not None:
            page["totalSize"] = total
        response = Mock(status_code=status, headers=headers or {})
        response.json.return_value = payload if payload is not None else {"code": 0, "data": {"page": page}}
        return response

    def request(self, response):
        with patch.object(service, "_post_listing", return_value=response):
            return service.request_benchmark_page(Mock(), Mock(), {}, "a", "views", 1, Mock())

    def query_pages(self, pages, metric="views"):
        with patch.object(service, "request_benchmark_page", side_effect=pages):
            return service.query_benchmark_sku(Mock(), Mock(), {}, "a", metric, Mock())

    def test_metrics_change_both_request_order_and_winner(self):
        rows = [listing("A", 300, 2), listing("B", 5, 99)]
        for metric, expected in (("views", "店铺A"), ("sales", "店铺B")):
            self.assertEqual(service.build_benchmark_query("child-SKU", metric, 2), {
                "searchType": "sku", "searchContent": "child-SKU", "inquireType": 0,
                "shopeeStatus": "live", "status": "active", "orderBy": metric, "desc": True,
                "pageNo": 2, "pageSize": 50, "timeType": "create_time", "startDateStr": "", "endDateStr": "",
            })
            self.assertEqual(self.query_pages([(rows, 2)], metric)["shop_name"], expected)

    def test_complete_pagination_uses_max_per_store_not_sum(self):
        pages = [([listing("A", 60, item="1"), listing("A", 70, item="2")], 4),
                 ([listing("B", 100, item="3"), listing("C", 10, item="4")], 4)]
        with patch.object(service, "request_benchmark_page", side_effect=pages) as request:
            result = service.query_benchmark_sku(Mock(), Mock(), {}, "a", "views", Mock())
        self.assertEqual(result["shop_name"], "店铺B")
        self.assertEqual(result["metric_value"], 100)
        self.assertEqual(result["shop_count"], 3)
        self.assertEqual([call.args[5] for call in request.call_args_list], [1, 2])

    def test_zero_valid_missing_metric_or_shop_is_failure(self):
        self.assertEqual(self.query_pages([([listing("A", 0)], 1)])["status"], "matched")
        self.assertEqual(service.metric_number("1,234"), 1234)
        for value in (None, True, "", "nan", "1,2", "5k", -1, 1.1, float("nan")):
            with self.subTest(value=value), self.assertRaises(service.BigSellerQueryError):
                self.query_pages([([listing("A", value)], 1)])
        with self.assertRaisesRegex(service.BigSellerQueryError, "店铺名称"):
            self.query_pages([([{"views": 2, "shopId": 1}], 1)])

    def test_sales_does_not_fall_back_to_views(self):
        with self.assertRaises(service.BigSellerQueryError):
            self.query_pages([([listing("A", 999, None)], 1)], "sales")

    def test_only_verified_empty_is_not_found(self):
        self.assertEqual(self.query_pages([([], 0)])["status"], "not_found")
        for payload in ({"code": 0}, {"code": 0, "data": {"page": {}}},
                        {"code": 0, "data": {"page": {"rows": None}}}, {"code": 500}):
            with self.subTest(payload=payload), self.assertRaises(service.BigSellerQueryError):
                self.request(self.response(payload=payload))

    def test_incomplete_repeated_changed_pages_fail_without_a_winner(self):
        first = [listing("A", 200)]
        for pages in ([(first, 2), ([], 2)], [(first, 2), (first, 2)],
                      [(first, 2), ([listing("B")], 3)], [(first, 0)]):
            with self.subTest(pages=pages), self.assertRaises(service.BigSellerQueryError):
                self.query_pages(pages)

    def test_no_total_full_page_fetches_next_and_short_page_stops(self):
        full = [listing(str(i), i, item=str(i)) for i in range(service.PAGE_SIZE)]
        result = self.query_pages([(full, None), ([listing("winner", 999)], None)])
        self.assertEqual(result["shop_name"], "店铺winner")
        self.assertEqual(result["shop_count"], 51)

    def test_store_key_does_not_use_variation_shop_sku_id(self):
        rows = [listing("A", 20, item="1"), listing("A", 30, item="2")]
        for i, row in enumerate(rows):
            del row["shopId"]
            row["shopSkuId"] = str(i)
        result = self.query_pages([(rows, 2)])
        self.assertEqual(result["shop_count"], 1)
        self.assertEqual(result["metric_value"], 30)

    def test_rate_limit_retries_respects_delay_and_closes_responses(self):
        limited = self.response(status=429, headers={"Retry-After": "7"})
        success = self.response([listing("A")], total=1)
        with patch.object(service, "_post_listing", side_effect=[limited, success]) as post, \
             patch.object(service.time, "sleep") as sleep:
            rows, total = service.request_benchmark_page(Mock(), Mock(), {}, "a", "sales", 1, Mock())
        self.assertEqual((len(rows), total), (1, 1))
        self.assertEqual(post.call_count, 2)
        sleep.assert_called_once_with(7)
        limited.close.assert_called_once()
        success.close.assert_called_once()

    def test_business_rate_limit_retries_and_persistent_limit_stops(self):
        response = self.response(payload={"code": 9, "msg": "请求过于频繁"})
        with patch.object(service, "_post_listing", return_value=response) as post, \
             patch.object(service.time, "sleep"), self.assertRaises(service.BigSellerRateLimitError):
            service.request_benchmark_page(Mock(), Mock(), {}, "a", "views", 1, Mock())
        self.assertEqual(post.call_count, service.REQUEST_ATTEMPTS)
        response = self.response(status=429, headers={"Retry-After": "120"})
        with patch.object(service, "_post_listing", return_value=response) as post, \
             patch.object(service.time, "sleep") as sleep, self.assertRaises(service.BigSellerRateLimitError):
            service.request_benchmark_page(Mock(), Mock(), {}, "a", "views", 1, Mock())
        post.assert_called_once()
        sleep.assert_not_called()

    def test_transient_timeout_retries_without_leaking_secrets(self):
        with patch.object(service, "_post_listing", side_effect=requests.Timeout("password=private")), \
             patch.object(service.time, "sleep"), self.assertRaises(service.BigSellerQueryError) as caught:
            service.request_benchmark_page(Mock(), Mock(), {}, "a", "views", 1, Mock())
        self.assertNotIn("private", str(caught.exception))
        self.assertTrue(caught.exception.__suppress_context__)

    def test_503_long_cooldown_stops_batch_instead_of_skipping_to_next_sku(self):
        response = self.response(status=503, headers={"Retry-After": "120"})
        with patch.object(service, "_post_listing", return_value=response) as post, \
             patch.object(service.time, "sleep") as sleep, self.assertRaises(service.BigSellerRateLimitError):
            service.request_benchmark_page(Mock(), Mock(), {}, "a", "views", 1, Mock())
        post.assert_called_once()
        sleep.assert_not_called()

    def test_request_gate_and_http_contract(self):
        session, util = Mock(), Mock()
        util.build_api_url.side_effect = lambda path: "https://example.invalid/api" + path
        with patch.object(service, "api_headers", return_value={}), \
             patch.object(service, "_LAST_REQUEST", 10.0), \
             patch.object(service.time, "monotonic", return_value=10.2), \
             patch.object(service.time, "sleep") as sleep:
            service._post_listing(session, util, {}, service.build_benchmark_query("a", "sales"))
        self.assertAlmostEqual(sleep.call_args.args[0], 0.55)
        kwargs = session.post.call_args.kwargs
        self.assertFalse(kwargs["allow_redirects"])
        self.assertEqual(kwargs["timeout"], (10, 45))
        self.assertEqual(kwargs["json"]["orderBy"], "sales")

    def test_validation_rejects_bad_metric_and_missing_sheet_before_login(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            payload = dict(username="user", password=" secret ", source_file=str(source), output_dir=directory)
            job = service.validate_bigseller_benchmark_payload(payload)
            self.assertEqual(job.metric, "views")
            self.assertEqual(job.sheet_name, "滞销")
            self.assertEqual(job.password, " secret ")
            self.assertNotIn("secret", repr(job))
            for change in ({"metric": "bad"}, {"sheet_name": "missing"}, {"source_file": ""}, {"output_dir": ""}):
                with self.subTest(change=change), self.assertRaises(ValueError):
                    service.validate_bigseller_benchmark_payload({**payload, **change})

    def test_run_separates_failures_preserves_source_and_exports_all_source_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            original = source.read_bytes()
            job = service.validate_bigseller_benchmark_payload(dict(
                username="user", password="secret", source_file=str(source), output_dir=directory, metric="sales"))
            matched = {"sku": "a", "shop_name": "店铺A", "metric_value": 0, "shop_count": 1,
                       "item_id": "00123456789012345678", "status": "matched", "message": ""}
            no_match = {"sku": "b", "shop_name": "", "metric_value": None, "shop_count": 0,
                        "item_id": "", "status": "not_found", "message": "无在售商品"}
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=Mock()), \
                 patch.object(service, "login_for_query", return_value=Mock()) as login, \
                 patch.object(service, "query_benchmark_sku", side_effect=[matched, no_match, RuntimeError("password=secret")]):
                result = service.run_bigseller_sku_benchmark(job, Mock())
            self.assertEqual([result[k] for k in ("total_rows", "sku_count", "matched_count", "not_found_count", "failed_count")], [4, 3, 1, 1, 1])
            self.assertFalse(result["is_complete"])
            self.assertEqual(result["confirmed_count"], 2)
            self.assertEqual(result["failed_skus"], ["c"])
            self.assertEqual(result["retry_rounds_used"], 3)
            self.assertNotIn("secret", json.dumps(result))
            self.assertEqual(source.read_bytes(), original)
            login.return_value.close.assert_called_once()
            book = load_workbook(result["output_file"])
            try:
                self.assertEqual(book["滞销"]["B3"].value, "旧店铺")
                self.assertEqual(book["滞销"]["C3"].value, "=1+2")
                detail = book["BigSeller销量对标明细"]
                self.assertEqual(detail.max_row, 4)
                self.assertEqual(detail["C2"].value, 0)
                self.assertEqual(detail["E2"].value, "00123456789012345678")
                self.assertEqual(detail["F4"].value, "查询失败")
            finally:
                book.close()

    def test_auth_refresh_once_and_permission_error_marks_unexecuted(self):
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            job = service.validate_bigseller_benchmark_payload(dict(
                username="user", password="secret", source_file=str(source), output_dir=directory))
            first, second = Mock(), Mock()
            with patch.object(service, "query_process_config", return_value={}), \
                 patch.object(service, "load_bigseller_request_util", return_value=Mock()), \
                 patch.object(service, "login_for_query", side_effect=[first, second]) as login, \
                 patch.object(service, "query_benchmark_sku", side_effect=[service.BigSellerAuthenticationError("过期"), service.BigSellerPermissionError("无权限")]) as query:
                result = service.run_bigseller_sku_benchmark(job, Mock())
            self.assertEqual(login.call_count, 2)
            self.assertEqual(query.call_count, 2)
            self.assertEqual(result["failed_count"], 3)
            self.assertEqual(result["not_found_count"], 0)
            self.assertFalse(result["is_complete"])
            self.assertEqual(result["retry_rounds_used"], 0)
            self.assertIn("未执行", result["rows"][1]["message"])
            first.close.assert_called_once()
            second.close.assert_called_once()

    def test_bridge_uses_bound_account_and_exposes_safe_context(self):
        bridge = AppBridge()
        bridge.config_store = Mock()
        bridge.config_store.load.return_value = AppSettings(
            accounts=[BoundAccount(id="bs", vendor="bigseller", name="BS", username="bound", password="private")],
            active_account_ids={"bigseller": "bs"}, captcha_username="captcha", captcha_password="private-captcha")
        bridge.tasks = Mock()
        job = SimpleNamespace(output_dir=Path("output"), source_file=Path("source.xlsx"), sheet_name="Sheet1", metric="sales")
        with patch("backend.app_bridge.validate_bigseller_benchmark_payload", return_value=job) as validate, \
             patch("backend.app_bridge.run_bigseller_sku_benchmark", return_value={"failed_count": 1}) as run:
            bridge.start_bigseller_sku_benchmark({"username": "override", "password": "override", "metric": "sales"})
            self.assertEqual(validate.call_args.args[0]["username"], "bound")
            self.assertEqual(validate.call_args.args[0]["captcha_password"], "private-captcha")
            context = bridge.tasks.start.call_args.kwargs
            self.assertEqual(context["tool"], "bigseller_sku_benchmark")
            self.assertNotIn("private", json.dumps(context))
            runner = bridge.tasks.start.call_args.args[1]
            self.assertEqual(runner(Mock()), {"failed_count": 1})
            run.assert_called_once()

    def test_bridge_blocks_unbound_account_and_inspects_real_sheet(self):
        bridge = AppBridge()
        bridge.config_store = Mock()
        bridge.config_store.load.return_value = AppSettings()
        bridge.tasks = Mock()
        self.assertTrue(bridge.start_bigseller_sku_benchmark({})["requires_account"])
        bridge.tasks.start.assert_not_called()
        with tempfile.TemporaryDirectory() as directory:
            source = self.source(directory)
            result = bridge.inspect_bigseller_benchmark_file(str(source))
            self.assertTrue(result["ok"])
            self.assertEqual(result["sheets"], ["滞销"])
            self.assertEqual(result["sheet_name"], "滞销")


class BenchmarkRecoveryTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.directory = Path(directory.name)
        source = BenchmarkTests().source(self.directory, ("a", "b"))
        self.job = service.validate_bigseller_benchmark_payload(dict(
            username="user", password="private", source_file=str(source), output_dir=str(self.directory)))

    def matched(self, sku):
        return {"sku": sku, "status": "matched", "shop_name": "店铺A", "shop_count": 1,
                "metric_value": 0, "item_id": "123"}

    def saved(self, path):
        data = service.load_workbook_input(self.job.source_file, self.job.sheet_name)
        return service.read_checkpoint(path, service._checkpoint_identity(self.job, data))

    @contextmanager
    def transport(self, query, wait=None):
        with ExitStack() as stack:
            stack.enter_context(patch.object(service, "query_process_config", return_value={}))
            stack.enter_context(patch.object(service, "load_bigseller_request_util", return_value=Mock()))
            login = stack.enter_context(patch.object(service, "login_for_query", return_value=Mock()))
            request = stack.enter_context(patch.object(service, "query_benchmark_sku", side_effect=query))
            delay = stack.enter_context(patch.object(service, "wait_before_retry", side_effect=wait))
            yield request, delay, login

    def test_deferred_retry_preserves_confirmed_skus_and_commits_before_wait(self):
        calls = []
        persisted = []

        def query(_session, _util, _config, sku, *_, **_kwargs):
            calls.append(sku)
            if sku == "b" and calls.count("b") < 3:
                raise service.BigSellerQueryError("分页未完成")
            return self.matched(sku)

        def wait(seconds, *_):
            if seconds > 0:
                checkpoint = next(self.directory.rglob("*.sqlite3"))
                persisted.append(self.saved(checkpoint))

        with self.transport(query, wait) as (_, delays, _):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertEqual(calls, ["a", "b", "b", "b"])
        self.assertEqual([c.args[0] for c in delays.call_args_list if c.args[0] > 0], [30, 60])
        self.assertEqual([rows["a"]["status"] for rows in persisted], ["matched", "matched"])
        self.assertEqual([rows["b"]["attempts"] for rows in persisted], [1, 2])
        self.assertTrue(result["is_complete"])
        self.assertEqual(result["confirmed_count"], 2)
        self.assertEqual(result["recovered_count"], 1)
        self.assertEqual(result["retry_rounds_used"], 2)

    def test_persistent_failure_is_never_reported_as_complete(self):
        def query(_session, _util, _config, sku, *_, **_kwargs):
            if sku == "a":
                return {"sku": sku, "status": "not_found", "shop_name": "", "shop_count": 0,
                        "metric_value": None, "item_id": ""}
            raise service.BigSellerQueryError("分页提前结束")

        logs = []
        with self.transport(query) as (requests, delays, _):
            result = service.run_bigseller_sku_benchmark(self.job, logs.append)
        self.assertEqual([c.args[3] for c in requests.call_args_list], ["a", "b", "b", "b", "b"])
        self.assertEqual([c.args[0] for c in delays.call_args_list if c.args[0] > 0], [30, 60, 120])
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["not_found_count"], 1)
        self.assertEqual(result["failed_skus"], ["b"])
        self.assertEqual(result["confirmed_count"], 1)
        self.assertFalse(any("[BS对标进度 2/2]" in line for line in logs))
        self.assertIn("结果不完整", Path(result["output_file"]).stem)
        self.assertEqual(self.saved(result["checkpoint_file"])["b"]["attempts"], 4)

    def test_cooldown_waits_before_any_next_sku_and_persists_server_deadline(self):
        events = []

        def query(_session, _util, _config, sku, *_, **_kwargs):
            events.append(sku)
            if len(events) == 1:
                raise service.BenchmarkCooldownError("限流", 120)
            return self.matched(sku)

        def wait(seconds, *_):
            if seconds > 0:
                rows = self.saved(next(self.directory.rglob("*.sqlite3")))
                self.assertTrue(all(row["retry_not_before"] == 2000000120 for row in rows.values()))
                events.append(seconds)

        with self.transport(query, wait), patch.object(service.time, "time", return_value=2000000000):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertEqual(events, ["a", 120, "a", "b"])
        self.assertTrue(result["is_complete"])

    def test_persistent_cooldown_does_not_query_next_sku_and_resume_honors_deadline(self):
        with self.transport(service.BenchmarkCooldownError("限流", 120)) as (query, _, _), \
             patch.object(service.time, "time", return_value=2000000000):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertEqual([c.args[3] for c in query.call_args_list], ["a"] * 4)
        self.assertEqual(result["confirmed_count"], 0)
        saved = self.saved(result["checkpoint_file"])
        self.assertEqual(saved["b"]["attempts"], 0)
        self.assertEqual(saved["b"]["retry_not_before"], 2000000120)
        events = []

        def retry(_session, _util, _config, sku, *_, **_kwargs):
            events.append(sku)
            return self.matched(sku)

        def wait(seconds, *_):
            if seconds > 0:
                events.append(seconds)

        with self.transport(retry, wait), patch.object(service.time, "time", return_value=2000000001):
            recovered = service.run_bigseller_sku_benchmark(
                replace(self.job, resume_checkpoint=Path(result["checkpoint_file"])), Mock())
        self.assertEqual(events, [119, "a", "b"])
        self.assertTrue(recovered["is_complete"])

    def test_interruption_keeps_prior_commits_and_resume_queries_only_pending(self):
        with self.transport([self.matched("a"), KeyboardInterrupt()]), self.assertRaises(KeyboardInterrupt):
            service.run_bigseller_sku_benchmark(self.job, Mock())
        previous = next(self.directory.rglob("*.sqlite3"))
        original = previous.read_bytes()
        self.assertEqual(self.saved(previous)["a"]["status"], "matched")
        with self.transport([self.matched("b")]) as (query, _, _):
            result = service.run_bigseller_sku_benchmark(replace(self.job, resume_checkpoint=previous), Mock())
        self.assertEqual([c.args[3] for c in query.call_args_list], ["b"])
        self.assertEqual(result["resumed_count"], 1)
        self.assertTrue(result["is_complete"])
        self.assertNotEqual(result["checkpoint_file"], str(previous))
        self.assertEqual(previous.read_bytes(), original)

    def test_export_failure_preserves_results_and_resume_reexports_without_requests(self):
        with self.transport([self.matched("a"), self.matched("b")]), \
             patch.object(service, "export_benchmark_workbook", side_effect=PermissionError("private")):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["failed_count"], 0)
        self.assertEqual(result["confirmed_count"], 2)
        self.assertEqual(result["output_file"], "")
        self.assertNotIn("private", json.dumps(result))
        with self.transport(AssertionError("must not query")) as (query, _, login):
            recovered = service.run_bigseller_sku_benchmark(
                replace(self.job, resume_checkpoint=Path(result["checkpoint_file"])), Mock())
        query.assert_not_called()
        login.assert_not_called()
        self.assertTrue(recovered["is_complete"])
        self.assertTrue(Path(recovered["output_file"]).is_file())

    def test_rejected_checkpoint_row_stops_queries_exports_commits_and_resumes_pending(self):
        source = BenchmarkTests().source(self.directory, ("a", "b", "c"))
        self.job = replace(self.job, source_file=source)
        create = service.create_checkpoint

        def reject_second_confirmed_row(*args):
            checkpoint = create(*args)
            # A real SQLite transaction rejection, including sensitive raw
            # diagnostic text that must never reach logs or the task result.
            checkpoint._connection.execute("""
                CREATE TRIGGER reject_result BEFORE UPDATE ON results
                WHEN NEW.sku = 'b' AND json_extract(NEW.row_json, '$.status') = 'matched'
                BEGIN SELECT RAISE(ABORT, 'private storage diagnostic'); END
            """)
            return checkpoint

        logs = []
        with self.transport([self.matched("a"), self.matched("b"), self.matched("c")]) as (query, _, _), \
             patch.object(service, "create_checkpoint", side_effect=reject_second_confirmed_row):
            result = service.run_bigseller_sku_benchmark(self.job, logs.append)
        self.assertEqual([call.args[3] for call in query.call_args_list], ["a", "b"])
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["confirmed_count"], 1)
        self.assertEqual(result["failed_skus"], ["b", "c"])
        self.assertEqual(result["retry_rounds_used"], 0)
        self.assertTrue(result["checkpoint_error"])
        self.assertIn(result["checkpoint_error"], result["completion_message"])
        saved = self.saved(result["checkpoint_file"])
        self.assertEqual(saved["a"]["status"], "matched")
        self.assertEqual(saved["b"]["status"], "failed")
        self.assertEqual(saved["b"]["attempts"], 0)
        self.assertTrue(Path(result["output_file"]).is_file())
        self.assertFalse(any("[BS对标进度 2/3]" in line for line in logs))
        self.assertNotIn("private", json.dumps(result) + "\n".join(logs))
        with self.transport([self.matched("b"), self.matched("c")]) as (query, _, _):
            recovered = service.run_bigseller_sku_benchmark(
                replace(self.job, resume_checkpoint=Path(result["checkpoint_file"])), Mock())
        self.assertEqual([call.args[3] for call in query.call_args_list], ["b", "c"])
        self.assertTrue(recovered["is_complete"])
        self.assertEqual(recovered["resumed_count"], 1)

    def test_checkpoint_and_export_failure_return_recovery_path_and_both_reasons(self):
        save = Checkpoint.save_result

        def reject_second(checkpoint, row):
            if row["sku"] == "b":
                raise service.CheckpointError("对标进度保存失败：磁盘空间不足（SQLITE_FULL）")
            return save(checkpoint, row)

        with self.transport([self.matched("a"), self.matched("b")]), \
             patch.object(Checkpoint, "save_result", reject_second), \
             patch.object(service, "export_benchmark_workbook", side_effect=PermissionError("private")):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["confirmed_count"], 1)
        self.assertEqual(result["output_file"], "")
        self.assertEqual(self.saved(result["checkpoint_file"])["a"]["status"], "matched")
        self.assertIn("SQLITE_FULL", result["completion_message"])
        self.assertIn("结果导出失败", result["completion_message"])
        self.assertNotIn("private", json.dumps(result))

    def test_final_checkpoint_save_failure_never_marks_task_complete(self):
        save = Checkpoint.save_results
        saves = 0

        def reject_final(checkpoint, rows):
            nonlocal saves
            saves += 1
            if saves == 4:  # initial batch, two committed SKU rows, final batch
                raise service.CheckpointError("对标进度保存失败：进度文件被其他程序占用（SQLITE_BUSY）")
            return save(checkpoint, rows)

        with self.transport([self.matched("a"), self.matched("b")]), \
             patch.object(Checkpoint, "save_results", reject_final):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["confirmed_count"], 2)
        self.assertEqual(result["failed_count"], 0)
        self.assertIn("SQLITE_BUSY", result["completion_message"])
        self.assertTrue(all(row["status"] == "matched" for row in self.saved(result["checkpoint_file"]).values()))
        self.assertTrue(Path(result["output_file"]).is_file())

    def test_initial_checkpoint_save_failure_keeps_original_resume_path_without_login(self):
        with self.transport([self.matched("a"), KeyboardInterrupt()]), self.assertRaises(KeyboardInterrupt):
            service.run_bigseller_sku_benchmark(self.job, Mock())
        previous = next(self.directory.rglob("*.sqlite3"))
        original = previous.read_bytes()
        with self.transport([]) as (query, _, login), \
             patch.object(Checkpoint, "save_results", side_effect=service.CheckpointError("对标进度保存失败：磁盘空间不足")):
            result = service.run_bigseller_sku_benchmark(replace(self.job, resume_checkpoint=previous), Mock())
        query.assert_not_called()
        login.assert_not_called()
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["checkpoint_file"], str(previous))
        self.assertEqual(result["confirmed_count"], 1)
        self.assertEqual(previous.read_bytes(), original)
        self.assertTrue(Path(result["output_file"]).is_file())

    def test_checkpoint_creation_failure_has_no_false_recovery_claim_or_login(self):
        with self.transport([]) as (query, _, login), \
             patch.object(service, "create_checkpoint", side_effect=service.CheckpointError("对标进度创建失败：目录不可写")):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        query.assert_not_called()
        login.assert_not_called()
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["confirmed_count"], 0)
        self.assertEqual(result["checkpoint_file"], "")
        self.assertIn("尚未建立可恢复进度", result["completion_message"])
        self.assertNotIn("进度已保存", result["completion_message"])

    def test_cooldown_save_failure_still_stops_without_queries_or_false_deadline(self):
        save = Checkpoint.save_results

        def reject_cooldown(checkpoint, rows):
            rows = list(rows)
            if any(row.get("retry_not_before", 0) for row in rows):
                raise service.CheckpointError("对标进度保存失败：目录不可写")
            return save(checkpoint, rows)

        def query(_session, _util, _config, sku, *_, on_cooldown):
            on_cooldown(2000000030)
            return self.matched(sku)

        with self.transport(query) as (request, delays, _), \
             patch.object(Checkpoint, "save_results", reject_cooldown), \
             self.assertRaisesRegex(service.BenchmarkCheckpointError, "无法保存服务冷却进度"):
            service.run_bigseller_sku_benchmark(self.job, Mock())
        request.assert_called_once()
        self.assertFalse(any(call.args[0] > 0 for call in delays.call_args_list))
        saved = self.saved(next(self.directory.rglob("*.sqlite3")))
        self.assertTrue(all(row["retry_not_before"] == 0 for row in saved.values()))
        self.assertTrue(all(row["status"] == "failed" for row in saved.values()))

    def test_resume_checks_account_metric_and_source_before_login(self):
        with self.transport([self.matched("a"), self.matched("b")]):
            result = service.run_bigseller_sku_benchmark(self.job, Mock())
        payload = dict(username="user", password="private", source_file=str(self.job.source_file),
                       output_dir=str(self.directory), resume_checkpoint=result["checkpoint_file"])
        for change in ({"username": "different"}, {"metric": "sales"}):
            with self.subTest(change=change), self.assertRaises(ValueError):
                service.validate_bigseller_benchmark_payload({**payload, **change})
        book = load_workbook(self.job.source_file)
        book.active["C3"] = "changed"
        book.save(self.job.source_file)
        book.close()
        with self.assertRaises(ValueError):
            service.validate_bigseller_benchmark_payload(payload)

    def test_login_failure_produces_incomplete_artifact_without_pointless_rounds(self):
        logs = []
        safe_reason = "BigSeller 账号或密码错误，请重新保存账号密码后继续补查"
        with self.transport([]) as (query, delays, login):
            login.side_effect = service.BigSellerAuthenticationError(safe_reason)
            result = service.run_bigseller_sku_benchmark(self.job, logs.append)
        query.assert_not_called()
        delays.assert_not_called()
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["failed_count"], 2)
        self.assertEqual(result["retry_rounds_used"], 0)
        self.assertTrue(Path(result["output_file"]).is_file())
        self.assertTrue(any(safe_reason in line for line in logs), logs)
        self.assertTrue(all(safe_reason in row["message"]
                            for row in self.saved(result["checkpoint_file"]).values()))
        self.assertNotIn(self.job.password, json.dumps(result, ensure_ascii=False) + "\n".join(logs))

    def test_interrupt_during_request_cooldown_already_has_durable_deadline(self):
        query = service.query_benchmark_sku

        def interrupt_wait(seconds, *_):
            if seconds > 0:
                raise SystemExit("application closed while waiting")

        for status in (429, 503):
            existing = set(self.directory.rglob("*.sqlite3"))
            response = BenchmarkTests().response(status=status, headers={"Retry-After": "30"})
            with self.subTest(status=status), self.transport(query, interrupt_wait), \
                 patch.object(service, "_post_listing", return_value=response), \
                 patch.object(service, "_COOLDOWN_UNTIL", 0.0), \
                 patch.object(service.time, "time", return_value=2000000000), \
                 self.assertRaises(SystemExit):
                service.run_bigseller_sku_benchmark(self.job, Mock())
            checkpoint = (set(self.directory.rglob("*.sqlite3")) - existing).pop()
            rows = self.saved(checkpoint)
            self.assertTrue(all(row["retry_not_before"] == 2000000030 for row in rows.values()))
            response.close.assert_called_once()

    def test_interrupt_waiting_another_tasks_cooldown_preserves_deadline(self):
        query = service.query_benchmark_sku

        def interrupt_wait(seconds, *_):
            if seconds > 0:
                raise SystemExit()

        with self.transport(query, interrupt_wait) as (_, _, login), \
             patch.object(service, "_COOLDOWN_UNTIL", 1030.0), \
             patch.object(service.time, "monotonic", return_value=1000.0), \
             patch.object(service.time, "time", return_value=2000000000), \
             self.assertRaises(SystemExit):
            service.run_bigseller_sku_benchmark(self.job, Mock())
        login.return_value.post.assert_not_called()
        rows = self.saved(next(self.directory.rglob("*.sqlite3")))
        self.assertTrue(all(row["retry_not_before"] == 2000000030 for row in rows.values()))


if __name__ == "__main__":
    unittest.main()
