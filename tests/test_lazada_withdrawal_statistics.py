from __future__ import annotations

import tempfile
import unittest
import logging
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from backend.services import lazada_withdrawal_statistics as service


class LazadaWithdrawalValidationTests(unittest.TestCase):
    def test_previous_month_crosses_year_boundary(self) -> None:
        self.assertEqual(
            service.previous_month_date_range(date(2026, 1, 8)),
            {"start_date": "2025-12-01", "end_date": "2025-12-31"},
        )

    def test_payload_defaults_to_previous_natural_month(self) -> None:
        job = service.validate_lazada_withdrawal_payload(
            {
                "username": "ziniao-user",
                "password": "secret",
                "country": "PH",
                "store_names": "店铺 A\n店铺 A\n店铺 B",
            },
            today=date(2026, 7, 15),
        )
        self.assertEqual(job.start_date, "2026-06-01")
        self.assertEqual(job.end_date, "2026-06-30")
        self.assertEqual(job.store_names, ("店铺 A", "店铺 B"))

    def test_payload_rejects_cross_month_range(self) -> None:
        with self.assertRaisesRegex(ValueError, "同一个自然月"):
            service.validate_lazada_withdrawal_payload(
                {
                    "username": "u",
                    "password": "p",
                    "country": "MY",
                    "start_date": "2026-05-31",
                    "end_date": "2026-06-01",
                    "store_names": "店铺",
                },
                today=date(2026, 7, 15),
            )

    def test_payload_rejects_future_end_date(self) -> None:
        with self.assertRaisesRegex(ValueError, "不得晚于今天"):
            service.validate_lazada_withdrawal_payload(
                {
                    "username": "u",
                    "password": "p",
                    "country": "TH",
                    "start_date": "2026-08-01",
                    "end_date": "2026-08-31",
                    "store_names": "店铺",
                },
                today=date(2026, 7, 15),
            )


class WithdrawalTableExtractionTests(unittest.TestCase):
    def test_english_source_columns_map_to_required_fields(self) -> None:
        record = service._map_withdrawal_row(
            ["Transaction Number", "Transaction Time", "Type", "Sub Type", "Amount", "Remarks"],
            [
                "502606281301398293700265",
                "29 Jun 2026 08:22:49",
                "Withdrawal",
                "Auto Withdrawal",
                "-7,850.47",
                "Paid | Bank Ref. TEST\nCredited in 2 days\n127.64USD",
            ],
        )
        self.assertEqual(record["transaction_time"], datetime(2026, 6, 29, 8, 22, 49))
        self.assertEqual(record["amount"], Decimal("-7850.47"))
        self.assertEqual(record["type"], "Withdrawal")
        self.assertEqual(record["remarks"].count("\n"), 2)

    def test_chinese_source_columns_are_supported(self) -> None:
        record = service._map_withdrawal_row(
            ["流水编号", "交易时间", "类型", "子类型", "金额", "备注"],
            ["1", "2026年06月29日 14:21:23", "提现", "自动提现", "(1,234.50)", "到账成功"],
        )
        self.assertEqual(record["transaction_time"], datetime(2026, 6, 29, 14, 21, 23))
        self.assertEqual(record["amount"], Decimal("-1234.50"))

    def test_missing_required_source_column_is_rejected(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "remarks"):
            service._withdrawal_column_map(["Transaction Time", "Type", "Amount"])

    def test_no_data_placeholder_is_not_a_business_row(self) -> None:
        with self.assertRaisesRegex(RuntimeError, "必要字段为空"):
            service._map_withdrawal_row(
                ["Transaction Time", "Type", "Amount", "Remarks"],
                ["没有数据"],
            )

    def test_collects_all_pages_and_deduplicates_by_transaction_number(self) -> None:
        headers = ["Transaction Number", "Transaction Time", "Type", "Sub Type", "Amount", "Remarks"]
        first = {
            "ready": True,
            "headers": headers,
            "rows": [{"cells": ["tx-1", "29 Jun 2026 08:22:49", "Withdrawal", "Auto", "-10.00", "one"]}],
            "nextFound": True,
            "nextDisabled": False,
            "signature": "1|tx-1|tx-1|1",
        }
        second = {
            "ready": True,
            "headers": headers,
            "rows": [
                {"cells": ["tx-1", "29 Jun 2026 08:22:49", "Withdrawal", "Auto", "-10.00", "one"]},
                {"cells": ["tx-2", "22 Jun 2026 08:22:49", "Withdrawal", "Auto", "-20.00", "two"]},
            ],
            "nextFound": True,
            "nextDisabled": True,
            "signature": "2|tx-1|tx-2|2",
        }
        with (
            patch.object(service, "_transaction_page_snapshot", side_effect=[first, second, second]),
            patch.object(service, "_click_transaction_next_page", return_value=True),
        ):
            records, pages = service._collect_withdrawal_transactions(object(), logging.getLogger("test"))
        self.assertEqual(pages, 2)
        self.assertEqual([record["transaction_number"] for record in records], ["tx-1", "tx-2"])


class WithdrawalWorkbookTests(unittest.TestCase):
    def test_workbook_has_exact_headers_and_typed_values(self) -> None:
        records = [
            {
                "transaction_number": "tx-1",
                "transaction_time": datetime(2026, 6, 29, 8, 22, 49),
                "type": "Withdrawal",
                "amount": Decimal("-7850.47"),
                "remarks": "Paid\nCredited in 2 days",
            }
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "store-withdrawal.xlsx"
            service._write_withdrawal_workbook(path, records, country_name="菲律宾", store_name="测试店铺")
            result = service.validate_withdrawal_workbook(
                path,
                records,
                country_name="菲律宾",
                store_name="测试店铺",
            )
            self.assertEqual(result["headers"], list(service.WITHDRAWAL_WORKBOOK_HEADERS))
            workbook = load_workbook(path)
            try:
                sheet = workbook["提现流水"]
                self.assertEqual(sheet.freeze_panes, "A2")
                self.assertEqual(sheet.auto_filter.ref, "A1:F2")
                self.assertIsInstance(sheet["A2"].value, datetime)
                self.assertEqual(sheet["C2"].value, -7850.47)
                self.assertTrue(sheet["D2"].alignment.wrap_text)
                self.assertGreater(sheet.row_dimensions[2].height, 22)
                self.assertEqual(sheet["E2"].value, "菲律宾")
                self.assertEqual(sheet["F2"].value, "测试店铺")
            finally:
                workbook.close()

    def test_no_data_workbook_contains_header_only(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "empty-withdrawal.xlsx"
            service._write_withdrawal_workbook(path, [], country_name="马来西亚", store_name="空店铺")
            result = service.validate_withdrawal_workbook(
                path,
                [],
                country_name="马来西亚",
                store_name="空店铺",
            )
            self.assertEqual(result["record_count"], 0)
            workbook = load_workbook(path)
            try:
                sheet = workbook["提现流水"]
                self.assertEqual(sheet.max_row, 1)
                self.assertEqual(sheet.auto_filter.ref, "A1:F1")
            finally:
                workbook.close()


class StatementWorkbookTests(unittest.TestCase):
    def _make_workbook(self, path: Path, statement: str, period: str) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Income Overview"
        sheet.append(["账单周期", "对账单编号 ", "交易日期", "费用名称"])
        sheet.append([period, statement, "05 Jul 2026", "货款"])
        workbook.save(path)

    def test_workbook_internal_identity_is_verified(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "statement.xlsx"
            self._make_workbook(path, "TH1K0MVRAO-2026-027", "29 Jun 2026 - 05 Jul 2026")
            actual = service.validate_statement_workbook(
                path,
                "TH1K0MVRAO-2026-027",
                "29 Jun 2026 - 05 Jul 2026",
            )
            self.assertEqual(actual["statement_number"], "TH1K0MVRAO-2026-027")

    def test_wrong_historical_export_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "statement.xlsx"
            self._make_workbook(path, "TH1K0MVRAO-2026-023", "01 Jun 2026 - 07 Jun 2026")
            with self.assertRaisesRegex(RuntimeError, "内部账单号不匹配"):
                service.validate_statement_workbook(
                    path,
                    "TH1K0MVRAO-2026-027",
                    "29 Jun 2026 - 05 Jul 2026",
                )

    def test_csv_statement_is_verified_when_platform_returns_csv(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "statement.csv"
            path.write_text(
                "账单周期;对账单编号 ;交易日期;费用名称\n"
                "15 Jun 2026 - 21 Jun 2026;TH1K0MVRAO-2026-025;21 Jun 2026;货款\n",
                encoding="utf-8-sig",
            )
            actual = service.validate_statement_file(
                path,
                "TH1K0MVRAO-2026-025",
                "15 Jun 2026 - 21 Jun 2026",
            )
            self.assertEqual(actual["sheet"], "CSV")


class ThailandExportBindingTests(unittest.TestCase):
    def test_download_clicks_same_new_export_id(self) -> None:
        class FakeWithdrawal:
            @staticmethod
            def _snapshot_download_files(_path):
                return {}

            @staticmethod
            def _open_thailand_statement_download_menu(_driver, _element):
                return True

            @staticmethod
            def _click_thailand_order_details_excel(_driver):
                return True

            @staticmethod
            def _wait_for_new_download_file(_path, _before, timeout):
                self.assertEqual(timeout, 120)
                return str(source)

            @staticmethod
            def _format_bill_period_part(start, end):
                return f"{start.strftime('%Y%m%d')}_{end.strftime('%Y%m%d')}"

            @staticmethod
            def _close_thailand_export_modal(_driver):
                return True

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            source = root / "platform.xlsx"
            source.write_bytes(b"xlsx")
            clicked: list[str] = []
            row = {
                "element": object(),
                "statement_number": "TH1K0MVRAO-2026-027",
                "period_range": (datetime(2026, 6, 29), datetime(2026, 7, 5)),
            }
            with (
                patch.object(
                    service,
                    "_export_rows",
                    return_value=[
                        {
                            "exportId": "18708755",
                            "status": "Finished",
                            "hasDownloadLink": True,
                        },
                        {
                            "exportId": "18708633",
                            "status": "Finished",
                            "hasDownloadLink": True,
                        },
                    ],
                ),
                patch.object(
                    service,
                    "_click_export_link_by_id",
                    side_effect=lambda _driver, export_id: clicked.append(export_id) or True,
                ),
            ):
                output, export_id = service._download_statement_exact(
                    object(),
                    row,
                    str(root),
                    root / "output",
                    FakeWithdrawal(),
                )
            self.assertEqual(export_id, "18708755")
            self.assertEqual(clicked, ["18708755"])
            self.assertEqual(
                Path(output).name,
                "TH1K0MVRAO-2026-027_20260629_20260705_Order_Details.xlsx",
            )


if __name__ == "__main__":
    unittest.main()
