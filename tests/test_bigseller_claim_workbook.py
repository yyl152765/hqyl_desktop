from __future__ import annotations

# Excel dates are deliberately naive wall-clock values.
# ruff: noqa: DTZ001
import hashlib
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image as PillowImage

from backend.services import bigseller_claim_workbook as service


class BigSellerClaimWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "新品认领表.xlsx"
        self.output = self.root / "output"
        book = Workbook()
        sheet = book.active
        sheet.title = "新sku上架"
        sheet.append(["SKU", "说明", "计算", "BS创建时间", "上架店铺", "上架时间"])
        sheet.append(["A-01", "商品一", "=1+1", "人工同名数据", "原店铺", datetime(2026, 1, 2)])
        sheet.append(["A-01", "重复 SKU"])
        sheet.append(["B-02\nC-03", "多个 SKU"])
        sheet.append(["1.13马来普杂整柜591件", "物流标题"])
        sheet.append(["以下柜子新品按小组产品切开", "分组说明"])
        sheet.append([123, "前导零 SKU"])
        sheet["A7"].number_format = "00000"
        sheet.append(["红色杯子-01", "中文 SKU"])
        sheet.append([None, "没有 SKU 的行"])
        sheet["D9"] = ""
        sheet.merge_cells("B9:C9")
        sheet["C2"].font = Font(name="Arial", bold=True, color="FF0000")
        sheet["E2"].fill = PatternFill("solid", fgColor="ABCDEF")
        sheet.column_dimensions["B"].width = 33
        sheet.row_dimensions[2].height = 42
        sheet.freeze_panes = "B2"
        validation = DataValidation(type="list", formula1='"甲,乙"')
        validation.add("B2:B8")
        sheet.add_data_validation(validation)
        other = book.create_sheet("原有汇总")
        other["A1"] = "='新sku上架'!C2"
        other["B1"] = "保留"
        other.sheet_state = "hidden"
        book.save(self.source)
        book.close()

    def book(self, path):
        book = load_workbook(path)
        self.addCleanup(book.close)
        return book

    def result(self, sku="A-01", **updates):
        data = {"sku": sku, "shop_name": "最早店铺", "item_id": "0012345678901234567",
                "created_time": "2024-02-01 08:09:10", "listed_time": "2024-02-02 09:10:11",
                "status": "matched", "message": ""}
        data.update(updates)
        return data

    def results(self, data):
        return [self.result(sku) for sku in data.skus]

    def owned_columns(self, sheet):
        return {key: next(cell.column for cell in sheet[1]
                          if cell.value == title and cell.comment and cell.comment.author == service.GENERATED_AUTHOR)
                for key, title in service.OUTPUT_HEADERS.items()}

    def test_parse_deduplicates_preserves_order_and_splits_multiline(self):
        data = service.load_claim_input(self.source)
        self.assertEqual(data.skus, ("A-01", "B-02", "C-03", "00123", "红色杯子-01"))
        self.assertEqual(data.row_numbers["A-01"], (2, 3))
        self.assertEqual(data.row_skus[4], ("B-02", "C-03"))
        self.assertEqual(data.total_rows, 5)
        self.assertEqual(data.sku_occurrence_count, 6)
        self.assertEqual([row["kind"] for row in data.skipped_rows], ["batch", "invalid"])
        self.assertEqual(data.fingerprint, hashlib.sha256(self.source.read_bytes()).hexdigest())
        self.assertTrue(any("多个换行 SKU" in message for message in data.warnings))

    def test_inspection_lists_sheets_and_reports_non_sku_sheet(self):
        result = service.inspect_claim_workbook(self.source)
        self.assertEqual(result["sheets"], ["新sku上架", "原有汇总"])
        self.assertEqual(result["sheet_name"], "新sku上架")
        self.assertEqual(result["sheet_details"][0]["sku_count"], 5)
        self.assertEqual(result["sheet_details"][0]["row_count"], 5)
        self.assertEqual(result["sheet_details"][0]["skipped_count"], 2)
        self.assertIn("未找到", result["sheet_details"][1]["error"])

    def test_parent_header_has_priority_and_whitespace_skus_are_not_removed(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["SKU", " Parent SKU "])
        sheet.append(["CHILD", "ABC 中文 1"])
        book.save(self.source)
        book.close()
        data = service.load_claim_input(self.source)
        self.assertEqual(data.skus, ("ABC 中文 1",))
        self.assertEqual(data.sku_column, 2)
        self.assertTrue(any("含空格" in message for message in data.warnings))

    def test_all_supported_headers(self):
        for header in ("主 SKU", "Parent SKU", "sku", "库存SKU", "库存SKU编号"):
            with self.subTest(header=header):
                book = Workbook()
                book.active.append(["标题"])
                book.active.append([header])
                book.active.append(["SKU-1"])
                book.save(self.source)
                book.close()
                data = service.load_claim_input(self.source)
                self.assertEqual((data.header_row, data.skus), (2, ("SKU-1",)))

    def test_invalid_formula_error_boolean_date_are_reported_not_queried(self):
        book = self.book(self.source)
        sheet = book["新sku上架"]
        for row, value in enumerate(("=1+1", "#REF!", True, datetime(2026, 1, 1)), 10):
            sheet.cell(row, 1, value)
        sheet["A14"] = "A-01\n以下新品按小组分组"
        sheet["A15"] = "MIEU31440527.14马来01整柜605件航利达"
        sheet["A16"] = "SKU"
        book.save(self.source)
        data = service.load_claim_input(self.source)
        self.assertEqual(data.row_numbers["A-01"], (2, 3, 14))
        skipped = {row["row_number"]: row for row in data.skipped_rows}
        for row in range(10, 15):
            self.assertEqual(skipped[row]["kind"], "invalid")
        self.assertEqual(skipped[15]["kind"], "batch")
        self.assertEqual(skipped[16]["kind"], "header")

    def test_incorrect_xml_dimension_does_not_truncate_rows(self):
        rebuilt = self.root / "dimension.xlsx"
        with ZipFile(self.source) as original, ZipFile(rebuilt, "w", ZIP_DEFLATED) as output:
            for name in original.namelist():
                data = original.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    import re
                    data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', data)
                output.writestr(name, data)
        data = service.load_claim_input(rebuilt)
        self.assertIn("红色杯子-01", data.skus)
        self.assertEqual(data.total_rows, 5)

    def test_limit_is_distinct_skus_and_excess_rejected(self):
        with patch.object(service, "MAX_SKUS", 4), self.assertRaisesRegex(ValueError, "超过 4"):
            service.load_claim_input(self.source, "新sku上架")
        with patch.object(service, "MAX_SKUS", 5):
            self.assertEqual(len(service.load_claim_input(self.source).skus), 5)

    def test_export_preserves_manual_cells_formulas_styles_other_sheets_and_images(self):
        image_path = self.root / "image.png"
        PillowImage.new("RGB", (8, 8), "red").save(image_path)
        original = self.book(self.source)
        original["新sku上架"].add_image(Image(str(image_path)), "B11")
        original["新sku上架"]["D9"] = ""
        original.save(self.source)
        original.close()
        service._preserve_empty_strings(self.source)
        self.assertEqual(self.book(self.source)["新sku上架"]["D9"].value, "")
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        data = service.load_claim_input(self.source)
        output = service.export_claim_workbook(data, self.results(data), self.output)
        result = self.book(output)
        sheet = result[data.sheet_name]
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())
        self.assertEqual(sheet["D2"].value, "人工同名数据")
        self.assertEqual(sheet["D9"].value, "")
        self.assertEqual(sheet["E2"].value, "原店铺")
        self.assertEqual(sheet["F2"].value, datetime(2026, 1, 2))
        self.assertEqual(sheet["C2"].value, "=1+1")
        self.assertEqual(sheet["C2"].font.name, "Arial")
        self.assertEqual(sheet["E2"].fill.fgColor.rgb, "00ABCDEF")
        self.assertEqual(sheet.column_dimensions["B"].width, 33)
        self.assertEqual(sheet.row_dimensions[2].height, 42)
        self.assertEqual(sheet.freeze_panes, "B2")
        self.assertIn("B9:C9", str(sheet.merged_cells))
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        self.assertEqual(len(sheet._images), 1)
        self.assertEqual(result["原有汇总"]["A1"].value, "='新sku上架'!C2")
        self.assertEqual(result["原有汇总"].sheet_state, "hidden")

    def test_single_dates_are_datetime_and_multiple_results_show_each_sku(self):
        data = service.load_claim_input(self.source)
        rows = self.results(data)
        rows[2] = self.result("C-03", shop_name="第二店铺", created_time="2025-03-04 12:13:14")
        output = service.export_claim_workbook(data, rows, self.output)
        book = self.book(output)
        sheet = book[data.sheet_name]
        columns = self.owned_columns(sheet)
        self.assertEqual(sheet.cell(2, columns["created_time"]).value, datetime(2024, 2, 1, 8, 9, 10))
        self.assertEqual(sheet.cell(2, columns["created_time"]).number_format, service.DATE_FORMAT)
        self.assertEqual(sheet.cell(3, columns["shop_name"]).value, "最早店铺")
        self.assertEqual(sheet.cell(4, columns["shop_name"]).value, "B-02：最早店铺\nC-03：第二店铺")
        self.assertEqual(sheet.cell(4, columns["created_time"]).value, "B-02：2024-02-01 08:09:10\nC-03：2025-03-04 12:13:14")
        self.assertEqual(sheet.cell(5, columns["status"]).value, "未查询")
        self.assertIn("物流标题", sheet.cell(5, columns["message"]).value)
        self.assertIn("人工核对", sheet.cell(6, columns["message"]).value)
        detail = book["BS新品认领查询明细"]
        self.assertEqual(detail.max_row, 9)
        self.assertEqual(detail["E2"].value, "0012345678901234567")
        self.assertEqual(detail["E2"].data_type, "s")
        self.assertEqual(detail["F2"].value, datetime(2024, 2, 1, 8, 9, 10))

    def test_missing_and_malformed_results_do_not_look_successful(self):
        data = service.load_claim_input(self.source)
        rows = [self.result("A-01", created_time="bad"),
                self.result("B-02", status="not_found"),
                self.result("C-03", status="unknown"),
                self.result("00123", listed_time="")]
        output = service.export_claim_workbook(data, rows, self.output)
        self.assertIn("结果不完整", output.name)
        sheet = self.book(output)[data.sheet_name]
        columns = self.owned_columns(sheet)
        self.assertEqual(sheet.cell(2, columns["status"]).value, "查询失败")
        self.assertIsNone(sheet.cell(2, columns["created_time"]).value)
        self.assertEqual(sheet.cell(2, columns["shop_name"]).value, "")
        self.assertEqual(sheet.cell(8, columns["status"]).value, "查询失败")
        self.assertEqual(sheet.cell(7, columns["status"]).value, "已匹配")
        self.assertIsNone(sheet.cell(7, columns["listed_time"]).value)
        self.assertIn("未提供", sheet.cell(7, columns["message"]).value)

    def test_remote_text_is_not_written_as_excel_formula(self):
        data = service.load_claim_input(self.source)
        rows = self.results(data)
        rows[0].update(shop_name="=HYPERLINK(\"https://example.invalid\")", message="=1+1")
        output = service.export_claim_workbook(data, rows, self.output)
        book = self.book(output)
        sheet = book[data.sheet_name]
        columns = self.owned_columns(sheet)
        self.assertEqual(sheet.cell(2, columns["shop_name"]).data_type, "s")
        self.assertEqual(sheet.cell(2, columns["message"]).data_type, "s")
        self.assertEqual(book["BS新品认领查询明细"]["D2"].data_type, "s")

    def test_repeat_export_updates_only_marked_columns_and_owned_detail(self):
        book = self.book(self.source)
        book.create_sheet("BS新品认领查询明细")["A1"] = "用户原有明细"
        book.save(self.source)
        data = service.load_claim_input(self.source)
        first = service.export_claim_workbook(data, self.results(data), self.output)
        repeated = service.load_claim_input(first, data.sheet_name)
        rows = [self.result(sku, shop_name="更新店铺") for sku in repeated.skus]
        second = service.export_claim_workbook(repeated, rows, self.output)
        result = self.book(second)
        sheet = result[data.sheet_name]
        columns = self.owned_columns(sheet)
        self.assertEqual(sheet.max_column, 11)
        self.assertEqual(sheet.cell(2, columns["shop_name"]).value, "更新店铺")
        self.assertEqual(sheet["D2"].value, "人工同名数据")
        self.assertEqual(result["BS新品认领查询明细"]["A1"].value, "用户原有明细")
        self.assertEqual(result.sheetnames.count("BS新品认领查询明细_2"), 1)
        self.assertEqual(len(result.sheetnames), 4)
        inspection = service.inspect_claim_workbook(second)
        self.assertNotIn("BS新品认领查询明细_2", inspection["sheets"])

    def test_user_same_named_column_without_ownership_is_never_overwritten(self):
        book = self.book(self.source)
        sheet = book["新sku上架"]
        for column, header in enumerate(service.OUTPUT_HEADERS.values(), 7):
            sheet.cell(1, column, header)
            sheet.cell(2, column, "用户维护")
        book.save(self.source)
        data = service.load_claim_input(self.source)
        output = service.export_claim_workbook(data, self.results(data), self.output)
        sheet = self.book(output)[data.sheet_name]
        self.assertEqual([sheet.cell(2, column).value for column in range(7, 12)], ["用户维护"] * 5)
        self.assertEqual(sheet.max_column, 16)

    def test_source_changes_after_import_prevent_export(self):
        data = service.load_claim_input(self.source)
        book = self.book(self.source)
        book[data.sheet_name]["A2"] = "CHANGED"
        book.save(self.source)
        with self.assertRaisesRegex(ValueError, "已变化"):
            service.export_claim_workbook(data, self.results(data), self.output)
        self.assertFalse(self.output.exists())

    def test_failed_save_cleans_temporary_file_and_does_not_create_result(self):
        data = service.load_claim_input(self.source)

        def broken_save(_book, filename):
            Path(filename).write_bytes(b"partial")
            raise OSError("disk full")

        with patch("openpyxl.workbook.workbook.Workbook.save", broken_save), self.assertRaisesRegex(OSError, "disk full"):
            service.export_claim_workbook(data, self.results(data), self.output)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_unique_exports_never_overwrite_source_or_previous_output(self):
        data = service.load_claim_input(self.source)
        first = service.export_claim_workbook(data, self.results(data), self.root)
        second = service.export_claim_workbook(data, self.results(data), self.root)
        self.assertNotEqual(first, second)
        self.assertNotEqual(first, self.source)
        self.assertTrue(first.is_file())
        self.assertTrue(second.is_file())
        self.assertEqual(list(self.root.glob(".bs-claim-*")), [])

    def test_no_valid_skus_and_unsupported_format_are_rejected(self):
        with self.assertRaisesRegex(ValueError, "仅支持 .xlsx"):
            service.load_claim_input(self.root / "source.xls")
        book = Workbook()
        book.active.append(["SKU"])
        book.active.append(["1.13整柜591件"])
        book.save(self.source)
        book.close()
        with self.assertRaisesRegex(ValueError, "没有可查询"):
            service.load_claim_input(self.source, "Sheet")
        self.assertEqual(service.inspect_claim_workbook(self.source)["sheet_name"], "")


if __name__ == "__main__":
    unittest.main()
