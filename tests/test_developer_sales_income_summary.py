from __future__ import annotations

import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from backend.core.mabang_client import MabangApiError
from backend.services.developer_sales_income_summary import (
    DeveloperIncomeSummary,
    IncomeDetailRow,
    PersonOption,
    UNASSIGNED_SALESPERSON,
    build_report_form_data,
    build_summary_workbook,
    parse_developer_names,
    parse_income_rows,
    parse_person_options,
    parse_table_foot_order_income,
    resolve_developer_options,
    summarize_all_people,
    validate_developer_sales_income_payload,
)


class DeveloperSalesIncomeSummaryTests(unittest.TestCase):
    def test_person_options_preserve_duplicate_suffix_and_independent_ids(self) -> None:
        html = """
        <label><input type="checkbox" name="developid[]" value="1117090">黄茵茵-开发&gt;1</label>
        <label><input type="checkbox" name="developid[]" value="1097736">A1018-黄茵茵-开发</label>
        <label><input type="checkbox" name="developid[]" value="0">无开发员</label>
        """

        options = parse_person_options(html, "developid[]")
        selected, missing = resolve_developer_options(["黄茵茵"], options)

        self.assertEqual(options[-1], PersonOption("0", "未分配开发员"))
        self.assertEqual([option.id for option in selected], ["1117090", "1097736"])
        self.assertEqual(
            [option.name for option in selected],
            ["黄茵茵-开发>1", "A1018-黄茵茵-开发"],
        )
        self.assertEqual(missing, ())

    def test_table_foot_uses_first_income_cell_as_order_income(self) -> None:
        html = (
            '<tr><th data-field="income">5,576.0500</th>'
            '<th data-field="income">92.1670</th>'
            '<th data-field="income">6,477.0550</th></tr>'
        )

        amount = parse_table_foot_order_income(html, context="黄茵茵-开发>1")

        self.assertEqual(amount, Decimal("5576.0500"))

    def test_missing_developer_names_are_returned_without_dropping_matches(self) -> None:
        selected, missing = resolve_developer_options(
            ["黄茵茵", "不存在的人"],
            [PersonOption("1", "黄茵茵-开发>1"), PersonOption("2", "A1018-黄茵茵-开发")],
        )

        self.assertEqual(len(selected), 2)
        self.assertEqual(missing, ("不存在的人",))

    def test_parse_developer_names_supports_one_name_per_line_and_deduplicates(self) -> None:
        self.assertEqual(
            parse_developer_names("黄茵茵\n\n 黄茵茵 \n刘琴铃"),
            ("黄茵茵", "刘琴铃"),
        )

    def test_validate_payload_uses_payment_date_range(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            job = validate_developer_sales_income_payload(
                {
                    "username": "user",
                    "password": "password",
                    "developer_names": ["黄茵茵"],
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-31",
                    "output_dir": temp_dir,
                }
            )

        self.assertEqual(job.developer_names, ("黄茵茵",))
        self.assertEqual(job.start_date, "2026-07-01")

    def test_validate_payload_allows_all_people_without_name_input(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            job = validate_developer_sales_income_payload(
                {
                    "username": "user",
                    "password": "password",
                    "developer_names": [],
                    "include_all": True,
                    "start_date": "2026-07-01",
                    "end_date": "2026-07-31",
                    "output_dir": temp_dir,
                }
            )

        self.assertTrue(job.include_all)

    def test_report_form_uses_us_dollars_payment_time_and_developer_id(self) -> None:
        payload = build_report_form_data("1117090", date(2026, 7, 1), date(2026, 7, 31))

        self.assertEqual(payload["timeKey"], "payTime")
        self.assertEqual(payload["moneyType"], "US")
        self.assertEqual(payload["developid[]"], "1117090")
        self.assertEqual(payload["timeStart"], "2026-07-01 00:00:00")
        self.assertEqual(payload["timeEnd"], "2026-07-31 23:59:59")

    def test_income_rows_skip_total_and_keep_unassigned_sales_income(self) -> None:
        content = (
            "SKU,开发员,销售员,收入-订单金额\n"
            "sku-1,李四-开发,张三,10.1234\n"
            "sku-2,李四-开发,--,2.5000\n"
            "合计,,,12.6234\n"
        ).encode("utf-8")

        rows = parse_income_rows(content, context="测试开发员")

        self.assertEqual(
            rows,
            [("张三", Decimal("10.1234")), (UNASSIGNED_SALESPERSON, Decimal("2.5000"))],
        )

    def test_all_people_summary_keeps_every_option_and_zero_amount_people(self) -> None:
        developer_options = [
            PersonOption("1", "黄茵茵-开发>1"),
            PersonOption("2", "A1018-黄茵茵-开发"),
            PersonOption("3", "零收入-开发"),
        ]
        salesperson_options = [PersonOption("10", "张三"), PersonOption("11", "零收入销售")]
        rows = [
            IncomeDetailRow("黄茵茵-开发>1", "张三", Decimal("10")),
            IncomeDetailRow("A1018-黄茵茵-开发", "张三", Decimal("20")),
        ]

        developers, salespeople = summarize_all_people(rows, developer_options, salesperson_options)

        self.assertEqual([item.developer.name for item in developers], [item.name for item in developer_options])
        self.assertEqual([item.amount for item in developers], [Decimal("10"), Decimal("20"), Decimal("0")])
        self.assertEqual(salespeople, {"张三": Decimal("30"), "零收入销售": Decimal("0")})

    def test_workbook_has_exactly_two_summary_sheets_and_separate_developers(self) -> None:
        summaries = [
            DeveloperIncomeSummary(PersonOption("1", "黄茵茵-开发>1"), Decimal("10.1234"), 2),
            DeveloperIncomeSummary(PersonOption("2", "A1018-黄茵茵-开发"), Decimal("2.5000"), 1),
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "summary.xlsx"
            build_summary_workbook(
                output_path,
                summaries,
                {"张三": Decimal("12.6234")},
            )
            workbook = load_workbook(output_path, data_only=True)

        self.assertEqual(workbook.sheetnames, ["开发员", "销售员"])
        self.assertEqual(workbook["开发员"]["A2"].value, "黄茵茵-开发>1")
        self.assertEqual(workbook["开发员"]["A3"].value, "A1018-黄茵茵-开发")
        self.assertAlmostEqual(workbook["开发员"]["B4"].value, 12.6234)
        self.assertAlmostEqual(workbook["销售员"]["B3"].value, 12.6234)
        workbook.close()

    def test_workbook_rejects_developer_salesperson_total_mismatch(self) -> None:
        summaries = [DeveloperIncomeSummary(PersonOption("1", "张三-开发"), Decimal("1"), 1)]
        with tempfile.TemporaryDirectory() as temp_dir:
            with self.assertRaisesRegex(MabangApiError, "合计不一致"):
                build_summary_workbook(
                    Path(temp_dir) / "summary.xlsx",
                    summaries,
                    {"李四": Decimal("2")},
                )

    def test_workbook_all_people_mode_keeps_independent_dimension_totals(self) -> None:
        summaries = [DeveloperIncomeSummary(PersonOption("1", "张三-开发"), Decimal("1"), 0)]
        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = Path(temp_dir) / "summary.xlsx"
            build_summary_workbook(
                output_path,
                summaries,
                {"李四": Decimal("2")},
                require_equal_totals=False,
            )
            workbook = load_workbook(output_path, data_only=True)

        self.assertEqual(workbook["开发员"]["B3"].value, 1)
        self.assertEqual(workbook["销售员"]["B3"].value, 2)
        workbook.close()


if __name__ == "__main__":
    unittest.main()
