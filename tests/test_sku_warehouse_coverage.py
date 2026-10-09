from __future__ import annotations

import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import MagicMock, Mock, patch

from openpyxl import load_workbook

from backend.services.sku_inventory_query import (
    DeveloperOption, MabangApiError, build_stock_query_payload,
    run_sku_inventory_query, validate_query_payload,
)
from backend.services.sku_warehouse_coverage import (
    coverage_record, export_coverage_records, iter_hot_stock_rows,
    parse_scope_warehouses, parse_warehouse_countries, southeast_asian_country,
    unclassified_warehouses,
)


DEVELOPER = DeveloperOption("101", "张三-开发")
EMPTY_STOCK_RESPONSE = {
    "success": True, "message": '<div class="alert"><h4>暂无内容！</h4></div>', "pageHtml": False,
}
CATALOGUE = '''
<label><input type="checkbox" name="developerIdIN[]" value="101">张三-开发</label>
<select name="defaultWarehouseId">
  <option value="">请选择</option>
  <option value="TH">泰国TH01仓</option><option value="TH2">泰国TH02仓</option>
  <option value="PH">PH8807-雅仓海外仓</option><option value="VN">寰球河内2仓-知意</option>
  <option value="TRANSIT">越南中转仓</option><option value="CN">广州仓</option>
  <option value="JP">日本海外仓</option><option value="SG">品牌仓</option>
  <option value="UNKNOWN">待核对品牌仓</option>
  <option value="1100808">Ola-Ola海外仓</option>
  <option value="1100602">T&amp;C 海外 2 号仓-T&amp;C海外仓</option>
</select>'''
COUNTRY_HTML = '''<table>
<tr><td><input class="checkWarehouse" data-id="SG"></td><td>品牌仓</td><td>自建本地仓</td><td>新加坡</td><td>15</td></tr>
<tr><td><input class="checkWarehouse" data-id="CN"></td><td>广州仓</td><td>自建本地仓</td><td>中国</td><td>20</td></tr>
</table>'''


def stock(sku="SKU-001", liveness="1", details=None):
    return {"id": sku, "stockSku": sku, "developerId": "101", "livenessType": liveness,
            "timeCreated": "2019-01-01", "nameCN": "商品", "stockWarehouseData": [] if details is None else details}


class SkuWarehouseCoverageTests(unittest.TestCase):
    def job(self, **overrides):
        return validate_query_payload({
            "username": "test-user", "password": "test-secret", "developer_ids": ["101"],
            "query_mode": "missing_warehouses", "output_dir": "unused", **overrides,
        })

    def test_all_history_and_all_statuses_with_separate_hot_filter(self):
        job = self.job(start_date="bad", end_date="2026-09-21")
        payload = dict(build_stock_query_payload(job, developer_id="101", page=3, liveness_type="2"))
        self.assertEqual(job.liveness_types, ("1", "2"))
        self.assertEqual((job.start_date, job.end_date), ("", ""))
        self.assertEqual(payload["status"], "")
        self.assertEqual(payload["timeCreatedStartTime"], "")
        self.assertEqual(payload["timeCreatedEndTime"], "")
        self.assertEqual(payload["developerIdM"], "101")
        self.assertEqual(payload["livenessType"], "2")
        self.assertEqual(payload["warehouseMold"], "all")
        for values in ([], ["3"], ["1", "bogus"]):
            with self.subTest(values=values), self.assertRaises(ValueError):
                self.job(liveness_types=values)

    def test_scope_uses_full_catalogue_country_metadata_and_excludes_transit(self):
        countries = parse_warehouse_countries(COUNTRY_HTML)
        warehouses = parse_scope_warehouses(CATALOGUE, countries)
        self.assertEqual({w["id"] for w in warehouses}, {"TH", "TH2", "PH", "VN", "SG"})
        self.assertEqual(unclassified_warehouses(CATALOGUE, countries), [{"id": "UNKNOWN", "name": "待核对品牌仓"}])
        self.assertEqual(southeast_asian_country("泰国中转仓", "泰国"), "")
        self.assertEqual(southeast_asian_country("泰国商品仓", "中国"), "")
        self.assertEqual(southeast_asian_country("ID8803-雅仓海外仓"), "印度尼西亚")
        self.assertEqual(southeast_asian_country("MY01 transfer warehouse"), "")
        self.assertEqual(southeast_asian_country("THINK仓"), "")
        self.assertEqual(southeast_asian_country("日本海外仓"), "")
        self.assertEqual(southeast_asian_country("品牌仓", "India"), "")
        for source in ("", "<p>权限不足</p>"):
            with self.assertRaises(MabangApiError):
                parse_warehouse_countries(source)
        with self.assertRaises(MabangApiError):
            parse_scope_warehouses("<select></select>", {})

    def test_association_not_stock_quantity_and_never_confuses_association_id(self):
        warehouses = parse_scope_warehouses(CATALOGUE, {})
        details = [
            {"warehouseId": "TH", "id": "association-1", "stockWarehouseAvailableInventory": 0},
            {"warehouseId": "PH", "id": "association-2", "stockWarehouseAvailableInventory": -2},
            {"warehouseId": "CN", "id": "TH2", "stockWarehouseAvailableInventory": 500},
        ]
        record = coverage_record(stock(details=details), DEVELOPER, warehouses)
        self.assertEqual(record["missing_warehouse_ids"], "TH2、VN")
        self.assertEqual(record["covered_warehouse_count"], 2)
        self.assertEqual(record["missing_warehouse_count"], 2)
        self.assertEqual(record["liveness_name"], "爆款")
        covered = coverage_record(stock(details=[{"warehouseId": w["id"]} for w in warehouses]), DEVELOPER, warehouses)
        self.assertEqual(covered["missing_warehouse_count"], 0)
        self.assertEqual(covered["missing_warehouse_names"], "")
        empty = coverage_record(stock(), DEVELOPER, warehouses)
        self.assertEqual(empty["missing_warehouse_count"], len(warehouses))
        for malformed in ({}, {**stock(), "stockWarehouseData": None}, stock(details=[{"id": "TH"}])):
            with self.assertRaises(MabangApiError):
                coverage_record(malformed, DEVELOPER, warehouses)

    def test_pagination_honors_total_pages_even_when_server_caps_page_size(self):
        responses = [
            {"success": True, "stockData": [stock("A")], "pageHtml": "1/2页"},
            {"success": True, "stockData": [stock("B")], "pageHtml": "2/2页"},
        ]
        with patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses) as post:
            rows = list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", None))
        self.assertEqual([row["stockSku"] for row in rows], ["A", "B"])
        self.assertEqual(post.call_count, 2)

    def test_verified_empty_first_page_is_valid_and_logs_developer_id(self):
        progress = Mock()
        with patch("backend.services.sku_warehouse_coverage.post_form_json", return_value=EMPTY_STOCK_RESPONSE) as post:
            rows = list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", progress))
        self.assertEqual(rows, [])
        post.assert_called_once()
        messages = "\n".join(call.args[0] for call in progress.call_args_list)
        for expected in (DEVELOPER.name, "ID 101", "爆款", "第 1 页", "暂无匹配 SKU"):
            self.assertIn(expected, messages)

    def test_populated_stock_data_takes_precedence_over_empty_message_template(self):
        response = {**EMPTY_STOCK_RESPONSE, "stockData": [stock()]}
        with patch("backend.services.sku_warehouse_coverage.post_form_json", return_value=response):
            rows = list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", None))
        self.assertEqual([row["stockSku"] for row in rows], ["SKU-001"])

    def test_unverified_missing_or_malformed_stock_data_is_not_treated_as_empty(self):
        variants = [
            {"stockData": None}, {"stockData": {}}, {"stockData": "[]"},
            {"hasData": False}, {"hasData": True}, {"success": False}, {"success": "true"},
            {"pageHtml": None}, {"pageHtml": "1/1页"}, {"pageHtml": "1/2页"},
            {"message": ""}, {"message": "暂无内容！请重新登录"}, {"message": "权限不足"},
        ]
        for variant in variants:
            response = {**EMPTY_STOCK_RESPONSE, **variant}
            with self.subTest(variant=variant), patch("backend.services.sku_warehouse_coverage.post_form_json", return_value=response):
                with self.assertRaises(MabangApiError):
                    list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", None))

    def test_declared_later_pages_cannot_be_silently_empty(self):
        final_responses = [
            EMPTY_STOCK_RESPONSE,
            {"success": True, "stockData": [], "pageHtml": "2/2页"},
            {"success": True, "stockData": []},
            {"success": True, "stockData": [], "pageHtml": False},
        ]
        for final_response in final_responses:
            responses = [
                {"success": True, "stockData": [stock("A")], "pageHtml": "1/2页"},
                final_response,
            ]
            with self.subTest(final_response=final_response), patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses):
                with self.assertRaises(MabangApiError):
                    list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", None))

    def test_pagination_remembers_total_when_intermediate_page_omits_it(self):
        responses = [
            {"success": True, "stockData": [stock("A")], "pageHtml": "1/3页"},
            {"success": True, "stockData": [stock("B")]},
            {"success": True, "stockData": [stock("C")], "pageHtml": "3/3页"},
        ]
        with patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses) as post:
            rows = list(iter_hot_stock_rows(Mock(), self.job(), DEVELOPER, "1", "https://test", None))
        self.assertEqual([row["stockSku"] for row in rows], ["A", "B", "C"])
        self.assertEqual(post.call_count, 3)

    def test_bad_filter_partial_and_repeated_pages_fail_without_exporting(self):
        cases = [
            [{"success": True, "stockData": [{**stock(), "developerId": "102"}]}],
            [{"success": True, "stockData": [stock(liveness="2")]}],
            [{"success": True, "stockData": [stock()], "pageHtml": "2/3页"}],
            [{"success": True, "stockData": [], "pageHtml": "1/2页"}],
            [{"success": False, "stockData": []}],
            [{"success": True, "stockData": [stock()]}, {"success": True, "stockData": [stock()]}],
        ]
        for responses in cases:
            with self.subTest(responses=responses), patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses):
                with self.assertRaises(MabangApiError):
                    list(iter_hot_stock_rows(Mock(), replace(self.job(), rows_per_page=1), DEVELOPER, "1", "https://test", None))

    def test_runner_queries_both_types_and_exports_all_rows_and_scope(self):
        client = MagicMock()
        client.__enter__.return_value = client
        responses = [
            {"success": True, "message": COUNTRY_HTML},
            {"success": True, "stockData": [stock("=FORMULA")], "pageHtml": "1/1页"},
            {"success": True, "stockData": [stock("B", "2")], "pageHtml": "1/1页"},
        ]
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch("backend.services.sku_warehouse_coverage.MabangClient", return_value=client),
            patch("backend.services.sku_warehouse_coverage.open_stock_page", return_value=("https://test/stock", CATALOGUE)),
            patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses) as post,
        ):
            result = run_sku_inventory_query(self.job(output_dir=tmp))
            self.assertEqual(result.query_mode, "missing_warehouses")
            self.assertEqual(result.sku_count, 2)
            self.assertEqual(result.warehouse_count, 5)
            self.assertEqual(result.missing_sku_count, 2)
            self.assertEqual([dict(c.args[2])["livenessType"] for c in post.call_args_list[1:]], ["1", "2"])
            book = load_workbook(result.output_file)
            try:
                self.assertEqual(book.sheetnames, ["SKU缺失仓库", "东南亚仓库范围", "查询口径"])
                self.assertEqual(book.worksheets[0].max_row, 3)
                self.assertEqual(book.worksheets[0]["B2"].value, "=FORMULA")
                self.assertEqual(book.worksheets[0]["B2"].data_type, "s")
                self.assertEqual(book.worksheets[1].max_row, 6)
                self.assertIn("待核对品牌仓", book.worksheets[2]["B8"].value)
            finally:
                book.close()

    def test_empty_developer_continues_to_both_types_of_distinct_id_with_same_name(self):
        client = MagicMock()
        client.__enter__.return_value = client
        catalogue = CATALOGUE + '<label><input type="checkbox" name="developerIdIN[]" value="102">张三-开发</label>'
        responses = [
            {"success": True, "message": COUNTRY_HTML},
            EMPTY_STOCK_RESPONSE, EMPTY_STOCK_RESPONSE,
            {**EMPTY_STOCK_RESPONSE, "stockData": [{**stock("B"), "developerId": "102"}]},
            {**EMPTY_STOCK_RESPONSE, "stockData": [{**stock("C", "2"), "developerId": "102"}]},
        ]
        progress = Mock()
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch("backend.services.sku_warehouse_coverage.MabangClient", return_value=client),
            patch("backend.services.sku_warehouse_coverage.open_stock_page", return_value=("https://test/stock", catalogue)),
            patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses) as post,
        ):
            result = run_sku_inventory_query(self.job(output_dir=tmp, developer_ids=["101", "102"]), progress)
            self.assertEqual(result.developer_count, 2)
            self.assertEqual(result.sku_count, 2)
            self.assertEqual({record["sku"] for record in result.records}, {"B", "C"})
            requests = [dict(call.args[2]) for call in post.call_args_list[1:]]
            self.assertEqual([(p["developerIdM"], p["livenessType"]) for p in requests],
                             [("101", "1"), ("101", "2"), ("102", "1"), ("102", "2")])
            self.assertTrue(result.output_file.is_file())
        messages = "\n".join(call.args[0] for call in progress.call_args_list)
        self.assertIn("张三-开发（ID 101）", messages)
        self.assertIn("张三-开发（ID 102）", messages)

    def test_runner_all_verified_empty_responses_export_headers_and_scope(self):
        client = MagicMock()
        client.__enter__.return_value = client
        responses = [{"success": True, "message": COUNTRY_HTML}, EMPTY_STOCK_RESPONSE, EMPTY_STOCK_RESPONSE]
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch("backend.services.sku_warehouse_coverage.MabangClient", return_value=client),
            patch("backend.services.sku_warehouse_coverage.open_stock_page", return_value=("https://test/stock", CATALOGUE)),
            patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses) as post,
        ):
            result = run_sku_inventory_query(self.job(output_dir=tmp))
            self.assertEqual(result.sku_count, 0)
            self.assertEqual(result.missing_sku_count, 0)
            self.assertEqual(result.warehouse_count, 5)
            self.assertEqual(post.call_count, 3)
            book = load_workbook(result.output_file)
            try:
                self.assertEqual(book.worksheets[0].max_row, 1)
                self.assertEqual(book.worksheets[1].max_row, 6)
            finally:
                book.close()

    def test_unexpected_response_aborts_without_export_and_redacts_diagnostics(self):
        client = MagicMock()
        client.__enter__.return_value = client
        responses = [
            {"success": True, "message": COUNTRY_HTML},
            {"success": True, "message": "<b>查询异常 test-user p&amp;ss p&ss</b>", "session": "private-session-value"},
        ]
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch("backend.services.sku_warehouse_coverage.MabangClient", return_value=client),
            patch("backend.services.sku_warehouse_coverage.open_stock_page", return_value=("https://test/stock", CATALOGUE)),
            patch("backend.services.sku_warehouse_coverage.post_form_json", side_effect=responses),
            patch("backend.services.sku_warehouse_coverage.export_coverage_records") as export,
        ):
            with self.assertRaises(MabangApiError) as error:
                run_sku_inventory_query(self.job(output_dir=tmp, password="p&ss"))
            export.assert_not_called()
            self.assertEqual(list(Path(tmp).iterdir()), [])
        message = str(error.exception)
        for expected in (DEVELOPER.name, "ID 101", "爆款", "第 1 页", "stockData", "缺失", "查询异常"):
            self.assertIn(expected, message)
        for secret in ("test-user", "p&ss", "p&amp;ss", "private-session-value", "<b>"):
            self.assertNotIn(secret, message)

    def test_empty_result_still_exports_headers_and_full_warehouse_scope(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = export_coverage_records([], [{"id": "1", "name": "泰国仓", "country": "泰国"}], Path(tmp))
            book = load_workbook(output)
            try:
                self.assertEqual(book.worksheets[0].max_row, 1)
                self.assertEqual(book.worksheets[1].max_row, 2)
            finally:
                book.close()


if __name__ == "__main__":
    unittest.main()
