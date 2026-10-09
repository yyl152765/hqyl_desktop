from __future__ import annotations

import hashlib
import tempfile
import unittest
from copy import copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from backend.services import bigseller_benchmark_workbook as service


class BigSellerBenchmarkWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "9月滞销sku.xlsx"
        self.output_dir = self.root / "导出"
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = "9月滞销 SKU"
        headers = ("库存SKU编号", "商品名", "浏览量TOP1", None, "浏览量指标值", "浏览量对标状态",
                   "浏览量对标说明", "销量第一店铺", "销量指标值", "销量对标状态", "销量对标说明", "计算")
        for column, header in enumerate(headers, 1):
            sheet.cell(2, column, header)
        for row, sku in enumerate(("sku-A", "sku-A", 123, None, "sku-failed", "sku-missing", "sku-absent"), 3):
            sheet.cell(row, 1, sku)
            sheet.cell(row, 2, "旧商品名")
            sheet.cell(row, 3, "旧浏览店铺")
            sheet.cell(row, 4, "空表头下仍有数据")
            sheet.cell(row, 5, 900)
            sheet.cell(row, 6, "旧成功状态")
            sheet.cell(row, 7, "旧浏览说明")
            sheet.cell(row, 8, "旧销量店铺")
            sheet.cell(row, 9, 800)
            sheet.cell(row, 10, "旧销量状态")
            sheet.cell(row, 11, "旧销量说明")
        sheet["A5"].number_format = "00000"
        sheet["L3"] = "=1+1"
        sheet["L3"].number_format = "0.00"
        sheet["L3"].font = Font(name="Arial", bold=True, color="FF0000")
        sheet["A2"].fill = PatternFill("solid", fgColor="ABCDEF")
        sheet["C3"].fill = PatternFill("solid", fgColor="FEDCBA")
        sheet.column_dimensions["B"].width = 32
        sheet.row_dimensions[3].height = 35
        sheet.freeze_panes = "C3"
        validation = DataValidation(type="list", formula1='"甲,乙"')
        sheet.add_data_validation(validation)
        validation.add("B3:B9")
        summary = book.create_sheet("原有汇总")
        summary["A1"] = "='Sheet1'!L3"
        summary["B1"] = "保留"
        summary.sheet_state = "hidden"
        book.save(self.source)
        book.close()

    def load(self, path):
        book = load_workbook(path)
        self.addCleanup(book.close)
        return book

    def result(self, sku="sku-A", **updates):
        result = {"sku": sku, "shop_name": "新店铺", "metric_value": 12,
                  "shop_count": 3, "item_id": "001234567890123456789",
                  "status": "matched", "message": ""}
        result.update(updates)
        return result

    def integrity_results(self, *, recovered=False):
        return [
            self.result(attempts=1), self.result("00123", metric_value=0, attempts=1),
            self.result("sku-failed", status="matched" if recovered else "failed",
                        message="" if recovered else "限流，等待后仍失败", attempts=7),
            self.result("sku-missing", status="not_found", attempts=1),
            self.result("sku-absent", attempts=1),
        ]

    def integrity_summary(self, *, recovered=False, **updates):
        summary = {
            "is_complete": recovered, "sku_count": 5, "total_rows": 6,
            "matched_count": 4 if recovered else 3, "not_found_count": 1,
            "failed_count": 0 if recovered else 1, "confirmed_count": 5 if recovered else 4,
            "retry_rounds_used": 2, "recovered_count": 1 if recovered else 0,
            "resumed_count": 4 if recovered else 0, "checkpoint_file": "内部进度，不应写入Excel.json",
        }
        summary.update(updates)
        return summary

    def test_second_row_header_duplicate_mapping_and_numeric_leading_zeros(self):
        info = service.load_workbook_input(self.source)
        self.assertEqual((info.sheet_name, info.header_row, info.sku_column), ("Sheet1", 2, 1))
        self.assertEqual(info.skus, ("sku-A", "00123", "sku-failed", "sku-missing", "sku-absent"))
        self.assertEqual(info.row_numbers["sku-A"], (3, 4))
        self.assertEqual(info.row_numbers["00123"], (5,))
        self.assertEqual(info.total_rows, 6)
        self.assertEqual(info.warnings, ())
        inspection = service.inspect_workbook(self.source)
        self.assertEqual(inspection["default_sheet"], "Sheet1")
        self.assertEqual(inspection["sheets"][0]["row_count"], 6)
        self.assertEqual(inspection["sheets"][0]["sku_count"], 5)
        self.assertIn("未找到", inspection["sheets"][1]["error"])

    def test_inventory_sku_header_wins_over_generic_and_searches_first_ten_rows(self):
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "SKU"
        sheet["B10"] = " 库存 SKU 编号 "
        sheet["B11"] = "001"
        sheet["B12"] = 0
        book.save(self.source)
        book.close()
        info = service.load_workbook_input(self.source)
        self.assertEqual((info.header_row, info.sku_column), (10, 2))
        self.assertEqual(info.skus, ("001", "0"))

    def test_generic_header_and_default_ignore_generated_detail(self):
        book = Workbook()
        book.active.title = "BigSeller浏览量对标明细"
        book.active.append(service.DETAIL_HEADERS)
        book.active.append(["detail-sku"])
        sheet = book.create_sheet("业务表")
        sheet.append(["sku"])
        sheet.append(["business-sku"])
        preferred = book.create_sheet("9月动销总表")
        preferred.append(["库存SKU编号"])
        preferred.append(["preferred-sku"])
        book.save(self.source)
        book.close()
        self.assertEqual(service.inspect_workbook(self.source)["default_sheet"], "9月动销总表")
        self.assertEqual(service.load_workbook_input(self.source, "业务表").skus, ("business-sku",))

    def test_export_preserves_source_formulas_formats_multisheets_and_blank_sku(self):
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        info = service.load_workbook_input(self.source)
        output = service.export_benchmark_workbook(info, [self.result(), self.result("00123", metric_value=0)], "views", self.output_dir)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), before)
        self.assertNotEqual(output, self.source)
        self.assertEqual(output.parent, self.output_dir)
        result = self.load(output)
        original = self.load(self.source)
        sheet = result["Sheet1"]
        for row in (3, 4):
            self.assertEqual(sheet.cell(row, 3).value, "新店铺")
            self.assertEqual(sheet.cell(row, 5).value, 12)
            self.assertEqual(sheet.cell(row, 6).value, "已匹配")
        self.assertEqual(sheet["E5"].value, 0)
        self.assertEqual(sheet["A5"].value, 123)
        self.assertEqual(sheet["A5"].number_format, "00000")
        for column in range(1, 13):
            self.assertEqual(sheet.cell(6, column).value, original["Sheet1"].cell(6, column).value)
        for row in range(3, 10):
            self.assertEqual(sheet.cell(row, 4).value, "空表头下仍有数据")
            self.assertEqual(sheet.cell(row, 8).value, "旧销量店铺")
            self.assertEqual(sheet.cell(row, 9).value, 800)
        self.assertIsNone(sheet["D2"].value)
        self.assertEqual(sheet["L3"].value, "=1+1")
        self.assertEqual(sheet["L3"]._style, original["Sheet1"]["L3"]._style)
        self.assertEqual(copy(sheet["C3"].fill), copy(original["Sheet1"]["C3"].fill))
        self.assertEqual(str(sheet.merged_cells), "A1:C1")
        self.assertEqual(sheet.freeze_panes, "C3")
        self.assertEqual(sheet.column_dimensions["B"].width, 32)
        self.assertEqual(sheet.row_dimensions[3].height, 35)
        self.assertEqual(str(sheet.data_validations.dataValidation[0].sqref), "B3:B9")
        self.assertEqual(result["原有汇总"]["A1"].value, "='Sheet1'!L3")
        self.assertEqual(result["原有汇总"].sheet_state, "hidden")
        detail = result["BigSeller浏览量对标明细"]
        self.assertEqual(detail.max_row, len(info.skus) + 1)
        self.assertEqual(detail["E2"].value, "001234567890123456789")
        self.assertEqual(detail["E2"].data_type, "s")
        self.assertEqual(detail["A3"].value, "00123")

    def test_failed_not_found_and_missing_results_clear_previous_success_values(self):
        info = service.load_workbook_input(self.source)
        results = [self.result("sku-failed", status="failed", message="接口返回失败"),
                   self.result("sku-missing", status="not_found", message="")]
        output = service.export_benchmark_workbook(info, results, "views", self.output_dir)
        sheet = self.load(output)["Sheet1"]
        for row, status in ((7, "查询失败"), (8, "未找到商品"), (9, "查询失败")):
            with self.subTest(row=row):
                self.assertIsNone(sheet.cell(row, 3).value)
                self.assertIsNone(sheet.cell(row, 5).value)
                self.assertEqual(sheet.cell(row, 6).value, status)
                self.assertTrue(sheet.cell(row, 7).value)
                self.assertEqual(sheet.cell(row, 8).value, "旧销量店铺")
        self.assertEqual(sheet["G7"].value, "接口返回失败")

    def test_sales_and_views_coexist_and_reruns_update_one_detail_sheet(self):
        info = service.load_workbook_input(self.source)
        views_path = service.export_benchmark_workbook(info, [self.result()], "views", self.output_dir)
        sales_input = service.load_workbook_input(views_path)
        sales_path = service.export_benchmark_workbook(sales_input, [self.result(shop_name="销量冠军", metric_value=99)], "sales", self.output_dir)
        sales_book = self.load(sales_path)
        self.assertEqual(sales_book["Sheet1"]["C3"].value, "新店铺")
        self.assertEqual(sales_book["Sheet1"]["H3"].value, "销量冠军")
        self.assertEqual(sales_book["Sheet1"]["I3"].value, 99)
        self.assertIn("BigSeller浏览量对标明细", sales_book.sheetnames)
        self.assertIn("BigSeller销量对标明细", sales_book.sheetnames)
        rerun = service.export_benchmark_workbook(service.load_workbook_input(sales_path),
                                                 [self.result(shop_name="更新店铺", metric_value=100)], "views", self.output_dir)
        rerun_book = self.load(rerun)
        self.assertEqual(rerun_book.sheetnames.count("BigSeller浏览量对标明细"), 1)
        self.assertEqual(rerun_book["BigSeller浏览量对标明细"].max_row, len(info.skus) + 1)
        self.assertEqual(rerun_book["BigSeller浏览量对标明细"]["B2"].value, "更新店铺")
        self.assertEqual(rerun_book["BigSeller销量对标明细"]["B2"].value, "销量冠军")

    def test_added_columns_append_after_data_even_with_empty_header(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["库存SKU编号", None, "原数据"])
        sheet.append(["a", "空表头有数据", "保持"])
        sheet["AO2"].fill = PatternFill("solid", fgColor="EEEEEE")
        book.save(self.source)
        book.close()
        output = service.export_benchmark_workbook(service.load_workbook_input(self.source), [self.result("a")], "sales", self.output_dir)
        sheet = self.load(output).active
        self.assertIsNone(sheet["B1"].value)
        self.assertEqual(sheet["B2"].value, "空表头有数据")
        self.assertEqual(sheet["C2"].value, "保持")
        self.assertEqual(sheet["D1"].value, "销量第一店铺")
        self.assertEqual(sheet["D2"].value, "新店铺")

    def test_new_columns_respect_merged_titles_beyond_the_table(self):
        book = Workbook()
        sheet = book.active
        sheet.merge_cells("A1:F1")
        sheet["A1"] = "合并标题"
        sheet["A2"] = "SKU"
        sheet["A3"] = "a"
        sheet["AO3"].fill = PatternFill("solid", fgColor="EEEEEE")
        book.save(self.source)
        book.close()
        output = service.export_benchmark_workbook(service.load_workbook_input(self.source), [self.result("a")], "views", self.output_dir)
        sheet = self.load(output).active
        self.assertEqual(sheet["G2"].value, "浏览量第一店铺")
        self.assertEqual(str(sheet.merged_cells), "A1:F1")

    def test_source_change_during_query_prevents_misaligned_export(self):
        info = service.load_workbook_input(self.source)
        book = self.load(self.source)
        book["Sheet1"]["A3"] = "用户在查询期间换了SKU"
        book.save(self.source)
        changed_bytes = self.source.read_bytes()
        with self.assertRaisesRegex(ValueError, "查询期间已变化"):
            service.export_benchmark_workbook(info, [self.result()], "views", self.output_dir)
        self.assertEqual(self.source.read_bytes(), changed_bytes)
        self.assertFalse(self.output_dir.exists())

    def test_lookup_contract_and_excel_formula_strings_are_exported_as_text(self):
        info = service.load_workbook_input(self.source)
        output = service.export_benchmark_workbook(info, {"sku-A": self.result(shop_name="=1+2", message="=HYPERLINK(\"x\")")}, "views", self.output_dir)
        book = self.load(output)
        for address in ("C3", "G3"):
            self.assertEqual(book["Sheet1"][address].data_type, "s")
        self.assertEqual(book["Sheet1"]["C3"].value, "=1+2")
        self.assertEqual(book["BigSeller浏览量对标明细"]["B2"].data_type, "s")

    def test_invalid_success_data_is_written_as_explicit_failure(self):
        info = service.load_workbook_input(self.source)
        output = service.export_benchmark_workbook(info, [self.result(metric_value=None)], "views", self.output_dir)
        sheet = self.load(output)["Sheet1"]
        self.assertIsNone(sheet["C3"].value)
        self.assertIsNone(sheet["E3"].value)
        self.assertEqual(sheet["F3"].value, "查询失败")
        self.assertIn("指标值", sheet["G3"].value)

    def test_repeated_exports_have_unique_paths_and_do_not_overwrite(self):
        info = service.load_workbook_input(self.source)
        first = service.export_benchmark_workbook(info, [self.result()], "views", self.output_dir)
        first_bytes = first.read_bytes()
        second = service.export_benchmark_workbook(info, [self.result(metric_value=20)], "views", self.output_dir)
        self.assertNotEqual(first, second)
        self.assertEqual(first.read_bytes(), first_bytes)
        self.assertEqual(self.load(second)["Sheet1"]["E3"].value, 20)

    def test_existing_unrelated_detail_named_sheet_is_preserved(self):
        book = self.load(self.source)
        unrelated = book.create_sheet("BigSeller浏览量对标明细")
        unrelated["A1"] = "用户自己的内容"
        book.save(self.source)
        output = service.export_benchmark_workbook(service.load_workbook_input(self.source), [self.result()], "views", self.output_dir)
        result = self.load(output)
        self.assertEqual(result["BigSeller浏览量对标明细"]["A1"].value, "用户自己的内容")
        self.assertEqual(result["BigSeller浏览量对标明细_2"]["B2"].value, "新店铺")

    def test_formula_sku_without_cached_result_is_an_explicit_input_error(self):
        book = self.load(self.source)
        book["Sheet1"]["A3"] = '=TEXT(123,"00000")'
        book.save(self.source)
        with self.assertRaisesRegex(ValueError, "A3.*没有缓存结果"):
            service.load_workbook_input(self.source, "Sheet1")
        self.assertIn("A3", service.inspect_workbook(self.source)["sheets"][0]["error"])

    def test_invalid_selection_and_metric_do_not_produce_output(self):
        with self.assertRaisesRegex(ValueError, "不存在"):
            service.load_workbook_input(self.source, "缺少的表")
        info = service.load_workbook_input(self.source)
        with self.assertRaisesRegex(ValueError, "浏览量或销量"):
            service.export_benchmark_workbook(info, [], "orders", self.output_dir)
        self.assertFalse(self.output_dir.exists())

    def test_incomplete_export_audits_counts_and_lists_only_failed_skus(self):
        source_before = self.source.read_bytes()
        info = service.load_workbook_input(self.source)
        output = service.export_benchmark_workbook(
            info, self.integrity_results(), "views", self.output_dir, summary=self.integrity_summary(),
        )
        self.assertTrue(output.stem.endswith("_结果不完整"))
        self.assertEqual(self.source.read_bytes(), source_before)
        book = self.load(output)
        summary = dict(book["BigSeller浏览量对标汇总"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["完成性"], "结果不完整")
        self.assertEqual(summary["原表行数"], 6)
        self.assertEqual(summary["去重SKU数"], 5)
        self.assertEqual(summary["成功匹配数"], 3)
        self.assertEqual(summary["确认未匹配数"], 1)
        self.assertEqual(summary["失败待补查数"], 1)
        self.assertEqual(summary["已确认SKU数"], 4)
        self.assertEqual(summary["补查轮数"], 2)
        pending = book["BigSeller浏览量待补查"]
        self.assertEqual(list(pending.values), [service.PENDING_HEADERS, ("sku-failed", "限流，等待后仍失败", 7, "7")])
        self.assertEqual(pending.auto_filter.ref, "A1:D2")
        self.assertEqual(book["Sheet1"]["F8"].value, "未找到商品")
        self.assertNotIn("内部进度", str(list(book["BigSeller浏览量对标汇总"].values)))

    def test_integrity_audit_rejects_missing_duplicate_extra_or_invalid_results(self):
        info = service.load_workbook_input(self.source)
        valid = self.integrity_results()
        cases = {
            "missing": valid[:-1],
            "duplicate": [*valid, valid[0]],
            "extra": [*valid, self.result("额外SKU")],
            "bad_status": [{**valid[0], "status": "unknown"}, *valid[1:]],
            "invalid_success": [{**valid[0], "metric_value": None}, *valid[1:]],
            "invalid_attempts": [{**valid[0], "attempts": -1}, *valid[1:]],
            "invalid_result": [None, *valid[1:]],
            "mismatched_mapping": {item["sku"]: {**item, "sku": "错位SKU"} for item in valid},
        }
        for name, results in cases.items():
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, "完整性校验失败"):
                service.export_benchmark_workbook(info, results, "views", self.output_dir, summary=self.integrity_summary())
        self.assertFalse(self.output_dir.exists())

    def test_integrity_audit_rejects_counts_or_completion_that_disagree_with_results(self):
        info = service.load_workbook_input(self.source)
        bad_updates = (
            {"sku_count": 6}, {"total_rows": 5}, {"matched_count": 4}, {"not_found_count": 0},
            {"failed_count": 0}, {"confirmed_count": 5}, {"is_complete": True},
            {"is_complete": 0}, {"failed_count": True}, {"retry_rounds_used": -1}, {"resumed_count": 6},
        )
        for updates in bad_updates:
            with self.subTest(updates=updates), self.assertRaisesRegex(ValueError, "完整性校验失败"):
                service.export_benchmark_workbook(
                    info, self.integrity_results(), "views", self.output_dir,
                    summary=self.integrity_summary(**updates),
                )
        with self.assertRaisesRegex(ValueError, "完成状态"):
            service.export_benchmark_workbook(
                info, self.integrity_results(recovered=True), "views", self.output_dir,
                summary=self.integrity_summary(recovered=True, is_complete=False),
            )
        self.assertFalse(self.output_dir.exists())

    def test_recovered_rerun_clears_owned_pending_but_preserves_custom_and_other_metric_sheets(self):
        book = self.load(self.source)
        custom = book.create_sheet("BigSeller浏览量待补查")
        custom.append(service.PENDING_HEADERS)
        custom.append(("用户SKU", "用户备注", 8, "12"))
        custom_summary = book.create_sheet("BigSeller浏览量对标汇总")
        custom_summary.append(service.SUMMARY_HEADERS)
        custom_summary.append(("用户自定义", "保留"))
        book.save(self.source)
        original_bytes = self.source.read_bytes()
        first = service.export_benchmark_workbook(
            service.load_workbook_input(self.source), self.integrity_results(), "views", self.output_dir,
            summary=self.integrity_summary(),
        )
        first_bytes = first.read_bytes()
        sales = service.export_benchmark_workbook(
            service.load_workbook_input(first), self.integrity_results(), "sales", self.output_dir,
            summary=self.integrity_summary(),
        )
        sales_bytes = sales.read_bytes()
        complete = service.export_benchmark_workbook(
            service.load_workbook_input(sales), self.integrity_results(recovered=True), "views", self.output_dir,
            summary=self.integrity_summary(recovered=True),
        )
        result = self.load(complete)
        self.assertFalse(complete.stem.endswith("_结果不完整"))
        self.assertEqual(list(result["BigSeller浏览量待补查"].values),
                         [service.PENDING_HEADERS, ("用户SKU", "用户备注", 8, "12")])
        self.assertEqual(list(result["BigSeller浏览量待补查_2"].values), [service.PENDING_HEADERS])
        self.assertEqual(result["BigSeller浏览量待补查_2"].auto_filter.ref, "A1:D1")
        self.assertEqual(list(result["BigSeller浏览量对标汇总"].values),
                         [service.SUMMARY_HEADERS, ("用户自定义", "保留")])
        summary = dict(result["BigSeller浏览量对标汇总_2"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["完成性"], "查询完整")
        self.assertEqual(summary["确认未匹配数"], 1)
        self.assertEqual(summary["失败待补查数"], 0)
        self.assertEqual(summary["补查恢复数"], 1)
        self.assertEqual(summary["沿用已确认结果数"], 4)
        prior = self.load(sales)
        for name in ("BigSeller销量对标汇总", "BigSeller销量待补查", "BigSeller销量对标明细"):
            self.assertEqual(list(result[name].values), list(prior[name].values))
            self.assertEqual(result[name]["A1"].comment, prior[name]["A1"].comment)
        self.assertEqual(result["Sheet1"]["F7"].value, "已匹配")
        self.assertEqual(self.source.read_bytes(), original_bytes)
        self.assertEqual(first.read_bytes(), first_bytes)
        self.assertEqual(sales.read_bytes(), sales_bytes)

    def test_complete_mapping_export_accepts_confirmed_not_found_and_has_no_pending_sheet(self):
        results = self.integrity_results(recovered=True)
        output = service.export_benchmark_workbook(
            service.load_workbook_input(self.source), {item["sku"]: item for item in results}, "sales", self.output_dir,
            summary=self.integrity_summary(recovered=True),
        )
        book = self.load(output)
        summary = dict(book["BigSeller销量对标汇总"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(summary["完成性"], "查询完整")
        self.assertEqual(summary["确认未匹配数"], 1)
        self.assertNotIn("BigSeller销量待补查", book.sheetnames)
        self.assertNotIn("_结果不完整", output.stem)


if __name__ == "__main__":
    unittest.main()
