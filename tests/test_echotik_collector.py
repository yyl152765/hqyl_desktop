from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from backend.services.echotik_collector import (
    CreatorFilters,
    EchoTikCollectJob,
    EchoTikClient,
    EchoTikHtmlResponseError,
    EchoTikPage,
    NumberRange,
    ProductFilters,
    collect_complete_pages,
    collect_complete_product_pages,
    creator_matches_filters,
    normalize_filter_categories,
    parse_compact_number,
    parse_keywords,
    product_matches_filters,
    run_echotik_collection,
    validate_echotik_payload,
)
from decimal import Decimal


def response_mock(
    *,
    status_code: int = 200,
    content_type: str = "application/json",
    text: str = "",
    payload: object | None = None,
) -> Mock:
    response = Mock()
    response.status_code = status_code
    response.headers = {"Content-Type": content_type}
    response.text = text
    if payload is None:
        response.json.side_effect = ValueError("not json")
    else:
        response.json.return_value = payload
    return response


class FakeEchoTikClient:
    def __init__(self, username: str, password: str) -> None:
        self.username = username
        self.password = password

    def login(self) -> None:
        return None

    def product_page(self, keyword: str, page: int, per_page: int):
        if page == 1:
            return [
                {
                    "product_id": "p1",
                    "product_name": "Lip Matte",
                    "cover_url": "https://example.test/p1.jpg",
                    "avg_price": "฿99",
                    "influencers_count": "2",
                    "videos_count": "7",
                    "seller": {"seller_id": "s1", "seller_name": "Shop A"},
                    "category": {"name": "美妆"},
                }
            ]
        return []

    def product_detail(self, product_id: str):
        raise AssertionError("低请求模式不应调用商品详情接口")

    def influencer_page(self, product_id: str, page: int, per_page: int):
        if page == 1:
            return [
                {
                    "influencer_id": "i1",
                    "unique_id": "creator_handle",
                    "influencer_name": "creator-one",
                    "avatar_url": "https://example.test/i1.jpg",
                    "region": {"id": "TH", "name": "泰国"},
                    "follower_count": "1万",
                    "heart_count": "20万",
                    "categories": ["美妆"],
                    "sales": "250",
                    "product_ifl_gmv_amt_fz": "฿88",
                    "product_ifl_gmv_amt": "88",
                    "related_video": "3",
                    "total_video_viewers": "10万",
                    "related_live": "1",
                    "total_live_viewers": "2万",
                }
            ]
        return []


class EchoTikCollectorTests(unittest.TestCase):
    def test_collect_complete_pages_recovers_rows_lost_to_page_reordering(self) -> None:
        calls = {1: 0, 2: 0}

        def fetch(page: int) -> EchoTikPage:
            calls[page] += 1
            responses = {
                (1, 1): [{"product_id": "p1"}, {"product_id": "p2"}],
                (2, 1): [{"product_id": "p2"}, {"product_id": "p3"}],
                (1, 2): [{"product_id": "p1"}, {"product_id": "p4"}],
            }
            return EchoTikPage(
                rows=responses.get((page, calls[page]), []),
                total=4,
                last_page=2,
                current_page=page,
            )

        logs: list[str] = []
        rows = collect_complete_pages(
            fetch,
            row_key=lambda row: str(row.get("product_id") or ""),
            page_size=2,
            max_pages=10,
            delay_seconds=0,
            description="商品列表",
            progress=logs.append,
        )

        self.assertEqual({row["product_id"] for row in rows}, {"p1", "p2", "p3", "p4"})
        self.assertTrue(any("开始补采缺失分页" in line for line in logs))

    def test_collect_complete_pages_can_keep_stable_partial_rows_with_warning(self) -> None:
        warnings: list[str] = []

        rows = collect_complete_pages(
            lambda page: EchoTikPage(
                rows=[{"influencer_id": "i1"}, {"influencer_id": "i2"}],
                total=3,
                last_page=1,
                current_page=page,
            ),
            row_key=lambda row: str(row.get("influencer_id") or ""),
            page_size=50,
            max_pages=10,
            delay_seconds=0,
            description="达人列表",
            progress=None,
            partial_warnings=warnings,
        )

        self.assertEqual([row["influencer_id"] for row in rows], ["i1", "i2"])
        self.assertEqual(len(warnings), 1)
        self.assertIn("缺失 1 条", warnings[0])

    def test_product_pages_recover_with_alternative_sort_and_stable_ranking(self) -> None:
        class ReorderingProductClient:
            def __init__(self) -> None:
                self.calls: list[tuple[int, str, str]] = []

            def product_page(
                self,
                keyword: str,
                page: int,
                per_page: int,
                filters: ProductFilters | None = None,
                *,
                order: str = "total_sale_nd_cnt",
                sort: str = "desc",
            ) -> EchoTikPage:
                self.calls.append((page, order, sort))
                responses = {
                    ("total_sale_nd_cnt", "desc", 1): [
                        {"product_id": "p4", "total_sale_nd_cnt": "40"},
                        {"product_id": "p3", "total_sale_nd_cnt": "30"},
                    ],
                    ("total_sale_nd_cnt", "desc", 2): [
                        {"product_id": "p3", "total_sale_nd_cnt": "30"},
                        {"product_id": "p2", "total_sale_nd_cnt": "20"},
                    ],
                    ("total_sale_nd_cnt", "asc", 1): [
                        {"product_id": "p0", "total_sale_nd_cnt": "0"},
                        {"product_id": "p1", "total_sale_nd_cnt": "10"},
                    ],
                    ("total_sale_nd_cnt", "asc", 2): [
                        {"product_id": "p2", "total_sale_nd_cnt": "20"},
                        {"product_id": "p3", "total_sale_nd_cnt": "30"},
                    ],
                }
                return EchoTikPage(
                    rows=responses.get((order, sort, page), []),
                    total=4,
                    last_page=2,
                    current_page=page,
                )

        client = ReorderingProductClient()
        job = EchoTikCollectJob(
            username="user",
            password="password",
            keywords=("keyword",),
            output_dir=Path("."),
            page_size=2,
            product_delay_seconds=0,
        )

        rows = collect_complete_product_pages(client, job, "keyword", progress=None)

        self.assertEqual([row["product_id"] for row in rows], ["p4", "p3", "p2", "p1"])
        self.assertIn((1, "total_sale_nd_cnt", "asc"), client.calls)

    def test_request_rebuilds_session_and_retries_html_response(self) -> None:
        client = EchoTikClient("user", "password")
        client.access_token = "token"
        client.session.request = Mock(
            side_effect=[
                response_mock(content_type="text/html", text="<!DOCTYPE html><title>可能遇到了一点问题</title>"),
                response_mock(payload={"code": 0, "data": {"ok": True}}),
            ]
        )

        with patch.object(client, "_renew_authenticated_session") as renew_session:
            result = client._request("GET", "/data/products/p1")

        self.assertEqual(result, {"ok": True})
        renew_session.assert_called_once()
        self.assertEqual(client.session.request.call_count, 2)

    def test_request_raises_typed_error_after_repeated_html_responses(self) -> None:
        client = EchoTikClient("user", "password")
        client.access_token = "token"
        client.session.request = Mock(
            side_effect=[
                response_mock(content_type="text/html", text="<!DOCTYPE html>"),
                response_mock(content_type="text/html", text="<!DOCTYPE html>"),
                response_mock(content_type="text/html", text="<!DOCTYPE html>"),
            ]
        )

        with (
            patch.object(client, "_renew_authenticated_session"),
            patch("backend.services.echotik_collector.sleep_seconds"),
        ):
            with self.assertRaises(EchoTikHtmlResponseError):
                client._request("GET", "/data/products/p1")

    def test_request_renews_session_before_observed_request_threshold(self) -> None:
        client = EchoTikClient("user", "password")
        client.access_token = "token"
        client.request_count = 450
        client.session.request = Mock(return_value=response_mock(payload={"code": 0, "data": []}))

        with patch.object(client, "_renew_authenticated_session") as renew_session:
            result = client._request("GET", "/data/products")

        self.assertEqual(result, [])
        renew_session.assert_called_once()

    def test_parse_keywords_supports_common_separators_and_deduplicates(self) -> None:
        self.assertEqual(
            parse_keywords("ลิปแมตต์สีชัด\nLip Matte，ลิปสติกติดทน;lip matte"),
            ("ลิปแมตต์สีชัด", "Lip Matte", "ลิปสติกติดทน"),
        )

    def test_validate_payload_allows_all_products_mode_without_keyword(self) -> None:
        job = validate_echotik_payload(
            {
                "username": "user",
                "password": "password",
                "keyword": "",
                "output_dir": "C:/Temp",
            }
        )

        self.assertEqual(job.keywords, ("",))
        self.assertEqual(job.max_products, 200)

    def test_validate_payload_accepts_and_bounds_max_products(self) -> None:
        job = validate_echotik_payload(
            {
                "username": "user",
                "password": "password",
                "output_dir": "C:/Temp",
                "max_products": 500,
            }
        )
        self.assertEqual(job.max_products, 500)

        for value in (0, 10001):
            with self.subTest(value=value), self.assertRaisesRegex(ValueError, "最多采集商品数"):
                validate_echotik_payload(
                    {
                        "username": "user",
                        "password": "password",
                        "output_dir": "C:/Temp",
                        "max_products": value,
                    }
                )

    def test_validate_payload_accepts_keyword_list(self) -> None:
        job = validate_echotik_payload(
            {
                "username": "user",
                "password": "password",
                "keywords": ["关键词一", "关键词二", "关键词一"],
                "output_dir": "C:/Temp",
            }
        )
        self.assertEqual(job.keywords, ("关键词一", "关键词二"))
        self.assertFalse(job.fetch_product_details)

    def test_validate_payload_can_enable_product_details_explicitly(self) -> None:
        job = validate_echotik_payload(
            {
                "username": "user",
                "password": "password",
                "keywords": ["lip"],
                "output_dir": "C:/Temp",
                "fetch_product_details": True,
            }
        )

        self.assertTrue(job.fetch_product_details)

    def test_validate_payload_parses_filter_ranges(self) -> None:
        job = validate_echotik_payload(
            {
                "username": "user",
                "password": "password",
                "keywords": ["lip"],
                "output_dir": "C:/Temp",
                "filters": {
                    "products": {
                        "category_ids": ["beauty"],
                        "category_names": ["美妆"],
                        "category_paths": [["root", "beauty"]],
                        "sales_period_days": 30,
                        "period_sales": {"min": "1万", "max": "2.5万"},
                    },
                    "creators": {
                        "category_names": ["美妆"],
                        "video_play_count": {"min": "1M", "max": None},
                    },
                },
            }
        )
        self.assertEqual(job.product_filters.sales_period_days, 30)
        self.assertEqual(job.product_filters.category_paths, (("root", "beauty"),))
        self.assertEqual(job.product_filters.period_sales.minimum, Decimal("10000"))
        self.assertEqual(job.product_filters.period_sales.maximum, Decimal("25000"))
        self.assertEqual(job.creator_filters.video_play_count.minimum, Decimal("1000000"))

    def test_validate_payload_rejects_reversed_range(self) -> None:
        with self.assertRaisesRegex(ValueError, "最小值不能大于最大值"):
            validate_echotik_payload(
                {
                    "username": "user",
                    "password": "password",
                    "keywords": ["lip"],
                    "output_dir": "C:/Temp",
                    "filters": {"creators": {"sales": {"min": 100, "max": 10}}},
                }
            )

    def test_compact_number_and_local_filters(self) -> None:
        self.assertEqual(parse_compact_number("20.35万"), Decimal("203500"))
        self.assertEqual(parse_compact_number("14.76亿"), Decimal("1476000000"))
        self.assertEqual(parse_compact_number("฿6,810.53"), Decimal("6810.53"))
        self.assertEqual(parse_compact_number("1.2M"), Decimal("1200000.0"))

        product_filters = ProductFilters(
            category_names=("美妆",),
            period_sales=NumberRange(minimum=Decimal("100")),
            video_views=NumberRange(minimum=Decimal("10000")),
        )
        self.assertTrue(
            product_matches_filters(
                {"categories": "美妆个护/美妆/口红", "total_sale_7d_cnt": "100", "view_count": "1万"},
                product_filters,
            )
        )

        creator_filters = CreatorFilters(
            category_names=("口红",),
            sales=NumberRange(minimum=Decimal("200")),
            video_play_count=NumberRange(minimum=Decimal("100000")),
        )
        self.assertEqual(
            creator_matches_filters(
                {"categories": ["口红"], "sales": "250", "total_video_viewers": "10万"},
                creator_filters,
            ),
            (True, False),
        )
        self.assertEqual(
            creator_matches_filters(
                {"categories": ["口红"], "sales": "250", "total_video_viewers": ""},
                creator_filters,
            ),
            (False, True),
        )

    def test_product_page_maps_server_side_filters(self) -> None:
        client = EchoTikClient("user", "password")
        filters = ProductFilters(
            category_ids=("beauty", "fashion"),
            sales_period_days=7,
            period_sales=NumberRange(minimum=Decimal("100")),
            total_sales=NumberRange(maximum=Decimal("500")),
        )
        with patch.object(client, "_request_page", return_value=EchoTikPage(rows=[])) as request_page:
            client.product_page("lip", 1, 50, filters)
        params = request_page.call_args.kwargs["params"]
        self.assertEqual(params["product_categories"], ["beauty", "fashion"])
        self.assertEqual(params["dateRange"], 7)
        self.assertEqual(params["sales"], ">99")
        self.assertEqual(params["total_sale_cnt"], "<501")

    def test_product_page_sends_cascaded_category_path_with_bracket_key(self) -> None:
        client = EchoTikClient("user", "password")
        filters = ProductFilters(category_paths=(("602118", "818696"),))
        with patch.object(client, "_request_page", return_value=EchoTikPage(rows=[])) as request_page:
            client.product_page("", 1, 50, filters)

        params = request_page.call_args.kwargs["params"]
        self.assertEqual(params["product_categories[]"], ["602118", "818696"])
        self.assertNotIn("product_categories", params)

    def test_product_page_accepts_recovery_sort(self) -> None:
        client = EchoTikClient("user", "password")
        with patch.object(client, "_request_page", return_value=EchoTikPage(rows=[])) as request_page:
            client.product_page(
                "lip",
                1,
                50,
                order="total_sale_cnt",
                sort="asc",
            )
        params = request_page.call_args.kwargs["params"]
        self.assertEqual(params["order"], "total_sale_cnt")
        self.assertEqual(params["sort"], "asc")

    def test_product_page_omits_blank_keyword_for_all_products_mode(self) -> None:
        client = EchoTikClient("user", "password")
        with patch.object(client, "_request_page", return_value=EchoTikPage(rows=[])) as request_page:
            client.product_page("", 1, 50)

        params = request_page.call_args.kwargs["params"]
        self.assertNotIn("keyword", params)

    def test_product_collection_stops_at_configured_limit(self) -> None:
        class AllProductsClient:
            def __init__(self) -> None:
                self.pages: list[int] = []

            def product_page(
                self,
                keyword: str,
                page: int,
                per_page: int,
                filters: ProductFilters | None = None,
                *,
                order: str = "total_sale_nd_cnt",
                sort: str = "desc",
            ) -> EchoTikPage:
                self.pages.append(page)
                start = (page - 1) * per_page
                return EchoTikPage(
                    rows=[
                        {
                            "product_id": f"p{index:05d}",
                            "total_sale_7d_cnt": str(10000 - index),
                        }
                        for index in range(start, start + per_page)
                    ],
                    total=10000,
                    last_page=200,
                    current_page=page,
                )

        client = AllProductsClient()
        job = EchoTikCollectJob(
            username="user",
            password="password",
            keywords=("",),
            output_dir=Path("."),
            page_size=50,
            max_products=75,
            product_delay_seconds=0,
        )

        rows = collect_complete_product_pages(client, job, "", progress=None)

        self.assertEqual(len(rows), 75)
        self.assertEqual(client.pages, [1, 2])

    def test_product_collection_merges_multiple_category_paths(self) -> None:
        class CategoryPathClient:
            def product_page(
                self,
                keyword: str,
                page: int,
                per_page: int,
                filters: ProductFilters | None = None,
                *,
                order: str = "total_sale_nd_cnt",
                sort: str = "desc",
            ) -> EchoTikPage:
                path = filters.category_paths[0] if filters and filters.category_paths else ()
                prefix = "a" if path[-1] == "child-a" else "b"
                return EchoTikPage(
                    rows=[
                        {"product_id": f"{prefix}{index}", "total_sale_7d_cnt": str(100 - index)}
                        for index in range(2)
                    ],
                    total=2,
                    last_page=1,
                    current_page=1,
                )

        job = EchoTikCollectJob(
            username="user",
            password="password",
            keywords=("",),
            output_dir=Path("."),
            max_products=2,
            product_filters=ProductFilters(
                category_names=("猫狗配件",),
                category_paths=(("root", "child-a"), ("root", "child-b")),
            ),
            product_delay_seconds=0,
        )

        rows = collect_complete_product_pages(CategoryPathClient(), job, "", progress=None)

        self.assertEqual(len(rows), 2)
        self.assertEqual({row["product_id"] for row in rows}, {"a0", "b0"})

    def test_normalize_filter_categories_preserves_secondary_categories(self) -> None:
        categories = normalize_filter_categories(
            [
                {
                    "value": "beauty",
                    "label": "美妆个护",
                    "children": [
                        {"value": "lip", "label": "唇部彩妆"},
                        {"value": "face", "label": "面部彩妆"},
                    ],
                },
                {"value": "all", "label": "全部"},
            ]
        )

        self.assertEqual(
            categories,
            [
                {
                    "id": "beauty",
                    "name": "美妆个护",
                    "children": [
                        {"id": "lip", "name": "唇部彩妆", "children": []},
                        {"id": "face", "name": "面部彩妆", "children": []},
                    ],
                }
            ],
        )

    def test_run_collection_paginates_and_exports_excel(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("ลิปแมตต์สีชัด", "ลิปสติกติดทน"),
                output_dir=Path(temp_dir),
                page_size=50,
                product_delay_seconds=0,
                creator_delay_seconds=0,
            )
            logs: list[str] = []
            with patch("backend.services.echotik_collector.EchoTikClient", FakeEchoTikClient):
                result = run_echotik_collection(job, logs.append)

            self.assertEqual(result.product_count, 1)
            self.assertEqual(result.product_candidate_count, 1)
            self.assertEqual(result.creator_source_count, 1)
            self.assertEqual(result.creator_count, 1)
            self.assertEqual(result.creator_filtered_count, 0)
            self.assertEqual(result.failed_product_count, 0)
            self.assertTrue(result.output_file.is_file())
            self.assertTrue(any("商品列表第 1 页" in line for line in logs))
            self.assertTrue(any("筛选并去重后 1 个" in line for line in logs))
            workbook = load_workbook(result.output_file, read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["Products", "Creators", "FilterSummary"])
                product_sheet = workbook["Products"]
                creator_sheet = workbook["Creators"]
                self.assertEqual(product_sheet.max_row, 2)
                self.assertEqual(creator_sheet.max_row, 2)
                self.assertEqual(product_sheet["A2"].value, "ลิปแมตต์สีชัด | ลิปสติกติดทน")
                self.assertEqual(creator_sheet["A2"].value, "ลิปแมตต์สีชัด | ลิปสติกติดทน")
                header = [cell.value for cell in creator_sheet[1]]
                self.assertEqual(creator_sheet.cell(row=2, column=header.index("creator_id") + 1).value, "creator_handle")
                self.assertEqual(creator_sheet.cell(row=2, column=header.index("creator_uid") + 1).value, "i1")
                self.assertEqual(creator_sheet.cell(row=2, column=header.index("creator_name") + 1).value, "creator-one")
                self.assertEqual(creator_sheet.cell(row=2, column=header.index("creator_sales") + 1).value, "250")
                self.assertEqual(creator_sheet.cell(row=2, column=header.index("creator_sales_value") + 1).value, 250)
            finally:
                workbook.close()

    def test_run_collection_labels_all_products_mode_in_logs_and_excel(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("",),
                output_dir=Path(temp_dir),
                product_delay_seconds=0,
                creator_delay_seconds=0,
            )
            logs: list[str] = []
            with patch("backend.services.echotik_collector.EchoTikClient", FakeEchoTikClient):
                result = run_echotik_collection(job, logs.append)

            self.assertTrue(any("全部商品（未指定关键词）" in line for line in logs))
            self.assertEqual(result.products_preview[0]["keyword"], "全部商品")
            self.assertEqual(result.creators_preview[0]["keyword"], "全部商品")
            self.assertIn("全部商品", result.output_file.name)
            workbook = load_workbook(result.output_file, read_only=True)
            try:
                self.assertEqual(workbook["Products"]["A2"].value, "全部商品")
                self.assertEqual(workbook["Creators"]["A2"].value, "全部商品")
                summary = {
                    workbook["FilterSummary"].cell(row=row, column=1).value:
                    workbook["FilterSummary"].cell(row=row, column=2).value
                    for row in range(2, workbook["FilterSummary"].max_row + 1)
                }
                self.assertEqual(summary["关键词"], "全部商品")
                self.assertEqual(summary["每个搜索范围最多商品数"], 200)
            finally:
                workbook.close()

    def test_run_collection_keeps_partial_creator_rows_and_marks_product(self) -> None:
        class PartialCreatorClient(FakeEchoTikClient):
            def influencer_page(self, product_id: str, page: int, per_page: int):
                return EchoTikPage(
                    rows=super().influencer_page(product_id, page, per_page),
                    total=2,
                    last_page=1,
                    current_page=page,
                )

        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("lip",),
                output_dir=Path(temp_dir),
                product_delay_seconds=0,
                creator_delay_seconds=0,
            )
            with patch("backend.services.echotik_collector.EchoTikClient", PartialCreatorClient):
                result = run_echotik_collection(job)

            self.assertEqual(result.creator_source_count, 1)
            self.assertEqual(result.creator_count, 1)
            self.assertEqual(result.failed_product_count, 1)
            self.assertIn("缺失 1 条", result.products_preview[0]["error"])

    def test_run_collection_can_opt_in_to_product_details(self) -> None:
        class DetailClient(FakeEchoTikClient):
            def product_detail(self, product_id: str):
                return {"product_id": product_id, "sale_cnt": "10", "gmv_amt_fz": "฿990"}

        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("lip",),
                output_dir=Path(temp_dir),
                product_delay_seconds=0,
                creator_delay_seconds=0,
                fetch_product_details=True,
            )
            with patch("backend.services.echotik_collector.EchoTikClient", DetailClient):
                result = run_echotik_collection(job)

            self.assertEqual(result.failed_product_count, 0)
            self.assertEqual(result.products_preview[0]["total_sales"], "10")

    def test_run_collection_opens_circuit_after_repeated_html_failures(self) -> None:
        class HtmlFailingClient:
            last_instance = None

            def __init__(self, username: str, password: str) -> None:
                self.influencer_calls = 0
                HtmlFailingClient.last_instance = self

            def login(self) -> None:
                return None

            def product_page(self, keyword: str, page: int, per_page: int):
                if page == 1:
                    return [
                        {"product_id": "p1", "product_name": "Product 1"},
                        {"product_id": "p2", "product_name": "Product 2"},
                        {"product_id": "p3", "product_name": "Product 3"},
                    ]
                return []

            def influencer_page(self, product_id: str, page: int, per_page: int):
                self.influencer_calls += 1
                raise EchoTikHtmlResponseError("EchoTik 返回 HTML 错误页")

        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("toy",),
                output_dir=Path(temp_dir),
                product_delay_seconds=0,
                creator_delay_seconds=0,
            )
            logs: list[str] = []
            with patch("backend.services.echotik_collector.EchoTikClient", HtmlFailingClient):
                result = run_echotik_collection(job, logs.append)

            self.assertEqual(result.product_count, 3)
            self.assertEqual(result.failed_product_count, 3)
            self.assertEqual(HtmlFailingClient.last_instance.influencer_calls, 2)
            self.assertTrue(any("已触发熔断" in line for line in logs))
            self.assertIn("未采集", result.products_preview[2]["error"])

    def test_run_collection_filters_creators_after_complete_pagination(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            job = EchoTikCollectJob(
                username="demo@example.com",
                password="secret",
                keywords=("lip",),
                output_dir=Path(temp_dir),
                product_delay_seconds=0,
                creator_delay_seconds=0,
                creator_filters=CreatorFilters(sales=NumberRange(minimum=Decimal("300"))),
            )
            with patch("backend.services.echotik_collector.EchoTikClient", FakeEchoTikClient):
                result = run_echotik_collection(job)

            self.assertEqual(result.product_count, 1)
            self.assertEqual(result.creator_source_count, 1)
            self.assertEqual(result.creator_count, 0)
            self.assertEqual(result.creator_filtered_count, 1)
            workbook = load_workbook(result.output_file, read_only=True)
            try:
                product_sheet = workbook["Products"]
                headers = [cell.value for cell in product_sheet[1]]
                self.assertEqual(product_sheet.cell(2, headers.index("source_creator_count") + 1).value, 1)
                self.assertEqual(product_sheet.cell(2, headers.index("matched_creator_count") + 1).value, 0)
                self.assertEqual(workbook["Creators"].max_row, 1)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
