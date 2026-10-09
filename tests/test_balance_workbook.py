from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import load_workbook
from PIL import Image

from backend.services.balance_workbook import write_balance_workbook


class BalanceWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.image = self.root / "proof.png"
        Image.new("RGB", (1600, 900), "white").save(self.image)

    def row(self, **changes):
        return {"store_name": "=HYPERLINK(\"https://example.invalid\")", "country": "TH", "currency": "THB",
                "status": "success", "captured_at": "2026-09-28T15:00:00+08:00", "notes": [], "field_errors": {},
                "values": {"total": "0.00", "pending": "-123.45", "income": "1.25", "balance": "0.00", "ads": "-2", "processing": None},
                "evidence": {field: str(self.image) for field in ("total", "pending", "income", "balance", "ads")}, **changes}

    def test_temu_layout_money_images_and_literal_store_names(self):
        target = self.root / "temu.xlsx"
        write_balance_workbook("temu", [self.row()], target, month="2026-08")
        book = load_workbook(target)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual([cell.value for cell in sheet[1]][:5], ["平台", "店铺", "账户总金额", "预估待结算销售额", "截图"])
        self.assertEqual(sheet["A2"].value, "temu")
        self.assertEqual(sheet["B2"].data_type, "s")
        self.assertEqual(sheet["C2"].value, 0)
        self.assertEqual(sheet["D2"].value, -123.45)
        self.assertEqual(sheet["K2"].value, "2026-08")
        self.assertIn(self.row()["store_name"], sheet["E2"].value)
        self.assertIn(self.row()["captured_at"], sheet["F2"].value)
        self.assertIn("非截图原文", sheet["E2"].value)
        self.assertEqual(sheet["E2"].data_type, "s")
        self.assertEqual(len(sheet._images), 2)
        self.assertEqual([image.anchor._from.col for image in sheet._images], [4, 5])
        for image in sheet._images:
            self.assertAlmostEqual(image.anchor.ext.cx / image.anchor.ext.cy, 1600 / 900)
            self.assertGreater(image.anchor._from.rowOff, 0)
        self.assertTrue(self.image.exists())
        with ZipFile(target) as archive:
            media = [name for name in archive.namelist() if name.startswith("xl/media/")]
            self.assertEqual(len(media), 2)
            for filename in media:
                self.assertEqual(archive.read(filename), self.image.read_bytes())

    def test_lazada_sheets_blank_zero_negative_and_no_cross_currency_totals(self):
        target = self.root / "lazada.xlsx"
        rows = [self.row(), self.row(store_name="另一个站点", country="PH", currency="PHP", values={"income": "12", "balance": "0", "ads": "0", "processing": "-8.1"}, evidence={"processing": str(self.image)})]
        write_balance_workbook("lazada", rows, target)
        book = load_workbook(target)
        self.addCleanup(book.close)
        self.assertEqual(book.sheetnames, ["余额统计", "Income截图", "Balance截图", "Ads截图", "processing"])
        sheet = book.active
        self.assertEqual([cell.value for cell in sheet[1]][:5], ["店铺名称", "income", "Balance", "Ads", "processing"])
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(sheet["C2"].value, 0)
        self.assertEqual(sheet["D2"].value, -2)
        self.assertIsNone(sheet["E2"].value)
        self.assertEqual(sheet["E3"].value, -8.1)
        self.assertEqual(sheet["G2"].value, "THB")
        self.assertEqual(sheet["G3"].value, "PHP")
        self.assertEqual(len(book["processing"]._images), 1)
        self.assertEqual(book["processing"]._images[0].anchor._from.row, 2)
        self.assertIn("另一个站点", book["processing"]["H3"].value)
        self.assertIn(rows[1]["captured_at"], book["processing"]["H3"].value)
        self.assertIsNone(book["processing"]["H2"].value)

    def test_failed_store_blank_amounts_errors_and_precision_retained(self):
        target = self.root / "partial.xlsx"
        write_balance_workbook("temu", [self.row(status="failed", values={}, evidence={}, message="=危险公式"), self.row(values={"total": "1234567890123456.78", "pending": "0"})], target)
        book = load_workbook(target)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertIsNone(sheet["C2"].value)
        self.assertEqual(sheet["L2"].data_type, "s")
        self.assertEqual(sheet["I2"].value, "失败")
        self.assertEqual(sheet["C3"].value, "1234567890123456.78")
        self.assertEqual(sheet["C3"].data_type, "s")

    def test_atomic_export_failure_keeps_previous_workbook(self):
        target = self.root / "keep.xlsx"
        target.write_bytes(b"existing workbook")
        with patch("backend.services.balance_workbook.Workbook.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                write_balance_workbook("temu", [self.row()], target)
        self.assertEqual(target.read_bytes(), b"existing workbook")
        self.assertEqual(list(self.root.glob(".*.xlsx")), [])


if __name__ == "__main__":
    unittest.main()
