from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from backend.services.sku_inventory_query import (
    OUTPUT_COLUMNS,
    DeveloperOption,
    build_stock_query_payload,
    export_records,
    flatten_stock_warehouse_rows,
    parse_pagination_state,
    parse_developer_options,
    query_inventory_records,
    row_matches_created_range,
    SkuInventoryQuery,
    validate_query_payload,
    warehouse_unshipped_count,
)


class SkuInventoryQueryTests(unittest.TestCase):
    def test_parse_developers_only_keeps_names_ending_in_development(self) -> None:
        html = """
        <label><input type="checkbox" name="developerIdIN[]" value="all">全部开发员</label>
        <label><input type="checkbox" name="developerIdIN[]" value="101">张三-开发</label>
        <label><input type="checkbox" name="developerIdIN[]" value="102">李四-欧美开发 &gt; 9</label>
        <label><input type="checkbox" name="developerIdIN[]" value="103">赵六-开发助理</label>
        <label><input type="checkbox" name="developerIdIN[]" value="104">王五-采购</label>
        <label><input type="checkbox" name="buyerIdIN[]" value="105">陈七-开发</label>
        """

        options = parse_developer_options(html)

        self.assertEqual(
            [(option.id, option.name) for option in options],
            [("101", "张三-开发"), ("102", "李四-欧美开发")],
        )

    def test_validate_query_payload_normalizes_developers_and_page_size(self) -> None:
        job = validate_query_payload(
            {
                "username": "user",
                "password": "secret",
                "developer_ids": ["101", "101", " 102 "],
                "start_date": "2026-08-01",
                "end_date": "2026-08-17",
                "output_dir": "D:/output",
                "rows_per_page": 5000,
            }
        )

        self.assertEqual(job.developer_ids, ("101", "102"))
        self.assertEqual(job.rows_per_page, 1000)

    def test_validate_query_payload_requires_a_developer(self) -> None:
        with self.assertRaisesRegex(ValueError, "至少选择一名开发人员"):
            validate_query_payload(
                {
                    "username": "user",
                    "password": "secret",
                    "developer_ids": [],
                    "start_date": "2026-08-01",
                    "end_date": "2026-08-17",
                    "output_dir": "D:/output",
                }
            )

    def test_stock_payload_uses_created_time_developer_and_all_warehouses(self) -> None:
        job = validate_query_payload(
            {
                "username": "user",
                "password": "secret",
                "developer_ids": ["101"],
                "start_date": "2026-08-01",
                "end_date": "2026-08-17",
                "output_dir": "D:/output",
            }
        )

        payload = dict(build_stock_query_payload(job, developer_id="101", page=3))

        self.assertEqual(payload["developerIdM"], "101")
        self.assertEqual(payload["queryTime"], "timeCreated")
        self.assertEqual(payload["timeCreatedStartTime"], "2026-08-01 00:00:00")
        self.assertEqual(payload["timeCreatedEndTime"], "2026-08-17 23:59:59")
        self.assertEqual(payload["warehouseMold"], "all")
        self.assertEqual(payload["defaultStockWarehouseDetailId"], "")
        self.assertEqual(payload["page"], "3")

    def test_parse_pagination_state_reads_current_and_total_pages(self) -> None:
        current_page, total_pages = parse_pagination_state(
            {
                "pageHtml": (
                    '<div class="static-text">共<span>2174</span>条 '
                    '<span class="semibold">5/5</span>页</div>'
                )
            }
        )

        self.assertEqual((current_page, total_pages), (5, 5))

    def test_query_stops_on_full_last_page_reported_by_mabang(self) -> None:
        job = SkuInventoryQuery(
            username="user",
            password="secret",
            developer_ids=("101",),
            start_date="2026-08-01",
            end_date="2026-08-17",
            output_dir=Path("D:/output"),
            rows_per_page=2,
        )
        response = {
            "success": True,
            "stockData": [{"stockId": "1"}, {"stockId": "2"}],
            "pageHtml": '<span class="semibold">1/1</span>页',
        }
        fake_client = unittest.mock.MagicMock()
        fake_client.__enter__.return_value = fake_client
        fake_client.__exit__.return_value = None

        with (
            patch("backend.services.sku_inventory_query.MabangClient", return_value=fake_client),
            patch(
                "backend.services.sku_inventory_query.open_stock_page",
                return_value=(
                    "https://example.test/stock",
                    '<label><input type="checkbox" name="developerIdIN[]" value="101">张三-开发</label>',
                ),
            ),
            patch("backend.services.sku_inventory_query.post_form_json", return_value=response) as post,
        ):
            records, developers = query_inventory_records(job)

        self.assertEqual(records, [])
        self.assertEqual([developer.id for developer in developers], ["101"])
        self.assertEqual(post.call_count, 1)

    def test_query_stops_when_mabang_repeats_last_full_page(self) -> None:
        job = SkuInventoryQuery(
            username="user",
            password="secret",
            developer_ids=("101",),
            start_date="2026-08-01",
            end_date="2026-08-17",
            output_dir=Path("D:/output"),
            rows_per_page=2,
        )
        responses = [
            {"success": True, "stockData": [{"stockId": "1"}, {"stockId": "2"}]},
            {"success": True, "stockData": [{"stockId": "3"}, {"stockId": "4"}]},
            {"success": True, "stockData": [{"stockId": "3"}, {"stockId": "4"}]},
        ]
        fake_client = unittest.mock.MagicMock()
        fake_client.__enter__.return_value = fake_client
        fake_client.__exit__.return_value = None

        with (
            patch("backend.services.sku_inventory_query.MabangClient", return_value=fake_client),
            patch(
                "backend.services.sku_inventory_query.open_stock_page",
                return_value=(
                    "https://example.test/stock",
                    '<label><input type="checkbox" name="developerIdIN[]" value="101">张三-开发</label>',
                ),
            ),
            patch("backend.services.sku_inventory_query.post_form_json", side_effect=responses) as post,
        ):
            records, developers = query_inventory_records(job)

        self.assertEqual(records, [])
        self.assertEqual([developer.id for developer in developers], ["101"])
        self.assertEqual(post.call_count, 3)

    def test_flatten_outputs_sku_warehouse_rows_with_positive_available_inventory(self) -> None:
        stock_row = {
            "stockSku": "SKU-001",
            "nameCN": "测试商品",
            "timeCreatedShowTimezone": "2026-08-03 10:20:30",
            "predictDailySales": "2.5",
            "stockWarehouseData": [
                {
                    "warehouseId": "W1",
                    "name": "一号仓",
                    "stockWarehouseAvailableInventory": "12",
                    "salesDays": "0",
                    "unshipped": "7",
                    "fbaWaitingQuantity": "2",
                    "transferred_warehouse_quantity": "3",
                    "manual_outbound_quantity": "4",
                },
                {
                    "warehouseId": "W2",
                    "name": "二号仓",
                    "availabledInventory": "7.5",
                    "daysales": "0",
                    "salesDays": "",
                },
                {
                    "warehouseId": "W3",
                    "name": "零库存仓",
                    "stockWarehouseAvailableInventory": "0",
                    "salesDays": "20",
                },
            ],
        }

        records = flatten_stock_warehouse_rows(
            stock_row,
            developer=DeveloperOption(id="101", name="张三-开发"),
        )

        self.assertEqual(len(records), 2)
        self.assertEqual(records[0]["developer_name"], "张三-开发")
        self.assertEqual(records[0]["sku"], "SKU-001")
        self.assertEqual(records[0]["warehouse_name"], "一号仓")
        self.assertEqual(records[0]["available_inventory"], 12)
        self.assertEqual(records[0]["forecast_daily_sales"], 2.5)
        self.assertEqual(records[0]["current_sales_days"], 0)
        self.assertEqual(records[0]["unshipped_count"], 16)
        self.assertEqual(records[0]["transit_inventory"], 0)
        self.assertEqual(records[1]["available_inventory"], 7.5)
        self.assertEqual(records[1]["forecast_daily_sales"], 0)
        self.assertEqual(records[1]["current_sales_days"], "--")
        self.assertEqual(records[1]["unshipped_count"], 0)

    def test_flatten_keeps_any_positive_inventory_unshipped_or_transit_quantity(self) -> None:
        cases = [
            ("positive_inventory", {"stockWarehouseAvailableInventory": "12"}, (12, 0, 0)),
            ("zero_inventory_unshipped", {"unshipped": "7"}, (0, 7, 0)),
            (
                "negative_inventory_unshipped",
                {"stockWarehouseAvailableInventory": "-3", "unshipped": "7"},
                (-3, 7, 0),
            ),
            ("fba_unshipped", {"fbaWaitingQuantity": "2"}, (0, 2, 0)),
            ("transfer_unshipped", {"transferred_warehouse_quantity": "3"}, (0, 3, 0)),
            ("manual_unshipped", {"manual_outbound_quantity": "4"}, (0, 4, 0)),
            ("zero_inventory_transit_100", {"allotShippingQuantity": "100"}, (0, 0, 100)),
            ("purchase_transit", {"shippingQuantity": "100"}, (0, 0, 100)),
            ("processing_transit", {"processingQuantity": "3"}, (0, 0, 3)),
            ("overseas_transit", {"hwc_in_transit_quantity": "4"}, (0, 0, 4)),
            ("manual_inbound_transit", {"manual_inbound_quantity": "5"}, (0, 0, 5)),
            (
                "five_part_transit_total_is_separate_from_unshipped",
                {
                    "shippingQuantity": "100",
                    "processingQuantity": "2",
                    "allotShippingQuantity": "3",
                    "hwc_in_transit_quantity": "4",
                    "manual_inbound_quantity": "5",
                    "transferred_warehouse_quantity": "7",
                },
                (0, 7, 114),
            ),
            (
                "negative_inventory_transit",
                {"stockWarehouseAvailableInventory": "-3", "allotShippingQuantity": "100"},
                (-3, 0, 100),
            ),
            ("formatted_transit", {"allotShippingQuantity": "1,200.5"}, (0, 0, 1200.5)),
            ("all_zero", {"unshipped": "0", "allotShippingQuantity": "0"}, None),
            ("negative_inventory_only", {"stockWarehouseAvailableInventory": "-3"}, None),
            ("missing_quantities", {}, None),
            ("processing_is_not_unshipped", {"processWaitingQuantity": "1000"}, None),
        ]
        for label, quantities, expected in cases:
            with self.subTest(label=label):
                records = flatten_stock_warehouse_rows(
                    {
                        "stockSku": "SKU-001",
                        "stockWarehouseData": [
                            {"warehouseId": "W1", "stockWarehouseAvailableInventory": "0", **quantities}
                        ],
                    },
                    developer=DeveloperOption(id="101", name="张三-开发"),
                )
                if expected is None:
                    self.assertEqual(records, [])
                else:
                    self.assertEqual(len(records), 1)
                    self.assertEqual(
                        tuple(records[0][key] for key in ("available_inventory", "unshipped_count", "transit_inventory")),
                        expected,
                    )

    def test_flatten_uses_each_warehouses_quantities_without_stock_level_fallback(self) -> None:
        records = flatten_stock_warehouse_rows(
            {
                "stockSku": "SKU-001",
                "stockWarehouseAvailableInventory": "999",
                "unshipped": "999",
                "shippingQuantity": "999",
                "processingQuantity": "999",
                "allotShippingQuantity": "999",
                "hwc_in_transit_quantity": "999",
                "manual_inbound_quantity": "999",
                "stockWarehouseData": [
                    {"warehouseId": "stock", "stockWarehouseAvailableInventory": "5"},
                    {"warehouseId": "unshipped", "unshipped": "7", "transferred_warehouse_quantity": "3"},
                    {"warehouseId": "transit", "shippingQuantity": "100"},
                    {
                        "warehouseId": "all",
                        "stockWarehouseAvailableInventory": "2",
                        "unshipped": "4",
                        "shippingQuantity": "2",
                        "processingQuantity": "3",
                        "allotShippingQuantity": "6",
                        "hwc_in_transit_quantity": "4",
                        "manual_inbound_quantity": "5",
                    },
                    {"warehouseId": "empty"},
                ],
            },
            developer=DeveloperOption(id="101", name="张三-开发"),
        )

        self.assertEqual(
            [
                (record["warehouse_id"], record["available_inventory"], record["unshipped_count"], record["transit_inventory"])
                for record in records
            ],
            [("stock", 5, 0, 0), ("unshipped", 0, 10, 0), ("transit", 0, 0, 100), ("all", 2, 4, 20)],
        )

    def test_unshipped_count_matches_mabang_four_part_total(self) -> None:
        self.assertEqual(
            warehouse_unshipped_count(
                {
                    "unshipped": "7",
                    "waitingQuantity": "99",
                    "fbaWaitingQuantity": "2",
                    "transferred_warehouse_quantity": "3",
                    "manual_outbound_quantity": "4",
                    "processWaitingQuantity": "1000",
                }
            ),
            16,
        )
        self.assertEqual(warehouse_unshipped_count({"waitingQuantity": "1,200"}), 1200)

    def test_created_time_range_is_checked_locally_when_field_exists(self) -> None:
        job = validate_query_payload(
            {
                "username": "user",
                "password": "secret",
                "developer_ids": ["101"],
                "start_date": "2026-08-01",
                "end_date": "2026-08-17",
                "output_dir": "D:/output",
            }
        )

        self.assertTrue(row_matches_created_range({"timeCreated": "2026-08-17 23:59:59"}, job))
        self.assertFalse(row_matches_created_range({"timeCreated": "2026-07-31 23:59:59"}, job))
        self.assertTrue(row_matches_created_range({}, job))

    def test_export_writes_complete_sku_warehouse_detail_sheet(self) -> None:
        records = [
            {
                "developer_name": "张三-开发",
                "sku": "SKU-001",
                "product_name": "测试商品",
                "created_at": "2026-08-03 10:20:30",
                "warehouse_name": "一号仓",
                "warehouse_id": "W1",
                "available_inventory": 12,
                "forecast_daily_sales": 2.5,
                "current_sales_days": 0,
                "unshipped_count": 16,
                "transit_inventory": 100,
            }
        ]
        with tempfile.TemporaryDirectory() as temporary_dir:
            output_file = export_records(records, Path(temporary_dir))
            workbook = load_workbook(output_file, data_only=True)
            try:
                worksheet = workbook["SKU仓库明细"]
                rows = list(worksheet.iter_rows(values_only=True))
                auto_filter_ref = worksheet.auto_filter.ref
            finally:
                workbook.close()

        self.assertEqual(rows[0], OUTPUT_COLUMNS)
        self.assertEqual(rows[0][-1], "在途量")
        self.assertEqual(rows[1][0:6], ("张三-开发", "SKU-001", "测试商品", "2026-08-03 10:20:30", "一号仓", "W1"))
        self.assertEqual(rows[1][6:], (12, 2.5, 0, 16, 100))
        self.assertEqual(auto_filter_ref, "A1:K2")


if __name__ == "__main__":
    unittest.main()
