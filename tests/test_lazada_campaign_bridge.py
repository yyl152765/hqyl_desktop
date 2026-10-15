from __future__ import annotations

import json
import tempfile
import time
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook

from backend.app_bridge import AppBridge
from backend.config_store import AppSettings


def write_prices(path: Path, *, recommended=False) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    sheet = book.active
    sheet.append(["SKU", "商品名称", "原价", "活动价", "推荐价"])
    for sku, base, price in (("0001", 100, 95), ("0002", 100, 94.99), ("0003", 0, 50)):
        sheet.append([sku, "商品", base, None if recommended else price, f"<={price}"])
    book.save(path)
    book.close()
    return path


class CampaignBridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = write_prices(self.root / "商品.xlsx")
        self.bridge = AppBridge()
        self.bridge.config_store = Mock()
        self.bridge.config_store.load.return_value = AppSettings(output_dir=str(self.root / "输出"))
        self.bridge._payload_with_account = Mock(side_effect=AssertionError("本地工具不能要求账号"))

    def payload(self, **overrides):
        return {"source_files": [str(self.source)], "mode": "plain", "threshold": "0.95", **overrides}

    def wait_task(self, task):
        deadline = time.monotonic() + 10
        while task["status"] in {"pending", "running"} and time.monotonic() < deadline:
            time.sleep(.02)
            task = self.bridge.get_task_status(task["id"])
        return task

    def assert_lock_released(self):
        self.assertTrue(self.bridge._campaign_operation_lock.acquire(blocking=False))
        self.bridge._campaign_operation_lock.release()

    def test_info_is_local_without_account_or_browser_resolution(self):
        with patch("backend.app_bridge.resolve_lazada_monthly_report_runtime_paths", side_effect=AssertionError("不访问紫鸟")):
            info = self.bridge.get_lazada_campaign_workbook_info()
        self.assertEqual(info["workflow"], "local_excel")
        self.assertEqual(info["max_files"], 100)
        self.assertNotIn("client_path", info)
        self.assertEqual([mode["value"] for mode in info["modes"]], ["auto", "plain", "invited"])

    def test_file_picker_is_multiple_and_cancel_is_explicit(self):
        window = Mock()
        window.create_file_dialog.return_value = (str(self.source), str(self.root / "第二.xlsx"))
        with patch.dict("sys.modules", {"webview": SimpleNamespace(OPEN_DIALOG=1, windows=[window])}):
            selected = self.bridge.choose_lazada_campaign_files()
            self.assertTrue(selected["ok"])
            self.assertEqual(len(selected["paths"]), 2)
            self.assertTrue(window.create_file_dialog.call_args.kwargs["allow_multiple"])
            window.create_file_dialog.return_value = None
            self.assertTrue(self.bridge.choose_lazada_campaign_files()["cancelled"])
            window.create_file_dialog.side_effect = RuntimeError("internal detail")
            self.assertNotIn("internal detail", self.bridge.choose_lazada_campaign_files()["error"])

    def test_inspection_reads_local_files_and_reports_partial_invalid_input(self):
        corrupt = self.root / "损坏.xlsx"
        corrupt.write_bytes(b"not an Excel")
        original = self.source.read_bytes()
        result = self.bridge.inspect_lazada_campaign_files(self.payload(source_files=[str(self.source), str(corrupt)]))
        self.assertTrue(result["ok"], result)
        self.assertEqual((result["success_count"], result["failed_count"]), (1, 1))
        self.assertEqual(result["files"][0]["mode"], "plain")
        self.assertEqual(self.source.read_bytes(), original)

    def test_export_blank_template_without_accounts_and_validate_directory(self):
        result = self.bridge.export_lazada_campaign_template()
        self.assertTrue(result["ok"], result)
        book = load_workbook(result["output_file"])
        try:
            self.assertEqual([cell.value for cell in book.active[1]], ["SKU", "商品名称", "原价", "活动价", "推荐价"])
            self.assertEqual(book.active.max_row, 1)
        finally:
            book.close()
        self.assertFalse(self.bridge.export_lazada_campaign_template({"output_dir": ""})["ok"])
        self.assertFalse(self.bridge.export_lazada_campaign_template([])["ok"])

    def test_real_batch_no_account_exact_threshold_restore_and_source_unchanged(self):
        original = self.source.read_bytes()
        task = self.bridge.start_lazada_campaign_workbook(self.payload(password="ignored"))
        self.assertTrue(task["ok"], task)
        task = self.wait_task(task)
        self.assertEqual(task["status"], "success", task)
        result = task["result"]
        self.assertEqual((result["kept_rows"], result["removed_rows"], result["unparseable_rows"]), (1, 1, 1))
        self.assertEqual(result["updated_rows"], 0)
        book = load_workbook(result["files"][0]["processed_file"])
        try:
            self.assertEqual(book.active["A2"].value, "0001")
            self.assertEqual(book.active["D2"].value, 95)
            self.assertEqual(book.active.max_row, 2)
        finally:
            book.close()
        self.assertEqual(self.source.read_bytes(), original)
        context = self.bridge.get_latest_task_status("lazada_campaign_workbook")["context"]
        self.assertEqual(context["workflow"], "local_excel")
        self.assertEqual(context["source_files"], [str(self.source)])
        self.assertNotIn("ignored", json.dumps(task))
        self.assertNotIn("account_id", context)
        self.bridge._payload_with_account.assert_not_called()
        self.assert_lock_released()

    def test_busy_invalid_or_old_download_payload_never_starts(self):
        self.bridge.tasks = Mock()
        self.bridge._campaign_operation_lock.acquire()
        try:
            self.assertFalse(self.bridge.start_lazada_campaign_workbook(self.payload())["ok"])
        finally:
            self.bridge._campaign_operation_lock.release()
        for payload in (self.payload(threshold="NaN"), {"store_names": ["old"], "campaign_url": "old"}, []):
            self.assertFalse(self.bridge.start_lazada_campaign_workbook(payload)["ok"])
        self.bridge.tasks.start.assert_not_called()

    def test_failed_runner_releases_operation_lock(self):
        with patch("backend.app_bridge.lazada_campaign_workbook.run_campaign_batch", side_effect=OSError("cannot write")):
            task = self.wait_task(self.bridge.start_lazada_campaign_workbook(self.payload()))
        self.assertEqual(task["status"], "failed", task)
        self.assert_lock_released()

    def test_startup_failures_release_lock_and_hide_internal_errors(self):
        self.bridge.tasks = Mock()
        self.bridge.tasks.start.side_effect = RuntimeError("private detail")
        result = self.bridge.start_lazada_campaign_workbook(self.payload())
        self.assertFalse(result["ok"])
        self.assertNotIn("private detail", result["error"])
        self.assert_lock_released()
        self.bridge.tasks.start.side_effect = None
        self.bridge.tasks.start.return_value = {"ok": False, "error": "not started"}
        self.assertFalse(self.bridge.start_lazada_campaign_workbook(self.payload())["ok"])
        self.assert_lock_released()


if __name__ == "__main__":
    unittest.main()
