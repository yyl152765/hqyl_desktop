from __future__ import annotations

import unittest
from contextlib import ExitStack
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from launcher.main import resolve_entry_html, self_check


class LauncherPageTests(unittest.TestCase):
    def test_bigseller_claim_query_has_development_entry(self) -> None:
        self.assertEqual(resolve_entry_html(["--page=bigseller-claim-query"]).name, "bigseller-claim-query.html")

    def test_bigseller_sku_benchmark_has_development_entry(self) -> None:
        self.assertEqual(resolve_entry_html(["--page=bigseller-sku-benchmark"]).name, "bigseller-sku-benchmark.html")

    def test_default_page_remains_dashboard(self) -> None:
        self.assertEqual(resolve_entry_html([]).name, "dashboard.html")

    def test_vietnam_income_page_has_separate_development_entry(self) -> None:
        page = resolve_entry_html(["--page=vietnam-income-reconciliation"])
        self.assertEqual(page.name, "vietnam-income-reconciliation.html")

    def test_mabang_warehouse_permission_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=mabang-warehouse-permission"])
        self.assertEqual(page.name, "mabang-warehouse-permission.html")

    def test_mabang_arrival_query_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=mabang-arrival-query"])
        self.assertEqual(page.name, "mabang-arrival-query.html")

    def test_sku_inventory_query_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=sku-inventory-query"])
        self.assertEqual(page.name, "sku-inventory-query.html")

    def test_mabang_income_expense_report_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=mabang-income-expense-report"])
        self.assertEqual(page.name, "mabang-income-expense-report.html")

    def test_temu_shipping_channel_has_development_entry_and_menu(self) -> None:
        page = resolve_entry_html(["--page=temu-shipping-channel"])
        self.assertEqual(page.name, "temu-shipping-channel.html")
        frontend = Path(__file__).resolve().parents[1] / "frontend" / "assets"
        self.assertIn('id: "temu_shipping_channel", href: "temu-shipping-channel.html"', (frontend / "common.js").read_text(encoding="utf-8"))
        self.assertIn('temu_shipping_channel: { href: "temu-shipping-channel.html"', (frontend / "pages" / "dashboard.js").read_text(encoding="utf-8"))

    def test_developer_sales_income_summary_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=developer-sales-income-summary"])
        self.assertEqual(page.name, "developer-sales-income-summary.html")

    def test_bigseller_item_id_query_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=bigseller-item-id-query"])
        self.assertEqual(page.name, "bigseller-item-id-query.html")

    def test_lazada_monthly_report_has_development_entry(self) -> None:
        page = resolve_entry_html(["--page=lazada-monthly-report"])
        self.assertEqual(page.name, "lazada-monthly-report.html")

    def test_sidebar_contains_vietnam_income_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("vietnam-income-reconciliation.html", source)
        self.assertIn("越南收支报表", source)

    def test_sidebar_contains_mabang_warehouse_permission_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("mabang-warehouse-permission.html", source)
        self.assertIn("仓库权限开通", source)

    def test_sidebar_contains_mabang_arrival_query_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("mabang-arrival-query.html", source)
        self.assertIn("到货查询", source)

    def test_sidebar_contains_sku_inventory_query_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("sku-inventory-query.html", source)
        self.assertIn("SKU库存与可售天数", source)

    def test_sidebar_contains_mabang_income_expense_report_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("mabang-income-expense-report.html", source)
        self.assertIn("收支报表", source)

    def test_sidebar_contains_developer_sales_income_summary_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("developer-sales-income-summary.html", source)
        self.assertIn("开发与销售收入汇总", source)

    def test_sidebar_contains_bigseller_item_id_query_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("bigseller-item-id-query.html", source)
        self.assertIn("商品ID查询", source)

    def test_sidebar_contains_lazada_monthly_report_entry(self) -> None:
        common_js = Path(__file__).resolve().parents[1] / "frontend" / "assets" / "common.js"
        source = common_js.read_text(encoding="utf-8")
        self.assertIn("lazada-monthly-report.html", source)
        self.assertIn("月度账单", source)

    def test_kec_reconciliation_has_development_entry_and_menu(self) -> None:
        page = resolve_entry_html(["--page=kec-reconciliation"])
        self.assertEqual(page.name, "kec-reconciliation.html")
        frontend = Path(__file__).resolve().parents[1] / "frontend"
        common_js = (frontend / "assets" / "common.js").read_text(encoding="utf-8")
        self.assertIn('id: "kec_reconciliation", href: "kec-reconciliation.html"', common_js)
        self.assertIn("KEC 对账", common_js)
        dashboard = (frontend / "assets" / "pages" / "dashboard.js").read_text(encoding="utf-8")
        self.assertIn('kec_reconciliation: { href: "kec-reconciliation.html"', dashboard)
        self.assertIn('"key": "kec_reconciliation"', (Path(__file__).resolve().parents[1] / "backend" / "app_bridge.py").read_text(encoding="utf-8"))

    def test_unknown_page_is_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "unknown desktop page"):
            resolve_entry_html(["--page=not-a-real-page"])


class LauncherSelfCheckTests(unittest.TestCase):
    def setUp(self) -> None:
        self.stack = ExitStack()
        self.addCleanup(self.stack.close)
        # All real configuration/network access is forbidden in these checks.
        for target in (
            "backend.config_store.ConfigStore.load",
            "backend.config_store.ConfigStore.save",
            "backend.services.group_sales_report._load_groups_from_mabang_config",
            "backend.services.sample_registration_config.config_path",
            "backend.services.bigseller_sync.load_reference_bigseller_config",
            "backend.services.bigseller_item_id_query.run_bigseller_item_id_query",
            "backend.services.bigseller_sku_benchmark.run_bigseller_sku_benchmark",
            "backend.services.bigseller_claim_query.run_bigseller_claim_query",
            "requests.sessions.Session.request",
        ):
            self.stack.enter_context(patch(target, side_effect=AssertionError(f"offline self-check called {target}")))
        self.stack.enter_context(patch(
            "backend.services.shopee_ads_recharge.get_site_module",
            return_value=SimpleNamespace(run_ads_recharge=lambda: None),
        ))

    def test_self_check_loads_bigseller_and_roundtrips_excel_without_user_data(self) -> None:
        from backend import app_bridge
        from backend.services import bigseller_item_id_query, bigseller_sync

        original_group_options = app_bridge.get_group_options
        util = bigseller_sync.load_bigseller_request_util()
        with patch.object(util, "login_bigseller_by_request", side_effect=AssertionError("must not log in")) as login, patch.object(
            bigseller_sync, "load_bigseller_request_util", wraps=bigseller_sync.load_bigseller_request_util,
        ) as loader, patch.object(
            bigseller_item_id_query, "export_item_ids", wraps=bigseller_item_id_query.export_item_ids,
        ) as export:
            self.assertEqual(self_check(), 0)
        loader.assert_called_once_with()
        login.assert_not_called()
        export.assert_called_once()
        temp_dir = export.call_args.args[0]
        self.assertTrue(temp_dir.name.startswith("hqyl-self-check-"))
        self.assertFalse(temp_dir.exists())
        self.assertIs(app_bridge.get_group_options, original_group_options)

    def test_self_check_rejects_missing_bigseller_export(self) -> None:
        from backend.services.bigseller_sync import load_bigseller_request_util

        util = load_bigseller_request_util()
        with patch.object(util, "build_api_url", None):
            with self.assertRaisesRegex(RuntimeError, "BigSeller.*build_api_url"):
                self_check()

    def test_self_check_roundtrips_benchmark_checkpoint_and_both_metrics(self) -> None:
        from backend.services import bigseller_benchmark_checkpoint, bigseller_benchmark_workbook

        with patch.object(
            bigseller_benchmark_checkpoint, "read_checkpoint", wraps=bigseller_benchmark_checkpoint.read_checkpoint,
        ) as read, patch.object(
            bigseller_benchmark_workbook, "export_benchmark_workbook",
            wraps=bigseller_benchmark_workbook.export_benchmark_workbook,
        ) as export:
            self.assertEqual(self_check(), 0)
        self.assertEqual(read.call_count, 2)
        self.assertEqual([call.args[2] for call in export.call_args_list], ["views", "sales"])
        for call in export.call_args_list:
            self.assertFalse(call.kwargs["summary"]["is_complete"])
            self.assertFalse(call.args[3].exists())

    def test_self_check_rejects_missing_recovered_benchmark_rows(self) -> None:
        with patch("backend.services.bigseller_benchmark_checkpoint.read_checkpoint", return_value={}):
            with self.assertRaisesRegex(RuntimeError, "benchmark checkpoint recovery"):
                self_check()

    def test_self_check_roundtrips_claim_checkpoint_and_preserves_multi_sku_workbook(self) -> None:
        from backend.services import (
            bigseller_claim_checkpoint,
            bigseller_claim_workbook,
        )

        with patch.object(
            bigseller_claim_checkpoint, "read_claim_checkpoint", wraps=bigseller_claim_checkpoint.read_claim_checkpoint,
        ) as read, patch.object(
            bigseller_claim_workbook, "load_claim_input", wraps=bigseller_claim_workbook.load_claim_input,
        ) as load, patch.object(
            bigseller_claim_workbook, "export_claim_workbook", wraps=bigseller_claim_workbook.export_claim_workbook,
        ) as export:
            self.assertEqual(self_check(), 0)
        read.assert_called_once()
        load.assert_called_once()
        export.assert_called_once()
        self.assertEqual(export.call_args.args[0].row_numbers, {"00001": (2, 3), "00002": (3,)})
        self.assertFalse(export.call_args.args[2].exists())

    def test_self_check_rejects_missing_recovered_claim_rows(self) -> None:
        with (
            patch("backend.services.bigseller_claim_checkpoint.read_claim_checkpoint", return_value=({}, 0.0)),
            self.assertRaisesRegex(RuntimeError, "claim checkpoint recovery"),
        ):
            self_check()

    def test_self_check_rejects_benchmark_export_with_missing_retry_row(self) -> None:
        from openpyxl import load_workbook
        from backend.services import bigseller_benchmark_workbook

        real_export = bigseller_benchmark_workbook.export_benchmark_workbook

        def corrupt_export(*args, **kwargs):
            output_file = real_export(*args, **kwargs)
            workbook = load_workbook(output_file)
            try:
                workbook["BigSeller浏览量待补查"].delete_rows(2)
                workbook.save(output_file)
            finally:
                workbook.close()
            return output_file

        with patch.object(bigseller_benchmark_workbook, "export_benchmark_workbook", side_effect=corrupt_export):
            with self.assertRaisesRegex(RuntimeError, "benchmark Excel integrity"):
                self_check()

    def test_self_check_rejects_query_contract_drift(self) -> None:
        from backend.services import bigseller_item_id_query

        valid_query = bigseller_item_id_query.build_listing_query("self-check SKU")
        for key, value in (("searchType", "parentSku"), ("inquireType", False), ("orderBy", "update_time"), ("desc", "true")):
            with self.subTest(key=key), patch.object(
                bigseller_item_id_query, "build_listing_query", return_value={**valid_query, key: value},
            ):
                with self.assertRaisesRegex(RuntimeError, "BigSeller Item ID query contract"):
                    self_check()

    def test_self_check_rejects_extra_excel_columns_and_formula_cells(self) -> None:
        from openpyxl import load_workbook
        from backend.services import bigseller_item_id_query

        real_export = bigseller_item_id_query.export_item_ids
        for corruption, expected_error in (("extra_column", "only SKU / Item ID"), ("formula", "text preservation")):
            def corrupt_export(output_dir, rows):
                output_file = real_export(output_dir, rows)
                workbook = load_workbook(output_file)
                try:
                    if corruption == "extra_column":
                        workbook.active["C1"] = "Unexpected"
                    else:
                        workbook.active["A2"] = "=1+1"
                    workbook.save(output_file)
                finally:
                    workbook.close()
                return output_file

            with self.subTest(corruption=corruption), patch.object(
                bigseller_item_id_query, "export_item_ids", side_effect=corrupt_export,
            ):
                with self.assertRaisesRegex(RuntimeError, expected_error):
                    self_check()


if __name__ == "__main__":
    unittest.main()
