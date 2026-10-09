from __future__ import annotations

import hashlib
import tempfile
import unittest
from datetime import date
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel

from backend.services.temu_shipping_workbook import read_temu_shipping_workbook


NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TODAY = date(2026, 9, 10)


class TemuShippingWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.path = Path(self.temp_dir.name) / "每日尺寸.xlsx"

    def write_book(self, rows=(), *, headers=None, title="Sheet1", header_row=1):
        book = Workbook()
        sheet = book.active
        sheet.title = title
        for _ in range(header_row - 1):
            sheet.append(["导入说明"])
        sheet.append(headers or ["长", "宽", "高", "尺寸比", "定量", "重量"])
        for row in rows:
            sheet.append(row)
        book.save(self.path)
        book.close()

    def patch_caches(self, values):
        """Add Excel-style formula caches without relying on an Excel installation."""
        with ZipFile(self.path) as source:
            members = {name: source.read(name) for name in source.namelist()}
        for sheet_id, caches in values.items():
            member = f"xl/worksheets/sheet{sheet_id}.xml"
            xml = ET.fromstring(members[member])
            for cell in xml.findall(".//m:c", NS):
                if cell.attrib["r"] not in caches:
                    continue
                cached = cell.find("m:v", NS)
                if cached is None:
                    cached = ET.SubElement(cell, f"{{{NS['m']}}}v")
                value = caches[cell.attrib["r"]]
                if isinstance(value, tuple):
                    value, cell_type = value
                    cell.set("t", cell_type)
                cached.text = str(value)
            members[member] = ET.tostring(xml, encoding="utf-8")
        with ZipFile(self.path, "w", ZIP_DEFLATED) as target:
            for name, content in members.items():
                target.writestr(name, content)

    def daily_book(self, *, cached_day=TODAY, with_caches=True, date_formula="=TODAY()", extra_dates=()):
        book = Workbook()
        sheet = book.active
        sheet.title = "Sheet1"
        sheet.append(["长", "宽", "高", "尺寸比", "定量", "重量"])
        for row in (2, 3):
            for col in "ABC":
                sheet[f"{col}{row}"] = f"=INDEX('每日尺寸配置'!${col}$6:${col}$261,MOD('每日尺寸配置'!$B$1+ROW()*37+178,256)+1)"
            sheet[f"F{row}"] = f"=A{row}*B{row}*C{row}/6"
        config = book.create_sheet("每日尺寸配置")
        config["A1"] = "更新日期"
        config["B1"] = date_formula
        config["B1"].number_format = "yyyy-mm-dd"
        for index, _day in enumerate(extra_dates, 2):
            config[f"B{index}"] = "=TODAY()"
            config[f"B{index}"].number_format = "yyyy-mm-dd"
        book.save(self.path)
        book.close()
        if with_caches:
            patches = {1: {"A2": 17, "B2": 12, "C2": 10, "F2": 340,
                           "A3": 16, "B3": 15, "C3": 11, "F3": 440}}
            if cached_day is not None:
                patches[2] = {"B1": to_excel(cached_day)}
            for index, day in enumerate(extra_dates, 2):
                if day is not None:
                    patches.setdefault(2, {})[f"B{index}"] = to_excel(day)
            self.patch_caches(patches)

    def read(self):
        return read_temu_shipping_workbook(self.path, today=TODAY)

    def test_daily_formula_caches_and_source_file_are_preserved(self):
        self.daily_book()
        before = self.path.read_bytes()
        data = self.read()
        self.assertEqual(data["rows"], [
            {"excel_row": 2, "length": "17", "width": "12", "height": "10", "weight": "340.00"},
            {"excel_row": 3, "length": "16", "width": "15", "height": "11", "weight": "440.00"},
        ])
        self.assertEqual(data["calculation_date"], TODAY.isoformat())
        self.assertTrue(data["uses_formula_cache"])
        self.assertEqual(data["source_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(self.path.read_bytes(), before)
        self.assertEqual(data["fingerprint"], self.read()["fingerprint"])

    def test_order_and_duplicates_are_preserved_and_fingerprint_is_value_based(self):
        self.write_book([[17, 12, 10, None, None, 340], [16, 15, 11, None, None, 440],
                         [17, 12, 10, None, None, 340]])
        first = self.read()
        self.assertEqual(first["row_count"], 3)
        self.assertEqual([row["excel_row"] for row in first["rows"]], [2, 3, 4])
        self.assertEqual([row["length"] for row in first["rows"]], ["17", "16", "17"])
        book = load_workbook(self.path)
        book.active.column_dimensions["A"].width = 55
        book.save(self.path)
        book.close()
        formatted = self.read()
        self.assertEqual(first["fingerprint"], formatted["fingerprint"])
        self.write_book([[16, 15, 11, None, None, 440], [17, 12, 10, None, None, 340],
                         [17, 12, 10, None, None, 340]])
        self.assertNotEqual(first["fingerprint"], self.read()["fingerprint"])

    def test_old_future_and_missing_date_caches_do_not_restrict_import(self):
        self.daily_book()
        baseline = self.read()
        for cached_day in (None, date(2026, 9, 9), date(2026, 9, 11)):
            with self.subTest(day=cached_day):
                self.daily_book(cached_day=cached_day)
                before = self.path.read_bytes()
                result = self.read()
                self.assertEqual(result["calculation_date"], cached_day.isoformat() if cached_day else None)
                self.assertEqual(result["import_date"], TODAY.isoformat())
                self.assertEqual(result["rows"], baseline["rows"])
                self.assertEqual(result["fingerprint"], baseline["fingerprint"])
                self.assertEqual(self.path.read_bytes(), before)

    def test_next_day_import_keeps_saved_values_and_reports_actual_dates(self):
        self.daily_book(cached_day=date(2026, 9, 10))
        before = self.path.read_bytes()
        original = read_temu_shipping_workbook(self.path, today="2026-09-10")
        next_day = read_temu_shipping_workbook(self.path, today="2026-09-11")
        self.assertEqual(next_day["calculation_date"], "2026-09-10")
        self.assertEqual(next_day["import_date"], "2026-09-11")
        self.assertEqual(next_day["rows"], original["rows"])
        self.assertEqual(next_day["fingerprint"], original["fingerprint"])
        self.assertEqual(next_day["source_sha256"], hashlib.sha256(before).hexdigest())
        self.assertEqual(self.path.read_bytes(), before)

    def test_composite_or_unparseable_date_formulas_leave_metadata_unknown(self):
        for formula in ("=TODAY()+1", "=IF(TODAY()>0,1,0)", '=TODAY("'):
            with self.subTest(formula=formula):
                self.daily_book(date_formula=formula)
                before = self.path.read_bytes()
                result = self.read()
                self.assertIsNone(result["calculation_date"])
                self.assertEqual(result["row_count"], 2)
                self.assertEqual(result["rows"][0]["weight"], "340.00")
                self.assertEqual(self.path.read_bytes(), before)

    def test_invalid_date_cache_leaves_metadata_unknown_but_keeps_data(self):
        for cached_value in (("not-a-date", "str"), ("#VALUE!", "e"), ("1", "b"), 0):
            with self.subTest(cached_value=cached_value):
                self.daily_book(cached_day=None)
                self.patch_caches({2: {"B1": cached_value}})
                before = self.path.read_bytes()
                result = self.read()
                self.assertIsNone(result["calculation_date"])
                self.assertEqual(result["rows"][0]["length"], "17")
                self.assertEqual(result["rows"][1]["weight"], "440.00")
                self.assertEqual(self.path.read_bytes(), before)

    def test_multiple_date_anchors_only_report_one_unambiguous_date(self):
        for extra_dates, expected in (((TODAY, TODAY), TODAY.isoformat()),
                                      ((date(2026, 9, 11),), None), ((None,), None)):
            with self.subTest(extra_dates=extra_dates):
                self.daily_book(extra_dates=extra_dates)
                before = self.path.read_bytes()
                result = self.read()
                self.assertEqual(result["calculation_date"], expected)
                self.assertEqual(result["row_count"], 2)
                self.assertEqual(self.path.read_bytes(), before)

    def test_actual_dimension_and_weight_formula_caches_are_still_required(self):
        for column, formula in ((0, "=17"), (1, "=12"), (2, "=10"), (5, "=340")):
            with self.subTest(column=column):
                values = [17, 12, 10, None, None, 340]
                values[column] = formula
                self.write_book([values])
                before = self.path.read_bytes()
                coordinate = f"{'ABCDEF'[column]}2"
                with self.assertRaisesRegex(ValueError, f"{coordinate}.*公式没有缓存结果"):
                    self.read()
                self.assertEqual(self.path.read_bytes(), before)

    def test_constant_rows_do_not_require_unrelated_today_cache(self):
        self.write_book([[17, 12, 10, None, None, 340]])
        book = load_workbook(self.path)
        book.create_sheet("其他内容")["A1"] = "=TODAY()"
        book.save(self.path)
        book.close()
        self.assertIsNone(self.read()["calculation_date"])

    def test_middle_blank_or_partial_row_does_not_shift_later_rows(self):
        for bad_row in ([None] * 6, [17, None, 10, None, None, 340]):
            with self.subTest(row=bad_row):
                self.write_book([[17, 12, 10, None, None, 340], bad_row,
                                 [16, 15, 11, None, None, 440]])
                with self.assertRaisesRegex(ValueError, "第 3 行.*缺少数值"):
                    self.read()

    def test_trailing_formatting_and_irrelevant_columns_do_not_add_rows(self):
        self.write_book([[17, 12, 10, None, None, 340], [None, None, None, "说明"]])
        book = load_workbook(self.path)
        book.active["A99"].number_format = "0.00"
        book.save(self.path)
        book.close()
        self.assertEqual(self.read()["row_count"], 1)

    def test_reordered_unit_headers_and_non_default_sheet(self):
        self.write_book([["18.2500", "125.005", "10", "12.50"]],
                        headers=[" 长度（cm） ", "申报重量(g)", "高", "宽（厘米）"],
                        title="每日数据", header_row=3)
        data = self.read()
        self.assertEqual(data["sheet_name"], "每日数据")
        self.assertEqual(data["header_row"], 3)
        self.assertEqual(data["rows"][0], {"excel_row": 4, "length": "18.25", "width": "12.5", "height": "10", "weight": "125.01"})

    def test_sheet1_takes_priority_over_another_matching_sheet(self):
        self.write_book([[17, 12, 10, None, None, 340]], title="其他")
        book = load_workbook(self.path)
        target = book.create_sheet("Sheet1")
        target.append(["长", "宽", "高", "重量"])
        target.append([16, 15, 11, 440])
        book.save(self.path)
        book.close()
        self.assertEqual(self.read()["rows"][0]["length"], "16")

    def test_bad_headers_duplicate_headers_and_no_rows_are_rejected(self):
        for headers, message in ((["A", "B", "C", "D"], "未找到完整表头"),
                                 (["长", "宽", "高", "重量(kg)"], "未找到完整表头"),
                                 (["长", "宽", "高", "重量", "重量"], "表头重复"),
                                 (["长", "宽", "高", "重量"], "没有长宽高和重量数据")):
            with self.subTest(headers=headers):
                self.write_book(headers=headers)
                with self.assertRaisesRegex(ValueError, message):
                    self.read()

    def test_invalid_numbers_are_rejected_with_source_location(self):
        for value in (0, -1, True, "NaN", "Infinity", "-Infinity", "不是数值", "#DIV/0!", "1e999999"):
            with self.subTest(value=value):
                self.write_book([[value, 12, 10, None, None, 340]])
                with self.assertRaisesRegex(ValueError, "第 2 行 A2"):
                    self.read()

    def test_weight_is_rounded_half_up_without_an_upper_limit(self):
        for weight, expected in (("1.005", "1.01"), ("499.994", "499.99"), ("0.005", "0.01"),
                                 (500, "500.00"), (501, "501.00"), ("499.995", "500.00"),
                                 ("499.999", "500.00"), ("1234.567", "1234.57"), (1000, "1000.00")):
            with self.subTest(weight=weight):
                self.write_book([[17, 12, 10, None, None, weight]])
                self.assertEqual(self.read()["rows"][0]["weight"], expected)
        self.write_book([[17, 12, 10, None, None, "0.004"]])
        with self.assertRaisesRegex(ValueError, "F2.*四舍五入到两位小数后必须大于 0"):
            self.read()

    def test_only_xlsx_and_readable_existing_files_are_accepted(self):
        with self.assertRaisesRegex(ValueError, "仅支持 .xlsx"):
            read_temu_shipping_workbook(self.path.with_suffix(".xls"))
        with self.assertRaisesRegex(ValueError, "存在的 Excel"):
            self.read()
        self.path.write_bytes(b"this is not an xlsx")
        with self.assertRaisesRegex(ValueError, "无法读取 Excel"):
            self.read()


if __name__ == "__main__":
    unittest.main()
