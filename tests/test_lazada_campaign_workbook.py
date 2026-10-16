from __future__ import annotations

import json
import tempfile
import unittest
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

from backend.services.lazada_campaign_workbook import (
    CampaignWorkbookJob, export_campaign_template, inspect_campaign_files, run_campaign_batch, validate_campaign_payload,
)


INVITED = ["Seller SKU", "Sku ID", "Shop Sku", "Product Name", "Current Price", "Current Stock", "Campaign Price", "Recommended Price"]
REGULAR = ["Seller SKU", "Sku ID", "Shop Sku", "Product Name", "Product Url", "Sales Price", "Campaign Price (Mandatory)", "Recommended Price"]
HELP = "Campaign Price is the price you submit to be eligible to join campaigns. If you have submitted the same deal, the lowest one will apply."


class CampaignWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def write(self, name="活动.xlsx", rows=(), *, mode="invited", title="商品", extra=None):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        sheet = book.active
        sheet.title = title
        sheet.append(INVITED if mode == "invited" else REGULAR)
        if mode == "regular":
            sheet.append([None] * 6 + [HELP, "Campaign Recommended Price is the price range that you may submit to join this campaign."])
            sheet.row_dimensions[2].height = 118
        for row in rows:
            sheet.append(row)
        if extra:
            extra(book)
        book.save(path)
        book.close()
        return path

    def run_job(self, sources, **kwargs):
        payload = {"source_files": [str(path) for path in sources], "output_dir": str(self.root / "结果"), **kwargs}
        return run_campaign_batch(validate_campaign_payload(payload))

    def output(self, result, index=0):
        book = load_workbook(result["files"][index]["processed_file"])
        self.addCleanup(book.close)
        return book

    def test_invited_fills_numeric_price_and_keeps_threshold_boundary(self):
        path = self.write(rows=[
            ["001", "1", "s1", "A", 100, 10, None, "<=95"],
            ["002", "2", "s2", "B", 100, 10, 96, "≤94.99"],
            ["003", "3", "s3", "C", "PHP 1,000.00", 10, 1, "<= ₱ 960.5"],
        ])
        original = path.read_bytes()
        result = self.run_job([path])
        self.assertTrue(result["is_complete"])
        self.assertEqual([result[key] for key in ("total_rows", "updated_rows", "removed_rows", "unparseable_rows", "kept_rows")], [3, 3, 1, 0, 2])
        sheet = self.output(result).active
        self.assertEqual(sheet["A2"].value, "001")
        self.assertEqual(sheet["G2"].value, 95)
        self.assertEqual(sheet["G3"].value, 960.5)
        self.assertEqual(path.read_bytes(), original)
        report = load_workbook(result["output_file"])
        self.addCleanup(report.close)
        audit = report["筛除与异常明细"]
        self.assertEqual(audit["D2"].value, 3)
        self.assertEqual(json.loads(audit["K2"].value)[6], 96)

    def test_document_invited_example_289_recommended_282_93_is_retained(self):
        path = self.write(rows=[["document-sku", "1", "s1", "A", 289, 10, None, "<=282.93"]])
        result = self.run_job([path])
        self.assertEqual(result["kept_rows"], 1)
        self.assertEqual(result["removed_rows"], 0)
        self.assertEqual(result["updated_rows"], 1)
        self.assertEqual(self.output(result).active["G2"].value, 282.93)

    def test_localized_invited_market_campaign_price_uses_recommendation_in_column_i(self):
        source = self.root / "中文特邀.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["Sku ID", "卖家SKU", "产品名称", "提醒", "当前价格", "当前存货",
                      "市场活动价", "流量加权价", "推荐价格", "Is Hero Product"])
        sheet.append(["1", "0001", "保留商品", None, 100, 5, None, 1, "<=95", "No"])
        sheet.append(["2", "0002", "剔除商品", None, 100, 5, 97, 999, "<=94.99", "No"])
        sheet.append(["3", "0003", "异常商品", None, 0, 5, None, 50, "<=95", "No"])
        sheet.append(["4", "0004", "已有活动价", None, 100, 5, 96, 20, None, "No"])
        book.save(source)
        book.close()
        original = source.read_bytes()
        inspection = inspect_campaign_files({"source_files": [str(source)]})
        self.assertTrue(inspection["files"][0]["ok"], inspection)
        self.assertEqual(inspection["files"][0]["sheets"][0]["columns"],
                         {"base_price": 5, "campaign_price": 7, "recommended_price": 9})
        result = self.run_job([source])
        self.assertTrue(result["is_complete"], result)
        self.assertEqual(result["files"][0]["mode"], "invited")
        self.assertEqual([result[key] for key in ("total_rows", "updated_rows", "kept_rows", "removed_rows", "unparseable_rows")], [4, 2, 2, 1, 1])
        output = self.output(result).active
        self.assertEqual((output["G1"].value, output["H1"].value, output["I1"].value), ("市场活动价", "流量加权价", "推荐价格"))
        self.assertEqual((output["G2"].value, output["H2"].value, output["I2"].value), (95, 1, "<=95"))
        self.assertEqual((output["B3"].value, output["G3"].value, output["H3"].value), ("0004", 96, 20))
        self.assertEqual(source.read_bytes(), original)

    def test_market_campaign_price_alias_keeps_ambiguity_checks_and_ignores_traffic_price(self):
        for title, campaign_headers, successful in (
            ("市场价格全称", ["市场活动价格（PHP）"], True),
            ("活动价重复", ["市场活动价", "Campaign Price"], False),
            ("只有流量价", ["流量加权价"], False),
        ):
            with self.subTest(title=title):
                source = self.root / f"{title}.xlsx"
                book = Workbook()
                book.active.append(["SKU", "当前价格", *campaign_headers, "推荐价格"])
                book.active.append(["0001", 100, *([95] * len(campaign_headers)), "<=95"])
                book.save(source)
                book.close()
                result = self.run_job([source])
                self.assertEqual(result["success_count"], int(successful), result)

    def test_actual_regular_sales_price_range_preserves_instruction_and_prefilled_price(self):
        path = self.write(mode="regular", rows=[
            ["001", "1", "s1", "A", "url", 158, 151.36, "0 < Price <= 151.36"],
            ["002", "2", "s2", "B", "url", 100, 94, "0 < Price <= 99"],
        ])
        result = self.run_job([path])
        self.assertEqual(result["files"][0]["mode"], "regular")
        self.assertEqual(result["total_rows"], 2)
        self.assertEqual(result["updated_rows"], 0)
        self.assertEqual(result["removed_rows"], 1)
        self.assertEqual(result["unparseable_rows"], 0)
        sheet = self.output(result).active
        self.assertEqual(sheet["G2"].value, HELP)
        self.assertEqual(sheet.row_dimensions[2].height, 118)
        self.assertEqual(sheet["G3"].value, 151.36)
        self.assertEqual(sheet["H3"].value, "0 < Price <= 151.36")

    def test_regular_missing_price_uses_recommended_upper_bound_with_maximum_column_priority(self):
        path = self.write(mode="regular", rows=[
            ["001", "1", "s1", "A", "url", 100, None, "0 < Price <= 99"],
            ["002", "2", "s2", "B", "url", 100, 96, "0 < Price <= 97"],
        ])
        result = self.run_job([path])
        self.assertEqual(result["unparseable_rows"], 0)
        self.assertEqual(result["kept_rows"], 2)
        self.assertEqual(result["updated_rows"], 1)
        self.assertEqual(self.output(result).active["G3"].value, 99)
        self.assertEqual(self.output(result).active["G4"].value, 96)
        book = load_workbook(path)
        book.active["I1"] = "Maximum Campaign Price (PHP)"
        book.save(path)
        book.close()
        maximum_blank = self.run_job([path])
        self.assertEqual(self.output(maximum_blank).active["G3"].value, 99)
        book = load_workbook(path)
        book.active["I3"] = 98
        book.save(path)
        book.close()
        filled = self.run_job([path])
        self.assertEqual(filled["unparseable_rows"], 0)
        self.assertEqual(filled["updated_rows"], 1)
        self.assertEqual(self.output(filled).active["G3"].value, 98)
        book = load_workbook(path)
        book.active["I3"] = 0
        book.save(path)
        book.close()
        maximum_invalid = self.run_job([path])
        self.assertEqual((maximum_invalid["unparseable_rows"], maximum_invalid["kept_rows"]), (1, 1))

    def test_regular_chinese_range_fills_activity_and_preserves_instruction_row(self):
        source = self.root / "范围价普通.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["卖家SKU", "Sku ID", "店铺sku", "商品名", "商品链接", "售价",
                      "大促价格 (必填)", "推荐价格", "库存"])
        help_text = "大促价格是你提到到这个活动的商品价格。如果你同一商品在相同的时间段提报了2个价格，提报的最低价会被应用。"
        sheet.append([None] * 6 + [help_text, "大促推荐价是你可以提报这个活动价格范围", "库存代表当前库存并显示在买方页面上"])
        sheet.row_dimensions[2].height = 118
        sheet.append(["0001", "1", "shop1", "保留", "url", 100, None, "0 < Price <= 95", 5])
        sheet.append(["0002", "2", "shop2", "剔除", "url", 100, None, "0 < Price <= 94.99", 5])
        sheet.append(["0003", "3", "shop3", "已有价", "url", 100, 96, "0 < Price <= 1", 5])
        sheet.append(["0004", "4", "shop4", "异常", "url", 0, None, "0 < Price <= 95", 5])
        book.save(source)
        book.close()
        before = source.read_bytes()
        for mode in ("auto", "invited", "plain"):
            with self.subTest(mode=mode):
                inspection = inspect_campaign_files({"source_files": [str(source)], "mode": mode})["files"][0]
                self.assertEqual(inspection["total_rows"], 4)
                self.assertEqual(inspection["sheets"][0]["first_data_row"], 3)
                result = self.run_job([source], mode=mode)
                self.assertTrue(result["is_complete"], result)
                output = self.output(result).active
                self.assertEqual(output["G2"].value, help_text)
                self.assertEqual(output.row_dimensions[2].height, 118)
                if mode == "plain":
                    self.assertEqual((result["kept_rows"], result["updated_rows"], result["unparseable_rows"]), (1, 0, 3))
                    self.assertEqual(output["A3"].value, "0003")
                else:
                    self.assertEqual(output["G3"].value, 95)
                    self.assertEqual(output["H3"].value, "0 < Price <= 95")
                    self.assertEqual(result["unparseable_rows"], 1)
                    self.assertEqual(result["updated_rows"], 3 if mode == "invited" else 2)
                    self.assertEqual(result["kept_rows"], 1 if mode == "invited" else 2)
        self.assertEqual(source.read_bytes(), before)

    def test_recommended_ranges_are_strict_and_malformed_rows_are_audited(self):
        valid = ("0 < Price <= 95", "0<price≤95", "<=95", "95")
        invalid = ("0 < Price < 95", "1 < Price <= 95", "0 < Price <= 95 or 99", "0 < Price <= NaN",
                   "0 < Price <= Infinity", "0 < Price <= 0", "0 < Price <= -95", "price 95", "=95", None)
        for mode in ("auto", "invited"):
            with self.subTest(mode=mode):
                source = self.write(name=f"范围校验_{mode}.xlsx", mode="regular", rows=[
                    [str(index), str(index), "sku", "商品", "url", 100, None, value]
                    for index, value in enumerate((*valid, *invalid))
                ])
                result = self.run_job([source], mode=mode)
                self.assertEqual(result["total_rows"], len(valid) + len(invalid))
                self.assertEqual((result["kept_rows"], result["updated_rows"], result["unparseable_rows"]),
                                 (len(valid), len(valid), len(invalid)))
                output = self.output(result).active
                self.assertEqual([output.cell(row, 7).value for row in range(3, 3 + len(valid))], [95] * len(valid))

    def test_instruction_detection_does_not_skip_incomplete_product_row(self):
        def extra(book):
            book.active.delete_rows(2)
        for index, first_row in enumerate((
            ["0001", "1", "s1", "商品", "url", None, HELP, "0 < Price <= 95"],
            ["库存代表当前商品SKU", None, None, None, None, None, HELP, "0 < Price <= 95"],
            ["0001", None, None, None, None, "请填写原价", "请填写活动价", None],
        )):
            with self.subTest(index=index):
                source = self.write(name=f"不完整商品{index}.xlsx", mode="regular", rows=[
                    first_row,
                    ["0002", "2", "s2", "商品", "url", 100, 95, "0 < Price <= 95"],
                ], extra=extra)
                result = self.run_job([source])
                self.assertEqual((result["total_rows"], result["unparseable_rows"], result["kept_rows"]), (2, 1, 1))

    def test_all_matching_sheets_and_other_sheets_are_preserved(self):
        def extra(book):
            regular = book.create_sheet("普通")
            regular.append(REGULAR)
            regular.append([None] * 6 + [HELP, "Campaign Recommended Price describes the range"])
            regular.append(["0002", "2", "s2", "B", "url", 100, 96, "0 < Price <= 96"])
            notes = book.create_sheet("说明")
            notes["A1"] = "保留说明内容"
        path = self.write(rows=[["0001", "1", "s1", "A", 100, 10, None, "<=95"]], extra=extra)
        result = self.run_job([path])
        self.assertEqual(result["total_rows"], 2)
        self.assertEqual(result["files"][0]["mode"], "mixed")
        self.assertEqual(len(result["files"][0]["sheets"]), 2)
        book = self.output(result)
        self.assertEqual(book.sheetnames, ["商品", "普通", "说明"])
        self.assertEqual(book["说明"]["A1"].value, "保留说明内容")

    def test_hidden_rows_styles_ids_nonprice_formulas_filter_and_table_ranges(self):
        def extra(book):
            sheet = book.active
            sheet["A4"].number_format = "@"
            sheet["A4"].font = Font(bold=True, color="AA0011")
            sheet["A4"].fill = PatternFill("solid", fgColor="FFFF00")
            sheet["B4"] = "=1+1"
            sheet.row_dimensions[3].hidden = True
            sheet.row_dimensions[4].height = 25
            sheet.auto_filter.ref = "A1:H4"
            table = Table(displayName="CampaignTable", ref="A1:H4")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            sheet.add_table(table)
        path = self.write(rows=[
            ["0001", "1", "s1", "A", 100, 10, None, "<=95"],
            ["0002", "2", "s2", "B", 100, 10, None, "<=90"],
            ["0003", "3", "s3", "C", 100, 10, None, "<=96"],
        ], extra=extra)
        result = self.run_job([path])
        sheet = self.output(result).active
        self.assertEqual(result["removed_rows"], 1)
        self.assertEqual(sheet["A3"].value, "0003")
        self.assertTrue(sheet["A3"].font.bold)
        self.assertEqual(sheet["A3"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(sheet["A3"].number_format, "@")
        self.assertEqual(sheet["B3"].value, "=1+1")
        self.assertEqual(sheet.auto_filter.ref, "A1:H3")
        self.assertEqual(sheet.tables["CampaignTable"].ref, "A1:H3")
        self.assertFalse(sheet.row_dimensions[3].hidden)
        self.assertEqual(sheet.row_dimensions[3].height, 25)

    def test_formula_and_unparseable_prices_are_excluded_with_original_formula_in_audit(self):
        path = self.write(rows=[
            ["001", "1", "s1", "A", "=50*2", 10, None, "<=95"],
            ["002", "2", "s2", "B", 100, 10, "=95", "<=96"],
            ["003", "3", "s3", "C", 100, 10, None, "<=NaN"],
            ["004", "4", "s4", "D", 100, 10, None, "<=price 97 is maximum"],
            ["005", "5", "s5", "E", -100, 10, None, "<=95"],
            ["006", "6", "s6", "F", 100, 10, None, "<=95"],
        ])
        result = self.run_job([path])
        self.assertEqual(result["unparseable_rows"], 5)
        self.assertEqual(result["kept_rows"], 1)
        sheet = self.output(result).active
        self.assertEqual(sheet["A2"].value, "006")
        report = load_workbook(result["output_file"])
        self.addCleanup(report.close)
        audit = report["筛除与异常明细"]
        self.assertIn("公式", audit["G2"].value)
        self.assertEqual(json.loads(audit["K2"].value)[4], "=50*2")
        self.assertEqual(audit["I2"].data_type, "s")
        self.assertEqual(audit["I2"].value, "=50*2")

    def test_duplicate_price_headers_fail_and_generic_headers_use_existing_price(self):
        valid = self.write("valid.xlsx", rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]])
        duplicate = self.write("duplicate.xlsx", rows=[])
        book = load_workbook(duplicate)
        book.active["I1"] = "Current Price (PHP)"
        book.save(duplicate)
        book.close()
        ambiguous = self.write("ambiguous.xlsx", rows=[])
        book = load_workbook(ambiguous)
        book.active["E1"] = "Original Price"
        book.active["F1"] = "Stock"
        book.save(ambiguous)
        book.close()
        result = self.run_job([valid, duplicate, ambiguous])
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["success_count"], 2)
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("歧义", result["files"][1]["error"])
        self.assertEqual(result["files"][2]["mode"], "plain")
        explicit = self.run_job([ambiguous], mode="invited")
        self.assertTrue(explicit["is_complete"])

    def test_header_position_and_chinese_currency_suffixes(self):
        path = self.root / "中文.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.append(["活动信息"])
        sheet.append(["SKU", "当前价格（PHP）", "活动价格（选填）", "推荐价格\n此处填写平台推荐值"])
        sheet.append(["0003", 100, None, "≤95"])
        book.save(path)
        book.close()
        result = self.run_job([path])
        self.assertEqual(result["kept_rows"], 1)
        self.assertEqual(self.output(result).active["C3"].value, 95)
        self.assertEqual(result["files"][0]["sheets"][0]["header_row"], 2)

    def test_inspection_matches_processing_and_keeps_source_unchanged(self):
        path = self.write(mode="regular", rows=[["001", "1", "s1", "A", "url", 100, None, "0 < Price <= 95"]])
        corrupt = self.root / "bad.xlsx"
        corrupt.write_bytes(b"this is not a workbook")
        before = path.read_bytes()
        result = inspect_campaign_files({"source_files": [str(path), str(corrupt)]})
        self.assertEqual(result["file_count"], 2)
        self.assertTrue(result["files"][0]["ok"])
        self.assertFalse(result["files"][1]["ok"])
        self.assertEqual(result["files"][0]["unparseable_rows"], 0)
        self.assertEqual(result["files"][0]["missing_price_rows"], 1)
        self.assertEqual(result["files"][0]["sheets"][0]["first_data_row"], 3)
        self.assertEqual(path.read_bytes(), before)
        self.assertFalse((self.root / "Lazada活动处理").exists())

    def test_same_named_files_unique_batches_archive_and_progress(self):
        first = self.write("one/shop.xlsx", rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]])
        second = self.write("two/shop.xlsx", rows=[["002", "2", "s2", "B", 100, 10, None, "<=96"]])
        job = validate_campaign_payload({"source_files": [first, second], "output_dir": self.root / "结果"})
        messages = []
        result = run_campaign_batch(job, messages.append)
        next_result = run_campaign_batch(job)
        self.assertNotEqual(result["output_dir"], next_result["output_dir"])
        self.assertNotEqual(result["files"][0]["processed_file"], result["files"][1]["processed_file"])
        self.assertEqual(len(messages), 4)
        with ZipFile(result["archive_file"]) as archive:
            self.assertEqual(len(archive.namelist()), 3)
            self.assertIn("Excel折扣筛选汇总与异常明细.xlsx", archive.namelist())

    def test_invalid_payloads_and_thresholds(self):
        path = self.write()
        payload = {"source_files": [str(path)]}
        bad = [None, [], {}, {"source_files": str(path)}, {"source_files": [None]},
               {"source_files": [str(self.root / "none.xlsx")]}, {**payload, "mode": []},
               {**payload, "mode": "wrong"}, {**payload, "output_dir": ""},
               {**payload, "source_files": [str(path)] * 101}]
        bad += [{**payload, "threshold": value} for value in (0, 1.1, -0.1, "NaN", "Infinity", True, None, "x", [])]
        for value in bad:
            with self.subTest(value=value), self.assertRaises(ValueError):
                validate_campaign_payload(value)
        self.assertEqual(validate_campaign_payload({**payload, "threshold": "1"}).threshold, Decimal("1"))
        self.assertEqual(len(validate_campaign_payload({"source_files": [str(path), path]}).source_files), 1)
        with patch("backend.services.lazada_campaign_workbook.MAX_FILE_BYTES", 1), self.assertRaises(ValueError):
            validate_campaign_payload(payload)

    def test_macro_binary_and_extension_are_preserved(self):
        source = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]])
        target = source.with_suffix(".xlsm")
        with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as macro:
            for item in original.infolist():
                macro.writestr(item, original.read(item.filename))
            macro.writestr("xl/vbaProject.bin", b"macro-preservation-test")
        result = self.run_job([target])
        self.assertTrue(result["is_complete"])
        output = Path(result["files"][0]["processed_file"])
        self.assertEqual(output.suffix, ".xlsm")
        with ZipFile(output) as bundle:
            self.assertEqual(bundle.read("xl/vbaProject.bin"), b"macro-preservation-test")

    def test_auxiliary_field_description_sheet_is_ignored(self):
        def extra(book):
            note = book.create_sheet("字段说明")
            note.append(["Field", "Description"])
            note.append(["Campaign Price", "Price submitted to a campaign"])
        path = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]], extra=extra)
        result = self.run_job([path])
        self.assertTrue(result["is_complete"])
        self.assertEqual(self.output(result)["字段说明"]["A2"].value, "Campaign Price")

    def test_reference_formula_with_deletions_fails_safely_and_source_stays_unchanged(self):
        def extra(book):
            book.active["B3"] = "=A3"
        path = self.write(rows=[
            ["001", "1", "s1", "A", 100, 10, None, "<=94"],
            ["002", "2", "s2", "B", 100, 10, None, "<=96"],
        ], extra=extra)
        original = path.read_bytes()
        result = self.run_job([path])
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("引用公式", result["files"][0]["error"])
        self.assertEqual(result["files"][0]["processed_file"], "")
        self.assertEqual(path.read_bytes(), original)

    def test_large_formatted_dimensions_fail_before_processing(self):
        def extra(book):
            book.active["A100001"] = "out of bounds"
        path = self.write(extra=extra)
        result = self.run_job([path])
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("100,000", result["files"][0]["error"])

    def test_price_precision_rejects_loss_and_decimal_threshold_is_exact(self):
        path = self.write(rows=[
            ["001", "1", "s1", "A", 100, 10, None, "<=95"],
            ["002", "2", "s2", "B", "10000000000000000", 10, None, "<=10000000000000000"],
        ])
        result = self.run_job([path], threshold="0.9500000000000000000000000000000001")
        self.assertEqual(result["removed_rows"], 1)
        self.assertEqual(result["unparseable_rows"], 1)
        self.assertEqual(result["kept_rows"], 0)

    def test_data_validation_ranges_move_with_removed_rows(self):
        def extra(book):
            validation = DataValidation(type="decimal", operator="greaterThan", formula1="0")
            book.active.add_data_validation(validation)
            validation.add("G3")
        path = self.write(rows=[
            ["001", "1", "s1", "A", 100, 10, None, "<=94"],
            ["002", "2", "s2", "B", 100, 10, None, "<=96"],
        ], extra=extra)
        result = self.run_job([path])
        self.assertEqual(str(self.output(result).active.data_validations.dataValidation[0].sqref), "G2")

    def test_report_export_failure_returns_created_templates_and_recovery_details(self):
        path = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]])
        with patch("backend.services.lazada_campaign_workbook._report", side_effect=PermissionError("locked report")):
            result = self.run_job([path])
        self.assertFalse(result["is_complete"])
        self.assertEqual(result["success_count"], 1)
        self.assertIn("汇总", result["completion_message"])
        self.assertTrue(Path(result["files"][0]["processed_file"]).is_file())
        self.assertTrue(Path(result["archive_file"]).is_file())
        self.assertEqual(result["output_file"], "")

    def test_zip_export_failure_returns_report_and_created_templates(self):
        path = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]])
        with patch("backend.services.lazada_campaign_workbook.ZipFile", side_effect=PermissionError("locked zip")):
            result = self.run_job([path])
        self.assertFalse(result["is_complete"])
        self.assertIn("ZIP", result["completion_message"])
        self.assertTrue(Path(result["output_file"]).is_file())
        self.assertEqual(result["files"][0]["audit_file"], result["output_file"])
        self.assertEqual(result["archive_file"], "")

    def test_valid_and_malformed_activity_sheets_fail_as_a_whole(self):
        def extra(book):
            malformed = book.create_sheet("缺少原价")
            malformed.append(["Seller SKU", "Campaign Price", "Recommended Price"])
            malformed.append(["unscreened-sku", 1, "<=1"])
        path = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=95"]], extra=extra)
        result = self.run_job([path])
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("缺少原价", result["files"][0]["error"])
        self.assertEqual(result["files"][0]["processed_file"], "")

    def test_unused_regular_recommended_reference_formula_blocks_unsafe_row_move(self):
        path = self.write(mode="regular", rows=[
            ["001", "1", "s1", "A", "url", 100, 94, "0 < Price <= 94"],
            ["002", "2", "s2", "B", "url", 100, 96, "=F4*0.97"],
        ])
        result = self.run_job([path])
        self.assertEqual(result["failed_count"], 1)
        self.assertIn("H4", result["files"][0]["error"])

    def test_removed_price_formula_does_not_block_other_valid_sheets(self):
        def extra(book):
            regular = book.create_sheet("普通")
            regular.append(REGULAR)
            regular.append([None] * 6 + [HELP, "Campaign Recommended Price describes the range"])
            regular.append(["002", "2", "s2", "B", "url", 100, "=F3*0.96", "0 < Price <= 96"])
        path = self.write(rows=[["001", "1", "s1", "A", 100, 10, None, "<=94"]], extra=extra)
        result = self.run_job([path])
        self.assertTrue(result["is_complete"])
        self.assertEqual(result["removed_rows"], 1)
        self.assertEqual(result["unparseable_rows"], 1)

    def test_large_alternating_exclusions_preserve_order_links_and_metadata(self):
        size = 1000
        def extra(book):
            sheet = book.active
            sheet.merge_cells("I1:J1")
            sheet["I1"] = "保留标题"
            sheet["I1"].font = Font(bold=True)
            sheet.merge_cells("I1000:J1000")
            sheet["I1000"] = "保留合并备注"
            sheet["I1000"].font = Font(bold=True)
            sheet["A1000"].hyperlink = "https://example.org/product/998"
            sheet["A1000"].font = Font(bold=True, color="FF0011")
            sheet["D1000"].comment = Comment("保留商品备注", "test")
            sheet["C1000"] = "=1+1"
            sheet.row_dimensions[1000].height = 33
            sheet.row_dimensions[1000].hidden = True
            sheet.auto_filter.ref = "A1:J1001"
            table = Table(displayName="LargeCampaignTable", ref="A1:H1001")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            sheet.add_table(table)
            validation = DataValidation(type="decimal", operator="greaterThan", formula1="0")
            sheet.add_data_validation(validation)
            validation.add("G1000")
        rows = [[f"SKU-{index:05d}", str(index), f"shop-{index}", "商品", 100, 10, None,
                 "<=95" if index % 2 == 0 else "<=90"] for index in range(size)]
        path = self.write(rows=rows, extra=extra)
        original = path.read_bytes()
        result = self.run_job([path])
        self.assertEqual(result["kept_rows"], 500)
        self.assertEqual(result["removed_rows"], 500)
        self.assertEqual(result["updated_rows"], 1000)
        sheet = self.output(result).active
        self.assertEqual([sheet.cell(row, 1).value for row in range(2, 502)],
                         [f"SKU-{index:05d}" for index in range(0, size, 2)])
        self.assertEqual(sheet["A501"].hyperlink.ref, "A501")
        self.assertEqual(sheet["A501"].hyperlink.target, "https://example.org/product/998")
        self.assertTrue(sheet["A501"].font.bold)
        self.assertEqual(sheet["D501"].comment.text, "保留商品备注")
        self.assertEqual(sheet["C501"].value, "=1+1")
        self.assertTrue(sheet.row_dimensions[501].hidden)
        self.assertEqual(sheet.row_dimensions[501].height, 33)
        self.assertEqual(sheet.auto_filter.ref, "A1:J501")
        self.assertEqual(sheet.tables["LargeCampaignTable"].ref, "A1:H501")
        self.assertEqual(str(sheet.data_validations.dataValidation[0].sqref), "G501")
        self.assertIn("I1:J1", {str(ref) for ref in sheet.merged_cells.ranges})
        self.assertEqual(sheet["I1"].value, "保留标题")
        self.assertIn("I501:J501", {str(ref) for ref in sheet.merged_cells.ranges})
        self.assertEqual(sheet["I501"].value, "保留合并备注")
        self.assertTrue(sheet["I501"].font.bold)
        self.assertEqual(path.read_bytes(), original)

    def test_blank_generic_template_can_be_filled_and_filtered_in_auto_mode(self):
        exported = export_campaign_template({"output_dir": self.root / "模板"})
        source = Path(exported["output_file"])
        self.assertTrue(exported["ok"])
        book = load_workbook(source)
        sheet = book.active
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual([cell.value for cell in sheet[1]], ["SKU", "商品名称", "原价", "活动价", "推荐价"])
        self.assertEqual(sheet.freeze_panes, "A2")
        sheet.append(["0001", "保留商品", 100, 95, "<=1"])
        sheet.append(["0002", "剔除商品", 100, 94.99, 100])
        book.save(source)
        book.close()
        before = source.read_bytes()
        result = self.run_job([source])
        self.assertTrue(result["is_complete"], result)
        self.assertEqual(result["files"][0]["mode"], "plain")
        self.assertEqual((result["kept_rows"], result["removed_rows"], result["updated_rows"]), (1, 1, 0))
        sheet = self.output(result).active
        self.assertEqual((sheet["A2"].value, sheet["D2"].value), ("0001", 95))
        self.assertIn("字段说明", self.output(result).sheetnames)
        self.assertEqual(source.read_bytes(), before)
        self.assertNotEqual(export_campaign_template({"output_dir": self.root / "模板"})["path"], str(source))

    def test_generic_english_and_chinese_aliases(self):
        for index, (base, price) in enumerate((("Base Price", "Discount Price"), ("List Price", "Sale Price"),
                                              ("商品原价", "折后价"), ("原售价", "促销价"))):
            with self.subTest(base=base, price=price):
                source = self.root / f"generic{index}.xlsx"
                book = Workbook()
                book.active.append(["SKU", base, price])
                book.active.append(["001", 100, 95])
                book.save(source)
                book.close()
                result = self.run_job([source])
                self.assertTrue(result["is_complete"], result)
                self.assertEqual(result["kept_rows"], 1)
                self.assertEqual(result["files"][0]["mode"], "plain")

    def test_plain_ignores_optional_recommended_headers_and_does_not_fill_missing(self):
        source = self.root / "已有价.xlsx"
        book = Workbook()
        book.active.append(["SKU", "原价", "活动价", "推荐价", "Recommended Price", "最高活动价"])
        book.active.append(["001", 100, 95, "无效推荐", "0 < Price <= 50", 99])
        book.active.append(["002", 100, None, 99, 99, 99])
        book.save(source)
        book.close()
        result = self.run_job([source], mode="plain")
        self.assertTrue(result["is_complete"], result)
        self.assertEqual((result["kept_rows"], result["unparseable_rows"], result["updated_rows"]), (1, 1, 0))
        self.assertEqual(self.output(result).active["C2"].value, 95)
        report = load_workbook(result["output_file"])
        self.addCleanup(report.close)
        self.assertIn("缺少活动价", report["筛除与异常明细"]["G2"].value)

    def test_explicit_recommended_mode_fills_generic_template(self):
        source = Path(export_campaign_template({"output_dir": self.root})["path"])
        book = load_workbook(source)
        book.active.append(["001", "商品", 100, None, "<=95"])
        book.save(source)
        book.close()
        result = self.run_job([source], mode="invited")
        self.assertEqual((result["kept_rows"], result["updated_rows"]), (1, 1))
        self.assertEqual(self.output(result).active["D2"].value, 95)

    def test_blank_template_has_no_sample_products_and_export_validation(self):
        source = Path(export_campaign_template({"output_dir": self.root})["path"])
        result = self.run_job([source])
        self.assertTrue(result["is_complete"], result)
        self.assertEqual((result["total_rows"], result["unparseable_rows"]), (0, 0))
        for payload in (None, [], {}, {"output_dir": ""}, {"output_dir": source}):
            with self.subTest(payload=payload), self.assertRaises(ValueError):
                export_campaign_template(payload)


if __name__ == "__main__":
    unittest.main()
