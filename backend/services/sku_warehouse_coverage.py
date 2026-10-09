"""Read-only hot SKU coverage against the inventory page's warehouse catalogue."""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.services.sku_inventory_query import (
    BASE_URL, MAX_PAGE_COUNT, STOCK_LIST_API_URL, TIME_CREATED_KEYS,
    DeveloperOption, MabangApiError, MabangClient, ProgressCallback,
    SkuInventoryQuery, SkuInventoryResult, build_stock_query_payload, first_value,
    normalize_created_at, open_stock_page, parse_developer_options,
    parse_pagination_state, post_form_json, safe_text, stock_row_identity,
)

LIVENESS_NAMES = {"1": "爆款", "2": "旺款"}
# Confirmed by the business owner on 2026-09-21; these brands have no country
# configured in Mabang and are outside the requested Southeast Asia scope.
NON_SEA_WAREHOUSE_IDS = frozenset({"1100808", "1100602"})
NON_SEA_WAREHOUSE_NAMES = frozenset({"Ola-Ola海外仓", "T&C 海外 2 号仓-T&C海外仓"})
SEA_COUNTRIES = {
    "泰国": ("泰国", "thailand", "TH"),
    "越南": ("越南", "vietnam", "viet nam", "VN"),
    "菲律宾": ("菲律宾", "philippines", "PH"),
    "马来西亚": ("马来西亚", "马来", "malaysia", "MY"),
    "印度尼西亚": ("印度尼西亚", "印尼", "indonesia", "ID"),
    "新加坡": ("新加坡", "singapore", "SG"),
    "缅甸": ("缅甸", "myanmar", "burma", "MM"),
    "柬埔寨": ("柬埔寨", "cambodia", "KH"),
    "老挝": ("老挝", "laos", "LA"),
    "文莱": ("文莱", "brunei", "BN"),
    "东帝汶": ("东帝汶", "timor-leste", "east timor", "TL"),
}


def southeast_asian_country(name: str, country: str = "") -> str:
    if "中转" in name or re.search(r"\b(?:transit|transfer)\b", name, re.I):
        return ""
    if country:
        return next((label for label, aliases in SEA_COUNTRIES.items()
                     if country.casefold() in {a.casefold() for a in aliases}), "")
    # Third-party warehouses often have no country configured in Mabang.
    for label, aliases in SEA_COUNTRIES.items():
        for alias in aliases:
            if len(alias) == 2 and alias.isascii():
                if re.search(rf"(?<![A-Za-z]){alias}(?=\d|[^A-Za-z]|$)", name, re.I):
                    return label
            elif alias.casefold() in name.casefold():
                return label
    if any(city in name for city in ("河内", "胡志明")):
        return "越南"
    return ""


def parse_warehouse_countries(fragment: str) -> dict[str, str]:
    countries = {}
    for row in BeautifulSoup(fragment, "html.parser").select("tr"):
        checkbox = row.select_one(".checkWarehouse[data-id]")
        if checkbox is None:
            continue
        cells = row.find_all("td", recursive=False)
        # With attributetype="", the endpoint omits the warehouse-code column:
        # checkbox, name, type, country, then numeric inventory columns.
        if len(cells) < 4:
            raise MabangApiError("仓库目录列结构异常，无法读取所属国家")
        countries[safe_text(checkbox.get("data-id"))] = cells[3].get_text(" ", strip=True)
    if not countries:
        raise MabangApiError("仓库资料未返回任何仓库，无法核对所属国家")
    return countries


def parse_scope_warehouses(stock_html: str, countries: dict[str, str]) -> list[dict[str, str]]:
    soup = BeautifulSoup(stock_html, "html.parser")
    selector = soup.select_one('select[name="defaultWarehouseId"]')
    if selector is None:
        raise MabangApiError("库存 SKU 页面缺少完整仓库列表，无法计算缺失仓库")
    warehouses = {}
    for option in selector.select("option[value]"):
        warehouse_id = safe_text(option.get("value"))
        name = option.get_text(" ", strip=True)
        if warehouse_id in NON_SEA_WAREHOUSE_IDS or name in NON_SEA_WAREHOUSE_NAMES:
            continue
        country = southeast_asian_country(name, countries.get(warehouse_id, ""))
        if warehouse_id and country:
            warehouses[warehouse_id] = {"id": warehouse_id, "name": name, "country": country}
    if not warehouses:
        raise MabangApiError("没有识别到东南亚非中转仓，请检查仓库资料和账号权限")
    return sorted(warehouses.values(), key=lambda item: (item["country"], item["name"], item["id"]))


def unclassified_warehouses(stock_html: str, countries: dict[str, str]) -> list[dict[str, str]]:
    result = []
    for option in BeautifulSoup(stock_html, "html.parser").select('select[name="defaultWarehouseId"] option[value]'):
        identity, name = safe_text(option.get("value")), option.get_text(" ", strip=True)
        if identity in NON_SEA_WAREHOUSE_IDS or name in NON_SEA_WAREHOUSE_NAMES:
            continue
        if not identity or countries.get(identity) or southeast_asian_country(name):
            continue
        if re.search(r"中转|测试|中国|广州|武汉|南宁|荆州|华东|日本|琦玉|俄罗斯|北俄|莫斯科|沙特|哈萨克|巴西|美洲|美国|欧洲|英国|德国|法国|西班牙|意大利|澳洲|澳大利亚|加拿大|\b(?:transit|transfer)\b", name, re.I):
            continue
        result.append({"id": identity, "name": name})
    return result


def _stock_response_message(data: dict[str, Any], job: SkuInventoryQuery) -> str:
    message = data.get("message") or data.get("msg")
    if not isinstance(message, str):
        return "马帮未返回具体原因"
    # Never log the complete response: it can include session credentials.
    for secret in (job.password, job.username):
        if secret:
            message = message.replace(secret, "[已隐藏]")
    message = BeautifulSoup(message, "html.parser").get_text(" ", strip=True)
    for secret in (job.password, job.username):
        if secret:
            message = message.replace(secret, "[已隐藏]")
    return message[:240] or "马帮未返回具体原因"


def _is_empty_stock_response(data: dict[str, Any], page: int) -> bool:
    # Verified against stock.getStockList: an empty first page omits stockData
    # and hasData, returns pageHtml=false and the "暂无内容！" HTML template.
    # The template also appears WITH populated stockData, so it alone is not
    # evidence of an empty result. Never accept this shape mid-pagination.
    message = data.get("message")
    return (
        page == 1
        and data.get("success") is True
        and "stockData" not in data
        and "hasData" not in data
        and data.get("pageHtml") is False
        and isinstance(message, str)
        and BeautifulSoup(message, "html.parser").get_text(" ", strip=True) == "暂无内容！"
    )


def iter_hot_stock_rows(client, job, developer, liveness_type, referer, progress):
    seen: set[str] = set()
    expected_total_pages: int | None = None
    for page in range(1, MAX_PAGE_COUNT + 1):
        context = f"{developer.name}（ID {developer.id}） · {LIVENESS_NAMES[liveness_type]} · 第 {page} 页"
        if progress:
            progress(f"{context}：正在查询...")
        data = post_form_json(client, STOCK_LIST_API_URL, build_stock_query_payload(
            job, developer_id=developer.id, page=page, liveness_type=liveness_type,
        ), referer)
        if not data.get("success"):
            raise MabangApiError(f"{context}：查询失败；{_stock_response_message(data, job)}")
        rows = data.get("stockData")
        if not isinstance(rows, list):
            if _is_empty_stock_response(data, page):
                if progress:
                    progress(f"{context}：暂无匹配 SKU，继续查询")
                return
            field_type = "缺失" if "stockData" not in data else type(rows).__name__
            raise MabangApiError(
                f"{context}：返回字段异常（stockData={field_type}）；{_stock_response_message(data, job)}"
            )
        current_page, total_pages = parse_pagination_state(data)
        if current_page is not None and current_page != page:
            raise MabangApiError("马帮分页返回了错误页码，已停止导出，避免遗漏 SKU")
        if total_pages is not None:
            if expected_total_pages is not None and total_pages != expected_total_pages:
                raise MabangApiError(f"{context}：马帮总页数发生变化，请重新查询，避免遗漏 SKU")
            expected_total_pages = total_pages
        if not rows and expected_total_pages is not None and (page > 1 or expected_total_pages > 1):
            raise MabangApiError(f"{context}：库存 SKU 分页提前返回空数据，请重新查询")
        new_count = 0
        for stock in rows:
            if not isinstance(stock, dict):
                raise MabangApiError("库存 SKU 返回记录格式异常")
            identity = stock_row_identity(stock)
            if identity in seen:
                continue
            seen.add(identity)
            new_count += 1
            # Verify the two server-side filters instead of silently labeling
            # unrelated rows as belonging to the requested developer/type.
            if safe_text(stock.get("developerId")) != developer.id:
                raise MabangApiError("库存 SKU 的开发员筛选结果不一致，请重新查询")
            if safe_text(stock.get("livenessType")) != liveness_type:
                raise MabangApiError("库存 SKU 的爆款／旺款筛选结果不一致，请重新查询")
            yield stock
        if rows and not new_count:
            raise MabangApiError("马帮返回重复分页，已停止导出，避免把不完整结果当作全部 SKU")
        if expected_total_pages is not None:
            if page >= expected_total_pages:
                return
        elif len(rows) < job.rows_per_page:
            return
    raise MabangApiError(f"库存 SKU 查询超过 {MAX_PAGE_COUNT} 页，请减少所选开发员")


def coverage_record(stock: dict[str, Any], developer: DeveloperOption,
                    warehouses: list[dict[str, str]]) -> dict[str, Any]:
    sku = safe_text(first_value(stock, ("stockSku", "stockSkuRaw", "sku")))
    details = stock.get("stockWarehouseData")
    if not sku or not isinstance(details, list):
        raise MabangApiError("SKU 或仓库关联数据缺失，无法可靠判断缺失仓库")
    present_ids = set()
    for detail in details:
        if not isinstance(detail, dict) or not safe_text(detail.get("warehouseId")):
            raise MabangApiError(f"SKU {sku} 的仓库关联数据异常")
        # `id` is a stock-warehouse association ID, not the warehouse ID.
        # Zero-stock / in-transit rows still represent an existing association.
        present_ids.add(safe_text(detail["warehouseId"]))
    missing = [warehouse for warehouse in warehouses if warehouse["id"] not in present_ids]
    return {
        "developer_name": developer.name,
        "sku": sku,
        "product_name": safe_text(first_value(stock, ("nameCN", "stockNameCN", "productName"))),
        "created_at": normalize_created_at(first_value(stock, TIME_CREATED_KEYS)),
        "liveness_name": LIVENESS_NAMES[safe_text(stock.get("livenessType"))],
        "missing_warehouse_count": len(missing),
        "missing_warehouse_names": "、".join(w["name"] for w in missing),
        "missing_warehouse_ids": "、".join(w["id"] for w in missing),
        "covered_warehouse_count": len(warehouses) - len(missing),
    }


def export_coverage_records(records, warehouses, output_dir: Path, unresolved=()) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / f"爆款旺款缺失仓库_{datetime.now():%Y%m%d_%H%M%S_%f}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU缺失仓库"
    fields = ("developer_name", "sku", "product_name", "liveness_name", "created_at",
              "covered_warehouse_count", "missing_warehouse_count", "missing_warehouse_names", "missing_warehouse_ids")
    sheet.append(["开发人员", "库存SKU", "商品中文名称", "款式类型", "SKU创建时间",
                  "已关联仓库数", "缺失仓库数", "缺失仓库", "缺失仓库ID"])
    for record in records:
        sheet.append([record[field] for field in fields])
    scope = workbook.create_sheet("东南亚仓库范围")
    scope.append(["仓库ID", "仓库名称", "所属国家"])
    for warehouse in warehouses:
        scope.append([warehouse["id"], warehouse["name"], warehouse["country"]])
    rules = workbook.create_sheet("查询口径")
    rules.append(["项目", "说明"])
    rules.append(["款式", "马帮库存 SKU 活跃度：爆款／旺款，以本次选择为准"])
    rules.append(["SKU范围", "所选开发员名下全部创建时间、全部商品状态的所选款式"])
    rules.append(["仓库范围", "库存 SKU 页面仓库目录中的东南亚仓，不含中转仓；国家为空时按仓库名称识别"])
    rules.append(["缺失口径", "该 SKU 未关联目标仓库；已关联但库存为 0 的仓库仍算已覆盖"])
    rules.append(["完整覆盖", "缺失仓库数为 0 时，表示已关联全部目标仓库"])
    rules.append(["明确排除", "Ola-Ola海外仓、T&C 海外 2 号仓-T&C海外仓不属于东南亚仓库"])
    if unresolved:
        rules.append(["待确认仓库", "以下仓库未填写国家且名称无法识别，尚未纳入范围：" + "、".join(w["name"] for w in unresolved)])
    for worksheet, widths in ((sheet, (22, 24, 38, 12, 22, 18, 18, 85, 60)),
                              (scope, (18, 45, 18)), (rules, (20, 95))):
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for index, width in enumerate(widths, 1):
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(output_file)
    workbook.close()
    return output_file


def run_warehouse_coverage_query(job: SkuInventoryQuery, progress: ProgressCallback | None = None) -> SkuInventoryResult:
    records = []
    with MabangClient(BASE_URL) as client:
        if progress:
            progress("正在登录马帮并读取开发员、东南亚仓库目录...")
        client.login(job.username, job.password)
        referer, html = open_stock_page(client)
        options = {developer.id: developer for developer in parse_developer_options(html)}
        if any(identity not in options for identity in job.developer_ids):
            raise MabangApiError("所选开发人员已失效，请重新加载开发人员")
        developers = [options[identity] for identity in job.developer_ids]
        data = post_form_json(client, urljoin(referer, "/index.php?mod=warehouse.dosearchwarehouse"), [
            ("status", ""), ("type", ""), ("orderby", ""), ("attributetype", ""), ("dataTpye", "public"),
        ], referer)
        if not data.get("success") or not isinstance(data.get("message"), str):
            raise MabangApiError("读取仓库所属国家失败，无法确定完整东南亚仓库范围")
        countries = parse_warehouse_countries(data["message"])
        warehouses = parse_scope_warehouses(html, countries)
        unresolved = unclassified_warehouses(html, countries)
        if progress:
            progress(f"本次对比 {len(warehouses)} 个东南亚非中转仓：" + "、".join(w["name"] for w in warehouses))
            if unresolved:
                progress("以下仓库国家未填写且名称无法识别，尚未纳入范围，需核对：" + "、".join(w["name"] for w in unresolved))
        for developer in developers:
            seen_skus = set()
            for liveness_type in job.liveness_types:
                for stock in iter_hot_stock_rows(client, job, developer, liveness_type, referer, progress):
                    record = coverage_record(stock, developer, warehouses)
                    if record["sku"] in seen_skus:
                        raise MabangApiError(f"SKU {record['sku']} 在查询期间重复或款式发生变化，请重新查询")
                    seen_skus.add(record["sku"])
                    records.append(record)
    records.sort(key=lambda r: (r["developer_name"], -r["missing_warehouse_count"], r["sku"]))
    output_file = export_coverage_records(records, warehouses, job.output_dir, unresolved)
    missing_count = sum(record["missing_warehouse_count"] > 0 for record in records)
    if progress:
        progress(f"查询完成：{len(records)} 个爆款／旺款 SKU，其中 {missing_count} 个存在缺失仓库")
    return SkuInventoryResult(
        records=tuple(records), developer_count=len(developers),
        sku_count=len({r["sku"] for r in records}), warehouse_count=len(warehouses),
        output_file=output_file, query_mode="missing_warehouses",
        scope_warehouses=tuple(warehouses), missing_sku_count=missing_count,
        unclassified_warehouses=tuple(unresolved),
    )
