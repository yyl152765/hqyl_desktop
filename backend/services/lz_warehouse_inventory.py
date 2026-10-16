"""Export native inventory-statistics pages into a separate file per warehouse.

Uses the same search and download endpoints as 商品 → 库存查询. Warehouse
IDs are discovered by the selected names, and every export uses its own search
token. Business rows (including zero stock) are never filtered or deduplicated.
"""

from __future__ import annotations

import hashlib
import json
import math
import re
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Sequence
from urllib.parse import urlencode, urljoin, urlparse
from uuid import uuid4

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import IllegalCharacterError
import xlrd

from backend.core.mabang_client import MabangApiError, MabangClient, MabangLoginError, decode_html_page
from backend.output_paths import default_output_directories


BASE_URL = "https://900853.private.mabangerp.com"
PRIVATE_BASE_URL = "https://private-amz.mabangerp.com"
INVENTORY_PAGE_PATH = "/index.php?mod=warehouse.inventorydetail"
SEARCH_PATH = "/index.php?mod=warehouse.searchwarehousestock"
PAGINATION_PATH = "/index.php?mod=warehouse.getSearchWarehouseStockPage"
EXPORT_PATH = "/index.php?mod=warehouse.doexportwarehousestock"
TARGET_WAREHOUSE_NAMES = ("泰国TH04仓", "泰国TH05仓", "泰国TH06仓")
EXPORT_PAGE_SIZE = 10000
MAX_EXPORT_PAGES = 1000
PREVIEW_ROW_LIMIT = 100
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class LzInventoryJob:
    username: str
    password: str = field(repr=False)
    output_dir: Path
    warehouse_names: tuple[str, ...]


@dataclass(frozen=True)
class LzWarehouseExport:
    warehouse_name: str
    output_file: Path
    record_count: int
    source_record_count: int
    export_page_count: int
    preview_columns: tuple[str, ...]
    preview_rows: tuple[dict[str, Any], ...]


@dataclass(frozen=True)
class LzInventoryResult:
    exports: tuple[LzWarehouseExport, ...]
    failures: tuple[dict[str, str], ...] = ()

    @property
    def output_files(self) -> tuple[Path, ...]:
        return tuple(item.output_file for item in self.exports)

    @property
    def output_file(self) -> Path | None:
        return self.output_files[0] if self.exports else None

    @property
    def record_count(self) -> int:
        return sum(item.record_count for item in self.exports)

    @property
    def source_record_count(self) -> int:
        return sum(item.source_record_count for item in self.exports)

    @property
    def export_page_count(self) -> int:
        return sum(item.export_page_count for item in self.exports)

    @property
    def warehouse_counts(self) -> dict[str, int]:
        return {item.warehouse_name: item.record_count for item in self.exports}

    @property
    def preview_columns(self) -> tuple[str, ...]:
        return self.exports[0].preview_columns if self.exports else ()

    @property
    def preview_rows(self) -> tuple[dict[str, Any], ...]:
        return self.exports[0].preview_rows if self.exports else ()

    @property
    def is_complete(self) -> bool:
        return not self.failures

    @property
    def completion_message(self) -> str:
        if not self.failures:
            return "所选仓库已全部导出"
        return f"已生成 {len(self.exports)} 个仓库文件，{len(self.failures)} 个仓库导出失败，请查看失败明细"


def validate_lz_inventory_payload(payload: dict[str, Any] | None) -> LzInventoryJob:
    data = dict(payload or {})
    username = str(data.get("username") or "").strip()
    password = str(data.get("password") or "")
    output_dir = str(data.get("output_dir") or default_output_directories("lz_warehouse_inventory")["output_dir"]).strip()
    if not username or not password:
        raise ValueError("请先绑定马帮账号和密码")
    if not output_dir:
        raise ValueError("请选择输出目录")
    raw_names = data.get("warehouse_names")
    if not isinstance(raw_names, (list, tuple)) or not raw_names:
        raise ValueError("请至少选择一个仓库")
    if any(not isinstance(name, str) or not name.strip() for name in raw_names):
        raise ValueError("仓库名称无效，请重新选择")
    warehouse_names = tuple(dict.fromkeys(name.strip() for name in raw_names))
    target = Path(output_dir).expanduser()
    if target.exists() and not target.is_dir():
        raise ValueError("输出目录不能是文件")
    return LzInventoryJob(username=username, password=password, output_dir=target, warehouse_names=warehouse_names)


def parse_inventory_warehouses(page_html: str) -> tuple[dict[str, str], ...]:
    soup = BeautifulSoup(page_html, "html.parser")
    options: dict[str, str] = {}
    for checkbox in soup.select('input[name="warehouseIds[]"]'):
        if checkbox.has_attr("disabled"):
            continue
        label = checkbox.find_parent("label")
        name = label.get_text(" ", strip=True) if label else ""
        warehouse_id = str(checkbox.get("value") or "").strip()
        if name and warehouse_id and warehouse_id != "all":
            if warehouse_id in options and options[warehouse_id] != name:
                raise MabangApiError("仓库目录存在冲突记录，请刷新后重新查询")
            options[warehouse_id] = name
    return tuple({"id": identity, "name": name} for identity, name in options.items())


def parse_target_warehouses(
    page_html: str, warehouse_names: Sequence[str] = TARGET_WAREHOUSE_NAMES,
) -> tuple[dict[str, str], ...]:
    if not warehouse_names:
        raise MabangApiError("请至少选择一个仓库")
    matches: dict[str, set[str]] = {name: set() for name in warehouse_names}
    for option in parse_inventory_warehouses(page_html):
        if option["name"] in matches:
            matches[option["name"]].add(option["id"])
    missing = [name for name, ids in matches.items() if not ids]
    if missing:
        raise MabangApiError("账号未找到可查询的仓库：" + "、".join(missing))
    if any(len(ids) != 1 for ids in matches.values()):
        raise MabangApiError("目标仓库存在同名记录，无法安全确定查询范围")
    result = tuple({"id": next(iter(ids)), "name": name} for name, ids in matches.items())
    if len({item["id"] for item in result}) != len(result):
        raise MabangApiError("目标仓库 ID 重复，无法安全确定查询范围")
    return result


def _login(client: MabangClient, username: str, password: str) -> None:
    try:
        client.login(username, password)
    except MabangLoginError as exc:
        raise MabangApiError(str(exc)) from None
    except httpx.TimeoutException:
        raise MabangApiError("马帮登录请求超时，请检查网络后重试") from None
    except httpx.HTTPStatusError as exc:
        raise MabangApiError(f"马帮登录请求失败（HTTP {exc.response.status_code}），请稍后重试") from None
    except httpx.RequestError:
        raise MabangApiError("马帮登录网络连接失败，请检查网络后重试") from None
    except (MabangApiError, httpx.HTTPError):
        raise MabangApiError("马帮登录未能完成认证，请检查当前账号的登录状态") from None


def fetch_lz_inventory_warehouses(username: str, password: str) -> tuple[dict[str, str], ...]:
    with MabangClient(BASE_URL) as client:
        _login(client, username, password)
        _, page_html = _open_inventory_page(client)
    warehouses = parse_inventory_warehouses(page_html)
    if not warehouses:
        raise MabangApiError("账号没有可查询的仓库，请检查商品仓库查看权限")
    return warehouses


def build_inventory_search_payload(warehouse_ids: Sequence[str]) -> list[tuple[str, str]]:
    # AdvanceSearch-form plus the range/paging defaults appended by the page.
    return [
        ("stockOrderby", ""), ("parentCategoryId", ""), ("categoryId", ""),
        ("third_category_id", ""),
        *(("warehouseIds[]", str(value)) for value in warehouse_ids),
        ("stockName", "nameCN"), ("stockNameValue", ""),
        ("inventoryAlertId", "0"), ("livenessType", ""), ("isNewType", ""),
        ("gridcodeStr", ""), ("stockSkuStr", ""),
        ("page", "1"), ("rowsPerPage", "50"), ("warehouseId", ""),
        ("startTime", ""), ("endTime", ""), ("isIdn", "1"),
        ("warehouseIdArr", ""), ("stockQuantitylt", ""), ("stockQuantitygt", ""),
        ("stockWarningQuantitylt", ""), ("stockWarningQuantitygt", ""),
        ("saleAvailableDayslt", ""), ("saleAvailableDaysgt", ""),
    ]


def parse_record_count(page_html: str) -> int:
    soup = BeautifulSoup(page_html, "html.parser")
    first_count = soup.select_one(".semibold")
    text = first_count.get_text(strip=True).replace(",", "") if first_count else ""
    if re.fullmatch(r"\d+", text):
        return int(text)
    match = re.search(r"共\s*([\d,]+)\s*条", soup.get_text(" ", strip=True))
    if match:
        return int(match.group(1).replace(",", ""))
    raise MabangApiError("库存查询未返回可核对的总条数，已停止导出")


def _request(client: MabangClient, method: str, url: str, *, context: str, **kwargs: Any) -> httpx.Response:
    try:
        response = client.client.request(method, url, **kwargs)
        response.raise_for_status()
        return response
    except httpx.HTTPStatusError as exc:
        raise MabangApiError(f"{context}失败（HTTP {exc.response.status_code}）") from None
    except httpx.RequestError:
        # HTTP exception URLs can contain authentication/search tokens.
        raise MabangApiError(f"{context}网络请求失败，请稍后重新运行") from None


def _post_json(client: MabangClient, path: str, payload: list[tuple[str, str]], referer: str) -> dict[str, Any]:
    response = _request(
        client, "POST", PRIVATE_BASE_URL + path, context="库存查询",
        content=urlencode(payload), headers={
            "Referer": referer, "Origin": PRIVATE_BASE_URL,
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "X-Requested-With": "XMLHttpRequest",
        },
    )
    try:
        data = response.json()
    except ValueError:
        raise MabangApiError("库存查询返回非 JSON 响应，请检查登录状态") from None
    if not isinstance(data, dict) or data.get("success") not in (True, 1):
        raise MabangApiError("库存查询失败，请检查账号的库存查询和导出权限")
    return data


def _open_inventory_page(client: MabangClient) -> tuple[str, str]:
    response = _request(client, "GET", BASE_URL + INVENTORY_PAGE_PATH, context="打开库存查询")
    page_html = decode_html_page(response)
    iframe = BeautifulSoup(page_html, "html.parser").find("iframe", id="iframeContent")
    if iframe and iframe.get("src"):
        iframe_url = urljoin(str(response.url), str(iframe["src"]))
        parsed = urlparse(iframe_url)
        if parsed.scheme != "https" or not (parsed.hostname or "").endswith(".mabangerp.com"):
            raise MabangApiError("库存查询页面返回非马帮地址")
        response = _request(client, "GET", iframe_url, context="加载库存查询", headers={"Referer": str(response.url)})
        page_html = decode_html_page(response)
    return str(response.url), page_html


def read_export_table(content: bytes) -> tuple[tuple[str, ...], list[tuple[Any, ...]]]:
    """Read the native spreadsheet, preserving every column and literal value."""
    try:
        if content.startswith(b"PK\x03\x04"):
            workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
            try:
                if len(workbook.worksheets) != 1:
                    raise MabangApiError("库存导出包含多个工作表，无法确认完整数据范围")
                values = list(workbook.worksheets[0].iter_rows(values_only=True))
            finally:
                workbook.close()
        elif content.startswith(OLE2_MAGIC):
            workbook_xls = xlrd.open_workbook(file_contents=content)
            try:
                if workbook_xls.nsheets != 1:
                    raise MabangApiError("库存导出包含多个工作表，无法确认完整数据范围")
                sheet = workbook_xls.sheet_by_index(0)
                values = [tuple(sheet.row_values(index)) for index in range(sheet.nrows)]
            finally:
                workbook_xls.release_resources()
        else:
            raise MabangApiError("库存导出未返回有效 Excel，可能登录失效或没有导出权限")
    except MabangApiError:
        raise
    except Exception:
        raise MabangApiError("库存导出 Excel 无法读取，文件可能不完整") from None
    if not values:
        raise MabangApiError("库存导出文件为空")
    headers = tuple(str(value or "").strip() for value in values[0])
    if not headers or any(not value for value in headers) or len(set(headers)) != len(headers):
        raise MabangApiError("库存导出表头存在空白或重复字段")
    if "仓库" not in headers or not any(name in headers for name in ("库存SKU编号", "库存SKU", "SKU")):
        raise MabangApiError("导出文件不是库存查询明细，缺少仓库或 SKU 字段")
    rows = [tuple(row) for row in values[1:] if any(value not in (None, "") for value in row)]
    if any(len(row) != len(headers) for row in rows):
        raise MabangApiError("库存导出明细列数与表头不一致")
    return headers, rows


def _append_row(sheet: Any, values: Sequence[Any], *, header: bool = False) -> None:
    cells = []
    for value in values:
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"
        if header:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="24364B")
        cells.append(cell)
    sheet.append(cells)


def _export_warehouse(
    client: MabangClient, job: LzInventoryJob, warehouse: dict[str, str],
    referer: str, page_size: int, progress: ProgressCallback | None,
) -> LzWarehouseExport:
    warehouse_name = warehouse["name"]

    def emit(message: str) -> None:
        if progress:
            progress(f"【{warehouse_name}】{message}")

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("lz各仓实时库存")
    sheet.freeze_panes = "A2"
    preview: list[dict[str, Any]] = []
    headers: tuple[str, ...] = ()
    fingerprints: set[str] = set()
    record_count = 0
    output_file: Path | None = None
    temporary_file: Path | None = None
    try:
        search = _post_json(client, SEARCH_PATH, build_inventory_search_payload([warehouse["id"]]), referer)
        token = str(search.get("searchToken") or "").strip()
        if not token:
            raise MabangApiError("库存查询未返回导出令牌，已停止以避免导出其他查询结果")
        pagination_html = str(search.get("pageHtml") or "")
        if not pagination_html.strip():
            pagination = _post_json(client, PAGINATION_PATH, [("page", "1"), ("rowsPerPage", "50")], referer)
            pagination_html = str(pagination.get("pageHtml") or "")
        source_count = parse_record_count(pagination_html)
        # Even an empty warehouse gets a header-only native export. Never
        # invent a schema when the upstream download fails.
        page_count = max(1, math.ceil(source_count / page_size))
        if page_count > MAX_EXPORT_PAGES or source_count > 1048575:
            raise MabangApiError("库存数据超出单个 Excel 工作表或分页上限")
        emit(f"查询共 {source_count:,} 条，需导出 {page_count} 页，每页最多 {page_size:,} 条")
        for page_number in range(1, page_count + 1):
            emit(f"正在下载导出分页 {page_number}/{page_count}")
            response = _request(
                client, "GET", PRIVATE_BASE_URL + EXPORT_PATH, context=f"库存第 {page_number} 页导出",
                params={"mod": "warehouse.doexportwarehousestock", "flag": "1", "page": str(page_number), "showRmbColumn": "0", "searchToken": token},
                headers={"Referer": referer},
            )
            page_headers, rows = read_export_table(response.content)
            expected = min(page_size, source_count - record_count)
            if len(rows) != expected:
                raise MabangApiError(f"第 {page_number} 页条数不完整：预期 {expected} 条，实际 {len(rows)} 条；请重新查询")
            signature = hashlib.sha256(json.dumps(rows, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8")).hexdigest()
            if signature in fingerprints:
                raise MabangApiError(f"第 {page_number} 页与已下载分页重复，已停止合并")
            fingerprints.add(signature)
            if not headers:
                headers = page_headers
                for index, name in enumerate(headers, start=1):
                    sheet.column_dimensions[get_column_letter(index)].width = 40 if "名称" in name else max(14, min(len(name) * 2 + 4, 26))
                _append_row(sheet, headers, header=True)
            elif page_headers != headers:
                raise MabangApiError(f"第 {page_number} 页表头与第一页不一致，已停止合并")
            warehouse_index = headers.index("仓库")
            for row in rows:
                name = str(row[warehouse_index] or "").strip()
                if name != warehouse_name:
                    raise MabangApiError(f"第 {page_number} 页包含非目标仓库，已停止合并")
                _append_row(sheet, row)
                if len(preview) < PREVIEW_ROW_LIMIT:
                    preview.append({key: value.isoformat(sep=" ") if isinstance(value, datetime) else value for key, value in zip(headers, row)})
            record_count += len(rows)
            emit(f"已合并第 {page_number}/{page_count} 页，本页 {len(rows):,} 条，累计 {record_count:,} 条")
        if record_count != source_count:
            raise MabangApiError("合并条数与查询总数不一致，未发布结果")
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{record_count + 1}"
        job.output_dir.mkdir(parents=True, exist_ok=True)
        safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", warehouse_name).strip(" .")[:80] or "仓库"
        output_file = job.output_dir / f"lz各仓实时库存_{safe_name}_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex[:8]}.xlsx"
        temporary_file = output_file.with_suffix(".tmp")
        emit("正在保存并核对合并表格")
        workbook.save(temporary_file)
        # Count the saved cells independently before exposing the final file.
        with temporary_file.open("rb") as saved_file:
            check = load_workbook(saved_file, read_only=True, data_only=False)
            try:
                saved_rows = check.worksheets[0].iter_rows(values_only=True)
                if tuple(next(saved_rows, ())) != headers or sum(1 for _ in saved_rows) != record_count:
                    raise MabangApiError("合并表格回读行数或表头校验失败")
            finally:
                check.close()
        temporary_file.replace(output_file)
        emit(f"完整导出 {record_count:,} 条、{len(headers)} 列")
        emit(f"已保存：{output_file}")
        return LzWarehouseExport(warehouse_name, output_file, record_count, source_count, page_count, headers, tuple(preview))
    finally:
        # openpyxl only removes write-only XML after a successful save. Do not
        # create a writer on login failure, and clean an existing failed stream.
        writer = getattr(sheet, "_writer", None)
        try:
            if writer is not None and not sheet.closed:
                sheet.close()
            workbook.close()
        except Exception:
            emit("库存临时写入流关闭失败")
        finally:
            if writer is not None and Path(writer.out).exists():
                try:
                    writer.cleanup()
                except OSError:
                    emit("库存临时文件清理失败，请检查临时目录写入权限")
        if temporary_file and temporary_file.exists():
            try:
                temporary_file.unlink()
            except OSError:
                emit("合并临时文件清理失败，请检查输出目录写入权限")


def run_lz_inventory_export(job: LzInventoryJob, progress: ProgressCallback | None = None) -> LzInventoryResult:
    def emit(message: str) -> None:
        if progress:
            progress(message)

    exports: list[LzWarehouseExport] = []
    failures: list[dict[str, str]] = []
    with MabangClient(BASE_URL) as client:
        emit("正在通过接口登录马帮")
        _login(client, job.username, job.password)
        referer, page_html = _open_inventory_page(client)
        warehouses = parse_target_warehouses(page_html, job.warehouse_names)
        page_size_match = re.search(r"var\s+pageSize\s*=\s*parseInt\(\s*(\d+)\s*\)", page_html)
        page_size = int(page_size_match.group(1)) if page_size_match else EXPORT_PAGE_SIZE
        if page_size < 1:
            raise MabangApiError("库存导出分页大小无效")
        emit(f"已选择 {len(warehouses)} 个仓库，每个仓库独立查询、合并分页并生成 Excel")
        for index, warehouse in enumerate(warehouses, start=1):
            emit(f"开始处理仓库 {index}/{len(warehouses)}：{warehouse['name']}")
            try:
                exports.append(_export_warehouse(client, job, warehouse, referer, page_size, progress))
            except (IllegalCharacterError, ValueError, TypeError):
                message = "库存导出包含 Excel 无法写入的字段，未生成该仓文件"
                failures.append({"warehouse_name": warehouse["name"], "error": message})
                emit(f"【{warehouse['name']}】导出失败：{message}")
            except (MabangApiError, OSError) as exc:
                # A failed warehouse never publishes a partial workbook, and
                # must not discard successfully verified files from others.
                failures.append({"warehouse_name": warehouse["name"], "error": str(exc)})
                emit(f"【{warehouse['name']}】导出失败：{exc}")
    result = LzInventoryResult(tuple(exports), tuple(failures))
    emit(result.completion_message)
    return result
