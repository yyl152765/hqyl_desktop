from __future__ import annotations

import copy
import csv
import re
import threading
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import StringIO
from pathlib import Path
from typing import Any, Callable, Iterable

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from backend.core.mabang_client import MabangApiError, MabangClient, decode_html_page


BASE_URL = "https://900853.private.mabangerp.com"
REPORT_PAGE_PATH = "/index.php?mod=reports.countryReports"
REPORT_URL = f"{BASE_URL}/index.php"
REPORT_PARAMS = {"mod": "reports.countryReports"}
DEFAULT_ROWS_PER_PAGE = "2000"
ALL_PEOPLE_CONCURRENCY = 20
DEVELOPER_INPUT_NAME = "developid[]"
SALESPERSON_INPUT_NAME = "saleid[]"
SUMMARY_EXPORT_FIELDS = (
    "salesSkuNewId",
    "sale",
    "develop",
    "income",
)
BASE_FORM_DATA: dict[str, Any] = {
    "historyType": "1",
    "platformOrShop": "country",
    "table": "country",
    "timeStart2": "2021-12-01 00:00:00",
    "timeEnd2": "2021-12-31 23:59:59",
    "timeStart3": "2019-12-01 00:00:00",
    "timeEnd3": "2019-12-31 23:59:59",
    "timeKey": "payTime",
    "searchType": "stockSku",
    "operate": "eq",
    "moneyType": "US",
    "page": "1",
    "rowsPerPage": DEFAULT_ROWS_PER_PAGE,
}
TOTAL_SKU_LABELS = {"合计", "总计", "total"}
UNASSIGNED_SALESPERSON = "未分配销售员"
UNASSIGNED_DEVELOPER = "未分配开发员"
MONEY_HEADER = "收入-订单金额（USD）"
MONEY_QUANTUM = Decimal("0.0001")

ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class PersonOption:
    id: str
    name: str


@dataclass(frozen=True)
class DeveloperSalesIncomeQuery:
    username: str
    password: str
    developer_names: tuple[str, ...]
    start_date: str
    end_date: str
    output_dir: Path
    include_all: bool = False


@dataclass(frozen=True)
class DeveloperIncomeSummary:
    developer: PersonOption
    amount: Decimal
    source_row_count: int


@dataclass(frozen=True)
class IncomeDetailRow:
    developer: str
    salesperson: str
    amount: Decimal


@dataclass(frozen=True)
class PersonIncomeTarget:
    role: str
    input_name: str
    option: PersonOption


@dataclass(frozen=True)
class DeveloperSalesIncomeResult:
    developer_count: int
    salesperson_count: int
    source_row_count: int
    missing_developer_names: tuple[str, ...]
    output_file: Path
    output_dir: Path


def safe_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def clean_person_label(value: object) -> str:
    text = safe_text(value).replace("□", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return re.sub(r"\s*>\s*(\d+)\s*$", r">\1", text)


def person_match_key(value: object) -> str:
    """Return a human-name key without collapsing distinct Mabang IDs."""
    text = clean_person_label(value)
    text = re.sub(r">\d+$", "", text).strip()
    text = re.sub(r"^[A-Za-z]+\d+\s*-\s*", "", text).strip()
    text = re.sub(
        r"\s*-\s*(?:[^-]*(?:开发|销售|采购|供应链)[^-]*)$",
        "",
        text,
    ).strip()
    return text.casefold()


def parse_person_options(page_html: str, input_name: str) -> list[PersonOption]:
    soup = BeautifulSoup(page_html or "", "html.parser")
    options: list[PersonOption] = []
    seen_ids: set[str] = set()
    for input_tag in soup.find_all("input"):
        if safe_text(input_tag.get("type")).lower() != "checkbox":
            continue
        if safe_text(input_tag.get("name")) != input_name:
            continue
        person_id = safe_text(input_tag.get("value"))
        label = input_tag.find_parent("label")
        name = clean_person_label(label.get_text(" ", strip=True) if label else "")
        if not person_id or person_id == "all" or person_id in seen_ids:
            continue
        if person_id == "0":
            name = UNASSIGNED_DEVELOPER if input_name == DEVELOPER_INPUT_NAME else UNASSIGNED_SALESPERSON
        if not name or name in {"全部开发员", "全部销售员"}:
            continue
        seen_ids.add(person_id)
        options.append(PersonOption(id=person_id, name=name))
    return options


def resolve_developer_options(
    requested_names: Iterable[str],
    available_options: Iterable[PersonOption],
) -> tuple[tuple[PersonOption, ...], tuple[str, ...]]:
    options = list(available_options)
    selected: list[PersonOption] = []
    selected_ids: set[str] = set()
    missing: list[str] = []
    requested_keys: set[str] = set()

    for requested_name in requested_names:
        display_name = clean_person_label(requested_name)
        match_key = person_match_key(display_name)
        if not match_key or match_key in requested_keys:
            continue
        requested_keys.add(match_key)
        matches = [option for option in options if person_match_key(option.name) == match_key]
        if not matches:
            missing.append(display_name)
            continue
        for option in matches:
            if option.id in selected_ids:
                continue
            selected_ids.add(option.id)
            selected.append(option)
    return tuple(selected), tuple(missing)


def parse_developer_names(value: object) -> tuple[str, ...]:
    if isinstance(value, str):
        raw_names = value.splitlines()
    elif isinstance(value, (list, tuple, set)):
        raw_names = [safe_text(item) for item in value]
    else:
        raw_names = []
    names: list[str] = []
    seen: set[str] = set()
    for item in raw_names:
        name = clean_person_label(item)
        key = person_match_key(name)
        if not name or not key or key in seen:
            continue
        seen.add(key)
        names.append(name)
    return tuple(names)


def parse_boolean(value: object) -> bool:
    if isinstance(value, bool):
        return value
    return safe_text(value).casefold() in {"1", "true", "yes", "on"}


def validate_developer_sales_income_payload(payload: dict[str, Any]) -> DeveloperSalesIncomeQuery:
    username = safe_text(payload.get("username"))
    password = safe_text(payload.get("password"))
    developer_names = parse_developer_names(payload.get("developer_names"))
    start_date = safe_text(payload.get("start_date"))
    end_date = safe_text(payload.get("end_date"))
    output_dir_text = safe_text(payload.get("output_dir"))
    include_all = parse_boolean(payload.get("include_all"))

    if not username:
        raise ValueError("请输入马帮账号")
    if not password:
        raise ValueError("请输入马帮密码")
    if not include_all and not developer_names:
        raise ValueError("请按一行一个姓名输入至少一名开发员")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", start_date):
        raise ValueError("付款开始日期格式必须是 YYYY-MM-DD")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", end_date):
        raise ValueError("付款结束日期格式必须是 YYYY-MM-DD")
    start_at = datetime.strptime(start_date, "%Y-%m-%d")
    end_at = datetime.strptime(end_date, "%Y-%m-%d")
    if start_at > end_at:
        raise ValueError("付款开始日期不能晚于结束日期")
    if not output_dir_text:
        raise ValueError("请选择输出目录")

    return DeveloperSalesIncomeQuery(
        username=username,
        password=password,
        developer_names=developer_names,
        start_date=start_date,
        end_date=end_date,
        output_dir=Path(output_dir_text).expanduser(),
        include_all=include_all,
    )


def fetch_report_person_options(client: MabangClient) -> tuple[list[PersonOption], list[PersonOption]]:
    response = client.client.get(client._url(REPORT_PAGE_PATH))
    response.raise_for_status()
    page_html = decode_html_page(response)
    developers = parse_person_options(page_html, DEVELOPER_INPUT_NAME)
    salespeople = parse_person_options(page_html, SALESPERSON_INPUT_NAME)
    if not developers:
        raise MabangApiError("未从商品销量报表获取到开发员选项，请确认账号权限")
    if not salespeople:
        raise MabangApiError("未从商品销量报表获取到销售员选项，请确认账号权限")
    return developers, salespeople


def build_report_form_data(developer_id: str | None, start_date: date, end_date: date) -> dict[str, Any]:
    form_data = copy.deepcopy(BASE_FORM_DATA)
    form_data.update(
        {
            "timeStart": f"{start_date:%Y-%m-%d} 00:00:00",
            "timeEnd": f"{end_date:%Y-%m-%d} 23:59:59",
            "timeKey": "payTime",
            "moneyType": "US",
        }
    )
    if safe_text(developer_id):
        form_data[DEVELOPER_INPUT_NAME] = safe_text(developer_id)
    return form_data


def decode_csv_content(content: bytes) -> str:
    if not content:
        raise MabangApiError("马帮商品销量报表导出返回空文件")
    if content.lstrip().startswith(b"<"):
        raise MabangApiError("马帮商品销量报表导出返回的不是 CSV 文件")
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8-sig", errors="replace")


def parse_money(value: object, *, context: str) -> Decimal:
    text = safe_text(value).rstrip("\t").replace(",", "").replace("$", "")
    if not text:
        return Decimal("0")
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise MabangApiError(f"{context} 收入-订单金额不是有效数字: {value!r}") from exc
    if not amount.is_finite():
        raise MabangApiError(f"{context} 收入-订单金额不是有效数字: {value!r}")
    return amount


def parse_income_detail_rows(content: bytes, *, context: str) -> list[IncomeDetailRow]:
    reader = csv.DictReader(StringIO(decode_csv_content(content)))
    if reader.fieldnames is None:
        raise MabangApiError(f"{context} 导出结果没有表头")
    reader.fieldnames = [
        safe_text(column).lstrip("\ufeff") if index == 0 else safe_text(column)
        for index, column in enumerate(reader.fieldnames)
    ]
    missing_headers = [
        header for header in ("SKU", "开发员", "销售员", "收入-订单金额") if header not in reader.fieldnames
    ]
    if missing_headers:
        raise MabangApiError(f"{context} 导出结果缺少必需列: {', '.join(missing_headers)}")

    rows: list[IncomeDetailRow] = []
    for row in reader:
        sku = safe_text(row.get("SKU")).lstrip("\ufeff")
        if not sku or sku.casefold() in TOTAL_SKU_LABELS:
            continue
        developer = clean_person_label(row.get("开发员"))
        if not developer or developer in {"--", "无开发员"}:
            developer = UNASSIGNED_DEVELOPER
        salesperson = clean_person_label(row.get("销售员"))
        if not salesperson or salesperson in {"--", "无销售员"}:
            salesperson = UNASSIGNED_SALESPERSON
        amount = parse_money(row.get("收入-订单金额"), context=f"{context} SKU={sku}")
        rows.append(IncomeDetailRow(developer=developer, salesperson=salesperson, amount=amount))
    return rows


def parse_income_rows(content: bytes, *, context: str) -> list[tuple[str, Decimal]]:
    return [
        (row.salesperson, row.amount)
        for row in parse_income_detail_rows(content, context=context)
    ]


def download_developer_income_rows(
    client: MabangClient,
    developer: PersonOption,
    start_date: date,
    end_date: date,
) -> list[tuple[str, Decimal]]:
    export_file = client.download_product_sales_report_csv(
        url=REPORT_URL,
        params=REPORT_PARAMS,
        form_data=build_report_form_data(developer.id, start_date, end_date),
        report_type="countryReports",
        platform_or_shop="country",
        referer=f"{BASE_URL}{REPORT_PAGE_PATH}",
        fields=SUMMARY_EXPORT_FIELDS,
    )
    return parse_income_rows(export_file.content, context=developer.name)


def parse_table_foot_order_income(table_foot: object, *, context: str) -> Decimal:
    soup = BeautifulSoup(safe_text(table_foot), "html.parser")
    income_cells = soup.find_all(attrs={"data-field": "income"})
    if not income_cells:
        return Decimal("0")
    return parse_money(income_cells[0].get_text(" ", strip=True), context=context)


def fetch_person_income_total(
    client: MabangClient,
    target: PersonIncomeTarget,
    start_date: date,
    end_date: date,
) -> Decimal:
    if target.input_name not in {DEVELOPER_INPUT_NAME, SALESPERSON_INPUT_NAME}:
        raise ValueError(f"不支持的人员筛选字段: {target.input_name}")
    form_data = build_report_form_data(None, start_date, end_date)
    form_data.update(
        {
            target.input_name: target.option.id,
            "page": "1",
            "rowsPerPage": "100",
        }
    )
    payload = client._post_sales_report(
        url=REPORT_URL,
        params=REPORT_PARAMS,
        form_data=form_data,
        referer=f"{BASE_URL}{REPORT_PAGE_PATH}",
    )
    if not payload.get("success"):
        raise MabangApiError(f"{target.role}{target.option.name}收入查询失败: {payload}")
    return parse_table_foot_order_income(
        payload.get("tableFoot"),
        context=f"{target.role}={target.option.name}",
    )


def collect_all_people_income_totals(
    *,
    username: str,
    password: str,
    developer_options: Iterable[PersonOption],
    salesperson_options: Iterable[PersonOption],
    start_date: date,
    end_date: date,
    progress: ProgressCallback | None = None,
    max_workers: int = ALL_PEOPLE_CONCURRENCY,
) -> tuple[list[DeveloperIncomeSummary], dict[str, Decimal]]:
    developers = list(developer_options)
    salespeople = list(salesperson_options)
    _unique_option_names(developers, role="开发员")
    _unique_option_names(salespeople, role="销售员")
    targets = [
        PersonIncomeTarget("开发员", DEVELOPER_INPUT_NAME, option)
        for option in developers
    ] + [
        PersonIncomeTarget("销售员", SALESPERSON_INPUT_NAME, option)
        for option in salespeople
    ]
    worker_count = max(1, min(int(max_workers or 1), len(targets)))
    thread_state = threading.local()
    clients: list[MabangClient] = []
    clients_lock = threading.Lock()

    def worker(target: PersonIncomeTarget) -> tuple[PersonIncomeTarget, Decimal]:
        client = getattr(thread_state, "client", None)
        if client is None:
            client = MabangClient(BASE_URL, timeout_seconds=180, verify_ssl=False)
            client.login(username, password)
            thread_state.client = client
            with clients_lock:
                clients.append(client)
        return target, fetch_person_income_total(client, target, start_date, end_date)

    amounts_by_target: dict[tuple[str, str], Decimal] = {}
    completed_count = 0
    progress_step = max(1, min(50, len(targets) // 20 or 1))
    executor = ThreadPoolExecutor(max_workers=worker_count, thread_name_prefix="income-summary")
    futures = [executor.submit(worker, target) for target in targets]
    try:
        for future in as_completed(futures):
            target, amount = future.result()
            amounts_by_target[(target.role, target.option.id)] = amount
            completed_count += 1
            if completed_count % progress_step == 0 or completed_count == len(targets):
                _emit(progress, f"全部人员遍历进度: {completed_count}/{len(targets)}")
    except Exception:
        for future in futures:
            future.cancel()
        raise
    finally:
        executor.shutdown(wait=True, cancel_futures=True)
        for client in clients:
            client.close()

    developer_summaries = [
        DeveloperIncomeSummary(
            developer=option,
            amount=amounts_by_target[("开发员", option.id)],
            source_row_count=0,
        )
        for option in developers
    ]
    salesperson_totals = {
        option.name: amounts_by_target[("销售员", option.id)]
        for option in salespeople
    }
    return developer_summaries, salesperson_totals


def download_all_income_rows(
    client: MabangClient,
    start_date: date,
    end_date: date,
) -> list[IncomeDetailRow]:
    export_file = client.download_product_sales_report_csv(
        url=REPORT_URL,
        params=REPORT_PARAMS,
        form_data=build_report_form_data(None, start_date, end_date),
        report_type="countryReports",
        platform_or_shop="country",
        referer=f"{BASE_URL}{REPORT_PAGE_PATH}",
        fields=SUMMARY_EXPORT_FIELDS,
    )
    return parse_income_detail_rows(export_file.content, context="全部开发员和销售员")


def summarize_all_people(
    rows: Iterable[IncomeDetailRow],
    developer_options: Iterable[PersonOption],
    salesperson_options: Iterable[PersonOption],
) -> tuple[list[DeveloperIncomeSummary], dict[str, Decimal]]:
    developers = list(developer_options)
    salespeople = list(salesperson_options)
    developer_names = _unique_option_names(developers, role="开发员")
    salesperson_names = _unique_option_names(salespeople, role="销售员")
    developer_amounts = {name: Decimal("0") for name in developer_names.values()}
    developer_row_counts = {name: 0 for name in developer_names.values()}
    salesperson_totals = {name: Decimal("0") for name in salesperson_names.values()}
    extra_developers: dict[str, PersonOption] = {}

    for row in rows:
        developer_key = clean_person_label(row.developer).casefold()
        developer_name = developer_names.get(developer_key)
        if developer_name is None:
            developer_name = clean_person_label(row.developer) or UNASSIGNED_DEVELOPER
            developer_key = developer_name.casefold()
            extra_developers.setdefault(
                developer_key,
                PersonOption(id=f"csv:{developer_key}", name=developer_name),
            )
            developer_amounts.setdefault(developer_name, Decimal("0"))
            developer_row_counts.setdefault(developer_name, 0)
        developer_amounts[developer_name] += row.amount
        developer_row_counts[developer_name] += 1

        salesperson_key = clean_person_label(row.salesperson).casefold()
        salesperson_name = salesperson_names.get(salesperson_key)
        if salesperson_name is None:
            salesperson_name = clean_person_label(row.salesperson) or UNASSIGNED_SALESPERSON
            salesperson_totals.setdefault(salesperson_name, Decimal("0"))
        salesperson_totals[salesperson_name] += row.amount

    summaries = [
        DeveloperIncomeSummary(
            developer=option,
            amount=developer_amounts[option.name],
            source_row_count=developer_row_counts[option.name],
        )
        for option in developers
    ]
    for key in sorted(extra_developers):
        option = extra_developers[key]
        summaries.append(
            DeveloperIncomeSummary(
                developer=option,
                amount=developer_amounts[option.name],
                source_row_count=developer_row_counts[option.name],
            )
        )
    return summaries, salesperson_totals


def _unique_option_names(options: Iterable[PersonOption], *, role: str) -> dict[str, str]:
    names: dict[str, str] = {}
    for option in options:
        key = clean_person_label(option.name).casefold()
        if key in names:
            raise MabangApiError(f"马帮{role}选项存在无法区分的同名记录: {option.name}")
        names[key] = option.name
    return names


def _decimal_cell_value(value: Decimal) -> float:
    return float(value.quantize(MONEY_QUANTUM))


def _style_summary_sheet(worksheet) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    total_fill = PatternFill("solid", fgColor="D9EAF7")
    thin_gray = Side(style="thin", color="D9E2F3")
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:B{max(1, worksheet.max_row - 1)}"
    worksheet.column_dimensions["A"].width = 34
    worksheet.column_dimensions["B"].width = 26
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=2):
        for cell in row:
            cell.border = Border(bottom=thin_gray)
            cell.alignment = Alignment(vertical="center")
        row[1].number_format = '$#,##0.0000;[Red]-$#,##0.0000'
    if worksheet.max_row >= 2:
        for cell in worksheet[worksheet.max_row]:
            cell.fill = total_fill
            cell.font = Font(bold=True)
    worksheet.sheet_view.showGridLines = False


def build_summary_workbook(
    output_path: Path,
    developer_summaries: Iterable[DeveloperIncomeSummary],
    salesperson_totals: dict[str, Decimal],
    *,
    require_equal_totals: bool = True,
) -> Path:
    developer_rows = list(developer_summaries)
    developer_total = sum((row.amount for row in developer_rows), Decimal("0"))
    salesperson_total = sum(salesperson_totals.values(), Decimal("0"))
    if require_equal_totals and developer_total != salesperson_total:
        raise MabangApiError(
            "开发员与销售员收入合计不一致: "
            f"开发员={developer_total}, 销售员={salesperson_total}"
        )

    workbook = Workbook()
    developer_sheet = workbook.active
    developer_sheet.title = "开发员"
    developer_sheet.append(["开发员", MONEY_HEADER])
    for summary in developer_rows:
        developer_sheet.append([summary.developer.name, _decimal_cell_value(summary.amount)])
    developer_sheet.append(["合计", _decimal_cell_value(developer_total)])

    salesperson_sheet = workbook.create_sheet("销售员")
    salesperson_sheet.append(["销售员", MONEY_HEADER])
    for salesperson in sorted(salesperson_totals, key=str.casefold):
        salesperson_sheet.append([salesperson, _decimal_cell_value(salesperson_totals[salesperson])])
    salesperson_sheet.append(["合计", _decimal_cell_value(salesperson_total)])

    _style_summary_sheet(developer_sheet)
    _style_summary_sheet(salesperson_sheet)
    workbook.properties.title = "开发与销售收入汇总导出"
    workbook.properties.subject = "马帮商品销量报表，付款日期口径，美元"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return output_path


def run_developer_sales_income_summary(
    job: DeveloperSalesIncomeQuery,
    progress: ProgressCallback | None = None,
) -> DeveloperSalesIncomeResult:
    start_date = datetime.strptime(job.start_date, "%Y-%m-%d").date()
    end_date = datetime.strptime(job.end_date, "%Y-%m-%d").date()
    _emit(progress, f"付款日期: {job.start_date} 至 {job.end_date}")
    _emit(progress, "金额口径: 收入-订单金额（USD）")
    _emit(progress, "正在登录马帮并读取商品销量报表人员选项...")

    developer_summaries: list[DeveloperIncomeSummary] = []
    salesperson_totals: dict[str, Decimal] = {}
    source_row_count = 0
    missing_names: tuple[str, ...] = ()

    with MabangClient(BASE_URL, timeout_seconds=300, verify_ssl=False) as client:
        client.login(job.username, job.password)
        developer_options, salesperson_options = fetch_report_person_options(client)
        _emit(
            progress,
            f"已读取 {len(developer_options)} 个开发员选项、{len(salesperson_options)} 个销售员选项",
        )
        if job.include_all:
            _emit(progress, "全部人员模式：正在按独立人员 ID 遍历开发员和销售员...")
            developer_summaries, salesperson_totals = collect_all_people_income_totals(
                username=job.username,
                password=job.password,
                developer_options=developer_options,
                salesperson_options=salesperson_options,
                start_date=start_date,
                end_date=end_date,
                progress=progress,
            )
            source_row_count = len(developer_summaries) + len(salesperson_totals)
            _emit(
                progress,
                f"全部人员汇总完成: {len(developer_summaries)} 个开发员、"
                f"{len(salesperson_totals)} 个销售员、{source_row_count} 次人员统计",
            )
        else:
            selected_developers, missing_names = resolve_developer_options(
                job.developer_names,
                developer_options,
            )
            if missing_names:
                _emit(progress, f"未匹配到开发员，已跳过: {'、'.join(missing_names)}")
            if not selected_developers:
                raise ValueError("输入姓名均未匹配到马帮商品销量报表中的开发员")
            _emit(progress, f"共匹配 {len(selected_developers)} 个独立开发员 ID，开始逐个统计")

            selected_salesperson_totals: defaultdict[str, Decimal] = defaultdict(lambda: Decimal("0"))
            for index, developer in enumerate(selected_developers, start=1):
                _emit(progress, f"[{index}/{len(selected_developers)}] 正在统计 {developer.name}")
                rows = download_developer_income_rows(client, developer, start_date, end_date)
                developer_amount = sum((amount for _, amount in rows), Decimal("0"))
                developer_summaries.append(
                    DeveloperIncomeSummary(
                        developer=developer,
                        amount=developer_amount,
                        source_row_count=len(rows),
                    )
                )
                source_row_count += len(rows)
                for salesperson, amount in rows:
                    selected_salesperson_totals[salesperson] += amount
                _emit(
                    progress,
                    f"[{index}/{len(selected_developers)}] {developer.name}: "
                    f"{len(rows)} 行，收入订单金额 ${developer_amount.quantize(MONEY_QUANTUM)}",
                )
            salesperson_totals = dict(selected_salesperson_totals)

    scope_label = "_全部人员" if job.include_all else ""
    output_file = job.output_dir / (
        f"开发与销售收入汇总{scope_label}_{start_date:%Y%m%d}_{end_date:%Y%m%d}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    )
    _emit(progress, "正在生成包含“开发员”和“销售员”两个子表的 Excel...")
    build_summary_workbook(
        output_file,
        developer_summaries,
        salesperson_totals,
        require_equal_totals=not job.include_all,
    )
    _emit(progress, f"汇总完成: {output_file}")
    return DeveloperSalesIncomeResult(
        developer_count=len(developer_summaries),
        salesperson_count=len(salesperson_totals),
        source_row_count=source_row_count,
        missing_developer_names=missing_names,
        output_file=output_file,
        output_dir=job.output_dir,
    )


def _emit(progress: ProgressCallback | None, message: str) -> None:
    if progress is not None:
        progress(message)
