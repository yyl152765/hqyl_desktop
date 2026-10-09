from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill


SHEET_NAME = "商品基础信息"
COLUMNS = (
    "店铺名", "商品标题", "SPU ID", "SKC ID", "SKU ID", "SKC货号", "SKU货号",
    "叶子类目名称", "经营站点", "商品状态", "规格1名称", "规格2名称", "申报价格站点",
    "申报价格(USD)", "申报价格状态", "库存", "创建时间",
)
IDENTIFIERS = frozenset({"SPU ID", "SKC ID", "SKU ID", "SKC货号", "SKU货号"})
REQUIRED = frozenset({"商品标题", "SPU ID", "SKC ID", "SKU ID", "商品状态"})


@dataclass(frozen=True)
class SourceInfo:
    path: Path
    headers: tuple[str, ...]
    row_count: int
    sha256: str

    @property
    def added_columns(self) -> list[str]:
        return [name for name in self.headers if name not in COLUMNS]

    @property
    def missing_columns(self) -> list[str]:
        return [name for name in COLUMNS[1:] if name not in self.headers]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sheet(book):
    if SHEET_NAME in book.sheetnames:
        return book[SHEET_NAME]
    candidates = []
    for sheet in book.worksheets:
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if REQUIRED.issubset({str(value or "").strip() for value in first}):
            candidates.append(sheet)
    if len(candidates) != 1:
        raise ValueError("无法唯一确定商品基础信息工作表，请检查真实导出模板")
    return candidates[0]


def _headers(sheet) -> tuple[str, ...]:
    row = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ()))
    while row and row[-1] is None:
        row.pop()
    names = tuple(str(value or "").strip() for value in row)
    if not names or any(not name for name in names):
        raise ValueError("原始文件存在空表头，不能安全对齐字段")
    if len(set(names)) != len(names):
        raise ValueError("原始文件有重复表头，不能安全对齐字段")
    if not REQUIRED.issubset(names):
        raise ValueError("原始文件缺少必要字段：" + "、".join(sorted(REQUIRED - set(names))))
    return names


def _identifier(cell) -> str | None:
    value = cell.value
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("标识符包含布尔值")
    if isinstance(value, float):
        if not value.is_integer():
            raise ValueError("标识符包含小数，无法无损转换")
        value = int(value)
    text = str(value)
    if isinstance(value, int) and re.fullmatch(r"0+", cell.number_format or ""):
        text = text.zfill(len(cell.number_format))
    return text


def _rows(sheet, headers: tuple[str, ...]) -> Iterator[dict[str, Any]]:
    # iter_rows includes filtered and hidden rows. Never use visible-row metadata.
    for number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
        if not any(cell.value is not None for cell in cells):
            continue
        if any(cell.value is not None for cell in cells[len(headers):]):
            raise ValueError(f"第 {number} 行有无表头数据")
        data = {}
        for name, cell in zip(headers, cells):
            if cell.data_type in {"f", "e"}:
                raise ValueError(f"第 {number} 行字段 {name} 包含公式或错误值，请使用平台原始导出")
            data[name] = _identifier(cell) if name in IDENTIFIERS else cell.value
        if str(data.get("商品状态") or "").strip() != "在售中":
            raise ValueError(f"第 {number} 行商品状态不是在售中，请重新确认页面查询条件")
        if not str(data.get("SKU ID") or "").strip():
            raise ValueError(f"第 {number} 行缺少 SKU ID，请检查模板")
        yield data


def inspect_source(path: str | Path) -> SourceInfo:
    path = Path(path).resolve()
    if path.suffix.lower() != ".xlsx" or not path.is_file() or not path.stat().st_size:
        raise ValueError("未获得完整可读的 .xlsx 原始文件")
    try:
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = _sheet(book)
            # Recalculate dimensions while streaming, rather than trusting a stale XML dimension.
            sheet.reset_dimensions()
            headers = _headers(sheet)
            count = sum(1 for _ in _rows(sheet, headers))
        finally:
            book.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("原始 Excel 无法读取，可能下载未完成或返回了错误页面") from exc
    return SourceInfo(path, headers, count, file_sha256(path))


def iter_source_rows(source: SourceInfo) -> Iterator[dict[str, Any]]:
    """Read validated rows without applying Excel visibility or price filters."""
    if file_sha256(source.path) != source.sha256:
        raise ValueError("原始文件在核对过程中发生变化")
    book = load_workbook(source.path, read_only=True, data_only=False)
    try:
        sheet = _sheet(book)
        sheet.reset_dimensions()
        if _headers(sheet) != source.headers:
            raise ValueError("原始文件表头发生变化")
        yield from _rows(sheet, source.headers)
    finally:
        book.close()


def merge_sources(sources: list[tuple[str, SourceInfo]], output: Path) -> dict[str, Any]:
    headers = list(COLUMNS)
    for _, source in sources:
        for name in source.headers:
            if name not in headers:
                headers.append(name)
    if sum(source.row_count for _, source in sources) > 1_048_575:
        raise ValueError("汇总超过 Excel 单工作表行数上限，请减少本批店铺数量")
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.{uuid4().hex}.xlsx")
    book = Workbook(write_only=True)
    sheet = book.create_sheet(SHEET_NAME)
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    for index, name in enumerate(headers, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(index)].width = 48 if name == "商品标题" else 24
    header_cells = []
    for name in headers:
        cell = WriteOnlyCell(sheet, name)
        cell.font = Font(name="Microsoft YaHei", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176B5B")
        header_cells.append(cell)
    sheet.append(header_cells)
    total = 0
    try:
        for store_name, source in sources:
            if not store_name.strip():
                raise ValueError("汇总缺少紫鸟店铺名")
            if file_sha256(source.path) != source.sha256:
                raise ValueError(f"{store_name} 的原始文件在汇总前发生变化")
            original = load_workbook(source.path, read_only=True, data_only=False)
            count = 0
            try:
                source_sheet = _sheet(original)
                source_sheet.reset_dimensions()
                if _headers(source_sheet) != source.headers:
                    raise ValueError("原始文件表头在汇总前发生变化")
                for values in _rows(source_sheet, source.headers):
                    values["店铺名"] = store_name
                    row = []
                    for name in headers:
                        value = values.get(name)
                        cell = WriteOnlyCell(sheet, value)
                        # Exported titles/SKUs starting with '=' remain text, never executable formulas.
                        if isinstance(value, str):
                            cell.data_type = "s"
                        if name in IDENTIFIERS:
                            cell.number_format = "@"
                        row.append(cell)
                    sheet.append(row)
                    count += 1
            finally:
                original.close()
            if count != source.row_count or file_sha256(source.path) != source.sha256:
                raise ValueError(f"{store_name} 的原始文件在汇总中发生变化")
            total += count
        from openpyxl.utils import get_column_letter
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total + 1}"
        book.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=False)
        try:
            rows = check[SHEET_NAME].iter_rows(values_only=True)
            if tuple(next(rows)) != tuple(headers) or sum(1 for _ in rows) != total:
                raise ValueError("汇总文件保存后行数或表头校验失败")
        finally:
            check.close()
        temporary.replace(output)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return {"output_file": str(output), "row_count": total, "columns": headers}
