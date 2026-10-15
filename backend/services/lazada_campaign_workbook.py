"""Filter local Excel prices and optionally fill supported invitation templates."""

from __future__ import annotations

import json
import os
import re
from bisect import bisect_left, bisect_right
from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, localcontext
from pathlib import Path
from typing import Any, Callable
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries

MAX_FILES = 100
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 200
MODES = {"auto", "plain", "invited", "regular"}
COUNTERS = ("total_rows", "updated_rows", "removed_rows", "unparseable_rows", "kept_rows")
_ALIASES = {
    "campaign_price": ("campaignprice", "discountprice", "saleprice", "活动价格", "活动价", "市场活动价格", "市场活动价", "大促价格", "大促价", "折后价", "促销价"),
    "base_price": ("currentproductprice", "currentprice", "originalprice", "salesprice", "sellingprice", "retailprice",
                   "baseprice", "listprice", "商品当前价格", "当前价格", "当前售价", "销售价格", "原售价", "商品原价", "原价", "售价"),
    "recommended_price": ("recommendedprice", "suggestedprice", "推荐价格", "推荐价", "建议价格", "建议价"),
    "maximum_price": ("maximumcampaignprice", "maxcampaignprice", "最高活动价格", "最高活动价", "活动门槛价"),
}
_INSTRUCTION = re.compile(r"required|optional|must|please|do\s+not|必填|选填|填写|说明|不可修改|请勿|请输入", re.I)


@dataclass(frozen=True)
class CampaignWorkbookJob:
    source_files: tuple[Path, ...]
    output_dir: Path
    mode: str = "auto"
    threshold: Decimal = Decimal("0.95")


def _threshold(value: Any) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise ValueError("折扣率必须是大于 0 且不大于 1 的数字，例如 0.95")
    try:
        number = Decimal(str(value).strip())
    except InvalidOperation as exc:
        raise ValueError("折扣率必须是大于 0 且不大于 1 的数字，例如 0.95") from exc
    if not number.is_finite() or not 0 < number <= 1:
        raise ValueError("折扣率必须大于 0 且不大于 1，例如 0.95 表示活动价至少为原价的 95%")
    return number


def validate_campaign_payload(payload: Any) -> CampaignWorkbookJob:
    if not isinstance(payload, Mapping):
        raise ValueError("Excel 折扣筛选参数必须是对象")
    source_files = payload.get("source_files")
    if not isinstance(source_files, (list, tuple)) or not 1 <= len(source_files) <= MAX_FILES:
        raise ValueError(f"请选择 1 至 {MAX_FILES} 份 Excel 表格")
    mode = payload.get("mode", "auto")
    if not isinstance(mode, str) or mode not in MODES:
        raise ValueError("模板类型必须是 auto、plain、invited 或 regular")
    files, seen = [], set()
    for value in source_files:
        if not isinstance(value, (str, Path)) or not str(value).strip():
            raise ValueError("Excel 表格路径必须是非空字符串")
        source = Path(value).expanduser().resolve()
        if source.suffix.lower() not in {".xlsx", ".xlsm"}:
            raise ValueError(f"仅支持 .xlsx 或 .xlsm Excel 表格：{source.name}")
        if not source.is_file():
            raise ValueError(f"Excel 表格不存在：{source.name}")
        if not 0 < source.stat().st_size <= MAX_FILE_BYTES:
            raise ValueError(f"Excel 表格必须非空且不超过 50 MiB：{source.name}")
        key = os.path.normcase(str(source))
        if key not in seen:
            files.append(source)
            seen.add(key)
    output = payload.get("output_dir", str(files[0].parent / "Excel折扣筛选"))
    if not isinstance(output, (str, Path)) or not str(output).strip():
        raise ValueError("请选择有效的结果输出目录")
    output_dir = Path(output).expanduser().resolve()
    if output_dir.exists() and not output_dir.is_dir():
        raise ValueError("结果输出目录不能是文件")
    return CampaignWorkbookJob(tuple(files), output_dir, mode, _threshold(payload.get("threshold", "0.95")))


def _header(value: Any) -> str:
    # Help paragraphs can follow headings in imported English/Chinese exports.
    first_line = str(value or "").lstrip("\ufeff").strip().split("\n", 1)[0]
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", first_line.lower())


def _field(value: Any) -> str | None:
    text = _header(value)
    for field, aliases in _ALIASES.items():
        for alias in aliases:
            if text == alias:
                return field
            if text.startswith(alias):
                tail = text[len(alias):]
                if re.fullmatch(r"(?:(?:mandatory|required|optional|php|usd|currency|必填|选填|菲律宾比索|比索|元|价格单位))+(?:\d+)?", tail):
                    return field
    return None


def _sheet_layout(sheet: Any, requested_mode: str) -> dict[str, Any] | None:
    first_row_values = {_header(cell.value) for cell in sheet[1]}
    field_guide = (bool(first_row_values & {"field", "字段", "字段名"})
                   and bool(first_row_values & {"description", "说明", "描述", "字段说明"}))
    for row in sheet.iter_rows(max_row=min(10, sheet.max_row), max_col=min(200, sheet.max_column)):
        found: dict[str, list[int]] = {}
        for cell in row:
            field = _field(cell.value)
            if field:
                found.setdefault(field, []).append(cell.column)
        if "campaign_price" not in found:
            continue
        if "base_price" not in found:
            # The template may include a separate Field/Description help sheet.
            if field_guide:
                continue
            raise ValueError(f"工作表「{sheet.title}」第 {row[0].row} 行缺少原价/当前价格/售价表头")
        duplicates = [key for key in ("base_price", "campaign_price") if len(found[key]) != 1]
        if duplicates:
            raise ValueError(f"工作表「{sheet.title}」表头重复或有歧义：{'、'.join(duplicates)}")
        columns = {key: values[0] for key, values in found.items()}
        mode = requested_mode
        if mode == "auto":
            base_header = _header(sheet.cell(row[0].row, columns["base_price"]).value)
            campaign_header = str(sheet.cell(row[0].row, columns["campaign_price"]).value or "").lower()
            regular_marker = (base_header.startswith(("salesprice", "sellingprice", "retailprice", "销售价格", "售价"))
                              and ("mandatory" in campaign_header or "必填" in campaign_header))
            invited_marker = base_header.startswith(("currentprice", "currentproductprice", "当前价格", "商品当前价格"))
            if regular_marker and invited_marker:
                raise ValueError(f"工作表「{sheet.title}」同时包含普通和特邀模板标记，请明确选择模板类型")
            if regular_marker:
                mode = "regular"
            elif invited_marker and "recommended_price" in columns:
                mode = "invited"
            else:
                mode = "plain"
        if mode == "invited" and "recommended_price" not in columns:
            raise ValueError(f"工作表「{sheet.title}」缺少特邀活动的推荐价格表头")
        used_fields = {"base_price", "campaign_price"}
        if mode == "invited":
            used_fields.add("recommended_price")
        elif mode == "regular" and "maximum_price" in columns:
            used_fields.add("maximum_price")
        duplicates = [key for key in used_fields if len(found[key]) != 1]
        if duplicates:
            raise ValueError(f"工作表「{sheet.title}」表头重复或有歧义：{'、'.join(duplicates)}")
        columns = {key: columns[key] for key in used_fields}
        header_row = row[0].row
        first_data_row = header_row + 1
        if mode == "regular" and first_data_row <= sheet.max_row:
            values = [sheet.cell(first_data_row, col).value for col in range(1, sheet.max_column + 1)]
            price_values = [sheet.cell(first_data_row, columns[key]).value for key in ("campaign_price", "base_price")]
            price_columns = {column for matches in found.values() for column in matches}
            other_values = [value for col, value in enumerate(values, 1) if col not in price_columns]
            campaign_help = str(price_values[0] or "")
            help_only_row = (not any(value is not None and str(value).strip() for value in other_values)
                             and not str(price_values[1] or "").strip()
                             and len(campaign_help) >= 40
                             and re.search(r"campaign\s+price\s+is|price.*(?:submit|campaign|eligible)|活动价.*(?:填写|是|说明)", campaign_help, re.I))
            typed_instruction_row = (sum(bool(_INSTRUCTION.search(str(value or ""))) for value in values) >= 2
                                     and all(isinstance(value, str) and _INSTRUCTION.search(value) for value in price_values))
            if help_only_row or typed_instruction_row:
                first_data_row += 1
        return {"worksheet": sheet.title, "mode": mode, "header_row": header_row,
                "first_data_row": first_data_row, "columns": columns}
    return None


def _layouts(book: Any, mode: str) -> list[dict[str, Any]]:
    for sheet in book.worksheets:
        if sheet.max_row > MAX_ROWS or sheet.max_column > MAX_COLUMNS:
            raise ValueError(f"工作表「{sheet.title}」超过 {MAX_ROWS:,} 行或 {MAX_COLUMNS} 列；请清除多余格式区域或拆分表格后导入")
    layouts = [layout for sheet in book.worksheets if (layout := _sheet_layout(sheet, mode))]
    if not layouts:
        raise ValueError("前 10 行未找到完整价格表头；需要原价和活动价，特邀模板还需要推荐价格")
    return layouts


def _currency(value: Any) -> Decimal:
    if value is None or isinstance(value, (bool, date, datetime)):
        raise ValueError("缺少有效价格")
    text = format(value, "f") if isinstance(value, Decimal) else str(value).strip()
    text = re.sub(r"^(?:PHP|USD|₱|\$|￥|¥)\s*", "", text, flags=re.I)
    text = re.sub(r"\s*(?:PHP|USD|比索|元)$", "", text, flags=re.I)
    if not re.fullmatch(r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", text):
        raise ValueError("价格不是有效正数")
    try:
        number = Decimal(text.replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError("价格不是有效正数") from exc
    if not number.is_finite() or number <= 0 or len(number.as_tuple().digits) > 15:
        raise ValueError("价格必须是有限的正数，且不超过 Excel 的 15 位有效数字")
    return number


def _read_price(cell: Any, *, recommended: bool = False) -> Decimal:
    if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("=")):
        raise ValueError(f"{cell.coordinate} 含价格公式，请在 Excel/WPS 中将价格粘贴为数值后再处理")
    if cell.data_type == "e":
        raise ValueError(f"{cell.coordinate} 含 Excel 错误")
    value = cell.value
    if recommended and isinstance(value, str):
        value = re.sub(r"^(?:<=|≤)\s*", "", value.strip())
    try:
        return _currency(value)
    except ValueError as exc:
        raise ValueError(f"{cell.coordinate} {exc}") from exc


def _data_rows(sheet: Any, layout: dict[str, Any]):
    for row in sheet.iter_rows(min_row=layout["first_data_row"]):
        if any(cell.value is not None and str(cell.value).strip() for cell in row):
            yield row[0].row


def _prices(sheet: Any, layout: dict[str, Any], row: int) -> tuple[Decimal, Decimal, bool]:
    columns = layout["columns"]
    base = _read_price(sheet.cell(row, columns["base_price"]))
    campaign_cell = sheet.cell(row, columns["campaign_price"])
    # Do not replace formula prices blindly, even if another column has a number.
    if campaign_cell.data_type == "f":
        _read_price(campaign_cell)
    selected_cell, recommended = campaign_cell, False
    if layout["mode"] == "invited":
        candidate = sheet.cell(row, columns["recommended_price"])
        if candidate.value is not None and str(candidate.value).strip():
            selected_cell, recommended = candidate, True
    elif layout["mode"] == "regular" and (campaign_cell.value is None or not str(campaign_cell.value).strip()):
        if "maximum_price" in columns:
            selected_cell = sheet.cell(row, columns["maximum_price"])
        else:
            raise ValueError(f"{campaign_cell.coordinate} 缺少平台最高活动价；请重新导出并勾选「Fill in with Maximum Campaign Price / 使用最高活动价格填充」")
    elif layout["mode"] == "plain" and (campaign_cell.value is None or not str(campaign_cell.value).strip()):
        raise ValueError(f"{campaign_cell.coordinate} 缺少活动价，请在源 Excel 中填写后再筛选")
    return base, _read_price(selected_cell, recommended=recommended), selected_cell is not campaign_cell


def _book(source: Path):
    try:
        return load_workbook(source, keep_links=True, keep_vba=source.suffix.lower() == ".xlsm")
    except Exception as exc:
        raise ValueError("无法读取 Excel 表格，请确认文件未损坏、未加密且是有效 Excel 文件") from exc


def _close_book(book: Any) -> None:
    book.close()
    if book.vba_archive is not None:
        book.vba_archive.close()


def _base_result(source: Path) -> dict[str, Any]:
    return {"source_file": str(source), "name": source.name, "status": "failed", "ok": False, "mode": "",
            "processed_file": "", "error": "", "sheets": [], **dict.fromkeys(COUNTERS, 0)}


def inspect_campaign_files(payload: Any) -> dict[str, Any]:
    job = validate_campaign_payload(payload)
    results = []
    for source in job.source_files:
        result = _base_result(source)
        book = None
        try:
            book = _book(source)
            for layout in _layouts(book, job.mode):
                sheet, invalid, missing, rows = book[layout["worksheet"]], 0, 0, 0
                for row in _data_rows(sheet, layout):
                    rows += 1
                    cell = sheet.cell(row, layout["columns"]["campaign_price"])
                    if cell.value is None or not str(cell.value).strip():
                        missing += 1
                    try:
                        _prices(sheet, layout, row)
                    except ValueError:
                        invalid += 1
                result["sheets"].append({**layout, "total_rows": rows, "unparseable_rows": invalid,
                                          "missing_price_rows": missing})
            result.update(status="success", ok=True, mode=_result_mode(result["sheets"]),
                          total_rows=sum(s["total_rows"] for s in result["sheets"]),
                          unparseable_rows=sum(s["unparseable_rows"] for s in result["sheets"]),
                          missing_price_rows=sum(s["missing_price_rows"] for s in result["sheets"]))
            result["warnings"] = (["无法计算的行将在结果模板中移出，并写入异常明细。"]
                                  if result["unparseable_rows"] else [])
        except Exception as exc:
            result["error"] = str(exc)
        finally:
            if book is not None:
                _close_book(book)
        results.append(result)
    success = sum(result["status"] == "success" for result in results)
    return {"file_count": len(results), "success_count": success,
            "failed_count": len(results) - success, "files": results}


def _result_mode(sheets: list[dict[str, Any]]) -> str:
    modes = {sheet["mode"] for sheet in sheets}
    return next(iter(modes)) if len(modes) == 1 else "mixed"


def _atomic_save(book: Any, output: Path) -> None:
    temporary = output.with_name(f".{output.stem}.{uuid4().hex}.tmp{output.suffix}")
    try:
        book.save(temporary)
        os.replace(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def export_campaign_template(payload: Any) -> dict[str, Any]:
    """Export an empty, immediately importable local Excel price template."""
    if not isinstance(payload, Mapping):
        raise ValueError("模板导出参数必须是对象")
    value = payload.get("output_dir")
    if not isinstance(value, (str, Path)) or not str(value).strip():
        raise ValueError("请选择有效的模板保存目录")
    output_dir = Path(value).expanduser().resolve()
    if output_dir.exists() and not output_dir.is_dir():
        raise ValueError("模板保存目录不能是文件")
    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"Excel折扣筛选模板_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex}.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "商品数据"
    sheet.append(["SKU", "商品名称", "原价", "活动价", "推荐价"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:E1"
    for column, width in {"A": 26, "B": 44, "C": 20, "D": 20, "E": 20}.items():
        sheet.column_dimensions[column].width = width
    sheet.column_dimensions["A"].number_format = "@"
    for column in ("C", "D", "E"):
        sheet.column_dimensions[column].number_format = "0.00"
    guide = book.create_sheet("字段说明")
    guide.append(["字段", "说明"])
    notes = [
        ("SKU", "选填。按文本填写，可保留前导零。"),
        ("商品名称", "选填。"),
        ("原价", "必填正数，也支持原售价、商品原价、Original Price、Base Price、List Price等表头。"),
        ("活动价", "必填正数，也支持折后价、促销价、Discount Price、Sale Price等表头。"),
        ("推荐价", "选填。通用筛选只比较原价和活动价，不会读取推荐价来覆盖或补填活动价。"),
        ("筛选规则", "折扣率=活动价÷原价。低于指定最低折扣率的行移出结果，等于最低值时保留。"),
        ("异常处理", "缺价、非正数或价格公式等行移出结果，记录到单独的异常明细中；源文件不改动。"),
        ("填写位置", "请从商品数据工作表第2行开始填写商品。本模板不含示例商品，字段说明不参与筛选。"),
    ]
    for note in notes:
        guide.append(note)
    guide.freeze_panes = "A2"
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 112
    for worksheet in book:
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B87")
    try:
        _atomic_save(book, output)
    finally:
        book.close()
    return {"ok": True, "output_file": str(output), "path": str(output)}


def _remove_rows(sheet: Any, rows: list[int]) -> None:
    rows = sorted(set(rows))
    old_dimensions = {row: dim for row, dim in sheet.row_dimensions.items()}
    removed = set(rows)
    # Repeated delete_rows rescans and shifts the entire remaining grid for
    # each interval. Rebuild openpyxl's sparse cell map once so alternating
    # exclusions cost O(stored cells * log(excluded rows)), preserving the
    # original Cell objects, styles, comments, and independent formulas.
    retained_cells = {}
    for (old_row, column), cell in sheet._cells.items():
        if old_row in removed:
            continue
        new_row = old_row - bisect_left(rows, old_row)
        cell.row = new_row
        if getattr(cell, "hyperlink", None) is not None:
            cell.hyperlink.ref = cell.coordinate
        retained_cells[(new_row, column)] = cell
    sheet._cells = retained_cells
    sheet._current_row = max((row for row, _ in retained_cells), default=0)
    # Cell movement also requires explicit row metadata and range updates.
    sheet.row_dimensions.clear()
    for old_row, dimension in old_dimensions.items():
        if old_row not in removed:
            new_row = old_row - bisect_left(rows, old_row)
            dimension.index = new_row
            sheet.row_dimensions[new_row] = dimension
    def shrink(ref: str) -> str:
        left, top, right, bottom = range_boundaries(ref)
        top -= bisect_left(rows, top)
        bottom -= bisect_right(rows, bottom)
        return f"{get_column_letter(left)}{top}:{get_column_letter(right)}{max(top, bottom)}"
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = shrink(sheet.auto_filter.ref)
    for table in sheet.tables.values():
        table.ref = shrink(table.ref)
        if table.autoFilter and table.autoFilter.ref:
            table.autoFilter.ref = shrink(table.autoFilter.ref)
    for validation in list(sheet.data_validations.dataValidation):
        refs = []
        for ref in validation.sqref.ranges:
            _, top, _, bottom = range_boundaries(str(ref))
            if bottom - top + 1 > bisect_right(rows, bottom) - bisect_left(rows, top):
                refs.append(shrink(str(ref)))
        if refs:
            validation.sqref = " ".join(refs)
        else:
            sheet.data_validations.dataValidation.remove(validation)
    # Preserve merged headings and adjust any surviving merged row ranges.
    merges = list(sheet.merged_cells.ranges)
    sheet.merged_cells.ranges.clear()
    for merged in merges:
        if not any(merged.min_row <= row <= merged.max_row for row in rows):
            sheet.merge_cells(shrink(str(merged)))


def _audit_row(source: Path, sheet: Any, row: int, status: str, reason: str,
               layout: dict[str, Any], base: Decimal | None = None, price: Decimal | None = None,
               source_values: list[Any] | None = None) -> list[Any]:
    values = source_values if source_values is not None else [cell.value for cell in sheet[row]]
    columns = layout["columns"]
    raw_price = sheet.cell(row, columns["campaign_price"]).value
    raw_base = sheet.cell(row, columns["base_price"]).value
    return [source.name, str(source), sheet.title, row, layout["mode"], status, reason,
            str(price) if price is not None else ("" if raw_price is None else str(raw_price)),
            str(base) if base is not None else ("" if raw_base is None else str(raw_base)),
            str(price / base) if price is not None and base is not None else "",
            json.dumps(values, ensure_ascii=False, default=str)]


def _process(source: Path, output: Path, mode: str, threshold: Decimal, audit: list[list[Any]]) -> dict[str, Any]:
    result = _base_result(source)
    book = _book(source)
    local_audit = []
    try:
        layouts = _layouts(book, mode)
        deletion_plans = {}
        for layout in layouts:
            sheet, remove = book[layout["worksheet"]], []
            column_count = sheet.max_column
            summary = {**layout, **dict.fromkeys(COUNTERS, 0)}
            for row in _data_rows(sheet, layout):
                summary["total_rows"] += 1
                source_values = [sheet.cell(row, column).value for column in range(1, column_count + 1)]
                try:
                    base, price, updated = _prices(sheet, layout, row)
                except ValueError as exc:
                    summary["unparseable_rows"] += 1
                    remove.append(row)
                    local_audit.append(_audit_row(source, sheet, row, "异常", str(exc), layout, source_values=source_values))
                    continue
                if updated:
                    sheet.cell(row, layout["columns"]["campaign_price"]).value = int(price) if price == price.to_integral() else float(price)
                    summary["updated_rows"] += 1
                with localcontext() as context:
                    context.prec = max(50, len(threshold.as_tuple().digits) + 20)
                    below_threshold = price < base * threshold
                if below_threshold:
                    summary["removed_rows"] += 1
                    remove.append(row)
                    local_audit.append(_audit_row(source, sheet, row, "折扣筛除", f"活动价/原价低于 {threshold}", layout, base, price, source_values))
                else:
                    summary["kept_rows"] += 1
            deletion_plans[sheet.title] = set(remove)
            result["sheets"].append(summary)
        if any(deletion_plans.values()):
            # Inspect stored cells to avoid allocating large blank dimension areas.
            references = [f"{sheet.title}!{cell.coordinate}" for sheet in book for cell in sheet._cells.values()
                          if cell.row not in deletion_plans.get(sheet.title, set()) and cell.data_type == "f"
                          and not re.fullmatch(r"=[\d\s.+*/^%()\-]+", str(cell.value))]
            if references:
                raise ValueError("表格含保留的引用公式，删行可能改变引用关系；请先将公式粘贴为数值后处理：" + "、".join(references[:5]))
        for sheet in book:
            if rows := deletion_plans.get(sheet.title):
                _remove_rows(sheet, list(rows))
        _atomic_save(book, output)
        audit.extend(local_audit)
        result.update(status="success", ok=True, processed_file=str(output), mode=_result_mode(result["sheets"]))
        for key in COUNTERS:
            result[key] = sum(sheet[key] for sheet in result["sheets"])
        return result
    finally:
        _close_book(book)


def _text_row(sheet: Any, values: list[Any]) -> None:
    sheet.append(values)
    # Audit source text can start with '='. Store it as text, never as a formula.
    row = sheet._current_row
    for column, value in enumerate(values, 1):
        if isinstance(value, str):
            sheet.cell(row, column).data_type = "s"


def _report(files: list[dict[str, Any]], audit: list[list[Any]], output: Path, threshold: Decimal) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "处理汇总"
    _text_row(sheet, ["源文件", "结果", "模板类型", "商品行", "填充行", "折扣筛除行", "异常行", "保留行", "结果文件", "错误", "最低活动价/原价"])
    for result in files:
        _text_row(sheet, [result["source_file"], result["status"], result["mode"],
                          *(result[key] for key in COUNTERS), result["processed_file"], result["error"], str(threshold)])
    detail = book.create_sheet("筛除与异常明细")
    _text_row(detail, ["源文件名", "源文件路径", "工作表", "原始行号", "模板类型", "处理分类", "原因", "活动价", "原价", "活动价/原价", "源行数据 JSON"])
    for row in audit:
        _text_row(detail, row)
    for worksheet in book:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="245B87")
        for column in range(1, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 22 if column < 7 else 35
    try:
        _atomic_save(book, output)
    finally:
        book.close()


def run_campaign_batch(job: CampaignWorkbookJob, progress: Callable[[str], None] | None = None) -> dict[str, Any]:
    # Recheck direct callers as well as bridge callers before creating outputs.
    job = validate_campaign_payload({"source_files": job.source_files, "output_dir": job.output_dir,
                                     "mode": job.mode, "threshold": job.threshold})
    output_dir = job.output_dir / f"Excel折扣筛选_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex[:8]}"
    output_dir.mkdir(parents=True, exist_ok=False)
    files, audit = [], []
    total = len(job.source_files)
    for index, source in enumerate(job.source_files, 1):
        if progress:
            progress(f"[{index}/{total}] 筛选 Excel 表格：{source.name}")
        output = output_dir / f"{index:03d}_{source.stem}_已处理{source.suffix.lower()}"
        try:
            result = _process(source, output, job.mode, job.threshold, audit)
        except Exception as exc:
            result = _base_result(source)
            result["error"] = str(exc)
        files.append(result)
        if progress:
            progress(f"[{index}/{total}] {source.name}：{'已完成' if result['status'] == 'success' else '处理失败'}")
    report = output_dir / "Excel折扣筛选汇总与异常明细.xlsx"
    output_files = [result["processed_file"] for result in files if result["status"] == "success"]
    completion_errors = []
    try:
        _report(files, audit, report, job.threshold)
        output_files.append(str(report))
        for result in files:
            result["audit_file"] = str(report)
    except Exception as exc:
        completion_errors.append(f"汇总与异常明细导出失败：{exc}")
    archive = output_dir / "Excel折扣筛选结果.zip"
    temporary = archive.with_name(f".{archive.name}.tmp")
    try:
        if output_files:
            with ZipFile(temporary, "w", ZIP_DEFLATED) as bundle:
                for path in output_files:
                    bundle.write(path, Path(path).name)
            os.replace(temporary, archive)
            output_files.append(str(archive))
    except Exception as exc:
        completion_errors.append(f"ZIP 打包失败：{exc}")
    finally:
        temporary.unlink(missing_ok=True)
    success = sum(result["status"] == "success" for result in files)
    return {"is_complete": success == total and not completion_errors, "file_count": total, "success_count": success,
            "failed_count": total - success, **{key: sum(result[key] for result in files) for key in COUNTERS},
            "output_dir": str(output_dir), "output_file": str(report) if report.exists() else "",
            "archive_file": str(archive) if archive.exists() else "", "completion_message": "；".join(completion_errors),
            "output_files": output_files, "files": files}
