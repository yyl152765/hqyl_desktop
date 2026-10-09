from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.services.vietnam_income_collection import load_manifest, report_is_completed, write_manifest


ProgressCallback = Callable[[str], None]

REQUIRED_MABANG_REPORTS = (
    "income_detail",
    "income_summary",
    "refund_part_01_10",
    "refund_part_11_20",
    "refund_part_21_end",
    "refunds_merged",
)

ZINIAO_REPORT_TYPES = ("income", "ads", "affiliate", "ads_credit")
DEFAULT_CHANNEL_RATE = Decimal("10")
STORE_REFUND_FIXED_ADD = Decimal("4")
ADJUSTMENT_DIVISOR = Decimal("0.92")

EXPENSE_COLUMNS = (
    "支出-成本",
    "支出-转帐费",
    "支出-平台费",
    "支出-交易手续费",
    "支出-佣金",
    "支出-服务费",
    "支出-广告费",
    "支出-头程费",
    "支出-FBA费用",
    "支出-优惠券",
    "支出-VAT税费",
    "支出-其他费用",
    "支出-其他支出",
    "支出-测评费",
)


@dataclass(frozen=True)
class VietnamFinalReconciliationJob:
    manifest_path: Path
    profit_rates: dict[str, Decimal]
    profit_rates_text: str = ""


def validate_vietnam_final_reconciliation_payload(
    payload: dict[str, Any] | None,
) -> VietnamFinalReconciliationJob:
    data = dict(payload or {})
    manifest_path = Path(str(data.get("manifest_path") or "").strip())
    manifest = load_manifest(manifest_path)
    if str(manifest.get("status") or "") not in {"mabang_ready", "final_ready"}:
        raise ValueError("请先完成紫鸟数据采集、马帮收支校验和退款合并")
    if str(((manifest.get("mabang") or {}).get("validation") or {}).get("status") or "") != "passed":
        raise ValueError("马帮收支校验未通过，不能生成最终核对文件")
    _assert_required_sources(manifest_path, manifest)

    profit_rates_text = str(data.get("profit_rates_text") or data.get("profit_rates") or "").strip()
    profit_rates = parse_profit_rates(profit_rates_text)

    # 回退：调用方未显式传入利润率文本时，读取 manifest 中已保存的利润率配置。
    # 前端将利润率保存到 config 后仅传 manifest_path，若此处不回退会导致“填了利润率仍被阻断生成”。
    if not profit_rates:
        config_rates = (manifest.get("config") or {}).get("profit_rates") or {}
        for store, value in config_rates.items():
            cleaned = _clean_store_name(store)
            if not cleaned:
                continue
            try:
                # config 中已是归一化后的小数（如 "0.125"），直接按小数解析，避免二次百分比换算。
                profit_rates[cleaned] = _decimal(value)
            except Exception:
                continue

    has_refund_stores = False
    mabang_reports = (manifest.get("mabang") or {}).get("reports") or {}
    refunds_path = _report_file(manifest_path, mabang_reports.get("refunds_merged", {}))
    if refunds_path and refunds_path.is_file():
        try:
            for row in read_csv_dicts(refunds_path):
                if _clean_store_name(row.get("店铺")):
                    has_refund_stores = True
                    break
        except:
            pass

    if not profit_rates and has_refund_stores:
        raise ValueError("请填写上月利润率，每行格式：店铺名称=利润率，例如 张三-越南001=12.5%")

    return VietnamFinalReconciliationJob(
        manifest_path=manifest_path,
        profit_rates=profit_rates,
        profit_rates_text=profit_rates_text,
    )


def normalize_store_summary(row: dict[str, Any]) -> dict[str, Any]:
    normalized = {
        "store_name": row.get("店铺名称", ""),
        "matched_ziniao_store": row.get("匹配紫鸟店铺", ""),
        "order_count": int(row.get("订单数") or 0),
        "adjusted_receivable": _decimal_text(_decimal(row.get("收入-应收货款-改"))),
        "all_expense": _decimal_text(_decimal(row.get("所有支出项"))),
        "shipping_adjusted": _decimal_text(_decimal(row.get("运费-改"))),
        "package_adjusted": _decimal_text(_decimal(row.get("支出-包材费-改"))),
        "actual_refund": _decimal_text(_decimal(row.get("实际退款金额RMB"))),
        "ads": _decimal_text(_decimal(row.get("广告"))),
        "actual_ads": _decimal_text(_decimal(row.get("实际广告"))),
        "affiliate": _decimal_text(_decimal(row.get("联盟"))),
        "actual_affiliate": _decimal_text(_decimal(row.get("实际联盟"))),
        "ads_credit": _decimal_text(_decimal(row.get("广告补贴"))),
        "actual_ads_credit": _decimal_text(_decimal(row.get("实际广告补贴"))),
        "final_ads": _decimal_text(_decimal(row.get("最终广告花费"))),
        "profit": _decimal_text(_decimal(row.get("利润"))),
        "profit_rate": _decimal_text(_decimal(row.get("利润率"))),
        "status": "normal",
        "exception_count": 0,
    }
    normalized["total_expense"] = _decimal_text(
        Decimal(normalized["all_expense"])
        + Decimal(normalized["actual_refund"])
        + Decimal(normalized["actual_affiliate"])
        + Decimal(normalized["final_ads"])
    )
    return normalized


def run_vietnam_final_reconciliation(
    job: VietnamFinalReconciliationJob,
    progress: ProgressCallback | None = None,
) -> dict[str, Any]:
    manifest = load_manifest(job.manifest_path)
    run_dir = job.manifest_path.parent
    output_dir = run_dir / "output" / "final"
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    final_section = manifest.setdefault("final", {})
    final_section["status"] = "running"
    final_section["started_at"] = _now_text()
    final_section["error_message"] = ""
    _touch_manifest(manifest)
    write_manifest(job.manifest_path, manifest)

    try:
        mabang_reports = (manifest.get("mabang") or {}).get("reports") or {}
        income_detail_path = _report_file(job.manifest_path, mabang_reports["income_detail"])
        refunds_path = _report_file(job.manifest_path, mabang_reports["refunds_merged"])
        if income_detail_path is None or refunds_path is None:
            raise RuntimeError("最终处理缺少马帮收支明细或退款合并文件")

        _emit(progress, "正在读取马帮收支明细并计算全表平均汇率")
        average_rate, income_headers, income_row_count = average_exchange_rate(income_detail_path)
        _emit(progress, f"全表平均汇率：{_decimal_text(average_rate)}，收支明细 {income_row_count} 行")

        _emit(progress, "正在合并紫鸟补贴订单和广告/联盟费用")
        subsidy_rows, subsidy_by_order = load_subsidy_rows(run_dir, manifest)
        ziniao_metrics = load_ziniao_store_metrics(run_dir, manifest)

        _emit(progress, "正在处理马帮退款并应用店铺利润率")
        refund_rows, refund_summary, missing_rate_stores = transform_refunds(refunds_path, job.profit_rates)
        if missing_rate_stores:
            sample = "、".join(missing_rate_stores[:12])
            suffix = "……" if len(missing_rate_stores) > 12 else ""
            raise ValueError(f"以下退款店铺未填写上月利润率：{sample}{suffix}")

        workbook_path = output_dir / f"越南Shopee收支报表_{manifest.get('period')}_{stamp}.xlsx"
        summary_path = output_dir / f"final_reconciliation_summary__{stamp}.json"
        unmatched_path = output_dir / f"unmatched_records__{stamp}.csv"
        log_path = output_dir / f"processing_log__{stamp}.json"

        _emit(progress, "正在生成最终核对工作簿")
        result = build_final_workbook(
            workbook_path,
            income_detail_path,
            income_headers,
            average_rate,
            subsidy_by_order,
            subsidy_rows,
            refund_rows,
            refund_summary,
            ziniao_metrics,
            manifest,
        )

        unmatched_rows = result["unmatched_rows"]
        write_dict_csv(unmatched_path, unmatched_rows)

        stores = []
        for r in result["store_summaries"]:
            norm = normalize_store_summary(r)
            store_exceptions = [
                ur for ur in unmatched_rows
                if _clean_store_name(ur.get("店铺名称")) == _clean_store_name(norm["store_name"])
            ]
            if store_exceptions:
                norm["status"] = "exception"
                norm["exception_count"] = len(store_exceptions)
            stores.append(norm)

        totals = {
            "order_count": sum(int(s["order_count"]) for s in stores),
            "adjusted_receivable": _decimal_text(sum(Decimal(s["adjusted_receivable"]) for s in stores)),
            "all_expense": _decimal_text(sum(Decimal(s["all_expense"]) for s in stores)),
            "actual_refund": _decimal_text(sum(Decimal(s["actual_refund"]) for s in stores)),
            "actual_affiliate": _decimal_text(sum(Decimal(s["actual_affiliate"]) for s in stores)),
            "final_ads": _decimal_text(sum(Decimal(s["final_ads"]) for s in stores)),
        }
        all_expense_dec = Decimal(totals["all_expense"])
        actual_refund_dec = Decimal(totals["actual_refund"])
        actual_affiliate_dec = Decimal(totals["actual_affiliate"])
        final_ads_dec = Decimal(totals["final_ads"])
        adjusted_receivable_dec = Decimal(totals["adjusted_receivable"])

        total_expense_dec = all_expense_dec + actual_refund_dec + actual_affiliate_dec + final_ads_dec
        profit_dec = adjusted_receivable_dec - total_expense_dec
        profit_rate_dec = profit_dec / adjusted_receivable_dec if adjusted_receivable_dec else Decimal("0")

        totals["total_expense"] = _decimal_text(total_expense_dec)
        totals["profit"] = _decimal_text(profit_dec)
        totals["profit_rate"] = _decimal_text(profit_rate_dec)

        affected_store_count = len(set(
            _clean_store_name(ur.get("店铺名称"))
            for ur in unmatched_rows
            if ur.get("店铺名称")
        ))
        exception_summary = {
            "total_count": len(unmatched_rows),
            "blocking_count": 0,
            "affected_store_count": affected_store_count,
            "amount": _decimal_text(sum(_decimal(ur.get("金额")) for ur in unmatched_rows)),
        }

        summary = {
            "generated_at": _now_text(),
            "period": manifest.get("period"),
            "manifest_path": str(job.manifest_path),
            "output_file": str(workbook_path),
            "average_exchange_rate": _decimal_text(average_rate),
            "income_row_count": income_row_count,
            "subsidy_row_count": len(subsidy_rows),
            "subsidy_matched_order_count": result["subsidy_matched_order_count"],
            "refund_row_count": len(refund_rows),
            "refund_store_count": len(refund_summary),
            "store_summary_count": len(result["store_summaries"]),
            "unmatched_count": len(unmatched_rows),
            "store_refund_fixed_add": _decimal_text(STORE_REFUND_FIXED_ADD),
            "exchange_rate_scope": "full_table_average",
            "channel_rate_rule": "extract N元/KG from 物流渠道; fallback 10",
            "profit_formula": "收入-应收货款-改 - 所有支出项 - 实际退款金额RMB - 实际联盟 - 最终广告花费",
            "totals": totals,
            "stores": stores,
            "exception_summary": exception_summary,
        }
        summary_path.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
        log_path.write_text(
            json.dumps(
                {
                    "generated_at": _now_text(),
                    "profit_rates": {key: _decimal_text(value) for key, value in sorted(job.profit_rates.items())},
                    "ziniao_metrics": {
                        key: {metric: _decimal_text(amount) for metric, amount in value.items()}
                        for key, value in sorted(ziniao_metrics.items())
                    },
                    "unmatched_rows": unmatched_rows[:500],
                    "unmatched_truncated": len(unmatched_rows) > 500,
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

        manifest = load_manifest(job.manifest_path)
        final_section = manifest.setdefault("final", {})
        final_section.update(
            {
                "status": "ready",
                "finished_at": _now_text(),
                "output_file": _relative(run_dir, workbook_path),
                "summary_file": _relative(run_dir, summary_path),
                "unmatched_file": _relative(run_dir, unmatched_path),
                "processing_log": _relative(run_dir, log_path),
                "average_exchange_rate": _decimal_text(average_rate),
                "refund_store_count": len(refund_summary),
                "unmatched_count": len(unmatched_rows),
            }
        )
        manifest["status"] = "final_ready"
        _touch_manifest(manifest)
        write_manifest(job.manifest_path, manifest)
        _emit(progress, f"最终收支报表已生成：{workbook_path}")
        return {
            "manifest_path": str(job.manifest_path),
            "run_dir": str(run_dir),
            "output_file": str(workbook_path),
            "output_dir": str(output_dir),
            "summary_file": str(summary_path),
            "unmatched_file": str(unmatched_path),
            "income_row_count": income_row_count,
            "store_summary_count": len(result["store_summaries"]),
            "refund_row_count": len(refund_rows),
            "unmatched_count": len(unmatched_rows),
            "failed_count": 0,
        }
    except Exception as exc:
        manifest = load_manifest(job.manifest_path)
        final_section = manifest.setdefault("final", {})
        final_section.update(
            {
                "status": "failed",
                "finished_at": _now_text(),
                "error_message": str(exc),
            }
        )
        _touch_manifest(manifest)
        write_manifest(job.manifest_path, manifest)
        raise


def parse_profit_rates(text: str) -> dict[str, Decimal]:
    result: dict[str, Decimal] = {}
    for raw_line in str(text or "").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" in line:
            name, value = line.split("=", 1)
        elif "\t" in line:
            name, value = line.rsplit("\t", 1)
        elif "," in line:
            name, value = line.rsplit(",", 1)
        elif "，" in line:
            name, value = line.rsplit("，", 1)
        else:
            raise ValueError(f"上月利润率格式不正确：{line}")
        store = _clean_store_name(name)
        if not store:
            continue
        result[store] = parse_profit_rate(value)
    return result


def parse_profit_rate(value: Any) -> Decimal:
    text = str(value or "").strip().replace("％", "%")
    if not text:
        raise ValueError("利润率不能为空")
    is_percent = text.endswith("%")
    text = text.rstrip("%").strip()
    number = _plain_decimal(text)
    if is_percent or abs(number) > Decimal("1"):
        number = number / Decimal("100")
    return number


def extract_channel_rate(channel: Any) -> Decimal:
    text = str(channel or "")
    match = re.search(
        r"(?<![\d.])([0-9]+(?:\.[0-9]+)?)\s*(?:元|块|¥|RMB)?\s*/?\s*(?:KG|kg|Kg|公斤|千克)",
        text,
    )
    if not match:
        return DEFAULT_CHANNEL_RATE
    value = _decimal(match.group(1))
    return value if value > 0 else DEFAULT_CHANNEL_RATE


def calculate_adjusted_shipping(order_number: Any, weight: Any, channel: Any) -> Decimal:
    if str(order_number or "").strip().endswith("_1"):
        return Decimal("0")
    return _money(_decimal(weight) / Decimal("1000") * extract_channel_rate(channel))


def average_exchange_rate(path: Path) -> tuple[Decimal, list[str], int]:
    total = Decimal("0")
    count = 0
    row_count = 0
    headers: list[str] = []
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        headers = list(reader.fieldnames or [])
        for row in reader:
            row_count += 1
            value = _decimal(row.get("汇率"))
            if value > 0:
                total += value
                count += 1
    return (total / count if count else Decimal("0")), headers, row_count


def load_subsidy_rows(
    run_dir: Path,
    manifest: dict[str, Any],
) -> tuple[list[dict[str, str]], dict[str, Decimal]]:
    rows: list[dict[str, str]] = []
    by_order: dict[str, Decimal] = defaultdict(Decimal)
    for store in manifest.get("stores") or []:
        report = (store.get("reports") or {}).get("income") or {}
        path = _path_from_report(run_dir, report)
        if path is None:
            continue
        for row in read_csv_dicts(path):
            item = dict(row)
            item.setdefault("店铺名称", str(store.get("store_name") or ""))
            rows.append(item)
            order = _clean_order(item.get("订单编号"))
            if order:
                by_order[order] += _decimal(item.get("补贴金额"))
    return rows, dict(by_order)


def load_ziniao_store_metrics(run_dir: Path, manifest: dict[str, Any]) -> dict[str, dict[str, Decimal]]:
    metrics: dict[str, dict[str, Decimal]] = {}
    for store in manifest.get("stores") or []:
        store_name = _clean_store_name(store.get("store_name"))
        reports = store.get("reports") or {}
        metrics[store_name] = {
            "ads": sum_expense_from_report(run_dir, reports.get("ads"), ("花费", "Expense", "expense")),
            "affiliate": sum_expense_from_report(
                run_dir,
                reports.get("affiliate"),
                ("expense", "Expense", "花费", "佣金", "商品佣金"),
            ),
            "ads_credit": sum_expense_from_report(run_dir, reports.get("ads_credit"), ("金额", "Amount", "amount")),
        }
    return metrics


def sum_expense_from_report(
    run_dir: Path,
    report: dict[str, Any] | None,
    column_candidates: tuple[str, ...],
) -> Decimal:
    path = _path_from_report(run_dir, report if isinstance(report, dict) else {})
    if path is None or path.stat().st_size == 0:
        return Decimal("0")
    rows = read_csv_rows(path)
    if not rows:
        return Decimal("0")
    header_index = -1
    column_index = -1
    candidate_keys = {_normalize_header(item) for item in column_candidates}
    for idx, row in enumerate(rows[:80]):
        for col_idx, cell in enumerate(row):
            if _normalize_header(cell) in candidate_keys:
                header_index = idx
                column_index = col_idx
                break
        if header_index >= 0:
            break
    if header_index < 0 or column_index < 0:
        return Decimal("0")
    total = Decimal("0")
    for row in rows[header_index + 1 :]:
        if column_index < len(row):
            total += _decimal(row[column_index])
    return total


def transform_refunds(
    path: Path,
    profit_rates: dict[str, Decimal],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Decimal]], list[str]]:
    rows: list[dict[str, Any]] = []
    summary: dict[str, dict[str, Decimal]] = defaultdict(lambda: {"refund": Decimal("0"), "count": Decimal("0")})
    seen_store: set[str] = set()
    missing: set[str] = set()
    for row in read_csv_dicts(path):
        item = dict(row)
        store = _clean_store_name(item.get("店铺"))
        refund_rmb = _decimal(item.get("退款金额RMB"))
        rate = profit_rates.get(store)
        if rate is None:
            missing.add(store or "未命名店铺")
            rate = Decimal("0")
        fixed_add = STORE_REFUND_FIXED_ADD if store not in seen_store else Decimal("0")
        seen_store.add(store)
        refund_without_fixed = _money(refund_rmb * rate)
        actual_refund = _money(refund_without_fixed + fixed_add)
        item["店铺"] = store
        item["上月利润率"] = _decimal_text(rate)
        item["实际退款金额RMB_不含固定加数"] = _decimal_text(refund_without_fixed)
        item["店铺固定加数"] = _decimal_text(fixed_add)
        item["实际退款金额RMB"] = _decimal_text(actual_refund)
        rows.append(item)
        summary[store]["refund"] += actual_refund
        summary[store]["count"] += Decimal("1")
    return rows, dict(summary), sorted(missing)


def build_final_workbook(
    workbook_path: Path,
    income_detail_path: Path,
    income_headers: list[str],
    average_rate: Decimal,
    subsidy_by_order: dict[str, Decimal],
    subsidy_rows: list[dict[str, str]],
    refund_rows: list[dict[str, Any]],
    refund_summary: dict[str, dict[str, Decimal]],
    ziniao_metrics: dict[str, dict[str, Decimal]],
    manifest: dict[str, Any],
) -> dict[str, Any]:
    workbook = Workbook(write_only=True)

    summary_sheet = workbook.create_sheet("核对汇总")
    income_sheet = workbook.create_sheet("收支报表")
    subsidy_sheet = workbook.create_sheet("补贴订单")
    refund_sheet = workbook.create_sheet("退款")
    unmatched_sheet = workbook.create_sheet("未匹配清单")
    apply_default_widths(summary_sheet, 18)
    apply_default_widths(income_sheet, len(income_headers) + 4)
    apply_default_widths(subsidy_sheet, 12)
    apply_default_widths(refund_sheet, 48)
    apply_default_widths(unmatched_sheet, 8)

    output_headers = build_income_output_headers(income_headers)
    income_sheet.append(output_headers)

    store_stats: dict[str, dict[str, Any]] = defaultdict(_empty_store_stats)
    subsidy_matched_orders: set[str] = set()
    with income_detail_path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        for row in reader:
            transformed = transform_income_row(row, output_headers, average_rate, subsidy_by_order)
            income_sheet.append([transformed.get(header, "") for header in output_headers])
            update_store_stats(store_stats, transformed)
            order = _clean_order(row.get("订单编号"))
            if order and order in subsidy_by_order:
                subsidy_matched_orders.add(order)

    write_rows_sheet(subsidy_sheet, subsidy_rows)
    write_rows_sheet(refund_sheet, refund_rows)

    unmatched_rows: list[dict[str, Any]] = []
    for order, amount in sorted(subsidy_by_order.items()):
        if order not in subsidy_matched_orders:
            unmatched_rows.append({"类型": "补贴订单未匹配", "店铺名称": "", "订单编号": order, "金额": _decimal_text(amount), "说明": "补贴订单号未在马帮收支明细中找到"})

    store_summaries = build_store_summaries(store_stats, refund_summary, ziniao_metrics, average_rate, unmatched_rows)
    write_summary_sheet(summary_sheet, store_summaries, manifest, average_rate)
    write_rows_sheet(unmatched_sheet, unmatched_rows)

    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(workbook_path)
    return {
        "store_summaries": store_summaries,
        "unmatched_rows": unmatched_rows,
        "subsidy_matched_order_count": len(subsidy_matched_orders),
    }


def build_income_output_headers(headers: list[str]) -> list[str]:
    output: list[str] = []
    for header in headers:
        output.append(header)
        if header == "收入-应收货款":
            output.extend(["补贴", "收入-应收货款-改"])
        if header == "预估运费":
            output.append("运费-改")
        if header == "支出-包材费":
            output.append("支出-包材费-改")
    return output


def transform_income_row(
    row: dict[str, Any],
    output_headers: list[str],
    average_rate: Decimal,
    subsidy_by_order: dict[str, Decimal],
) -> dict[str, Any]:
    order = _clean_order(row.get("订单编号"))
    subsidy = subsidy_by_order.get(order, Decimal("0"))
    item_total = _decimal(row.get("收入-应收货款"))
    adjusted_item_total = _money(item_total + subsidy * average_rate)
    shipping_adjusted = calculate_adjusted_shipping(order, row.get("订单重量"), row.get("物流渠道"))
    package_adjusted = _money(_decimal(row.get("支出-包材费")) * _decimal(row.get("销量")))
    transformed = {header: row.get(header, "") for header in output_headers}
    transformed["补贴"] = _decimal_text(subsidy)
    transformed["收入-应收货款-改"] = _decimal_text(adjusted_item_total)
    transformed["运费-改"] = _decimal_text(shipping_adjusted)
    transformed["支出-包材费-改"] = _decimal_text(package_adjusted)
    return transformed


def _empty_store_stats() -> dict[str, Any]:
    return {
        "orders": set(),
        "income_adjusted": Decimal("0"),
        "all_expense": Decimal("0"),
        "shipping_adjusted": Decimal("0"),
        "package_adjusted": Decimal("0"),
    }


def update_store_stats(stats: dict[str, dict[str, Any]], row: dict[str, Any]) -> None:
    store = _clean_store_name(row.get("店铺名称"))
    values = stats[store]
    order = _clean_order(row.get("订单编号"))
    if order:
        values["orders"].add(order)
    values["income_adjusted"] += _decimal(row.get("收入-应收货款-改"))
    shipping = _decimal(row.get("运费-改"))
    package = _decimal(row.get("支出-包材费-改"))
    values["shipping_adjusted"] += shipping
    values["package_adjusted"] += package
    total_expense = shipping + package
    for column in EXPENSE_COLUMNS:
        total_expense += _decimal(row.get(column))
    values["all_expense"] += total_expense


def build_store_summaries(
    store_stats: dict[str, dict[str, Any]],
    refund_summary: dict[str, dict[str, Decimal]],
    ziniao_metrics: dict[str, dict[str, Decimal]],
    average_rate: Decimal,
    unmatched_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    used_ziniao: set[str] = set()
    for store in sorted(store_stats):
        values = store_stats[store]
        ziniao_store = match_ziniao_store(store, ziniao_metrics)
        if ziniao_store:
            used_ziniao.add(ziniao_store)
        metrics = ziniao_metrics.get(ziniao_store or "", {})
        ads = metrics.get("ads", Decimal("0"))
        affiliate = metrics.get("affiliate", Decimal("0"))
        ads_credit = metrics.get("ads_credit", Decimal("0"))
        actual_ads = _money(ads * average_rate / ADJUSTMENT_DIVISOR) if ADJUSTMENT_DIVISOR else Decimal("0")
        actual_affiliate = _money(affiliate * average_rate / ADJUSTMENT_DIVISOR) if ADJUSTMENT_DIVISOR else Decimal("0")
        actual_ads_credit = _money(ads_credit * average_rate / ADJUSTMENT_DIVISOR) if ADJUSTMENT_DIVISOR else Decimal("0")
        final_ads = _money(actual_ads - actual_ads_credit)
        actual_refund = (refund_summary.get(store) or {}).get("refund", Decimal("0"))
        income_adjusted = values["income_adjusted"]
        all_expense = values["all_expense"]
        profit = _money(income_adjusted - all_expense - actual_refund - actual_affiliate - final_ads)
        profit_rate = profit / income_adjusted if income_adjusted else Decimal("0")
        rows.append(
            {
                "店铺名称": store,
                "匹配紫鸟店铺": ziniao_store or "",
                "订单数": len(values["orders"]),
                "收入-应收货款-改": _decimal_text(income_adjusted),
                "所有支出项": _decimal_text(all_expense),
                "运费-改": _decimal_text(values["shipping_adjusted"]),
                "支出-包材费-改": _decimal_text(values["package_adjusted"]),
                "实际退款金额RMB": _decimal_text(actual_refund),
                "广告": _decimal_text(ads),
                "实际广告": _decimal_text(actual_ads),
                "联盟": _decimal_text(affiliate),
                "实际联盟": _decimal_text(actual_affiliate),
                "广告补贴": _decimal_text(ads_credit),
                "实际广告补贴": _decimal_text(actual_ads_credit),
                "最终广告花费": _decimal_text(final_ads),
                "利润": _decimal_text(profit),
                "利润率": _decimal_text(profit_rate),
            }
        )
    for ziniao_store in sorted(set(ziniao_metrics) - used_ziniao):
        metrics = ziniao_metrics[ziniao_store]
        if any(value for value in metrics.values()):
            unmatched_rows.append(
                {
                    "类型": "紫鸟店铺未匹配马帮店铺",
                    "店铺名称": ziniao_store,
                    "订单编号": "",
                    "金额": _decimal_text(sum(metrics.values(), Decimal("0"))),
                    "说明": "广告/联盟/广告补贴费用未能匹配到马帮收支明细店铺",
                }
            )
    return rows


def match_ziniao_store(mabang_store: str, ziniao_metrics: dict[str, dict[str, Decimal]]) -> str:
    normalized_mabang = _normalize_store_for_match(mabang_store)
    exact = [store for store in ziniao_metrics if _normalize_store_for_match(store) == normalized_mabang]
    if len(exact) == 1:
        return exact[0]
    candidates = []
    for store in ziniao_metrics:
        normalized = _normalize_store_for_match(store)
        if len(normalized) >= 4 and (normalized in normalized_mabang or normalized_mabang in normalized):
            candidates.append(store)
    return candidates[0] if len(candidates) == 1 else ""


def write_summary_sheet(sheet: Any, rows: list[dict[str, Any]], manifest: dict[str, Any], average_rate: Decimal) -> None:
    sheet.append(["越南 Shopee 收支报表"])
    sheet.append(["核对月份", manifest.get("period", ""), "全表平均汇率", _decimal_text(average_rate)])
    sheet.append(["物流渠道规则", "从物流渠道文本提取 N元/KG；缺失按 10", "退款固定加数", "每个店铺 +4"])
    sheet.append([])
    write_rows_sheet(sheet, rows)


def write_rows_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    if not rows:
        sheet.append(["无数据"])
        return
    headers = list(rows[0].keys())
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])


def apply_default_widths(sheet: Any, column_count: int) -> None:
    for col_idx in range(1, min(int(column_count or 0), 80) + 1):
        sheet.column_dimensions[get_column_letter(col_idx)].width = 16 if col_idx <= 4 else 12


def read_csv_dicts(path: Path) -> list[dict[str, str]]:
    rows = read_csv_rows(path)
    if not rows:
        return []
    header_index = 0
    while header_index < len(rows) and not any(str(cell or "").strip() for cell in rows[header_index]):
        header_index += 1
    if header_index >= len(rows):
        return []
    headers = [str(cell or "").strip() for cell in rows[header_index]]
    result: list[dict[str, str]] = []
    for row in rows[header_index + 1 :]:
        if not any(str(cell or "").strip() for cell in row):
            continue
        item = {headers[idx]: str(row[idx] if idx < len(row) else "") for idx in range(len(headers)) if headers[idx]}
        result.append(item)
    return result


def read_csv_rows(path: Path) -> list[list[str]]:
    if not path.exists() or path.stat().st_size == 0:
        return []
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    return [list(row) for row in csv.reader(text.splitlines())]


def write_dict_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers = list(rows[0].keys()) if rows else ["类型", "店铺名称", "订单编号", "金额", "说明"]
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def _assert_required_sources(manifest_path: Path, manifest: dict[str, Any]) -> None:
    stores = manifest.get("stores") or []
    if not stores:
        raise ValueError("批次没有店铺信息")
    for store in stores:
        reports = store.get("reports") or {}
        for report_type in ZINIAO_REPORT_TYPES:
            if not report_is_completed(manifest_path, reports.get(report_type)):
                raise ValueError(f"紫鸟数据未完成：{store.get('store_name')} / {report_type}")
    mabang_reports = (manifest.get("mabang") or {}).get("reports") or {}
    for report_type in REQUIRED_MABANG_REPORTS:
        if not report_is_completed(manifest_path, mabang_reports.get(report_type)):
            raise ValueError(f"马帮数据未完成：{report_type}")


def _report_file(manifest_path: Path, report: dict[str, Any]) -> Path | None:
    file_text = str(report.get("file") or "").strip()
    if not file_text:
        return None
    path = Path(file_text)
    if not path.is_absolute():
        path = manifest_path.parent / path
    return path if path.is_file() else None


def _path_from_report(run_dir: Path, report: dict[str, Any] | None) -> Path | None:
    if not isinstance(report, dict):
        return None
    file_text = str(report.get("file") or "").strip()
    if not file_text:
        return None
    path = Path(file_text)
    if not path.is_absolute():
        path = run_dir / path
    return path if path.is_file() else None


def _clean_order(value: Any) -> str:
    return str(value or "").strip().lstrip("\t").strip()


def _clean_store_name(value: Any) -> str:
    return " ".join(str(value or "").strip().lstrip("\t").strip().split())


def _normalize_store_for_match(value: Any) -> str:
    text = _clean_store_name(value).casefold()
    text = re.sub(r"[\s\-_./()（）]+", "", text)
    return text


def _normalize_header(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "")).strip().casefold()


def _decimal(value: Any) -> Decimal:
    text = str(value or "").strip().lstrip("\t").strip()
    if text in {"", "-", "--", "nan", "None"}:
        return Decimal("0")
    text = text.replace(",", "")
    text = re.sub(r"[^0-9.\-%]", "", text)
    if not text or text in {"-", ".", "%"}:
        return Decimal("0")
    is_percent = text.endswith("%")
    text = text.rstrip("%")
    try:
        result = Decimal(text or "0")
    except InvalidOperation:
        return Decimal("0")
    return result / Decimal("100") if is_percent else result


def _plain_decimal(value: Any) -> Decimal:
    text = str(value or "").strip().replace(",", "")
    text = re.sub(r"[^0-9.\-]", "", text)
    if not text or text in {"-", "."}:
        return Decimal("0")
    try:
        return Decimal(text)
    except InvalidOperation:
        return Decimal("0")


def _money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def _decimal_text(value: Any) -> str:
    number = _decimal(value) if not isinstance(value, Decimal) else value
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _relative(run_dir: Path, path: Path) -> str:
    return str(path.relative_to(run_dir))


def _now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _touch_manifest(manifest: dict[str, Any]) -> None:
    manifest["updated_at"] = _now_text()


def _emit(progress: ProgressCallback | None, message: str) -> None:
    if progress is not None:
        progress(message)
