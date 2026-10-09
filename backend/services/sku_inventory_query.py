from __future__ import annotations

import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence
from urllib.parse import urlencode, urljoin

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.core.mabang_client import MabangApiError, MabangClient, decode_html_page


BASE_URL = "https://900853.private.mabangerp.com"
PRIVATE_BASE_URL = "https://private-amz.mabangerp.com"
STOCK_LIST_PAGE_PATH = "/index.php?mod=stock.list&searchStatus=3"
STOCK_LIST_API_URL = f"{PRIVATE_BASE_URL}/index.php?mod=stock.getStockList"
DEFAULT_ROWS_PER_PAGE = 500
MAX_PAGE_COUNT = 500
PREVIEW_ROW_LIMIT = 500
OUTPUT_COLUMNS = (
    "开发人员",
    "库存SKU",
    "商品中文名称",
    "SKU创建时间",
    "仓库",
    "仓库ID",
    "可用库存量",
    "预测日销量",
    "当前可售天数",
    "未发货数",
    "在途量",
)
TIME_CREATED_KEYS = (
    "timeCreatedShowTimezone",
    "timeCreated",
    "Stock_timeCreated",
    "createTime",
    "createdAt",
)
FORECAST_DAILY_SALES_KEYS = (
    "predictDailySales",
    "predictDailySale",
    "forecastDailySales",
    "forecastDaySale",
    "daysales",
    "daySales",
    "dailySales",
    "avgDailySales",
)
UNSHIPPED_COMPONENT_KEYS = (
    ("unshipped", "waitingQuantity"),
    ("fbaWaitingQuantity",),
    ("transferred_warehouse_quantity", "transferredWarehouseQuantity"),
    ("manual_outbound_quantity", "manualOutboundQuantity"),
)
TRANSIT_COMPONENT_KEYS = (
    "shippingQuantity",
    "processingQuantity",
    "allotShippingQuantity",
    "hwc_in_transit_quantity",
    "manual_inbound_quantity",
)

ProgressCallback = Callable[[str], None]


@dataclass(frozen=True)
class DeveloperOption:
    id: str
    name: str

    def to_dict(self) -> dict[str, str]:
        return {"id": self.id, "name": self.name}


@dataclass(frozen=True)
class SkuInventoryQuery:
    username: str
    password: str
    developer_ids: tuple[str, ...]
    start_date: str
    end_date: str
    output_dir: Path
    rows_per_page: int = DEFAULT_ROWS_PER_PAGE
    query_mode: str = "inventory"
    liveness_types: tuple[str, ...] = ("1", "2")


@dataclass(frozen=True)
class SkuInventoryResult:
    records: tuple[dict[str, Any], ...]
    developer_count: int
    sku_count: int
    warehouse_count: int
    output_file: Path
    query_mode: str = "inventory"
    scope_warehouses: tuple[dict[str, str], ...] = ()
    missing_sku_count: int = 0
    unclassified_warehouses: tuple[dict[str, str], ...] = ()


def safe_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def first_value(mapping: dict[str, Any], keys: Sequence[str]) -> object:
    lower_map = {safe_text(key).lower(): value for key, value in mapping.items()}
    for key in keys:
        if key in mapping and mapping.get(key) not in (None, ""):
            return mapping.get(key)
        value = lower_map.get(key.lower())
        if value not in (None, ""):
            return value
    return ""


def stock_row_identity(row: dict[str, Any]) -> str:
    stock_key = safe_text(first_value(row, ("stockId", "id", "stock_id"))) or safe_text(
        first_value(row, ("stockSku", "sku"))
    )
    if stock_key:
        return stock_key
    serialized = json.dumps(row, ensure_ascii=False, sort_keys=True, default=str, separators=(",", ":"))
    return f"json:{hashlib.sha256(serialized.encode('utf-8')).hexdigest()}"


def parse_pagination_state(data: dict[str, Any]) -> tuple[int | None, int | None]:
    page_html = safe_text(data.get("pageHtml"))
    if not page_html:
        return None, None
    page_text = BeautifulSoup(page_html, "html.parser").get_text(" ", strip=True)
    match = re.search(r"(\d+)\s*/\s*(\d+)\s*页", page_text)
    if not match:
        return None, None
    current_page = int(match.group(1))
    total_pages = int(match.group(2))
    if current_page < 1 or total_pages < 1:
        return None, None
    return current_page, total_pages


def parse_number(value: object) -> float:
    text = safe_text(value).replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except (TypeError, ValueError):
        return 0.0


def normalize_number(value: object) -> int | float:
    number = parse_number(value)
    return int(number) if number.is_integer() else number


def optional_number(value: object) -> int | float | str:
    text = safe_text(value)
    if not text or text in {"--", "None", "null"}:
        return "--"
    normalized = text.replace(",", "")
    try:
        number = float(normalized)
    except (TypeError, ValueError):
        return text
    return int(number) if number.is_integer() else number


def warehouse_unshipped_count(warehouse_row: dict[str, Any]) -> int | float:
    """Return the warehouse-level total shown in Mabang's 未发货 column."""
    total = sum(
        parse_number(first_value(warehouse_row, component_keys))
        for component_keys in UNSHIPPED_COMPONENT_KEYS
    )
    return normalize_number(total)


def warehouse_transit_inventory(warehouse_row: dict[str, Any]) -> int | float:
    """Return the five-part warehouse total shown in Mabang's 在途量 column."""
    total = sum(parse_number(first_value(warehouse_row, (key,))) for key in TRANSIT_COMPONENT_KEYS)
    return normalize_number(total)


def parse_datetime(value: object) -> datetime | None:
    text = re.sub(r"\s+", " ", safe_text(value)).replace("/", "-")
    if not text:
        return None
    match = re.search(r"\d{4}-\d{1,2}-\d{1,2}(?:\s+\d{1,2}:\d{1,2}(?::\d{1,2})?)?", text)
    if not match:
        return None
    normalized = match.group(0)
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(normalized, fmt)
        except ValueError:
            continue
    return None


def normalize_created_at(value: object) -> str:
    parsed = parse_datetime(value)
    if parsed is None:
        return safe_text(value)
    if parsed.hour == 0 and parsed.minute == 0 and parsed.second == 0:
        return parsed.strftime("%Y-%m-%d")
    return parsed.strftime("%Y-%m-%d %H:%M:%S")


def normalize_developer_name(value: object) -> str:
    text = re.sub(r"\s+", " ", safe_text(value).replace("□", " ")).strip()
    text = re.sub(r"\s*>\s*\d+\s*$", "", text).strip()
    return text


def is_developer_name(value: object) -> bool:
    return normalize_developer_name(value).endswith("开发")


def parse_developer_options(stock_page_html: str) -> list[DeveloperOption]:
    soup = BeautifulSoup(stock_page_html, "html.parser")
    options: list[DeveloperOption] = []
    seen_ids: set[str] = set()
    for input_tag in soup.find_all("input"):
        if safe_text(input_tag.get("type")).lower() != "checkbox":
            continue
        if safe_text(input_tag.get("name")) != "developerIdIN[]":
            continue
        developer_id = safe_text(input_tag.get("value"))
        label = input_tag.find_parent("label")
        name = normalize_developer_name(label.get_text(" ", strip=True) if label else "")
        if not developer_id or developer_id in {"0", "all"} or developer_id in seen_ids:
            continue
        if not name or "全部开发员" in name or "无开发员" in name or not is_developer_name(name):
            continue
        seen_ids.add(developer_id)
        options.append(DeveloperOption(id=developer_id, name=name))
    options.sort(key=lambda item: item.name.casefold())
    return options


def open_stock_page(client: MabangClient) -> tuple[str, str]:
    response = client.client.get(client._url(STOCK_LIST_PAGE_PATH))
    response.raise_for_status()
    page_url = str(response.url)
    page_html = decode_html_page(response)
    soup = BeautifulSoup(page_html, "html.parser")
    iframe = soup.find("iframe", id="iframeContent") or soup.find("iframe")
    iframe_src = safe_text(iframe.get("src") if iframe else "")
    if not iframe_src:
        return page_url, page_html
    iframe_url = urljoin(page_url, iframe_src)
    iframe_response = client.client.get(iframe_url, headers={"Referer": page_url})
    iframe_response.raise_for_status()
    return str(iframe_response.url), decode_html_page(iframe_response)


def fetch_developer_options(username: str, password: str) -> list[DeveloperOption]:
    with MabangClient(BASE_URL) as client:
        client.login(username, password)
        _, stock_page_html = open_stock_page(client)
    options = parse_developer_options(stock_page_html)
    if not options:
        raise MabangApiError("未获取到名称以“开发”结尾的开发人员，请确认账号拥有库存 SKU 页面权限")
    return options


def validate_query_payload(payload: dict[str, Any]) -> SkuInventoryQuery:
    username = safe_text(payload.get("username"))
    password = safe_text(payload.get("password"))
    start_date = safe_text(payload.get("start_date"))
    end_date = safe_text(payload.get("end_date"))
    output_dir_text = safe_text(payload.get("output_dir"))
    developer_ids = _coerce_string_list(payload.get("developer_ids"))
    query_mode = safe_text(payload.get("query_mode")) or "inventory"
    if query_mode not in {"inventory", "missing_warehouses"}:
        raise ValueError("不支持的 SKU 查询类型")
    liveness_types = _coerce_string_list(payload.get("liveness_types", ["1", "2"]))
    if query_mode == "missing_warehouses" and (not liveness_types or set(liveness_types) - {"1", "2"}):
        raise ValueError("请至少选择爆款或旺款")
    if not username:
        raise ValueError("请输入马帮账号")
    if not password:
        raise ValueError("请输入马帮密码")
    if not developer_ids:
        raise ValueError("请至少选择一名开发人员")
    if query_mode == "missing_warehouses":
        start_date = end_date = ""
    elif not re.fullmatch(r"\d{4}-\d{2}-\d{2}", start_date):
        raise ValueError("SKU创建时间的开始日期格式必须是 YYYY-MM-DD")
    if query_mode == "inventory" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", end_date):
        raise ValueError("SKU创建时间的结束日期格式必须是 YYYY-MM-DD")
    if query_mode == "inventory":
        start_at = datetime.strptime(start_date, "%Y-%m-%d")
        end_at = datetime.strptime(end_date, "%Y-%m-%d")
        if start_at > end_at:
            raise ValueError("SKU创建时间的开始日期不能晚于结束日期")
    if not output_dir_text:
        raise ValueError("请选择输出目录")
    try:
        rows_per_page = int(payload.get("rows_per_page") or DEFAULT_ROWS_PER_PAGE)
    except (TypeError, ValueError):
        rows_per_page = DEFAULT_ROWS_PER_PAGE
    return SkuInventoryQuery(
        username=username,
        password=password,
        developer_ids=developer_ids,
        start_date=start_date,
        end_date=end_date,
        output_dir=Path(output_dir_text).expanduser(),
        rows_per_page=max(50, min(rows_per_page, 1000)),
        query_mode=query_mode,
        liveness_types=liveness_types,
    )


def build_stock_query_payload(
    job: SkuInventoryQuery,
    *,
    developer_id: str,
    page: int,
    liveness_type: str = "",
) -> list[tuple[str, str]]:
    return [
        ("searchKey", "Stock_stockSku"),
        ("operate", "likeStart"),
        ("orderBys[]", ""),
        ("Stock_stockSku", ""),
        ("Stock_nameCN", ""),
        ("Stock_nameEN", ""),
        ("Stock_defaultRetailNameCn", ""),
        ("StockPlus_financial", ""),
        ("search-content", "库存SKU"),
        ("searchValue", ""),
        ("status", "" if job.query_mode == "missing_warehouses" else "3"),
        ("parentCategoryId", ""),
        ("categoryId", ""),
        ("third_category_id", ""),
        ("parentBrandId", ""),
        ("list-brandId", ""),
        ("labelId", ""),
        ("buyerId", ""),
        ("developerIdM", developer_id),
        ("dev_assistant", ""),
        ("artDesignerId", ""),
        ("salesId", ""),
        ("defaultStockWarehouseDetailId", ""),
        ("livenessType", liveness_type),
        ("isNewType", ""),
        ("stock_type", ""),
        ("isMachining", ""),
        ("showstart", "1"),
        ("isCloud", ""),
        ("isGift", ""),
        ("exceptionDeclaration", ""),
        ("singleWarehouseType", ""),
        ("isGoogsExpireManageSearch", ""),
        ("queryTime", "timeCreated"),
        ("timeCreatedStartTime", f"{job.start_date} 00:00:00" if job.start_date else ""),
        ("timeCreatedEndTime", f"{job.end_date} 23:59:59" if job.end_date else ""),
        ("page", str(page)),
        ("rowsPerPage", str(job.rows_per_page)),
        ("stockOrderby", ""),
        ("warehouseMold", "all"),
        ("isBatchSearch", "0"),
    ]


def row_matches_created_range(row: dict[str, Any], job: SkuInventoryQuery) -> bool:
    created_at = parse_datetime(first_value(row, TIME_CREATED_KEYS))
    if created_at is None:
        # 马帮已在服务端按创建时间过滤；部分账号的返回数据不带创建时间字段。
        return True
    start_at = datetime.strptime(job.start_date, "%Y-%m-%d")
    end_at = datetime.strptime(job.end_date, "%Y-%m-%d").replace(hour=23, minute=59, second=59)
    return start_at <= created_at <= end_at


def flatten_stock_warehouse_rows(
    stock_row: dict[str, Any],
    *,
    developer: DeveloperOption,
) -> list[dict[str, Any]]:
    warehouse_rows = stock_row.get("stockWarehouseData") or []
    if not isinstance(warehouse_rows, list):
        return []
    sku = safe_text(first_value(stock_row, ("stockSku", "stockSkuRaw", "sku")))
    if not sku:
        return []
    product_name = safe_text(first_value(stock_row, ("nameCN", "stockNameCN", "productName", "goodsName")))
    created_at = normalize_created_at(first_value(stock_row, TIME_CREATED_KEYS))
    stock_forecast = first_value(stock_row, FORECAST_DAILY_SALES_KEYS)
    records: list[dict[str, Any]] = []
    for warehouse_row in warehouse_rows:
        if not isinstance(warehouse_row, dict):
            continue
        available_value = first_value(
            warehouse_row,
            ("stockWarehouseAvailableInventory", "availabledInventory", "availableInventory"),
        )
        unshipped_count = warehouse_unshipped_count(warehouse_row)
        transit_inventory = warehouse_transit_inventory(warehouse_row)
        if not (parse_number(available_value) > 0 or unshipped_count > 0 or transit_inventory > 0):
            continue
        warehouse_id = safe_text(first_value(warehouse_row, ("warehouseId", "id")))
        warehouse_name = safe_text(first_value(warehouse_row, ("name", "warehouseName")))
        if not warehouse_name:
            warehouse_name = f"仓库 {warehouse_id}" if warehouse_id else "未命名仓库"
        forecast_value = first_value(warehouse_row, FORECAST_DAILY_SALES_KEYS)
        if forecast_value in (None, ""):
            forecast_value = stock_forecast
        records.append(
            {
                "developer_name": developer.name,
                "sku": sku,
                "product_name": product_name,
                "created_at": created_at,
                "warehouse_name": warehouse_name,
                "warehouse_id": warehouse_id,
                "available_inventory": normalize_number(available_value),
                "forecast_daily_sales": optional_number(forecast_value),
                "current_sales_days": optional_number(first_value(warehouse_row, ("salesDays", "saleAvailableDays"))),
                "unshipped_count": unshipped_count,
                "transit_inventory": transit_inventory,
            }
        )
    return records


def query_inventory_records(
    job: SkuInventoryQuery,
    progress: ProgressCallback | None = None,
) -> tuple[list[dict[str, Any]], list[DeveloperOption]]:
    records: list[dict[str, Any]] = []
    seen_details: set[tuple[str, str, str]] = set()
    with MabangClient(BASE_URL) as client:
        _emit(progress, "正在登录马帮...")
        client.login(job.username, job.password)
        _emit(progress, "正在读取名称以“开发”结尾的开发人员...")
        referer, stock_page_html = open_stock_page(client)
        available_options = parse_developer_options(stock_page_html)
        option_by_id = {item.id: item for item in available_options}
        missing_ids = [developer_id for developer_id in job.developer_ids if developer_id not in option_by_id]
        if missing_ids:
            raise MabangApiError("所选开发人员已失效或不符合名称规则，请重新加载开发人员")
        selected_options = [option_by_id[developer_id] for developer_id in job.developer_ids]
        for developer_index, developer in enumerate(selected_options, start=1):
            page = 1
            developer_record_count = 0
            seen_stock_rows: set[str] = set()
            previous_page_signature: tuple[str, ...] | None = None
            while True:
                _emit(
                    progress,
                    f"[{developer_index}/{len(selected_options)}] {developer.name}：正在查询第 {page} 页...",
                )
                payload = build_stock_query_payload(job, developer_id=developer.id, page=page)
                data = post_form_json(client, STOCK_LIST_API_URL, payload, referer)
                if not data.get("success"):
                    raise MabangApiError(
                        f"{developer.name} 的库存 SKU 查询失败：{safe_text(data.get('message')) or data}"
                    )
                rows = data.get("stockData") or []
                if not isinstance(rows, list):
                    raise MabangApiError(f"{developer.name} 的库存 SKU 返回字段异常")
                page_signature = tuple(
                    stock_row_identity(stock_row)
                    for stock_row in rows
                    if isinstance(stock_row, dict)
                )
                if rows and previous_page_signature is not None and page_signature == previous_page_signature:
                    _emit(
                        progress,
                        f"[{developer_index}/{len(selected_options)}] {developer.name}：马帮返回了重复分页，已停止继续翻页",
                    )
                    break
                previous_page_signature = page_signature
                new_stock_row_count = 0
                for stock_row in rows:
                    if not isinstance(stock_row, dict):
                        continue
                    stock_key = stock_row_identity(stock_row)
                    if stock_key and stock_key in seen_stock_rows:
                        continue
                    if stock_key:
                        seen_stock_rows.add(stock_key)
                    new_stock_row_count += 1
                    if not row_matches_created_range(stock_row, job):
                        continue
                    for record in flatten_stock_warehouse_rows(stock_row, developer=developer):
                        warehouse_key = record["warehouse_id"] or record["warehouse_name"]
                        detail_key = (developer.id, record["sku"], safe_text(warehouse_key))
                        if detail_key in seen_details:
                            continue
                        seen_details.add(detail_key)
                        records.append(record)
                        developer_record_count += 1
                current_page, total_pages = parse_pagination_state(data)
                if current_page is not None and total_pages is not None and current_page >= total_pages:
                    _emit(
                        progress,
                        f"[{developer_index}/{len(selected_options)}] {developer.name}：已到最后一页（{current_page}/{total_pages}）",
                    )
                    break
                if rows and new_stock_row_count == 0:
                    _emit(
                        progress,
                        f"[{developer_index}/{len(selected_options)}] {developer.name}：当前分页没有新增 SKU，已停止继续翻页",
                    )
                    break
                if not rows or len(rows) < job.rows_per_page:
                    break
                page += 1
                if page > MAX_PAGE_COUNT:
                    raise MabangApiError(f"{developer.name} 的分页超过 {MAX_PAGE_COUNT} 页，请缩小时间范围")
            _emit(progress, f"[{developer_index}/{len(selected_options)}] {developer.name}：匹配 {developer_record_count} 条符合条件的仓库记录")
    records.sort(
        key=lambda item: (
            safe_text(item.get("developer_name")).casefold(),
            safe_text(item.get("sku")).casefold(),
            safe_text(item.get("warehouse_name")).casefold(),
        )
    )
    return records, selected_options


def post_form_json(
    client: MabangClient,
    url: str,
    payload: list[tuple[str, str]],
    referer: str,
) -> dict[str, Any]:
    response = client.client.post(
        url,
        content=urlencode(payload, doseq=True),
        headers={
            "Referer": referer,
            "Origin": PRIVATE_BASE_URL,
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "X-Requested-With": "XMLHttpRequest",
        },
    )
    response.raise_for_status()
    try:
        data = response.json()
    except ValueError as exc:
        raise MabangApiError(f"库存 SKU 查询返回非 JSON 响应：{url}") from exc
    if not isinstance(data, dict):
        raise MabangApiError(f"库存 SKU 查询返回格式异常：{url}")
    if data.get("success") is False and "登录信息已超时" in safe_text(data.get("message")):
        raise MabangApiError("马帮登录信息已超时，请重新运行查询")
    return data


def export_records(records: Iterable[dict[str, Any]], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / f"SKU库存与可售天数_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SKU仓库明细"
    worksheet.append(list(OUTPUT_COLUMNS))
    field_names = (
        "developer_name",
        "sku",
        "product_name",
        "created_at",
        "warehouse_name",
        "warehouse_id",
        "available_inventory",
        "forecast_daily_sales",
        "current_sales_days",
        "unshipped_count",
        "transit_inventory",
    )
    record_count = 0
    for record in records:
        worksheet.append([record.get(field_name, "") for field_name in field_names])
        record_count += 1
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(OUTPUT_COLUMNS))}{max(1, record_count + 1)}"
    for index, width in enumerate((18, 22, 34, 21, 30, 15, 15, 15, 17, 13, 13), start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    workbook.save(output_file)
    workbook.close()
    return output_file


def run_sku_inventory_query(
    job: SkuInventoryQuery,
    progress: ProgressCallback | None = None,
) -> SkuInventoryResult:
    if job.query_mode == "missing_warehouses":
        from backend.services.sku_warehouse_coverage import run_warehouse_coverage_query

        return run_warehouse_coverage_query(job, progress)
    records, selected_options = query_inventory_records(job, progress)
    _emit(progress, "正在生成 Excel 文件...")
    output_file = export_records(records, job.output_dir)
    sku_count = len({safe_text(record.get("sku")) for record in records if safe_text(record.get("sku"))})
    warehouse_count = len(
        {
            safe_text(record.get("warehouse_id")) or safe_text(record.get("warehouse_name"))
            for record in records
            if safe_text(record.get("warehouse_id")) or safe_text(record.get("warehouse_name"))
        }
    )
    _emit(progress, f"查询完成：{sku_count} 个 SKU，{len(records)} 条 SKU＋仓库记录")
    return SkuInventoryResult(
        records=tuple(records),
        developer_count=len(selected_options),
        sku_count=sku_count,
        warehouse_count=warehouse_count,
        output_file=output_file,
    )


def preview_rows(records: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    return [dict(record) for record in records[:PREVIEW_ROW_LIMIT]]


def _coerce_string_list(value: object) -> tuple[str, ...]:
    if isinstance(value, str):
        source = [value]
    elif isinstance(value, (list, tuple, set)):
        source = list(value)
    else:
        source = []
    result: list[str] = []
    seen: set[str] = set()
    for item in source:
        text = safe_text(item)
        if not text or text in seen:
            continue
        seen.add(text)
        result.append(text)
    return tuple(result)


def _emit(progress: ProgressCallback | None, message: str) -> None:
    if progress:
        progress(message)
