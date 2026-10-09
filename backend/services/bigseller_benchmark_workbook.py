"""Read SKU workbooks and write benchmark results without changing the source."""

from __future__ import annotations

import hashlib
import math
import re
from collections.abc import Mapping, Sequence
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_SEARCH_ROWS = 10
SKU_HEADERS = ("库存SKU编号", "库存SKU", "SKU")
METRIC_LABELS = {"views": "浏览量", "sales": "销量"}
STATUS_LABELS = {"matched": "已匹配", "not_found": "未找到商品", "failed": "查询失败"}
DETAIL_HEADERS = ("SKU", "店铺", "指标值", "店铺数", "Item ID", "状态", "说明")
SUMMARY_HEADERS = ("汇总项目", "结果")
PENDING_HEADERS = ("SKU", "原因", "查询次数", "源行号")
GENERATED_SHEET_AUTHOR = "HQYL BigSeller benchmark"
XLS_WARNING = ".xls 输出会转换为 .xlsx，仅保留单元格值和基本格式；公式采用缓存值。需要完整保留公式和格式时，请先用 Excel 另存为 .xlsx。"


@dataclass(frozen=True)
class WorkbookInput:
    source_path: Path
    sheet_name: str
    header_row: int
    sku_column: int
    skus: tuple[str, ...]
    row_numbers: dict[str, tuple[int, ...]]
    warnings: tuple[str, ...] = ()
    source_sha256: str = ""

    @property
    def total_rows(self) -> int:
        return sum(len(rows) for rows in self.row_numbers.values())


def _source_path(path: str | Path) -> Path:
    source = Path(path).expanduser().resolve()
    if not source.is_file():
        raise ValueError("请选择存在的 Excel 文件")
    if source.suffix.lower() not in {".xlsx", ".xls"}:
        raise ValueError("支持 .xlsx 或 .xls 文件；请将其他格式另存为 .xlsx")
    return source


def _read_xls(source: Path) -> Workbook:
    """Explicit, value-based legacy conversion; callers surface XLS_WARNING."""
    import xlrd

    original = xlrd.open_workbook(str(source), formatting_info=True)
    book = Workbook()
    book.remove(book.active)
    try:
        for old_sheet in original.sheets():
            sheet = book.create_sheet(old_sheet.name)
            if old_sheet.visibility:
                sheet.sheet_state = "hidden" if old_sheet.visibility == 1 else "veryHidden"
            for row_index in range(old_sheet.nrows):
                for col_index in range(old_sheet.ncols):
                    old_cell = old_sheet.cell(row_index, col_index)
                    value = old_cell.value
                    if old_cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    elif old_cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate.xldate_as_datetime(value, original.datemode)
                    elif old_cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif old_cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#VALUE!")
                    cell = sheet.cell(row_index + 1, col_index + 1, value)
                    if old_cell.ctype == xlrd.XL_CELL_TEXT:
                        cell.data_type = "s"
                    xf = original.xf_list[old_cell.xf_index]
                    number_format = original.format_map.get(xf.format_key)
                    if number_format:
                        cell.number_format = number_format.format_str
                    old_font = original.font_list[xf.font_index]
                    cell.font = Font(name=old_font.name, size=old_font.height / 20,
                                     bold=bool(old_font.bold), italic=bool(old_font.italic))
            for index, info in old_sheet.colinfo_map.items():
                sheet.column_dimensions[get_column_letter(index + 1)].width = info.width / 256
                sheet.column_dimensions[get_column_letter(index + 1)].hidden = bool(info.hidden)
            for index, info in old_sheet.rowinfo_map.items():
                sheet.row_dimensions[index + 1].height = info.height / 20
                sheet.row_dimensions[index + 1].hidden = bool(info.hidden)
            for low_row, high_row, low_col, high_col in old_sheet.merged_cells:
                sheet.merge_cells(start_row=low_row + 1, end_row=high_row,
                                  start_column=low_col + 1, end_column=high_col)
    finally:
        original.release_resources()
    return book


def _open_source(source: Path, *, data_only: bool = False) -> Workbook:
    if source.suffix.lower() == ".xls":
        return _read_xls(source)
    return load_workbook(source, data_only=data_only, keep_links=True)


def _header_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").lstrip("\ufeff")).upper()


def _find_sku_header(sheet: Any) -> tuple[int, int]:
    # Prefer the inventory SKU header even if a generic SKU occurs earlier.
    candidates = [
        (SKU_HEADERS.index(_header_text(cell.value)), cell.row, cell.column)
        for row in sheet.iter_rows(max_row=min(HEADER_SEARCH_ROWS, sheet.max_row))
        for cell in row if _header_text(cell.value) in SKU_HEADERS
    ]
    if not candidates:
        raise ValueError(f"前 {HEADER_SEARCH_ROWS} 行未找到「库存SKU编号」或「SKU」表头")
    _, row, column = min(candidates)
    return row, column


def _sku_value(cell: Any, cached_cell: Any = None) -> str:
    value = cell.value
    if cell.data_type == "f":
        value = cached_cell.value if cached_cell is not None else None
        if value is None:
            raise ValueError(f"SKU 单元格 {cell.coordinate} 是没有缓存结果的公式，请用 Excel 计算并保存后重试")
    if value is None:
        return ""
    if cell.data_type == "e" or (cached_cell is not None and cached_cell.data_type == "e"):
        raise ValueError(f"SKU 单元格 {cell.coordinate} 含 Excel 错误，请修正后重试")
    if isinstance(value, bool):
        raise ValueError(f"SKU 单元格 {cell.coordinate} 是布尔值，请使用文本 SKU")
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError(f"SKU 单元格 {cell.coordinate} 不是有效数字")
        if value == int(value):
            text = str(int(value))
            # Excel's 000000 identifier format carries significant leading zeros.
            mask = str(cell.number_format or "").split(";", 1)[0].strip()
            if value >= 0 and re.fullmatch(r"0+", mask):
                text = text.zfill(len(mask))
        else:
            text = str(value)
    else:
        text = str(value).strip().lstrip("\ufeff").strip()
    if ILLEGAL_CHARACTERS_RE.search(text) or len(text) > 32767:
        raise ValueError(f"SKU 单元格 {cell.coordinate} 含无效字符或文本过长")
    return text


def _sheet_input(source: Path, sheet: Any, cached_sheet: Any = None, *, source_sha256: str = "") -> WorkbookInput:
    header_row, sku_column = _find_sku_header(sheet)
    rows: dict[str, list[int]] = {}
    for row in sheet.iter_rows(min_row=header_row + 1, min_col=sku_column, max_col=sku_column):
        cell = row[0]
        cached_cell = cached_sheet.cell(cell.row, sku_column) if cell.data_type == "f" and cached_sheet else None
        sku = _sku_value(cell, cached_cell)
        if sku:
            rows.setdefault(sku, []).append(cell.row)
    if not rows:
        raise ValueError("所选子表没有非空 SKU 数据")
    return WorkbookInput(
        source_path=source, sheet_name=sheet.title, header_row=header_row,
        sku_column=sku_column, skus=tuple(rows),
        row_numbers={sku: tuple(numbers) for sku, numbers in rows.items()},
        warnings=(XLS_WARNING,) if source.suffix.lower() == ".xls" else (),
        source_sha256=source_sha256,
    )


def _default_sheet(sheets: list[dict[str, Any]]) -> str:
    valid = [sheet for sheet in sheets if not sheet["error"] and "对标明细" not in sheet["name"]]
    if not valid:
        valid = [sheet for sheet in sheets if not sheet["error"]]
    preferred = next((sheet for sheet in valid if "动销总表" in sheet["name"]), None)
    return (preferred or (valid[0] if valid else {})).get("name", "")


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    source = _source_path(path)
    book = _open_source(source)
    cached = _open_source(source, data_only=True)
    try:
        sheets = []
        for sheet in book.worksheets:
            item = {"name": sheet.title, "header_row": None, "sku_column": None,
                    "sku_count": 0, "row_count": 0, "error": ""}
            try:
                info = _sheet_input(source, sheet, cached[sheet.title])
                item.update(header_row=info.header_row, sku_column=info.sku_column,
                            sku_count=len(info.skus), row_count=info.total_rows)
            except ValueError as exc:
                item["error"] = str(exc)
            sheets.append(item)
        return {"source_path": str(source), "sheets": sheets,
                "sheet_names": book.sheetnames, "default_sheet": _default_sheet(sheets),
                "warnings": [XLS_WARNING] if source.suffix.lower() == ".xls" else []}
    finally:
        book.close()
        cached.close()


def load_workbook_input(path: str | Path, sheet_name: str | None = None) -> WorkbookInput:
    source = _source_path(path)
    if not sheet_name:
        sheet_name = inspect_workbook(source)["default_sheet"]
        if not sheet_name:
            raise ValueError("工作簿中未找到包含 SKU 数据的子表")
    source_sha256 = _file_sha256(source)
    book = _open_source(source)
    cached = _open_source(source, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError("所选子表不存在，请重新选择 Excel 文件")
        if _file_sha256(source) != source_sha256:
            raise ValueError("源工作簿在读取过程中已变化，请重新选择文件")
        return _sheet_input(source, book[sheet_name], cached[sheet_name], source_sha256=source_sha256)
    finally:
        book.close()
        cached.close()


def _file_sha256(source: Path) -> str:
    digest = hashlib.sha256()
    with source.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_value(cell: Any, value: Any) -> None:
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        # Remote shop names and user SKUs must stay text, including '=...' values.
        cell.data_type = "s"
    else:
        cell.value = value


def _result_lookup(results: Sequence[Mapping[str, Any]] | Mapping[str, Mapping[str, Any]]) -> dict[str, Mapping[str, Any]]:
    if isinstance(results, Mapping):
        return {str(sku): result for sku, result in results.items()}
    return {str(result.get("sku", "")): result for result in results}


def _normalized_result(sku: str, result: Mapping[str, Any] | None) -> dict[str, Any]:
    data = dict(result or {})
    status = str(data.get("status") or "failed")
    if status not in STATUS_LABELS:
        status = "failed"
        data["message"] = "结果状态异常，请重新查询"
    shop = str(data.get("shop_name") or "").strip()
    value = data.get("metric_value")
    if status == "matched" and (not shop or isinstance(value, bool) or
                                not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0):
        status = "failed"
        data["message"] = "结果缺少有效店铺或指标值，请重新查询"
    message = str(data.get("message") or "")
    if not result:
        message = "本次未取得查询结果，请重新查询"
    elif not message and status == "failed":
        message = "查询失败，请重新查询"
    elif not message and status == "not_found":
        message = "未找到符合查询条件的商品"
    return {"sku": sku, "shop_name": shop if status == "matched" else None,
            "metric_value": value if status == "matched" else None,
            "shop_count": data.get("shop_count") if status == "matched" else None,
            "item_id": str(data.get("item_id") or "") if status == "matched" else None,
            "status": status, "message": message, "attempts": data.get("attempts", 0)}


def _audited_results(
    source: WorkbookInput,
    results: Sequence[Mapping[str, Any]] | Mapping[str, Mapping[str, Any]],
    summary: Mapping[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    """Require a complete accounting before claiming that an export is complete."""
    if not isinstance(summary, Mapping):
        raise ValueError("完整性校验失败：汇总格式无效")
    if isinstance(results, Mapping):
        pairs = list(results.items())
    else:
        pairs = []
        for result in results:
            if not isinstance(result, Mapping):
                raise ValueError("完整性校验失败：查询结果格式无效")
            pairs.append((result.get("sku", ""), result))
    lookup: dict[str, Mapping[str, Any]] = {}
    for raw_sku, result in pairs:
        sku = str(raw_sku)
        if sku in lookup:
            raise ValueError(f"完整性校验失败：SKU {sku} 有重复查询结果")
        if not isinstance(result, Mapping) or ("sku" in result and str(result["sku"]) != sku):
            raise ValueError(f"完整性校验失败：SKU {sku} 的查询结果不一致")
        lookup[sku] = result
    if set(lookup) != set(source.skus) or len(lookup) != len(source.skus):
        raise ValueError("完整性校验失败：查询结果缺少输入 SKU 或包含额外 SKU")
    normalized = []
    for sku in source.skus:
        result = lookup[sku]
        item = _normalized_result(sku, result)
        status = result.get("status")
        if not isinstance(status, str) or status not in STATUS_LABELS or item["status"] != status:
            raise ValueError(f"完整性校验失败：SKU {sku} 的状态或指标数据无效")
        if type(item["attempts"]) is not int or item["attempts"] < 0:
            raise ValueError(f"完整性校验失败：SKU {sku} 的查询次数无效")
        normalized.append(item)
    counts = {
        "sku_count": len(source.skus), "total_rows": source.total_rows,
        "matched_count": sum(item["status"] == "matched" for item in normalized),
        "not_found_count": sum(item["status"] == "not_found" for item in normalized),
        "failed_count": sum(item["status"] == "failed" for item in normalized),
    }
    counts["confirmed_count"] = counts["matched_count"] + counts["not_found_count"]
    for key, expected in counts.items():
        if type(summary.get(key)) is not int or summary[key] != expected:
            raise ValueError(f"完整性校验失败：汇总 {key} 与实际查询结果不一致")
    is_complete = counts["failed_count"] == 0
    if type(summary.get("is_complete")) is not bool or summary["is_complete"] != is_complete:
        raise ValueError("完整性校验失败：完成状态与失败待补查数不一致")
    audited: dict[str, Any] = {**counts, "is_complete": is_complete}
    for key in ("retry_rounds_used", "recovered_count", "resumed_count"):
        value = summary.get(key, 0)
        if type(value) is not int or value < 0 or (key != "retry_rounds_used" and value > len(source.skus)):
            raise ValueError(f"完整性校验失败：汇总 {key} 无效")
        audited[key] = value
    return normalized, audited


def _owned_integrity_sheet(
    book: Workbook, source_sheet: str, label: str, kind: str,
    headers: tuple[str, ...], *, create: bool = True,
) -> Any:
    # Header text alone is not ownership evidence: users may have a same-named list.
    marker = f"hqyl_desktop.bigseller_benchmark_workbook:v1:{label}:{kind}:{source_sheet}"
    base = f"BigSeller{label}{kind}"
    first_free = None
    for index in range(1, len(book.sheetnames) + 2):
        name = base if index == 1 else f"{base}_{index}"
        if name not in book.sheetnames:
            if first_free is None:
                first_free = name
            continue
        sheet = book[name]
        comment = sheet["A1"].comment
        existing_headers = tuple(sheet.cell(1, column).value for column in range(1, len(headers) + 1))
        if name != source_sheet and comment and comment.author == GENERATED_SHEET_AUTHOR and comment.text == marker and existing_headers == headers:
            if sheet.max_row > 1:
                sheet.delete_rows(2, sheet.max_row - 1)
            return sheet
    if not create:
        return None
    sheet = book.create_sheet(first_free)
    for index, title in enumerate(headers, 1):
        cell = sheet.cell(1, index)
        _write_value(cell, title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D87")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A1"].comment = Comment(marker, GENERATED_SHEET_AUTHOR)
    sheet.freeze_panes = "A2"
    return sheet


def _write_integrity_sheets(
    book: Workbook, source: WorkbookInput, label: str,
    normalized: list[dict[str, Any]], summary: Mapping[str, Any],
) -> None:
    sheet = _owned_integrity_sheet(book, source.sheet_name, label, "对标汇总", SUMMARY_HEADERS)
    rows = (
        ("完成性", "查询完整" if summary["is_complete"] else "结果不完整"),
        ("对标指标", label), ("来源子表", source.sheet_name),
        ("原表行数", summary["total_rows"]), ("去重SKU数", summary["sku_count"]),
        ("成功匹配数", summary["matched_count"]), ("确认未匹配数", summary["not_found_count"]),
        ("失败待补查数", summary["failed_count"]), ("已确认SKU数", summary["confirmed_count"]),
        ("补查轮数", summary["retry_rounds_used"]), ("补查恢复数", summary["recovered_count"]),
        ("沿用已确认结果数", summary["resumed_count"]),
        ("计数口径", "原表行数为非空SKU行数；其余查询数量按去重SKU统计。确认未匹配表示查询成功但未找到符合条件的商品，不计为失败。"),
    )
    for row_number, values in enumerate(rows, 2):
        for column, value in enumerate(values, 1):
            _write_value(sheet.cell(row_number, column), value)
            sheet.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 90
    sheet.row_dimensions[len(rows) + 1].height = 44
    failed = [item for item in normalized if item["status"] == "failed"]
    pending = _owned_integrity_sheet(book, source.sheet_name, label, "待补查", PENDING_HEADERS, create=bool(failed))
    if pending is None:
        return
    for row_number, item in enumerate(failed, 2):
        values = (item["sku"], item["message"], item["attempts"],
                  ", ".join(map(str, source.row_numbers[item["sku"]])))
        for column, value in enumerate(values, 1):
            _write_value(pending.cell(row_number, column), value)
            pending.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate((26, 85, 15, 30), 1):
        pending.column_dimensions[get_column_letter(index)].width = width
    pending.auto_filter.ref = f"A1:D{len(failed) + 1}"


def _output_columns(sheet: Any, source: WorkbookInput, label: str) -> dict[str, tuple[int, ...]]:
    aliases = {
        "shop_name": (f"{label}第一店铺", f"{label}第一", f"{label}TOP1"),
        "metric_value": (f"{label}指标值",),
        "status": (f"{label}对标状态",),
        "message": (f"{label}对标说明",),
    }
    widths = {"shop_name": 28, "metric_value": 18, "status": 18, "message": 65}
    columns = {}
    # Formatting can extend far beyond the real table. Reuse only that empty tail,
    # never an interior blank header, data column, formula or merged title area.
    last_used_column = max(
        (cell.column for row in sheet.iter_rows() for cell in row if cell.value is not None),
        default=0,
    )
    last_used_column = max(last_used_column, max((area.max_col for area in sheet.merged_cells.ranges), default=0))
    next_column = last_used_column + 1
    for key, headers in aliases.items():
        found = tuple(cell.column for cell in sheet[source.header_row]
                      if _header_text(cell.value) in tuple(_header_text(name) for name in headers))
        if found:
            columns[key] = found
            continue
        index = next_column
        next_column += 1
        if index > 16384:
            raise ValueError("子表列数已达 Excel 上限，无法添加对标结果列")
        cell = sheet.cell(source.header_row, index)
        cell._style = copy(sheet.cell(source.header_row, source.sku_column)._style)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        _write_value(cell, headers[0])
        sheet.column_dimensions[get_column_letter(index)].width = widths[key]
        columns[key] = (index,)
    return columns


def _detail_sheet(book: Workbook, source_sheet: str, label: str) -> Any:
    base = f"BigSeller{label}对标明细"
    index = 1
    while True:
        name = base if index == 1 else f"{base}_{index}"
        if name not in book.sheetnames:
            return book.create_sheet(name)
        sheet = book[name]
        existing_headers = tuple(sheet.cell(1, column).value for column in range(1, len(DETAIL_HEADERS) + 1))
        if name != source_sheet and existing_headers == DETAIL_HEADERS:
            # Reuse the same sheet object so its position, references and name survive.
            if sheet.max_row > 1:
                sheet.delete_rows(2, sheet.max_row - 1)
            return sheet
        index += 1


def export_benchmark_workbook(
    input_data: WorkbookInput,
    results: Sequence[Mapping[str, Any]] | Mapping[str, Mapping[str, Any]],
    metric: str,
    output_dir: str | Path,
    *,
    summary: Mapping[str, Any] | None = None,
) -> Path:
    if metric not in METRIC_LABELS:
        raise ValueError("对标指标必须选择浏览量或销量")
    destination = Path(output_dir).expanduser().resolve()
    if destination.exists() and not destination.is_dir():
        raise ValueError("输出目录不能是文件")
    label = METRIC_LABELS[metric]
    source_path = _source_path(input_data.source_path)
    if input_data.source_sha256 and _file_sha256(source_path) != input_data.source_sha256:
        raise ValueError("源工作簿在查询期间已变化，请重新选择文件并运行，避免结果写入错误的 SKU 行")
    book = _open_source(source_path)
    try:
        if input_data.source_sha256 and _file_sha256(source_path) != input_data.source_sha256:
            raise ValueError("源工作簿在读取过程中已变化，请重新选择文件并运行")
        if input_data.sheet_name not in book.sheetnames:
            raise ValueError("源工作簿已变化，请重新选择文件")
        sheet = book[input_data.sheet_name]
        if _find_sku_header(sheet) != (input_data.header_row, input_data.sku_column):
            raise ValueError("源工作簿表头已变化，请重新选择文件")
        audited_summary = None
        if summary is not None:
            normalized, audited_summary = _audited_results(input_data, results, summary)
        else:
            lookup = _result_lookup(results)
            normalized = [_normalized_result(sku, lookup.get(sku)) for sku in input_data.skus]
        columns = _output_columns(sheet, input_data, label)
        for result in normalized:
            for row_number in input_data.row_numbers[result["sku"]]:
                for key, indices in columns.items():
                    value = STATUS_LABELS[result[key]] if key == "status" else result[key]
                    for index in indices:
                        cell = sheet.cell(row_number, index)
                        _write_value(cell, value)
                        if key == "metric_value":
                            cell.number_format = "#,##0.########"
        detail = _detail_sheet(book, input_data.sheet_name, label)
        for index, title in enumerate(DETAIL_HEADERS, 1):
            cell = detail.cell(1, index)
            _write_value(cell, title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="275D87")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for row_number, result in enumerate(normalized, 2):
            values = (result["sku"], result["shop_name"], result["metric_value"],
                      result["shop_count"], result["item_id"],
                      STATUS_LABELS[result["status"]], result["message"])
            for column, value in enumerate(values, 1):
                cell = detail.cell(row_number, column)
                _write_value(cell, value)
                cell.number_format = "@" if column in {1, 5} else "General"
        for index, width in enumerate((26, 30, 18, 12, 24, 18, 72), 1):
            detail.column_dimensions[get_column_letter(index)].width = width
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = f"A1:G{len(normalized) + 1}"
        if audited_summary is not None:
            _write_integrity_sheets(book, input_data, label, normalized, audited_summary)
        destination.mkdir(parents=True, exist_ok=True)
        source_stem = input_data.source_path.stem
        if audited_summary is not None:
            source_stem = re.sub(
                rf"(_BigSeller{label}对标_\d{{8}}_\d{{6}}_[0-9a-f]{{8}})_结果不完整$", r"\1", source_stem,
            )
        stem = re.sub(r'[\\/:*?"<>|]', "_", source_stem)[:90]
        sheet_name = re.sub(r'[\\/:*?"<>|]', "_", input_data.sheet_name)[:35]
        incomplete = "_结果不完整" if audited_summary is not None and not audited_summary["is_complete"] else ""
        filename = f"{stem}_{sheet_name}_BigSeller{label}对标_{datetime.now():%Y%m%d_%H%M%S}_{uuid4().hex[:8]}{incomplete}.xlsx"
        output_path = destination / filename
        # Exclusive creation provides collision safety, even for concurrent runs.
        with output_path.open("xb") as stream:
            book.save(stream)
        return output_path
    finally:
        book.close()
