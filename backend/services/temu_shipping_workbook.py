"""Freeze TEMU dimensions in worksheet order without recalculating or editing Excel."""

from __future__ import annotations

import hashlib
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


HEADER_SEARCH_ROWS = 10
FIELDS = ("length", "width", "height", "weight")
FIELD_LABELS = {"length": "长", "width": "宽", "height": "高", "weight": "重量"}
_HEADERS = {
    "length": {"长", "长度", "长cm", "长度cm", "长厘米", "长度厘米"},
    "width": {"宽", "宽度", "宽cm", "宽度cm", "宽厘米", "宽度厘米"},
    "height": {"高", "高度", "高cm", "高度cm", "高厘米", "高度厘米"},
    "weight": {"重量", "重量g", "重量克", "申报重量", "申报重量g", "申报重量克"},
}
_RECALCULATE = "请用 Excel 或 WPS 打开文件，重新计算并保存后再导入"


def _header_text(value: Any) -> str:
    return re.sub(r"[\s()（）\[\]【】/／]", "", str(value or "").lstrip("\ufeff")).lower()


def _find_header(sheet: Any) -> tuple[int, dict[str, int]] | None:
    for row in sheet.iter_rows(max_row=min(HEADER_SEARCH_ROWS, sheet.max_row)):
        columns: dict[str, list[int]] = {field: [] for field in FIELDS}
        for cell in row:
            text = _header_text(cell.value)
            for field, aliases in _HEADERS.items():
                if text in aliases:
                    columns[field].append(cell.column)
        if all(columns.values()):
            duplicates = [FIELD_LABELS[field] for field, found in columns.items() if len(found) != 1]
            if duplicates:
                raise ValueError(f"工作表「{sheet.title}」第 {row[0].row} 行的表头重复：{'、'.join(duplicates)}")
            return row[0].row, {field: found[0] for field, found in columns.items()}
    return None


def _today_date(value: date | datetime | str | None) -> date:
    if value is None:
        return date.today()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("当前日期必须使用 YYYY-MM-DD 格式") from exc


def _calculation_date(formulas: Any, cached: Any) -> str | None:
    """Read optional saved date metadata without restricting cached data imports.

    Old or future dates are retained as recorded in the workbook. Missing,
    ambiguous or composite date anchors simply leave this metadata unknown;
    only the dimensions and weight cells determine whether import is valid.
    """
    dates: set[date] = set()
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                formula = str(cell.value)
                if "TODAY" not in formula.upper():
                    continue
                if not re.fullmatch(r"=\s*(?:_xlfn\.)?TODAY\s*\(\s*\)\s*", formula, flags=re.I):
                    return None
                value = cached[sheet.title][cell.coordinate].value
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    try:
                        value = from_excel(value, formulas.epoch)
                    except (ValueError, OverflowError):
                        return None
                if isinstance(value, datetime):
                    value = value.date()
                if not isinstance(value, date):
                    return None
                dates.add(value)
    return next(iter(dates)).isoformat() if len(dates) == 1 else None


def _number(cell: Any, cached_cell: Any, field: str, sheet_name: str) -> Decimal:
    location = f"工作表「{sheet_name}」第 {cell.row} 行 {cell.coordinate}（{FIELD_LABELS[field]}）"
    value = cached_cell.value if cell.data_type == "f" else cell.value
    if cell.data_type == "f" and value is None:
        raise ValueError(f"{location} 的公式没有缓存结果；{_RECALCULATE}")
    if value is None or (isinstance(value, str) and not value.strip()):
        raise ValueError(f"{location} 缺少数值，请补全后导入，避免渠道对应错位")
    if cell.data_type == "e" or cached_cell.data_type == "e":
        raise ValueError(f"{location} 含 Excel 错误 {value}，请修正后导入")
    if isinstance(value, (bool, date, datetime)):
        raise ValueError(f"{location} 必须是正数")
    try:
        number = Decimal(str(value).strip())
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"{location} 不是有效数字") from exc
    if not number.is_finite() or number <= 0:
        raise ValueError(f"{location} 必须是有限的正数")
    # Reject unreasonable text exponents before expanding them into JSON strings.
    if abs(number.adjusted()) > 100 or len(number.as_tuple().digits) > 100:
        raise ValueError(f"{location} 的数值位数过多，请检查单元格")
    if field == "weight":
        with localcontext() as context:
            context.prec = 110
            number = number.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
        if number <= 0:
            raise ValueError(f"{location} 四舍五入到两位小数后必须大于 0 g")
    return number


def _decimal_text(number: Decimal, *, weight: bool = False) -> str:
    text = format(number, "f")
    return text if weight or "." not in text else text.rstrip("0").rstrip(".")


def read_temu_shipping_workbook(
    path: str | Path, *, today: date | datetime | str | None = None,
) -> dict[str, Any]:
    """Return an immutable-by-value import snapshot with original Excel row numbers.

    Required headings are 长/宽/高 (cm) and 重量 (g). Formula values must have
    been calculated and saved by Excel/WPS. Repeated dimensions remain separate
    rows, and blank rows inside the data region are errors rather than skipped.
    """
    source = Path(path).expanduser().resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("仅支持 .xlsx 文件，请用 Excel 或 WPS 另存为 .xlsx 后导入")
    if not source.is_file():
        raise ValueError("请选择存在的 Excel 文件")
    current_day = _today_date(today)
    try:
        # Both views use identical bytes even if Excel saves again during import.
        content = source.read_bytes()
        formulas = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError("无法读取 Excel 文件，请确认文件未损坏、未加密且为 .xlsx 格式") from exc
    cached = None
    try:
        try:
            cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
        except Exception as exc:
            raise ValueError("无法读取 Excel 的缓存数据，请重新计算并保存后导入") from exc
        candidates = sorted(formulas.worksheets, key=lambda sheet: sheet.title != "Sheet1")
        selection = next(((sheet, header) for sheet in candidates if (header := _find_header(sheet))), None)
        if selection is None:
            raise ValueError(f"工作表前 {HEADER_SEARCH_ROWS} 行未找到完整表头：长、宽、高、重量（尺寸 cm，重量 g）")
        sheet, (header_row, columns) = selection
        cached_sheet = cached[sheet.title]
        last_row = max(
            (row for row in range(header_row + 1, sheet.max_row + 1)
             if any(sheet.cell(row, col).value is not None for col in columns.values())),
            default=header_row,
        )
        if last_row == header_row:
            raise ValueError(f"工作表「{sheet.title}」没有长宽高和重量数据")
        has_formulas = any(
            sheet.cell(row, col).data_type == "f"
            for row in range(header_row + 1, last_row + 1) for col in columns.values()
        )
        calculation_date = _calculation_date(formulas, cached) if has_formulas else None
        rows = []
        for row_number in range(header_row + 1, last_row + 1):
            item: dict[str, Any] = {"excel_row": row_number}
            for field, column in columns.items():
                number = _number(sheet.cell(row_number, column), cached_sheet.cell(row_number, column), field, sheet.title)
                item[field] = _decimal_text(number, weight=field == "weight")
            rows.append(item)
        frozen_json = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        warnings = []
        if has_formulas:
            warnings.append("公式采用文件中已保存的计算结果，本次导入已固定数值；源文件不会被修改。")
        return {
            "input_file": str(source),
            "sheet_name": sheet.title,
            "header_row": header_row,
            "row_count": len(rows),
            "rows": rows,
            "fingerprint": hashlib.sha256(frozen_json.encode("utf-8")).hexdigest(),
            "source_sha256": hashlib.sha256(content).hexdigest(),
            "warnings": warnings,
            "calculation_date": calculation_date,
            "import_date": current_day.isoformat(),
            "uses_formula_cache": has_formulas,
            "dimension_unit": "cm",
            "weight_unit": "g",
        }
    finally:
        formulas.close()
        if cached is not None:
            cached.close()
