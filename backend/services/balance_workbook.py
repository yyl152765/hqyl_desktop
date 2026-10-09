"""Excel deliverables for TEMU/Lazada balance snapshots."""
from __future__ import annotations

import os
import math
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.units import pixels_to_EMU


FIELDS = {"temu": ("total", "pending"), "lazada": ("income", "balance", "ads", "processing")}
STATUS_NAMES = {"success": "成功", "partial": "部分完成", "failed": "失败", "pending": "待处理", "running": "采集中"}


def _text(sheet, row: int, column: int, value) -> None:
    cell = sheet.cell(row, column, str(value or ""))
    # Assigning a leading '=' ordinarily creates a formula. All labels, names,
    # URLs and errors are untrusted strings and must stay literal text.
    cell.data_type = "s"
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def _amount(sheet, row: int, column: int, value) -> None:
    if value is None or value == "":
        return
    if isinstance(value, bool):
        raise ValueError("余额必须为有效数字")
    try:
        amount = Decimal(str(value))
    except InvalidOperation as exc:
        raise ValueError("余额必须为有效数字") from exc
    if not amount.is_finite():
        raise ValueError("余额必须为有限数字")
    cell = sheet.cell(row, column)
    # Excel has only 15 significant digits. Preserve unusual high-precision
    # balances as text rather than silently rounding the captured amount.
    if len(amount.as_tuple().digits) > 15:
        _text(sheet, row, column, format(amount, "f"))
        cell.number_format = "@"
    else:
        cell.value = amount
        cell.number_format = '#,##0.00########;[Red]-#,##0.00########;0.00'
        cell.alignment = Alignment(vertical="top", horizontal="right")


def _image(sheet, row: int, column: int, path: str, store: dict) -> None:
    if not path:
        return
    image = Image(str(path))
    scale = min(480 / image.width, 320 / image.height, 1)
    image.width *= scale
    image.height *= scale
    # Keep provenance outside the PNG. Native ZiNiao chrome can abbreviate its
    # store chip even when the complete native window title was matched.
    lines = ["采集记录（软件附注，非截图原文）", f"完整店铺名：{store.get('store_name') or ''}", f"采集时间：{store.get('captured_at') or '未记录'}"]
    caption = "\n".join(lines)
    _text(sheet, row, column, caption)
    sheet.cell(row, column).font = Font(size=10, color="214C68")
    line_count = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1 for char in line) / 62)) for line in lines)
    caption_height = line_count * 17 + 10
    image.anchor = OneCellAnchor(
        _from=AnchorMarker(col=column - 1, row=row - 1, rowOff=pixels_to_EMU(caption_height)),
        ext=XDRPositiveSize2D(pixels_to_EMU(image.width), pixels_to_EMU(image.height)),
    )
    sheet.add_image(image)
    sheet.column_dimensions[get_column_letter(column)].width = 69
    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 24, (image.height + caption_height) * .75 + 10)


def _header(sheet, headers: list[str]) -> None:
    for column, label in enumerate(headers, 1):
        _text(sheet, 1, column, label)
        cell = sheet.cell(1, column)
        cell.fill = PatternFill("solid", fgColor="214C68")
        cell.font = Font(color="FFFFFF", bold=True)
        sheet.column_dimensions[get_column_letter(column)].width = 20
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"


def _note(store: dict) -> str:
    values = [str(store.get("message") or "")]
    values += [str(item) for item in store.get("notes", [])]
    values += [f"{field}: {error}" for field, error in store.get("field_errors", {}).items()]
    labels = {"withdrawal_success": "最新提现已成功，按流程留空", "no_withdrawal": "无提现记录，按流程留空", "cross_border_unavailable": "跨境店余额不可用，按流程填 0"}
    values += [f"{field}: {labels.get(reason, reason)}" for field, reason in store.get("evidence_exemptions", {}).items()]
    return "；".join(dict.fromkeys(item for item in values if item))


def write_balance_workbook(platform: str, stores: list[dict], output_file: str | Path, *, month: str = "") -> str:
    """Write atomically; retain collector images and leave blanks distinct from 0."""
    platform = str(platform).lower()
    if platform not in FIELDS:
        raise ValueError("不支持的余额统计平台")
    target = Path(output_file).resolve()
    target.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    summary = book.active
    summary.title = "余额统计"
    if platform == "temu":
        _header(summary, ["平台", "店铺", "账户总金额", "预估待结算销售额", "截图", "截图（待结算）", "站点", "币种", "状态", "采集时间", "订单创建月份", "说明"])
    else:
        _header(summary, ["店铺名称", "income", "Balance", "Ads", "processing", "站点", "币种", "状态", "采集时间", "说明"])
    for index, store in enumerate(stores, 2):
        values, evidence = store.get("values", {}), store.get("evidence", {})
        if platform == "temu":
            _text(summary, index, 1, "temu")
            _text(summary, index, 2, store.get("store_name"))
            for column, field in ((3, "total"), (4, "pending")):
                _amount(summary, index, column, values.get(field))
                _image(summary, index, column + 2, evidence.get(field, ""), store)
            metadata = [store.get("country"), store.get("currency"), STATUS_NAMES.get(store.get("status"), store.get("status")), store.get("captured_at"), month, _note(store)]
            start = 7
        else:
            _text(summary, index, 1, store.get("store_name"))
            for column, field in enumerate(FIELDS[platform], 2):
                _amount(summary, index, column, values.get(field))
            metadata = [store.get("country"), store.get("currency"), STATUS_NAMES.get(store.get("status"), store.get("status")), store.get("captured_at"), _note(store)]
            start = 6
        for column, value in enumerate(metadata, start):
            _text(summary, index, column, value)
        if store.get("status") != "success":
            for cell in summary[index]:
                cell.fill = PatternFill("solid", fgColor="FFF0DD")
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["B" if platform == "temu" else "A"].width = 34
    summary.column_dimensions["L" if platform == "temu" else "J"].width = 50

    if platform == "lazada":
        for field, name in zip(FIELDS[platform], ("Income截图", "Balance截图", "Ads截图", "processing")):
            sheet = book.create_sheet(name)
            _header(sheet, ["店铺名称", field, "站点", "币种", "状态", "采集时间", "说明", "截图"])
            sheet.column_dimensions["A"].width = 34
            sheet.column_dimensions["G"].width = 45
            for index, store in enumerate(stores, 2):
                _text(sheet, index, 1, store.get("store_name"))
                _amount(sheet, index, 2, store.get("values", {}).get(field))
                for column, value in enumerate([store.get("country"), store.get("currency"), STATUS_NAMES.get(store.get("status"), store.get("status")), store.get("captured_at"), store.get("field_errors", {}).get(field, "")], 3):
                    _text(sheet, index, column, value)
                _image(sheet, index, 8, store.get("evidence", {}).get(field, ""), store)
            sheet.auto_filter.ref = sheet.dimensions

    temporary = target.with_name(f".{target.stem}.{uuid4().hex}.xlsx")
    try:
        book.save(temporary)
        with temporary.open("r+b") as stream:
            os.fsync(stream.fileno())
        os.replace(temporary, target)
    finally:
        book.close()
        temporary.unlink(missing_ok=True)
    return str(target)
