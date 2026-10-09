"""Query Mabang warehouse-transfer arrivals and export one workbook per remark.

The workflow in this module deliberately uses the authenticated ``httpx``
session owned by :class:`backend.core.mabang_client.MabangClient`.  Mabang's
export screen is still an HTML form, but no browser automation is required:
the available fields and hidden defaults are discovered from ``form#theform``
at run time and then submitted through the v2 export endpoints.
"""

from __future__ import annotations

import re
import time
import zipfile
from dataclasses import dataclass, replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence
from urllib.parse import unquote, urlencode, urljoin, urlparse

from bs4 import BeautifulSoup, Tag
import httpx
from openpyxl import load_workbook
import xlrd

from backend.core.mabang_client import MabangApiError, MabangClient, decode_html_page


ProgressCallback = Callable[[str], None]

MABANG_BASE_URL = "https://900853.private.mabangerp.com"
SHIPMENTS_PAGE_PATH = (
    "/index.php?mod=allocationwarehouse.shipments&platform=other&version=1"
)
SEARCH_PATH = "/index.php?mod=warehouseallocation.searchallocation"
EXPORT_ENTRY_PATH = "/index.php?platform=other&version=1"
EXPORT_ACTION_PATH = "/index.php?mod=export.doAllocationWarehouseExportFile"

DEFAULT_ROWS_PER_PAGE = 100
DEFAULT_MAX_PAGES = 500
DEFAULT_STEP_RETRIES = 6
DEFAULT_POLL_INTERVAL_SECONDS = 0.5
DEFAULT_POLL_TIMEOUT_SECONDS = 300.0
DEFAULT_CONTENT_RETRIES = 2

OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
EXCEL_SUFFIXES = {".xls", ".xlsx"}
BATCH_CODE_HEADERS = {"批次编号", "批次号", "调拨批次", "调拨批次号"}
ARRIVAL_OUTPUT_DIR_NAME = "到货查询"
ARRIVAL_FILENAME_PREFIX = "到货查询"
MAX_SEARCH_FILENAME_LENGTH = 80


@dataclass(frozen=True, slots=True)
class ArrivalQueryJob:
    username: str
    password: str
    remarks: tuple[str, ...]
    output_dir: Path
    base_url: str = MABANG_BASE_URL
    rows_per_page: int = DEFAULT_ROWS_PER_PAGE
    max_pages: int = DEFAULT_MAX_PAGES
    step_retries: int = DEFAULT_STEP_RETRIES
    poll_interval_seconds: float = DEFAULT_POLL_INTERVAL_SECONDS
    poll_timeout_seconds: float = DEFAULT_POLL_TIMEOUT_SECONDS


@dataclass(frozen=True, slots=True)
class ArrivalBatch:
    allocation_id: str
    batch_code: str
    remark: str
    matched_remarks: tuple[str, ...] = ()

    @property
    def data_code(self) -> str:
        """Alias for the source page's ``data-code`` attribute."""

        return self.batch_code


@dataclass(frozen=True, slots=True)
class ArrivalRemarkExport:
    remark: str
    matched_batch_count: int
    output_file: Path
    batches: tuple[ArrivalBatch, ...]


@dataclass(frozen=True, slots=True)
class ArrivalQueryResult:
    remark_count: int
    matched_remark_count: int
    matched_batch_count: int
    exported_batch_count: int
    output_files: tuple[Path, ...]
    output_dir: Path
    output_root: Path
    unmatched_remarks: tuple[str, ...]
    batches: tuple[ArrivalBatch, ...]
    exports: tuple[ArrivalRemarkExport, ...]

    @property
    def output_file(self) -> Path:
        """Compatibility alias for consumers that still expect one file."""

        return self.output_files[0]


@dataclass(frozen=True, slots=True)
class ExportField:
    code: str
    label: str


@dataclass(frozen=True, slots=True)
class ExportTemplate:
    page_url: str
    hidden_defaults: tuple[tuple[str, str], ...]
    fields: tuple[ExportField, ...]


def parse_remark_lines(value: Any) -> tuple[str, ...]:
    """Normalize multi-line remarks, dropping blanks and stable-deduplicating."""

    if isinstance(value, str):
        raw_values: Iterable[Any] = (value,)
    elif isinstance(value, (list, tuple, set)):
        raw_values = value
    elif value is None:
        raw_values = ()
    else:
        raw_values = (value,)

    remarks: list[str] = []
    seen: set[str] = set()
    for raw_value in raw_values:
        if raw_value is None:
            continue
        for raw_line in str(raw_value).splitlines():
            remark = raw_line.strip()
            if not remark or remark in seen:
                continue
            seen.add(remark)
            remarks.append(remark)
    return tuple(remarks)


def validate_arrival_query_payload(payload: dict[str, Any] | None) -> ArrivalQueryJob:
    """Validate bridge data and return an immutable arrival-query job."""

    data = dict(payload or {})
    username = _safe_text(data.get("username"))
    password = str(data.get("password") or "")
    raw_remarks = data.get("remarks")
    if raw_remarks is None:
        raw_remarks = data.get("remark_text")
    if raw_remarks is None:
        raw_remarks = data.get("remarks_text")
    remarks = parse_remark_lines(raw_remarks)
    output_text = _safe_text(data.get("output_dir"))
    base_url = (_safe_text(data.get("base_url")) or MABANG_BASE_URL).rstrip("/")

    if not username:
        raise ValueError("请输入马帮账号")
    if not password:
        raise ValueError("请输入马帮密码")
    if not remarks:
        raise ValueError("请至少输入一个备注，每行一个")
    if not output_text:
        raise ValueError("请选择输出目录")
    parsed_base = urlparse(base_url)
    if parsed_base.scheme not in {"http", "https"} or not parsed_base.netloc:
        raise ValueError("马帮地址格式无效")

    rows_per_page = _bounded_int(
        data.get("rows_per_page"),
        default=DEFAULT_ROWS_PER_PAGE,
        minimum=1,
        maximum=500,
    )
    max_pages = _bounded_int(
        data.get("max_pages"),
        default=DEFAULT_MAX_PAGES,
        minimum=1,
        maximum=5000,
    )
    step_retries = _bounded_int(
        data.get("step_retries"),
        default=DEFAULT_STEP_RETRIES,
        minimum=1,
        maximum=30,
    )
    poll_interval = _bounded_float(
        data.get("poll_interval_seconds"),
        default=DEFAULT_POLL_INTERVAL_SECONDS,
        minimum=0.0,
        maximum=60.0,
    )
    poll_timeout = _bounded_float(
        data.get("poll_timeout_seconds"),
        default=DEFAULT_POLL_TIMEOUT_SECONDS,
        minimum=0.01,
        maximum=3600.0,
    )

    return ArrivalQueryJob(
        username=username,
        password=password,
        remarks=remarks,
        output_dir=Path(output_text).expanduser(),
        base_url=base_url,
        rows_per_page=rows_per_page,
        max_pages=max_pages,
        step_retries=step_retries,
        poll_interval_seconds=poll_interval,
        poll_timeout_seconds=poll_timeout,
    )


def build_arrival_search_payload(
    remark: str,
    *,
    page: int = 1,
    rows_per_page: int = DEFAULT_ROWS_PER_PAGE,
) -> dict[str, str]:
    """Build the form used by the shipment list's remark search."""

    return {
        "orderBys[]": "",
        "warehouseMold": "all",
        "startWarehouseIdStr": "",
        "targetWarehouseIdStr": "",
        "search-content1": "remark",
        "search-content-text1": _safe_text(remark),
        "tablebase": "",
        "Orderby": "",
        "third_in_status": "",
        "third_out_status": "",
        "page": str(max(1, int(page))),
        "rowsPerPage": str(max(1, min(int(rows_per_page), 500))),
        "type": "1",
        # Empty is the page's “全部” status, not the pending-signature value 2.
        "allocationstatus": "",
        "startwarhouseId": "",
        "targetwarhouseId": "",
        "timetype": "",
        "datepickerfrom": "",
        "datepickerto": "",
        "orderbysVal": "",
        "orderbydac": "",
        "auditStatus": "",
        "labelId": "",
        "freight_set": "",
        "transportType": "",
    }


def parse_arrival_batches(message_html: Any) -> list[ArrivalBatch]:
    """Parse allocation IDs, ``data-code`` values, and shipment remarks."""

    if message_html is False or message_html is None:
        return []
    html_text = str(message_html or "")
    if not html_text:
        return []

    soup = BeautifulSoup(html_text, "html.parser")
    batches: list[ArrivalBatch] = []
    seen: set[str] = set()
    for checkbox in soup.select('input[name="allot[]"]'):
        allocation_id = _safe_text(checkbox.get("value"))
        batch_code = _safe_text(checkbox.get("data-code"))
        if not allocation_id:
            continue

        allocation_list = checkbox.find_parent("ul")
        container: Tag | None = None
        if allocation_list is not None and isinstance(allocation_list.parent, Tag):
            container = allocation_list.parent
        elif isinstance(checkbox.parent, Tag):
            container = checkbox.parent

        if not batch_code and container is not None:
            batch_link = container.select_one(
                "a[href*='mod=warehouseallocation.editallocation']"
            )
            if batch_link is not None:
                batch_code = _compact_text(batch_link.get_text(" ", strip=True))

        remark_node: Tag | None = None
        if allocation_list is not None:
            selected = allocation_list.select_one("textarea.fhRemark")
            remark_node = selected if isinstance(selected, Tag) else None
        if remark_node is None and container is not None:
            selected = container.select_one("textarea.fhRemark")
            remark_node = selected if isinstance(selected, Tag) else None
        remark = ""
        if remark_node is not None:
            remark = _compact_text(
                _safe_text(remark_node.get_text(" ", strip=True) or remark_node.get("value"))
            )

        key = allocation_id or batch_code
        if key in seen:
            continue
        seen.add(key)
        batches.append(
            ArrivalBatch(
                allocation_id=allocation_id,
                batch_code=batch_code,
                remark=remark,
            )
        )
    return batches


def parse_arrival_total_pages(page_html: Any) -> int:
    """Return zero for an empty result and otherwise parse ``current/total``."""

    if page_html is False or page_html is None:
        return 0
    text = str(page_html or "").strip()
    if not text:
        return 0
    soup_text = BeautifulSoup(text, "html.parser").get_text(" ", strip=True)
    match = re.search(r"[\d,]+\s*/\s*([\d,]+)\s*(?:页)?", soup_text)
    if not match:
        return 1
    return max(1, int(match.group(1).replace(",", "")))


def query_arrival_batches(
    client: MabangClient,
    remarks: Sequence[str],
    *,
    referer: str,
    rows_per_page: int = DEFAULT_ROWS_PER_PAGE,
    max_pages: int = DEFAULT_MAX_PAGES,
    progress: ProgressCallback | None = None,
) -> tuple[tuple[ArrivalBatch, ...], tuple[str, ...]]:
    """Query every remark serially, paging each one before starting the next."""

    by_allocation_id: dict[str, ArrivalBatch] = {}
    unmatched: list[str] = []
    total_remarks = len(remarks)

    for remark_index, remark in enumerate(remarks, start=1):
        _emit(progress, f"[{remark_index}/{total_remarks}] 查询备注: {remark}")
        matched_this_remark = False
        page = 1
        while True:
            payload = build_arrival_search_payload(
                remark,
                page=page,
                rows_per_page=rows_per_page,
            )
            endpoint = urljoin(referer, SEARCH_PATH)
            response = client.client.post(
                endpoint,
                content=urlencode(payload),
                headers={
                    "Referer": referer,
                    "Origin": _origin(endpoint),
                    "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
                    "X-Requested-With": "XMLHttpRequest",
                },
            )
            response.raise_for_status()
            response_payload = _json_object(response, context=f"备注“{remark}”第 {page} 页")
            if response_payload.get("success") is False:
                raise MabangApiError(
                    f"备注“{remark}”第 {page} 页查询失败: {response_payload}"
                )

            message_html = _search_result_html(response_payload)
            page_batches = parse_arrival_batches(message_html)
            if page_batches:
                matched_this_remark = True
            elif message_html and not _is_no_data_fragment(message_html):
                raise MabangApiError(
                    f"备注“{remark}”第 {page} 页返回了无法解析的调拨列表"
                )

            for batch in page_batches:
                existing = by_allocation_id.get(batch.allocation_id)
                if existing is None:
                    by_allocation_id[batch.allocation_id] = replace(
                        batch,
                        matched_remarks=(remark,),
                    )
                    continue
                matched_remarks = existing.matched_remarks
                if remark not in matched_remarks:
                    matched_remarks += (remark,)
                by_allocation_id[batch.allocation_id] = ArrivalBatch(
                    allocation_id=existing.allocation_id,
                    batch_code=existing.batch_code or batch.batch_code,
                    remark=existing.remark or batch.remark,
                    matched_remarks=matched_remarks,
                )

            total_pages = parse_arrival_total_pages(_search_page_html(response_payload))
            if not message_html or page >= max(1, total_pages):
                break
            if total_pages > max_pages or page >= max_pages:
                raise MabangApiError(
                    f"备注“{remark}”分页超过安全上限 {max_pages} 页"
                )
            page += 1

        if not matched_this_remark:
            unmatched.append(remark)
            _emit(progress, f"未查询到备注: {remark}")

    return tuple(by_allocation_id.values()), tuple(unmatched)


def build_export_entry_payload(
    client: MabangClient,
    allocation_ids: Sequence[str],
) -> list[tuple[str, str]]:
    """Build the POST that opens the allocation-warehouse export template."""

    normalized_ids = _stable_nonempty(allocation_ids)
    if not normalized_ids:
        raise ValueError("至少需要一个调拨批次才能导出")
    ids_text = ",".join(normalized_ids) + ","
    return [
        ("mod", "export.exportTemplate"),
        ("datasOpen", "2"),
        ("data", ids_text),
        ("type", "1"),
        ("menu", "allocationWarehouse"),
        ("exportUrl", client._url(EXPORT_ACTION_PATH)),
        ("mainMenu", ""),
        ("showRmbColumn", "0"),
    ]


def parse_export_template(
    page_html: str,
    *,
    page_url: str = "",
) -> ExportTemplate:
    """Discover hidden defaults and every export field from ``form#theform``."""

    soup = BeautifulSoup(page_html or "", "html.parser")
    form = soup.find("form", id="theform")
    if not isinstance(form, Tag):
        raise MabangApiError("到货导出页面缺少 form#theform")

    hidden_defaults: list[tuple[str, str]] = []
    for input_node in form.find_all("input"):
        if not isinstance(input_node, Tag):
            continue
        if _safe_text(input_node.get("type")).lower() != "hidden":
            continue
        if input_node.has_attr("disabled"):
            continue
        name = _safe_text(input_node.get("name"))
        if name:
            hidden_defaults.append((name, str(input_node.get("value") or "")))

    fields: list[ExportField] = []
    seen_codes: set[str] = set()
    for input_node in form.select('input[name="fieldlabel"]'):
        code = _safe_text(input_node.get("value"))
        if not code or code == "all" or code in seen_codes:
            continue
        label = _export_field_label(input_node)
        if not label:
            raise MabangApiError(f"到货导出字段缺少名称: {code}")
        seen_codes.add(code)
        fields.append(ExportField(code=code, label=label))

    if not fields:
        raise MabangApiError("到货导出页面没有可用字段")
    return ExportTemplate(
        page_url=page_url,
        hidden_defaults=tuple(hidden_defaults),
        fields=tuple(fields),
    )


def build_export_step1_payload(
    template: ExportTemplate,
    allocation_ids: Sequence[str],
) -> list[tuple[str, str]]:
    """Select every discovered field and disable “合并共有项”."""

    normalized_ids = _stable_nonempty(allocation_ids)
    if not normalized_ids:
        raise ValueError("至少需要一个调拨批次才能导出")

    overridden = {
        "orderIds",
        "fieldlabel",
        "map-uq[]",
        "map-name[]",
        "map-text[]",
        "tableBase",
        "isMerage",
        "version",
        "step",
        "sn",
        "taskId",
        "taskid",
        "sub_no",
        "1",
    }
    payload = [
        (key, value)
        for key, value in template.hidden_defaults
        if key not in overridden
    ]
    payload.append(("orderIds", ",".join(normalized_ids) + ","))
    for field in template.fields:
        payload.extend(
            [
                ("fieldlabel", field.code),
                ("map-uq[]", field.code),
                ("map-name[]", field.label),
                ("map-text[]", ""),
            ]
        )
    payload.extend(
        [
            ("tableBase", ""),
            ("isMerage", "2"),
            ("version", "v2"),
            ("step", "1"),
        ]
    )
    return payload


def validate_excel_bytes(content: bytes, *, content_type: str = "") -> str:
    """Reject error pages and return the verified Excel file suffix.

    The allocation-warehouse exporter currently returns an OLE2 ``.xls`` file,
    while other Mabang tenants can return an OOXML ``.xlsx`` archive.  Both are
    genuine Excel formats and must keep their real extension on disk.
    """

    if not content:
        raise MabangApiError("马帮导出文件为空")
    normalized_type = _safe_text(content_type).lower()
    prefix = content[:4096].lstrip(b"\xef\xbb\xbf\x00\t\r\n ").lower()
    if "html" in normalized_type or prefix.startswith(
        (b"<!doctype html", b"<html", b"<head", b"<body", b"<form", b"<script")
    ):
        raise MabangApiError("马帮导出返回了 HTML 页面，而不是 Excel 文件")

    stream = BytesIO(content)
    if zipfile.is_zipfile(stream):
        stream.seek(0)
        try:
            workbook = load_workbook(stream, read_only=True, data_only=True)
            try:
                if not workbook.sheetnames:
                    raise MabangApiError("马帮导出 XLSX 不包含工作表")
            finally:
                workbook.close()
        except MabangApiError:
            raise
        except Exception as exc:
            raise MabangApiError(f"马帮导出 XLSX 无法打开: {exc}") from exc
        return ".xlsx"

    if content.startswith(OLE2_MAGIC):
        if (
            len(content) < 512
            or content[28:30] != b"\xfe\xff"
            or int.from_bytes(content[30:32], "little") not in {9, 12}
        ):
            raise MabangApiError("马帮导出 XLS 文件头无效")
        return ".xls"

    raise MabangApiError("马帮导出文件不是有效的 XLSX 或 XLS 文件")


def validate_arrival_export_bytes(
    content: bytes,
    *,
    expected_batch_codes: Sequence[Any],
    content_type: str = "",
) -> str:
    """Validate that an arrival workbook contains this export's business rows.

    A workbook with only column headings is a valid Excel container, but it is
    not a successful arrival export.  Mabang's legacy XLS writer can also emit
    a mildly inconsistent compound-file directory, so ``xlrd`` is opened in
    its documented corruption-tolerant mode before the rows are checked.
    """

    suffix = validate_excel_bytes(content, content_type=content_type)
    normalized_expected = [_safe_text(code) for code in expected_batch_codes]
    if not normalized_expected or any(not code for code in normalized_expected):
        raise MabangApiError("本次查询结果缺少批次编号，无法校验到货导出")
    expected = set(normalized_expected)

    try:
        if suffix == ".xlsx":
            summaries = _xlsx_arrival_sheet_summaries(content)
        else:
            summaries = _xls_arrival_sheet_summaries(content)
    except MabangApiError:
        raise
    except Exception as exc:
        raise MabangApiError(f"马帮到货导出 {suffix.upper()} 内容无法读取: {exc}") from exc

    header_found = any(summary[0] for summary in summaries)
    data_row_count = sum(summary[1] for summary in summaries)
    exported_codes = {
        code
        for _header_found, _row_count, sheet_codes in summaries
        for code in sheet_codes
    }
    if not header_found:
        raise MabangApiError("马帮到货导出缺少“批次编号”列")
    if data_row_count <= 0:
        raise MabangApiError("马帮到货导出无数据（文件只有表头）")

    missing = expected - exported_codes
    if missing:
        raise MabangApiError(
            f"马帮到货导出内容与本次查询不一致，缺少 {len(missing)} 个批次编号"
        )
    unexpected = exported_codes - expected
    if unexpected:
        raise MabangApiError(
            f"马帮到货导出混入其他查询结果，包含 {len(unexpected)} 个额外批次编号"
        )
    return suffix


def _xlsx_arrival_sheet_summaries(
    content: bytes,
) -> list[tuple[bool, int, set[str]]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        return [
            _arrival_sheet_summary(sheet.iter_rows(values_only=True))
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def _xls_arrival_sheet_summaries(
    content: bytes,
) -> list[tuple[bool, int, set[str]]]:
    workbook = xlrd.open_workbook(
        file_contents=content,
        on_demand=True,
        ignore_workbook_corruption=True,
    )
    try:
        return [
            _arrival_sheet_summary(
                sheet.row_values(row_index) for row_index in range(sheet.nrows)
            )
            for sheet in (workbook.sheet_by_index(index) for index in range(workbook.nsheets))
        ]
    finally:
        workbook.release_resources()


def _arrival_sheet_summary(
    rows: Iterable[Sequence[Any]],
) -> tuple[bool, int, set[str]]:
    batch_column: int | None = None
    data_row_count = 0
    batch_codes: set[str] = set()
    for row in rows:
        values = list(row)
        if batch_column is None:
            headers = [_excel_cell_text(value) for value in values]
            batch_column = next(
                (
                    index
                    for index, header in enumerate(headers)
                    if header in BATCH_CODE_HEADERS
                ),
                None,
            )
            continue

        if not any(_excel_cell_text(value) for value in values):
            continue
        data_row_count += 1
        if batch_column < len(values):
            batch_code = _excel_cell_text(values[batch_column])
            if batch_code:
                batch_codes.add(batch_code)
    return batch_column is not None, data_row_count, batch_codes


def _excel_cell_text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _safe_text(value)


def validate_xlsx_bytes(content: bytes, *, content_type: str = "") -> None:
    """Backward-compatible validator name; accepts both Excel formats."""

    validate_excel_bytes(content, content_type=content_type)


def run_mabang_arrival_query(
    job: ArrivalQueryJob | dict[str, Any],
    progress: ProgressCallback | None = None,
) -> ArrivalQueryResult:
    """Run the complete query-and-export request workflow."""

    if isinstance(job, dict):
        job = validate_arrival_query_payload(job)
    if not isinstance(job, ArrivalQueryJob):
        raise TypeError("job 必须是 ArrivalQueryJob 或请求字典")

    _emit(progress, "正在登录马帮...")
    with MabangClient(job.base_url, timeout_seconds=300, verify_ssl=False) as client:
        client.login(job.username, job.password)
        _emit(progress, "正在打开分仓调拨发货页面...")
        shipments_referer = _open_shipments_page(client)
        output_dir = _arrival_output_dir(job.output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        run_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        exports: list[ArrivalRemarkExport] = []
        batches_by_id: dict[str, ArrivalBatch] = {}
        unmatched: list[str] = []

        # Mabang keeps the latest allocation search in the authenticated PHP
        # session.  Export each remark immediately after its search; querying
        # all remarks first leaves only the final search context active and can
        # produce header-only workbooks for every earlier remark.
        for export_index, remark in enumerate(job.remarks, start=1):
            _emit(
                progress,
                f"[{export_index}/{len(job.remarks)}] 查询备注: {remark}",
            )
            remark_batches, _remark_unmatched = query_arrival_batches(
                client,
                (remark,),
                referer=shipments_referer,
                rows_per_page=job.rows_per_page,
                max_pages=job.max_pages,
                progress=None,
            )
            if not remark_batches:
                unmatched.append(remark)
                _emit(progress, f"未查询到备注: {remark}")
                continue

            _emit(
                progress,
                f"[{export_index}/{len(job.remarks)}] 备注“{remark}”命中 "
                f"{len(remark_batches)} 个批次，正在导出独立 Excel...",
            )
            content = b""
            excel_suffix = ""
            for content_attempt in range(1, DEFAULT_CONTENT_RETRIES + 1):
                if content_attempt > 1:
                    _emit(progress, f"备注“{remark}”导出未成功，正在重新查询并重试...")
                    retry_batches, _retry_unmatched = query_arrival_batches(
                        client,
                        (remark,),
                        referer=shipments_referer,
                        rows_per_page=job.rows_per_page,
                        max_pages=job.max_pages,
                        progress=None,
                    )
                    if not retry_batches:
                        raise MabangApiError(
                            f"备注“{remark}”重新查询后未找到调拨批次，无法重试导出"
                        )
                    remark_batches = retry_batches

                try:
                    allocation_ids = [batch.allocation_id for batch in remark_batches]
                    template = _open_export_template(
                        client,
                        allocation_ids,
                        referer=shipments_referer,
                    )
                    final_payload = _run_export_steps(
                        client,
                        template=template,
                        allocation_ids=allocation_ids,
                        step_retries=job.step_retries,
                        poll_interval_seconds=job.poll_interval_seconds,
                        poll_timeout_seconds=job.poll_timeout_seconds,
                    )
                    _suggested_name, content, content_type = _download_export(
                        client,
                        final_payload,
                        referer=template.page_url,
                    )
                    excel_suffix = validate_arrival_export_bytes(
                        content,
                        expected_batch_codes=[
                            batch.batch_code for batch in remark_batches
                        ],
                        content_type=content_type,
                    )
                    break
                except (MabangApiError, httpx.HTTPError) as exc:
                    if content_attempt >= DEFAULT_CONTENT_RETRIES:
                        raise MabangApiError(
                            f"备注“{remark}”导出失败（已重新查询并重试）: {exc}"
                        ) from exc

            output_name = build_arrival_output_filename(
                remark,
                timestamp=run_timestamp,
                suffix=excel_suffix,
            )
            output_file = _unique_output_path(output_dir, output_name)
            output_file.write_bytes(content)
            exports.append(
                ArrivalRemarkExport(
                    remark=remark,
                    matched_batch_count=len(remark_batches),
                    output_file=output_file,
                    batches=remark_batches,
                )
            )
            for batch in remark_batches:
                existing = batches_by_id.get(batch.allocation_id)
                if existing is None:
                    batches_by_id[batch.allocation_id] = batch
                    continue
                matched_remarks = existing.matched_remarks
                if remark not in matched_remarks:
                    matched_remarks += (remark,)
                batches_by_id[batch.allocation_id] = ArrivalBatch(
                    allocation_id=existing.allocation_id,
                    batch_code=existing.batch_code or batch.batch_code,
                    remark=existing.remark or batch.remark,
                    matched_remarks=matched_remarks,
                )
            _emit(progress, f"已导出备注“{remark}”: {output_file.name}")

    batches = tuple(batches_by_id.values())
    if not batches:
        raise MabangApiError("所有备注均未查询到调拨批次，无法导出")

    _emit(progress, f"查询完成，共匹配 {len(batches)} 个唯一调拨批次")
    output_files = tuple(item.output_file for item in exports)
    _emit(progress, f"到货查询完成，共生成 {len(output_files)} 个 Excel 文件: {output_dir}")
    return ArrivalQueryResult(
        remark_count=len(job.remarks),
        matched_remark_count=len(exports),
        matched_batch_count=len(batches),
        exported_batch_count=sum(item.matched_batch_count for item in exports),
        output_files=output_files,
        output_dir=output_dir,
        output_root=job.output_dir,
        unmatched_remarks=tuple(unmatched),
        batches=batches,
        exports=tuple(exports),
    )


def _open_shipments_page(client: MabangClient) -> str:
    outer_response = client.client.get(client._url(SHIPMENTS_PAGE_PATH))
    outer_response.raise_for_status()
    outer_url = str(outer_response.url)
    soup = BeautifulSoup(decode_html_page(outer_response), "html.parser")
    iframe = soup.find("iframe", id="iframeContent") or soup.find("iframe")
    iframe_src = _safe_text(iframe.get("src") if isinstance(iframe, Tag) else "")
    if not iframe_src:
        return outer_url
    iframe_url = urljoin(outer_url, iframe_src)
    iframe_response = client.client.get(iframe_url, headers={"Referer": outer_url})
    iframe_response.raise_for_status()
    return str(iframe_response.url)


def _open_export_template(
    client: MabangClient,
    allocation_ids: Sequence[str],
    *,
    referer: str,
) -> ExportTemplate:
    entry_payload = build_export_entry_payload(client, allocation_ids)
    response = client.client.post(
        client._url(EXPORT_ENTRY_PATH),
        content=urlencode(entry_payload, doseq=True),
        headers={
            "Referer": referer,
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        },
    )
    response.raise_for_status()
    page_url = str(response.url)
    page_html = decode_html_page(response)
    soup = BeautifulSoup(page_html, "html.parser")
    if isinstance(soup.find("form", id="theform"), Tag):
        return parse_export_template(page_html, page_url=page_url)

    iframe = soup.find("iframe")
    iframe_src = _safe_text(iframe.get("src") if isinstance(iframe, Tag) else "")
    if not iframe_src:
        raise MabangApiError("无法定位到货导出模板 form#theform")
    iframe_url = urljoin(page_url, iframe_src)
    iframe_response = client.client.get(iframe_url, headers={"Referer": page_url})
    iframe_response.raise_for_status()
    return parse_export_template(
        decode_html_page(iframe_response),
        page_url=str(iframe_response.url),
    )


def _run_export_steps(
    client: MabangClient,
    *,
    template: ExportTemplate,
    allocation_ids: Sequence[str],
    step_retries: int,
    poll_interval_seconds: float,
    poll_timeout_seconds: float,
) -> dict[str, Any]:
    action_url = client._url(EXPORT_ACTION_PATH)
    step1 = _post_export_action(
        client,
        action_url=action_url,
        payload=build_export_step1_payload(template, allocation_ids),
        referer=template.page_url,
    )
    if not step1.get("success"):
        raise MabangApiError(f"到货导出 step1 失败: {step1}")
    if _extract_file_url(step1):
        return step1

    success_type = _safe_text(step1.get("success_type"))
    if success_type == "1":
        raise MabangApiError(f"到货导出同步响应缺少下载地址: {step1}")
    if success_type != "2":
        raise MabangApiError(f"到货导出返回未知 success_type: {step1}")

    sn = _safe_text(step1.get("sn"))
    try:
        subtask_num = int(step1.get("subtask_num") or 0)
    except (TypeError, ValueError):
        subtask_num = 0
    if not sn or subtask_num <= 0:
        raise MabangApiError(f"到货导出异步任务参数无效: {step1}")

    retry_count = max(1, int(step_retries))
    wait_interval = max(0.1, float(poll_interval_seconds))
    stage_timeout = max(wait_interval, float(poll_timeout_seconds))
    for sub_no in range(1, subtask_num + 1):
        completed = False
        last_step2: dict[str, Any] = {}
        failure_count = 0
        stage_deadline = time.monotonic() + stage_timeout
        while time.monotonic() < stage_deadline:
            last_step2 = _post_export_action(
                client,
                action_url=action_url,
                payload=_export_step_payload(sn, step=2, sub_no=sub_no),
                referer=template.page_url,
            )
            message = _payload_message(last_step2).lower()
            if "请选择填写要导出的数据" in message:
                raise MabangApiError(
                    f"到货导出 step2 未收到可导出数据: sn={sn}, sub_no={sub_no}"
                )
            if last_step2.get("success") and not any(
                marker in message for marker in ("处理中", "processing", "稍后")
            ):
                completed = True
                break
            if not last_step2.get("success"):
                failure_count += 1
                if failure_count >= retry_count:
                    break
            _sleep(wait_interval)
        if not completed:
            raise MabangApiError(
                f"到货导出 step2 失败: sn={sn}, sub_no={sub_no}, response={last_step2}"
            )

    step3: dict[str, Any] = {}
    task_id = ""
    task_ready = False
    failure_count = 0
    stage_deadline = time.monotonic() + stage_timeout
    while time.monotonic() < stage_deadline:
        step3 = _post_export_action(
            client,
            action_url=action_url,
            payload=_export_step_payload(sn, step=3),
            referer=template.page_url,
        )
        if step3.get("success") and _extract_file_url(step3):
            return step3
        task_id = _safe_text(step3.get("taskId") or step3.get("taskid"))
        if step3.get("success") and task_id:
            task_ready = True
            break
        if not step3.get("success"):
            failure_count += 1
            if failure_count >= retry_count:
                break
        _sleep(wait_interval)
    if not task_ready:
        raise MabangApiError(f"到货导出 step3 失败: {step3}")

    deadline = time.monotonic() + max(0.01, float(poll_timeout_seconds))
    last_step4: dict[str, Any] = {}
    while time.monotonic() < deadline:
        last_step4 = _post_export_action(
            client,
            action_url=action_url,
            payload=_export_step_payload(sn, step=4, task_id=task_id),
            referer=template.page_url,
        )
        if (
            last_step4.get("success")
            and _extract_file_url(last_step4)
            and _export_state_is_complete(last_step4)
        ):
            return last_step4
        _sleep(wait_interval)
    raise MabangApiError(
        f"到货导出 step4 超时: sn={sn}, taskId={task_id}, response={last_step4}"
    )


def _post_export_action(
    client: MabangClient,
    *,
    action_url: str,
    payload: Sequence[tuple[str, str]],
    referer: str,
) -> dict[str, Any]:
    response = client.client.post(
        action_url,
        content=urlencode(list(payload), doseq=True),
        headers={
            "Referer": referer,
            "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
            "X-Requested-With": "XMLHttpRequest",
        },
    )
    response.raise_for_status()
    return _json_object(response, context="到货导出")


def _export_step_payload(
    sn: str,
    *,
    step: int,
    sub_no: int | None = None,
    task_id: str = "",
) -> list[tuple[str, str]]:
    payload = [
        ("tableBase", ""),
        ("isMerage", "2"),
        ("version", "v2"),
        ("sn", sn),
    ]
    if sub_no is not None:
        payload.append(("sub_no", str(sub_no)))
    payload.append(("step", str(step)))
    if task_id:
        payload.append(("taskId", task_id))
    payload.append(("1", "1"))
    return payload


def _download_export(
    client: MabangClient,
    payload: dict[str, Any],
    *,
    referer: str,
) -> tuple[str, bytes, str]:
    file_url = _extract_file_url(payload)
    if not file_url:
        raise MabangApiError(f"到货导出响应缺少下载地址: {payload}")
    resolved_url = urljoin(referer or client.base_url.rstrip("/") + "/", file_url)
    response = client.client.get(resolved_url, headers={"Referer": referer})
    response.raise_for_status()
    content_type = response.headers.get("Content-Type") or ""
    name = _download_filename(
        response.headers.get("Content-Disposition"),
        resolved_url,
        content_type=content_type,
        content=response.content,
    )
    return name, response.content, content_type


def _extract_file_url(payload: Any) -> str:
    if not isinstance(payload, dict):
        return ""
    for key in ("file_url", "gourl", "fileUrl", "download_url", "downloadUrl"):
        value = _safe_text(payload.get(key))
        if value:
            return value
    nested = payload.get("data")
    if isinstance(nested, dict):
        return _extract_file_url(nested)
    return ""


def _export_state_is_complete(payload: Any) -> bool:
    if not isinstance(payload, dict):
        return False
    if "state" in payload:
        state = payload.get("state")
        return state is True or _safe_text(state).lower() in {"1", "true"}
    nested = payload.get("data")
    if isinstance(nested, dict):
        return _export_state_is_complete(nested)
    # Some tenants omit ``state`` from the terminal response.  An explicit
    # false value must still block download; a missing value is allowed because
    # the downloaded workbook is subsequently checked for rows and exact batch
    # codes before it can be saved.
    return True


def _download_filename(
    content_disposition: str | None,
    file_url: str,
    *,
    content_type: str = "",
    content: bytes = b"",
) -> str:
    header = content_disposition or ""
    match = re.search(r"filename\*=UTF-8''([^;]+)", header, flags=re.IGNORECASE)
    if match:
        name = unquote(match.group(1)).strip('"')
    else:
        match = re.search(r'filename="?([^";]+)"?', header, flags=re.IGNORECASE)
        name = unquote(match.group(1)).strip('"') if match else ""
    if not name:
        name = Path(urlparse(file_url).path).name
    name = _safe_filename(name)
    if Path(name).suffix.lower() not in EXCEL_SUFFIXES:
        name = _ensure_excel_filename(
            name,
            _infer_excel_suffix(content, content_type=content_type),
        )
    return name


def resolve_arrival_output_dir(output_root: str | Path) -> Path:
    """Return the fixed module directory without nesting it twice."""

    root = Path(output_root).expanduser()
    if root.name.casefold() == ARRIVAL_OUTPUT_DIR_NAME.casefold():
        return root
    return root / ARRIVAL_OUTPUT_DIR_NAME


def build_arrival_output_filename(
    remark: str,
    *,
    timestamp: str | datetime,
    suffix: str,
) -> str:
    """Build ``模块_时间_搜索名.ext`` using a Windows-safe search name."""

    if isinstance(timestamp, datetime):
        timestamp_text = timestamp.strftime("%Y%m%d_%H%M%S")
    else:
        timestamp_text = re.sub(r"[^0-9_-]+", "", _safe_text(timestamp))
    if not timestamp_text:
        timestamp_text = datetime.now().strftime("%Y%m%d_%H%M%S")
    normalized_suffix = suffix.lower() if suffix.lower() in EXCEL_SUFFIXES else ".xlsx"
    search_name = _safe_search_filename_component(remark)
    return (
        f"{ARRIVAL_FILENAME_PREFIX}_{timestamp_text}_{search_name}"
        f"{normalized_suffix}"
    )


def _safe_search_filename_component(value: Any) -> str:
    text = _compact_text(value)
    text = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", text).strip(" ._")
    if not text:
        text = "未命名"
    text = text[:MAX_SEARCH_FILENAME_LENGTH].rstrip(" ._") or "未命名"
    reserved = {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{index}" for index in range(1, 10)),
        *(f"LPT{index}" for index in range(1, 10)),
    }
    if text.upper() in reserved:
        text = f"_{text}"
    return text


def _arrival_output_dir(output_root: str | Path) -> Path:
    return resolve_arrival_output_dir(output_root)


def _unique_output_path(output_dir: Path, suggested_name: str) -> Path:
    safe_name = _safe_filename(suggested_name)
    if Path(safe_name).suffix.lower() not in EXCEL_SUFFIXES:
        safe_name = f"马帮到货查询_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    candidate = output_dir / safe_name
    counter = 2
    while candidate.exists():
        candidate = output_dir / (
            f"{Path(safe_name).stem}_{counter}{Path(safe_name).suffix}"
        )
        counter += 1
    return candidate


def _ensure_excel_filename(name: str, suffix: str) -> str:
    safe_name = _safe_filename(name)
    normalized_suffix = suffix.lower() if suffix.lower() in EXCEL_SUFFIXES else ".xlsx"
    current_suffix = Path(safe_name).suffix
    if current_suffix.lower() == normalized_suffix:
        return safe_name
    stem = Path(safe_name).stem if current_suffix else safe_name
    return f"{stem or '马帮到货查询'}{normalized_suffix}"


def _infer_excel_suffix(content: bytes, *, content_type: str = "") -> str:
    if content.startswith(OLE2_MAGIC):
        return ".xls"
    stream = BytesIO(content)
    if content and zipfile.is_zipfile(stream):
        return ".xlsx"
    normalized_type = _safe_text(content_type).lower()
    if "openxmlformats" in normalized_type:
        return ".xlsx"
    if "vnd.ms-excel" in normalized_type:
        return ".xls"
    return ".xlsx"


def _export_field_label(input_node: Tag) -> str:
    label_node = input_node.find_parent("label")
    if isinstance(label_node, Tag):
        label = _compact_text(label_node.get_text(" ", strip=True))
        if label:
            return label
    sibling = input_node.find_next_sibling(["span", "em", "strong"])
    if isinstance(sibling, Tag):
        label = _compact_text(sibling.get_text(" ", strip=True))
        if label:
            return label
    return _compact_text(
        _safe_text(
            input_node.get("data-name")
            or input_node.get("data-label")
            or input_node.get("title")
        )
    )


def _search_result_html(payload: dict[str, Any]) -> str:
    for key in ("message", "tableContent", "html"):
        value = payload.get(key)
        if value is not False and value is not None and str(value):
            return str(value)
    nested = payload.get("data")
    if isinstance(nested, dict):
        return _search_result_html(nested)
    return ""


def _search_page_html(payload: dict[str, Any]) -> Any:
    for key in ("pageHtml", "page_html", "pagination"):
        if key in payload:
            return payload.get(key)
    nested = payload.get("data")
    if isinstance(nested, dict):
        return _search_page_html(nested)
    return ""


def _is_no_data_fragment(fragment: str) -> bool:
    soup = BeautifulSoup(fragment or "", "html.parser")
    if soup.select_one(".group-nodata, .no-data, .nodata") is not None:
        return True
    text = _compact_text(soup.get_text(" ", strip=True)).lower()
    return any(marker in text for marker in ("暂无数据", "暂无记录", "无数据", "no data"))


def _json_object(response: Any, *, context: str) -> dict[str, Any]:
    try:
        payload = response.json()
    except Exception as exc:
        snippet = _compact_text(decode_html_page(response))[:500]
        raise MabangApiError(f"{context}返回非 JSON 响应: {snippet}") from exc
    if not isinstance(payload, dict):
        raise MabangApiError(f"{context}返回 JSON 结构无效: {payload}")
    return payload


def _payload_message(payload: dict[str, Any]) -> str:
    message = _safe_text(payload.get("message") or payload.get("\x00*\x00message"))
    if message:
        return message
    nested = payload.get("data")
    if isinstance(nested, dict):
        return _payload_message(nested)
    return ""


def _stable_nonempty(values: Sequence[Any]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = _safe_text(value)
        if text and text not in seen:
            seen.add(text)
            result.append(text)
    return result


def _safe_filename(value: Any) -> str:
    name = Path(_safe_text(value)).name
    name = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', "_", name).strip(" .")
    return name or f"马帮到货查询_{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def _origin(url: str) -> str:
    parsed = urlparse(url)
    return f"{parsed.scheme}://{parsed.netloc}"


def _compact_text(value: Any) -> str:
    return re.sub(r"\s+", " ", _safe_text(value)).strip()


def _safe_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _bounded_int(
    value: Any,
    *,
    default: int,
    minimum: int,
    maximum: int,
) -> int:
    try:
        parsed = int(value if value not in (None, "") else default)
    except (TypeError, ValueError):
        parsed = default
    return max(minimum, min(parsed, maximum))


def _bounded_float(
    value: Any,
    *,
    default: float,
    minimum: float,
    maximum: float,
) -> float:
    try:
        parsed = float(value if value not in (None, "") else default)
    except (TypeError, ValueError):
        parsed = default
    return max(minimum, min(parsed, maximum))


def _sleep(seconds: float) -> None:
    if seconds > 0:
        time.sleep(seconds)


def _emit(progress: ProgressCallback | None, message: str) -> None:
    if progress is not None:
        progress(message)


__all__ = [
    "ArrivalBatch",
    "ArrivalQueryJob",
    "ArrivalQueryResult",
    "ArrivalRemarkExport",
    "ExportField",
    "ExportTemplate",
    "build_arrival_search_payload",
    "build_arrival_output_filename",
    "build_export_entry_payload",
    "build_export_step1_payload",
    "parse_arrival_batches",
    "parse_arrival_total_pages",
    "parse_export_template",
    "parse_remark_lines",
    "query_arrival_batches",
    "resolve_arrival_output_dir",
    "run_mabang_arrival_query",
    "validate_arrival_query_payload",
    "validate_arrival_export_bytes",
    "validate_excel_bytes",
    "validate_xlsx_bytes",
]
