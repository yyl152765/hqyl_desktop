"""Lazada 提现统计：PH/MY 流水转 Excel，TH 收入账单下载并上传钉盘。"""

from __future__ import annotations

import csv
import json
import logging
import os
import re
import shutil
import sys
import threading
import time
import unicodedata
import warnings
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable


ProgressCallback = Callable[[str], None]

COUNTRIES: dict[str, dict[str, str]] = {
    "PH": {
        "name": "菲律宾",
        "domain": "sellercenter.lazada.com.ph",
        "drive_root_id": "228929555721",
    },
    "MY": {
        "name": "马来西亚",
        "domain": "sellercenter.lazada.com.my",
        "drive_root_id": "228929487440",
    },
    "TH": {
        "name": "泰国",
        "domain": "sellercenter.lazada.co.th",
        "drive_root_id": "228929395164",
    },
}
DRIVE_SPACE_ID = "28932857521"
MIN_CONCURRENT_STORES = 1
MAX_CONCURRENT_STORES = 4
DEFAULT_CONCURRENT_STORES = 2

_SUPERBROWSER_ROOT: str | None = None
_INVALID_FILE_CHARS = re.compile(r'[\\/:*?"<>|]+')
_STORE_NAME_CHAR_FOLD = str.maketrans(
    {
        "萬": "万",
        "賽": "赛",
        "諾": "诺",
        "鋒": "锋",
        "廣": "广",
        "為": "为",
        "爲": "为",
        "硯": "砚",
        "國": "国",
        "營": "营",
    }
)


@dataclass(frozen=True)
class LazadaWithdrawalStatisticsQuery:
    company: str
    username: str
    password: str
    country: str
    start_date: str
    end_date: str
    store_names: tuple[str, ...]
    max_concurrent_stores: int = DEFAULT_CONCURRENT_STORES
    browser_window_mode: str = "normal"
    client_path: str = ""
    webdriver_path: str = ""
    output_root: str = ""
    dingtalk_app_key: str = ""
    dingtalk_app_secret: str = ""
    dingtalk_user_id: str = ""


@dataclass
class StoreResult:
    country: str
    requested_store_name: str
    store_name: str
    start_date: str
    end_date: str
    status: str = "running"
    message: str = ""
    filter_verified: bool = False
    date_verified: bool = False
    no_data: bool = False
    local_files: list[str] | None = None
    drive_paths: list[str] | None = None
    record_count: int = 0
    downloaded_count: int = 0
    uploaded_count: int = 0
    failed_items: list[str] | None = None
    started_at: str = ""
    finished_at: str = ""

    def __post_init__(self) -> None:
        self.local_files = list(self.local_files or [])
        self.drive_paths = list(self.drive_paths or [])
        self.failed_items = list(self.failed_items or [])
        if not self.started_at:
            self.started_at = _now_text()


class _ProgressHandler(logging.Handler):
    def __init__(self, callback: ProgressCallback | None) -> None:
        super().__init__(logging.INFO)
        self.callback = callback

    def emit(self, record: logging.LogRecord) -> None:
        if self.callback is not None:
            self.callback(self.format(record))


def previous_month_date_range(today: date | None = None) -> dict[str, str]:
    current = today or date.today()
    first_this_month = current.replace(day=1)
    last_previous_month = first_this_month - timedelta(days=1)
    first_previous_month = last_previous_month.replace(day=1)
    return {
        "start_date": first_previous_month.isoformat(),
        "end_date": last_previous_month.isoformat(),
    }


def country_options() -> list[dict[str, str]]:
    return [{"code": code, "name": item["name"]} for code, item in COUNTRIES.items()]


def parse_store_names(value: Any) -> list[str]:
    if isinstance(value, (list, tuple, set)):
        raw_values = [str(item or "").strip() for item in value]
    else:
        text = str(value or "").replace("，", "\n").replace(",", "\n").replace(";", "\n")
        raw_values = [line.strip() for line in text.splitlines()]
    result: list[str] = []
    seen: set[str] = set()
    for item in raw_values:
        if item and item not in seen:
            result.append(item)
            seen.add(item)
    return result


def validate_lazada_withdrawal_payload(
    payload: dict[str, Any],
    *,
    default_output_root: str = "",
    today: date | None = None,
) -> LazadaWithdrawalStatisticsQuery:
    request = dict(payload or {})
    username = str(request.get("username") or "").strip()
    password = str(request.get("password") or "")
    if not username or not password:
        raise ValueError("请先绑定并选择紫鸟账号")

    country = str(request.get("country") or "").strip().upper()
    if country not in COUNTRIES:
        raise ValueError("请选择国家：菲律宾、马来西亚或泰国")

    defaults = previous_month_date_range(today)
    start_text = str(request.get("start_date") or defaults["start_date"]).strip()
    end_text = str(request.get("end_date") or defaults["end_date"]).strip()
    try:
        start_value = date.fromisoformat(start_text)
        end_value = date.fromisoformat(end_text)
    except ValueError as exc:
        raise ValueError("日期格式必须为 YYYY-MM-DD") from exc
    if start_value > end_value:
        raise ValueError("开始日期不得晚于结束日期")
    if (start_value.year, start_value.month) != (end_value.year, end_value.month):
        raise ValueError("开始日期和结束日期必须属于同一个自然月")
    if end_value > (today or date.today()):
        raise ValueError("结束日期不得晚于今天")

    stores = parse_store_names(request.get("store_names"))
    if not stores:
        raise ValueError("请输入店铺名称，每行一个")

    try:
        concurrent = int(request.get("max_concurrent_stores") or DEFAULT_CONCURRENT_STORES)
    except (TypeError, ValueError):
        concurrent = DEFAULT_CONCURRENT_STORES
    if not MIN_CONCURRENT_STORES <= concurrent <= MAX_CONCURRENT_STORES:
        raise ValueError("并发店铺数必须在 1～4 之间")

    window_mode = str(request.get("browser_window_mode") or "normal").strip().lower()
    if window_mode not in {"normal", "background"}:
        raise ValueError("浏览器窗口模式无效")

    return LazadaWithdrawalStatisticsQuery(
        company=str(request.get("company") or "").strip(),
        username=username,
        password=password,
        country=country,
        start_date=start_value.isoformat(),
        end_date=end_value.isoformat(),
        store_names=tuple(stores),
        max_concurrent_stores=concurrent,
        browser_window_mode=window_mode,
        client_path=str(request.get("client_path") or "").strip(),
        webdriver_path=str(request.get("webdriver_path") or "").strip(),
        output_root=str(request.get("output_root") or default_output_root or "").strip(),
        dingtalk_app_key=str(request.get("dingtalk_app_key") or "").strip(),
        dingtalk_app_secret=str(request.get("dingtalk_app_secret") or ""),
        dingtalk_user_id=str(request.get("dingtalk_user_id") or "").strip(),
    )


def _ensure_superbrowser_path() -> None:
    global _SUPERBROWSER_ROOT
    if _SUPERBROWSER_ROOT is not None:
        return
    if getattr(sys, "frozen", False):
        _SUPERBROWSER_ROOT = "<bundled>"
        return
    project_root = Path(__file__).resolve().parents[2]
    candidates = [
        Path(os.environ["SUPERBROWSER_PROCESS_ROOT"]).expanduser()
        if os.environ.get("SUPERBROWSER_PROCESS_ROOT")
        else None,
        project_root.parent / "superbrowser_process",
    ]
    candidate = next((path.resolve() for path in candidates if path and (path / "main").is_dir()), None)
    if candidate is None:
        checked = "、".join(str(path) for path in candidates if path)
        raise ModuleNotFoundError(f"未找到 superbrowser_process（已检查：{checked}）")
    _SUPERBROWSER_ROOT = str(candidate)
    if _SUPERBROWSER_ROOT not in sys.path:
        sys.path.insert(0, _SUPERBROWSER_ROOT)


def _reference_modules() -> tuple[Any, Any, Any, Any]:
    _ensure_superbrowser_path()
    from implement.lazada import lazada_balance_withdrawal as withdrawal
    from main.lazada import lazada_balance_withdrawal_operator_service as operator_service
    from main.super_browser_desktop import LoadSuperBrowser
    from util.webdriver_util import WebdriverUtil

    return operator_service, withdrawal, LoadSuperBrowser, WebdriverUtil


def resolve_lazada_runtime_paths() -> dict[str, str]:
    operator_service, _, _, _ = _reference_modules()
    process_config, _ = operator_service.load_runtime_configs()
    browser = process_config.get("ziniao", {}).get("browser", {}) or {}
    client_path = str(browser.get("client_path") or "")
    webdriver_path = str(browser.get("webdriver_path") or "")
    try:
        from util.ziniao_runtime_util import resolve_webdriver_path, resolve_ziniao_client_path

        client_path = resolve_ziniao_client_path(client_path, browser.get("version", "v6"))
        webdriver_path = resolve_webdriver_path(webdriver_path)
    except Exception:
        pass
    return {"client_path": client_path, "webdriver_path": webdriver_path}


def _safe_name(value: Any) -> str:
    text = _INVALID_FILE_CHARS.sub("_", str(value or "").strip()).rstrip(". ")
    return text[:180] or "store"


def _normalize_store_name(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).translate(_STORE_NAME_CHAR_FOLD)
    return "".join(char for char in text.casefold() if char.isalnum())


def _store_match_score(requested: str, actual: str) -> int:
    if requested == actual:
        return 100
    left = _normalize_store_name(requested)
    right = _normalize_store_name(actual)
    if left and left == right:
        return 90
    if min(len(left), len(right)) >= 8 and (left in right or right in left):
        return 70
    return 0


def _requested_store_for_browser(browser_name: str, requested_names: tuple[str, ...]) -> str:
    scores = [(name, _store_match_score(name, browser_name)) for name in requested_names]
    best_score = max((score for _, score in scores), default=0)
    best = [name for name, score in scores if score == best_score and score > 0]
    return best[0] if len(best) == 1 else browser_name


def _now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _make_logger(run_dir: Path, progress: ProgressCallback | None) -> tuple[logging.Logger, Path]:
    logger = logging.getLogger(f"lazada_withdrawal_statistics.{time.time_ns()}")
    logger.setLevel(logging.DEBUG)
    logger.propagate = False
    log_file = run_dir / "run.log"
    formatter = logging.Formatter("%(asctime)s - %(levelname)s - %(message)s")
    file_handler = logging.FileHandler(log_file, encoding="utf-8")
    file_handler.setFormatter(formatter)
    logger.addHandler(file_handler)
    callback_handler = _ProgressHandler(progress)
    callback_handler.setFormatter(logging.Formatter("%(levelname)s - %(message)s"))
    logger.addHandler(callback_handler)
    return logger, log_file


def _month_folder(start_date: str) -> str:
    return f"{date.fromisoformat(start_date).month}月"


def _create_drive_client(job: LazadaWithdrawalStatisticsQuery) -> Any:
    _ensure_superbrowser_path()
    from util.dingtalk_doc_api import APP_KEY, APP_SECRET, USER_ID
    from util.dingtalk_drive_util import DingDriveClient

    app_key = job.dingtalk_app_key or os.environ.get("LAZADA_WITHDRAWAL_DINGTALK_APP_KEY") or APP_KEY
    app_secret = (
        job.dingtalk_app_secret
        or os.environ.get("LAZADA_WITHDRAWAL_DINGTALK_APP_SECRET")
        or APP_SECRET
    )
    user_id = job.dingtalk_user_id or os.environ.get("LAZADA_WITHDRAWAL_DINGTALK_USER_ID") or USER_ID
    return DingDriveClient(app_key, app_secret, user_id, DRIVE_SPACE_ID)


def _upload_file(
    job: LazadaWithdrawalStatisticsQuery,
    store_name: str,
    local_file: str,
    logger: logging.Logger,
) -> str:
    client = _create_drive_client(job)
    country = COUNTRIES[job.country]
    folder = client.ensure_path(
        [_month_folder(job.start_date), store_name],
        parent_id=country["drive_root_id"],
    )
    file_name = Path(local_file).name
    dentry = client.upload_file(local_file, folder.id, file_name=file_name, overwrite=True)
    drive_path = (
        f"提现统计/lazada{country['name'].replace('马来西亚', '马来')}/"
        f"{_month_folder(job.start_date)}/{store_name}/{file_name}"
    )
    logger.info("钉盘上传完成：%s（dentryId=%s）", drive_path, dentry.id)
    return drive_path


def _dismiss_balance_update_modal(driver: Any) -> bool:
    script = r"""
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim().toLowerCase();
const dialog = Array.from(document.querySelectorAll('[role="dialog"],.next-dialog,.next-overlay-inner'))
  .find((el) => visible(el) && /seller balance updated/i.test(el.innerText || el.textContent || ''));
if (!dialog) return false;
const action = Array.from(dialog.querySelectorAll('button,a,[role="button"]'))
  .find((el) => visible(el) && /checkout it out|got it|close|关闭|知道了/i.test(clean(el.innerText || el.textContent || el.getAttribute('aria-label'))));
const target = action || dialog.querySelector('.next-dialog-close,[aria-label="Close"],[aria-label="close"]');
if (!target) return false;
target.click();
return true;
"""
    try:
        dismissed = bool(driver.execute_script(script))
        if dismissed:
            time.sleep(0.8)
        return dismissed
    except Exception:
        return False


def _dismiss_guided_tours(driver: Any, max_rounds: int = 4) -> int:
    """关闭会遮住筛选区的 Lazada 新手引导，不操作任何业务按钮。"""
    dismissed = 0
    for _ in range(max_rounds):
        try:
            clicked = bool(
                driver.execute_script(
                    r"""
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const safeDismiss = /^(skip|跳过|got it|知道了|关闭|close)$/i;
const candidates = Array.from(document.querySelectorAll('button,a,[role="button"]'))
  .filter((el) => visible(el) && safeDismiss.test(clean(el.innerText || el.textContent || el.getAttribute('aria-label'))));
if (!candidates.length) return false;
const preferred = candidates.find((el) => /^(skip|跳过)$/i.test(clean(el.innerText || el.textContent))) || candidates[0];
preferred.click();
return true;
"""
                )
            )
        except Exception:
            clicked = False
        if not clicked:
            break
        dismissed += 1
        time.sleep(0.8)
    return dismissed


def _scroll_to_balance_transactions(driver: Any) -> bool:
    try:
        moved = bool(
            driver.execute_script(
                r"""
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const nodes = Array.from(document.querySelectorAll('h1,h2,h3,h4,div,span')).filter(visible);
const heading = nodes.find((el) => /^(balance transactions|余额流水)$/i.test(clean(el.innerText || el.textContent)))
  || nodes.find((el) => /balance transactions|余额流水/i.test(clean(el.innerText || el.textContent)));
if (!heading) return false;
heading.scrollIntoView({block: 'start', inline: 'nearest'});
window.scrollBy(0, -180);
return true;
"""
            )
        )
        if moved:
            time.sleep(0.7)
        return moved
    except Exception:
        return False


def _withdrawal_button_state(driver: Any) -> dict[str, Any]:
    return driver.execute_script(
        r"""
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const button = Array.from(document.querySelectorAll('button,[role="button"]'))
  .find((el) => visible(el) && /^(Withdrawal|提现)$/i.test(clean(el.innerText || el.textContent)));
if (!button) return {found: false};
const style = window.getComputedStyle(button);
return {
  found: true,
  text: clean(button.innerText || button.textContent),
  className: clean(button.className),
  ariaPressed: clean(button.getAttribute('aria-pressed')),
  ariaSelected: clean(button.getAttribute('aria-selected')),
  backgroundColor: style.backgroundColor,
  color: style.color,
};
"""
    ) or {"found": False}


def _select_withdrawal_filter(driver: Any) -> tuple[bool, dict[str, Any], dict[str, Any]]:
    _scroll_to_balance_transactions(driver)
    before = _withdrawal_button_state(driver)
    clicked = bool(
        driver.execute_script(
            r"""
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const button = Array.from(document.querySelectorAll('button,[role="button"]'))
  .find((el) => visible(el) && /^(Withdrawal|提现)$/i.test(clean(el.innerText || el.textContent)));
if (!button) return false;
button.scrollIntoView({block: 'center', inline: 'nearest'});
button.click();
return true;
"""
        )
    )
    if clicked:
        time.sleep(1.5)
    after = _withdrawal_button_state(driver)
    selected = clicked and (
        after.get("ariaPressed") == "true"
        or after.get("ariaSelected") == "true"
        or "primary" in str(after.get("className") or "").casefold()
        or before.get("backgroundColor") != after.get("backgroundColor")
    )
    return selected, before, after


def _transaction_table_state(driver: Any, timeout: int = 30) -> dict[str, Any]:
    deadline = time.time() + timeout
    last: dict[str, Any] = {}
    while time.time() < deadline:
        last = driver.execute_script(
            r"""
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const visible = (el) => {
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const nodes = Array.from(document.querySelectorAll('h1,h2,h3,h4,div,section')).filter(visible);
const heading = nodes.find((el) => /^(balance transactions|余额流水)$/i.test(clean(el.innerText || el.textContent)))
  || nodes.find((el) => /balance transactions|余额流水/i.test(clean(el.innerText || el.textContent)));
let container = heading;
while (container && container !== document.body) {
  const text = clean(container.innerText || container.textContent);
  const rect = container.getBoundingClientRect();
  const hasFilter = /withdrawal|提现/i.test(text);
  const hasTable = Boolean(container.querySelector('table,[role="table"],.next-table'));
  if (rect.width >= 650 && hasFilter && hasTable) break;
  container = container.parentElement;
}
if (!container || container === document.body) return {ready: false, reason: 'transaction container not found'};
const headers = Array.from(container.querySelectorAll('th,[role="columnheader"]')).map((el) => clean(el.innerText || el.textContent)).filter(Boolean);
const text = clean(container.innerText || container.textContent);
const rowCount = container.querySelectorAll('tbody tr,[role="rowgroup"] [role="row"]').length;
const noData = /no data|no records|暂无数据|没有数据|无数据/i.test(text) || rowCount === 0;
return {ready: headers.length >= 4, headers, rowCount, noData, text: text.slice(0, 1600)};
"""
        ) or {}
        if last.get("ready"):
            return last
        time.sleep(1)
    return last


WITHDRAWAL_WORKBOOK_HEADERS = ("交易时间", "类型", "金额", "备注", "国家", "店铺")
_WITHDRAWAL_HEADER_ALIASES = {
    "transactionnumber": "transaction_number",
    "流水编号": "transaction_number",
    "交易编号": "transaction_number",
    "transactiontime": "transaction_time",
    "交易时间": "transaction_time",
    "type": "type",
    "类型": "type",
    "amount": "amount",
    "金额": "amount",
    "remarks": "remarks",
    "remark": "remarks",
    "备注": "remarks",
}
_ENGLISH_MONTHS = {
    name: index
    for index, name in enumerate(
        ("jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"),
        start=1,
    )
}


def _transaction_page_snapshot(driver: Any) -> dict[str, Any]:
    return driver.execute_script(
        r"""
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const multiline = (value) => String(value || '')
  .replace(/\u00a0/g, ' ')
  .split(/\r?\n/)
  .map((line) => line.replace(/[ \t]+/g, ' ').trim())
  .filter(Boolean)
  .join('\n');
const visible = (el) => {
  if (!el) return false;
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const nodes = Array.from(document.querySelectorAll('h1,h2,h3,h4,div,section')).filter(visible);
const heading = nodes.find((el) => /^(balance transactions|余额流水)$/i.test(clean(el.innerText || el.textContent)))
  || nodes.find((el) => /balance transactions|余额流水/i.test(clean(el.innerText || el.textContent)));
let container = heading;
while (container && container !== document.body) {
  const text = clean(container.innerText || container.textContent);
  const rect = container.getBoundingClientRect();
  if (rect.width >= 650 && /withdrawal|提现/i.test(text) && container.querySelector('table,[role="table"],.next-table')) break;
  container = container.parentElement;
}
if (!container || container === document.body) return {ready: false, reason: 'transaction container not found'};
const loading = Array.from(container.querySelectorAll('.next-loading,[aria-busy="true"]')).some(visible);
const headers = Array.from(container.querySelectorAll('th,[role="columnheader"]'))
  .filter(visible)
  .map((el) => clean(el.innerText || el.textContent))
  .filter(Boolean);
const rawRows = Array.from(container.querySelectorAll('tbody tr'))
  .filter(visible)
  .map((row) => ({
    cells: Array.from(row.querySelectorAll('td,[role="gridcell"]')).map((cell) => multiline(cell.innerText || cell.textContent)),
  }))
  .filter((row) => row.cells.some(Boolean));
const noDataRow = rawRows.some((row) => row.cells.length === 1 && /no data|no records|暂无数据|没有数据|无数据/i.test(row.cells[0] || ''));
const rows = rawRows.filter((row) => row.cells.length >= headers.length && headers.length > 0);
const nextButton = Array.from(container.querySelectorAll('.next-pagination-item.next-next,button[aria-label*="Next page"],button[aria-label*="下一页"]'))
  .find(visible) || null;
const currentButton = Array.from(container.querySelectorAll('.next-pagination-item.next-current,[aria-current="page"]')).find(visible) || null;
const nextDisabled = !nextButton || Boolean(
  nextButton.disabled
  || nextButton.getAttribute('disabled') !== null
  || nextButton.getAttribute('aria-disabled') === 'true'
  || /disabled/i.test(nextButton.className || '')
);
const text = clean(container.innerText || container.textContent);
const noData = rows.length === 0 && (noDataRow || /no data|no records|暂无数据|没有数据|无数据/i.test(text));
const firstKey = rows.length ? String(rows[0].cells[0] || '') : '';
const lastKey = rows.length ? String(rows[rows.length - 1].cells[0] || '') : '';
const currentPage = currentButton ? clean(currentButton.innerText || currentButton.textContent) : '';
return {
  ready: !loading && headers.length >= 4 && (rows.length > 0 || noData),
  loading,
  headers,
  rows,
  noData,
  nextFound: Boolean(nextButton),
  nextDisabled,
  currentPage,
  signature: [currentPage, firstKey, lastKey, String(rows.length)].join('|'),
};
"""
    ) or {}


def _click_transaction_next_page(driver: Any) -> bool:
    return bool(
        driver.execute_script(
            r"""
const visible = (el) => {
  if (!el) return false;
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const button = Array.from(document.querySelectorAll('.next-pagination-item.next-next,button[aria-label*="Next page"],button[aria-label*="下一页"]'))
  .find(visible);
if (!button || button.disabled || button.getAttribute('disabled') !== null || button.getAttribute('aria-disabled') === 'true' || /disabled/i.test(button.className || '')) return false;
button.scrollIntoView({block: 'center', inline: 'nearest'});
button.click();
return true;
"""
        )
    )


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", unicodedata.normalize("NFKC", str(value or "")).casefold())


def _withdrawal_column_map(headers: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        field = _WITHDRAWAL_HEADER_ALIASES.get(_normalize_header(header))
        if field and field not in mapping:
            mapping[field] = index
    missing = [field for field in ("transaction_time", "type", "amount", "remarks") if field not in mapping]
    if missing:
        raise RuntimeError(f"提现流水表缺少必要列：{', '.join(missing)}；实际表头={headers}")
    return mapping


def _parse_transaction_time(value: Any) -> datetime:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    english = re.fullmatch(r"(\d{1,2})\s+([A-Za-z]{3})\s+(\d{4})\s+(\d{1,2}):(\d{2}):(\d{2})", text)
    if english:
        month = _ENGLISH_MONTHS.get(english.group(2).casefold())
        if month:
            return datetime(
                int(english.group(3)),
                month,
                int(english.group(1)),
                int(english.group(4)),
                int(english.group(5)),
                int(english.group(6)),
            )
    normalized = text.replace("年", "-").replace("月", "-").replace("日", " ").replace("/", "-")
    normalized = re.sub(r"\s+", " ", normalized).strip()
    for pattern in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M:%S"):
        try:
            return datetime.strptime(normalized, pattern)
        except ValueError:
            continue
    raise RuntimeError(f"无法解析提现交易时间：{text!r}")


def _parse_amount(value: Any) -> Decimal:
    text = unicodedata.normalize("NFKC", str(value or "")).strip()
    negative_parentheses = text.startswith("(") and text.endswith(")")
    compact = re.sub(r"[^0-9,\.\-+]", "", text)
    if "," in compact and "." in compact:
        compact = compact.replace(",", "")
    elif "," in compact:
        tail = compact.rsplit(",", 1)[-1]
        compact = compact.replace(",", "." if len(tail) != 3 else "")
    compact = compact.lstrip("+")
    if negative_parentheses and not compact.startswith("-"):
        compact = f"-{compact}"
    try:
        amount = Decimal(compact)
    except (InvalidOperation, ValueError) as exc:
        raise RuntimeError(f"无法解析提现金额：{text!r}") from exc
    if not amount.is_finite():
        raise RuntimeError(f"提现金额不是有限数值：{text!r}")
    return amount


def _map_withdrawal_row(headers: list[str], cells: list[str]) -> dict[str, Any]:
    columns = _withdrawal_column_map(headers)

    def cell(field: str) -> str:
        index = columns.get(field)
        return str(cells[index] if index is not None and index < len(cells) else "").strip()

    transaction_time = cell("transaction_time")
    transaction_type = cell("type")
    amount = cell("amount")
    if not transaction_time or not transaction_type or not amount:
        raise RuntimeError(f"提现流水存在必要字段为空的行：{cells}")
    return {
        "transaction_number": cell("transaction_number"),
        "transaction_time": _parse_transaction_time(transaction_time),
        "type": transaction_type,
        "amount": _parse_amount(amount),
        "remarks": cell("remarks"),
    }


def _collect_withdrawal_transactions(
    driver: Any,
    logger: logging.Logger,
    *,
    max_pages: int = 50,
    page_timeout: int = 30,
) -> tuple[list[dict[str, Any]], int]:
    records: list[dict[str, Any]] = []
    seen_records: set[str] = set()
    seen_pages: set[str] = set()
    page_number = 0
    snapshot: dict[str, Any] = {}

    while page_number < max_pages:
        deadline = time.time() + page_timeout
        while time.time() < deadline:
            snapshot = _transaction_page_snapshot(driver)
            if snapshot.get("ready"):
                break
            time.sleep(0.5)
        if not snapshot.get("ready"):
            raise RuntimeError(f"提现流水当前页未稳定加载：{snapshot}")

        signature = str(snapshot.get("signature") or "")
        if signature in seen_pages:
            raise RuntimeError(f"提现流水分页出现重复页，已停止以避免重复数据：{signature}")
        seen_pages.add(signature)
        page_number += 1

        headers = [str(item or "").strip() for item in snapshot.get("headers") or []]
        _withdrawal_column_map(headers)
        for raw_row in snapshot.get("rows") or []:
            record = _map_withdrawal_row(headers, list(raw_row.get("cells") or []))
            identity = record["transaction_number"] or "|".join(
                (
                    record["transaction_time"].isoformat(),
                    record["type"],
                    str(record["amount"]),
                    record["remarks"],
                )
            )
            if identity in seen_records:
                continue
            seen_records.add(identity)
            records.append(record)
        logger.info("提现流水第 %s 页读取完成：本页 %s 行，累计 %s 行", page_number, len(snapshot.get("rows") or []), len(records))

        if not snapshot.get("nextFound") or snapshot.get("nextDisabled"):
            return records, page_number
        if not _click_transaction_next_page(driver):
            raise RuntimeError("提现流水下一页可用，但点击失败")

        previous_signature = signature
        deadline = time.time() + page_timeout
        while time.time() < deadline:
            next_snapshot = _transaction_page_snapshot(driver)
            if next_snapshot.get("ready") and next_snapshot.get("signature") != previous_signature:
                snapshot = next_snapshot
                break
            time.sleep(0.5)
        else:
            raise RuntimeError("点击提现流水下一页后，表格内容未发生变化")

    raise RuntimeError(f"提现流水超过最大分页限制 {max_pages} 页，已停止")


def _write_withdrawal_workbook(
    output_file: Path,
    records: list[dict[str, Any]],
    *,
    country_name: str,
    store_name: str,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "提现流水"
    sheet.append(list(WITHDRAWAL_WORKBOOK_HEADERS))
    for record in records:
        sheet.append(
            [
                record["transaction_time"],
                record["type"],
                float(record["amount"]),
                record["remarks"],
                country_name,
                store_name,
            ]
        )

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_gray = Side(style="thin", color="D9E2F3")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=thin_gray)
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{max(1, len(records) + 1)}"
    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 16
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 76
    sheet.column_dimensions["E"].width = 13
    sheet.column_dimensions["F"].width = 28
    for row_index in range(2, len(records) + 2):
        sheet.cell(row_index, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        sheet.cell(row_index, 3).number_format = "#,##0.00;[Red]-#,##0.00"
        for column_index in range(1, 7):
            sheet.cell(row_index, column_index).alignment = Alignment(
                vertical="top",
                wrap_text=column_index == 4,
            )
        remarks = str(sheet.cell(row_index, 4).value or "")
        visual_lines = 0
        for logical_line in remarks.splitlines() or [""]:
            visual_units = sum(2 if ord(char) > 127 else 1 for char in logical_line)
            visual_lines += max(1, (visual_units + 74) // 75)
        sheet.row_dimensions[row_index].height = min(180, max(24, visual_lines * 17 + 8))
    sheet.sheet_view.showGridLines = False

    temporary = output_file.with_name(f".{output_file.stem}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, output_file)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)


def validate_withdrawal_workbook(
    workbook_path: Path,
    records: list[dict[str, Any]],
    *,
    country_name: str,
    store_name: str,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    if not workbook_path.is_file() or workbook_path.stat().st_size == 0:
        raise RuntimeError("提现流水 Excel 未生成或为空")
    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    try:
        sheet = workbook["提现流水"]
        headers = tuple(sheet.cell(1, index).value for index in range(1, 7))
        if headers != WITHDRAWAL_WORKBOOK_HEADERS:
            raise RuntimeError(f"提现流水 Excel 表头不匹配：{headers}")
        actual_count = max(0, sheet.max_row - 1)
        if actual_count != len(records):
            raise RuntimeError(f"提现流水 Excel 行数不匹配：期望 {len(records)}，实际 {actual_count}")
        for row_index in range(2, sheet.max_row + 1):
            if sheet.cell(row_index, 1).data_type != "d" or not isinstance(sheet.cell(row_index, 1).value, datetime):
                raise RuntimeError(f"提现流水 Excel 第 {row_index} 行交易时间不是日期值")
            if sheet.cell(row_index, 3).data_type != "n":
                raise RuntimeError(f"提现流水 Excel 第 {row_index} 行金额不是数值")
            if sheet.cell(row_index, 5).value != country_name or sheet.cell(row_index, 6).value != store_name:
                raise RuntimeError(f"提现流水 Excel 第 {row_index} 行国家或店铺不匹配")
        return {"sheet": sheet.title, "record_count": actual_count, "headers": list(headers)}
    finally:
        workbook.close()


def _process_ph_my(
    driver: Any,
    job: LazadaWithdrawalStatisticsQuery,
    store_name: str,
    store_dir: Path,
    logger: logging.Logger,
) -> StoreResult:
    _, withdrawal, _, WebdriverUtil = _reference_modules()
    result = StoreResult(job.country, store_name, store_name, job.start_date, job.end_date)
    try:
        driver._lazada_target_country = job.country
        util = WebdriverUtil(driver)
        withdrawal._ensure_lazada_login(util)
        driver.get(f"https://{COUNTRIES[job.country]['domain']}/")
        withdrawal._wait_page_ready(driver, timeout=30)
        time.sleep(3)
        if not withdrawal.open_lazada_finance_page(util):
            if withdrawal._is_lazada_balance_unavailable(driver):
                result.status = "page_unavailable"
                result.message = "该店铺尚未开放 My Balance（平台试点限制）"
                return result
            raise RuntimeError("无法进入 My Balance 页面")
        withdrawal.close_lazada_popups(driver)
        withdrawal.open_lazada_withdrawal_page(util)
        withdrawal.close_lazada_popups(driver)
        _dismiss_balance_update_modal(driver)
        _dismiss_guided_tours(driver)

        selected, before, after = _select_withdrawal_filter(driver)
        if not selected:
            raise RuntimeError(f"提现筛选未通过选中态校验：before={before}, after={after}")
        result.filter_verified = True

        if not withdrawal.set_trade_date_range(driver, job.start_date, job.end_date):
            raise RuntimeError("日期范围提交失败")
        _dismiss_balance_update_modal(driver)
        if not withdrawal._verify_balance_date_range(driver, job.start_date, job.end_date):
            visible_range = withdrawal._get_balance_date_range_text(driver)
            raise RuntimeError(f"日期范围回读校验失败：{visible_range}")
        result.date_verified = True
        _dismiss_guided_tours(driver)
        _scroll_to_balance_transactions(driver)

        table = _transaction_table_state(driver)
        if not table.get("ready"):
            raise RuntimeError(f"提现流水表未加载或表头不完整：{table}")
        records, page_count = _collect_withdrawal_transactions(driver, logger)
        result.record_count = len(records)
        result.no_data = not records
        workbook_file = store_dir / f"{_safe_name(store_name)}-withdrawal.xlsx"
        _write_withdrawal_workbook(
            workbook_file,
            records,
            country_name=COUNTRIES[job.country]["name"],
            store_name=store_name,
        )
        validation = validate_withdrawal_workbook(
            workbook_file,
            records,
            country_name=COUNTRIES[job.country]["name"],
            store_name=store_name,
        )
        logger.info(
            "提现流水 Excel 校验通过：%s 页、%s 行、表头=%s",
            page_count,
            validation["record_count"],
            validation["headers"],
        )
        result.local_files.append(str(workbook_file))
        result.drive_paths.append(_upload_file(job, store_name, str(workbook_file), logger))
        result.uploaded_count = 1
        result.status = "no_data" if result.no_data else "success"
        result.message = (
            "无提现流水，空表头 Excel 已上传"
            if result.no_data
            else f"提现流水 Excel 已上传，共 {result.record_count} 行"
        )
        return result
    except Exception as exc:
        result.status = "failed"
        result.message = str(exc)
        logger.exception("%s 处理失败", store_name)
        return result
    finally:
        result.finished_at = _now_text()


def _export_rows(driver: Any) -> list[dict[str, Any]]:
    return driver.execute_script(
        r"""
const visible = (el) => {
  if (!el) return false;
  const rect = el.getBoundingClientRect();
  const style = window.getComputedStyle(el);
  return rect.width > 0 && rect.height > 0 && style.display !== 'none' && style.visibility !== 'hidden';
};
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
const dialog = Array.from(document.querySelectorAll('[role="dialog"],.next-dialog,.ant-modal,.next-overlay-inner'))
  .find((el) => visible(el) && /导出记录|Download Result File/i.test(clean(el.innerText || el.textContent)));
if (!dialog) return [];
return Array.from(dialog.querySelectorAll('tr')).slice(1).map((row) => {
  const cells = Array.from(row.querySelectorAll('td')).map((cell) => clean(cell.innerText || cell.textContent));
  const link = Array.from(row.querySelectorAll('a')).find((item) => clean(item.innerText || item.textContent) === 'Download Result File');
  return {exportId: cells[0] || '', exportTime: cells[1] || '', fileName: cells[2] || '', status: cells[3] || '', hasDownloadLink: Boolean(link)};
}).filter((row) => /^\d+$/.test(row.exportId));
"""
    ) or []


def _click_export_link_by_id(driver: Any, export_id: str) -> bool:
    return bool(
        driver.execute_script(
            r"""
const wanted = String(arguments[0] || '');
const clean = (value) => String(value || '').replace(/\s+/g, ' ').trim();
for (const row of Array.from(document.querySelectorAll('tr'))) {
  const cells = Array.from(row.querySelectorAll('td'));
  if (!cells.length || clean(cells[0].innerText || cells[0].textContent) !== wanted) continue;
  const link = Array.from(row.querySelectorAll('a')).find((item) => clean(item.innerText || item.textContent) === 'Download Result File');
  if (!link) return false;
  link.click();
  return true;
}
return false;
""",
            export_id,
        )
    )


def _download_statement_exact(
    driver: Any,
    row: dict[str, Any],
    browser_download_dir: str,
    destination: Path,
    withdrawal: Any,
    *,
    timeout: int = 210,
) -> tuple[str, str]:
    before = withdrawal._snapshot_download_files(browser_download_dir)
    if not withdrawal._open_thailand_statement_download_menu(driver, row.get("element")):
        raise RuntimeError("无法打开账单下载菜单")
    if not withdrawal._click_thailand_order_details_excel(driver):
        raise RuntimeError("无法点击 Order Details (excel)")

    deadline = time.time() + timeout
    export_id = ""
    while time.time() < deadline:
        rows = _export_rows(driver)
        if rows and not export_id:
            export_id = str(rows[0].get("exportId") or "")
        target = next((item for item in rows if str(item.get("exportId") or "") == export_id), None)
        if target and "finished" in str(target.get("status") or "").casefold() and target.get("hasDownloadLink"):
            if not _click_export_link_by_id(driver, export_id):
                raise RuntimeError(f"无法点击本次导出记录 {export_id} 的下载链接")
            source = withdrawal._wait_for_new_download_file(browser_download_dir, before, timeout=120)
            if not source:
                raise RuntimeError(f"导出记录 {export_id} 已完成，但文件未落地")
            destination.mkdir(parents=True, exist_ok=True)
            period_start, period_end = row.get("period_range") or (None, None)
            period_part = withdrawal._format_bill_period_part(period_start, period_end)
            statement_number = _safe_name(row.get("statement_number") or "statement")
            extension = Path(source).suffix.lower()
            if extension not in {".xlsx", ".csv"}:
                raise RuntimeError(f"平台返回的不是 xlsx/csv 账单文件：{Path(source).name}")
            output = destination / f"{statement_number}_{period_part}_Order_Details{extension}"
            shutil.copy2(source, output)
            withdrawal._close_thailand_export_modal(driver)
            return str(output), export_id
        time.sleep(1)
    withdrawal._close_thailand_export_modal(driver)
    raise RuntimeError(f"等待本次导出完成超时，export_id={export_id or '未识别'}")


def validate_statement_workbook(
    file_path: str | Path,
    expected_statement_number: str,
    expected_period: str,
) -> dict[str, str]:
    from openpyxl import load_workbook

    path = Path(file_path)
    if not path.is_file() or path.stat().st_size <= 0 or path.suffix.lower() != ".xlsx":
        raise RuntimeError(f"无效的泰国账单文件：{path}")
    # Lazada 的文件常带错误的 worksheet dimension（A1:A...）。read_only=True
    # 会因此只暴露 A 列；普通模式会按实际单元格恢复完整的 19 列。
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style.*", category=UserWarning)
        workbook = load_workbook(path, read_only=False, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        normalized = [re.sub(r"\s+", "", value).casefold() for value in headers]
        statement_index = next(
            (
                index
                for index, value in enumerate(normalized)
                if "对账单编号" in value or "statementnumber" in value or "statementno" in value
            ),
            None,
        )
        period_index = next(
            (
                index
                for index, value in enumerate(normalized)
                if "账单周期" in value or "billingperiod" in value or value == "period"
            ),
            None,
        )
        if statement_index is None or period_index is None:
            raise RuntimeError(f"Excel 缺少账单号或账期列：{headers}")
        actual_statement = ""
        actual_period = ""
        for values in rows:
            statement = str(values[statement_index] or "").strip() if statement_index < len(values) else ""
            period = str(values[period_index] or "").strip() if period_index < len(values) else ""
            if statement or period:
                actual_statement, actual_period = statement, period
                break
        if actual_statement != expected_statement_number:
            raise RuntimeError(
                f"Excel 内部账单号不匹配：期望 {expected_statement_number}，实际 {actual_statement or '空'}"
            )
        clean_expected_period = re.sub(r"\s+", " ", expected_period).strip()
        clean_actual_period = re.sub(r"\s+", " ", actual_period).strip()
        if clean_actual_period != clean_expected_period:
            raise RuntimeError(
                f"Excel 内部账期不匹配：期望 {clean_expected_period}，实际 {clean_actual_period or '空'}"
            )
        return {"statement_number": actual_statement, "period": actual_period, "sheet": sheet.title}
    finally:
        workbook.close()


def validate_statement_csv(
    file_path: str | Path,
    expected_statement_number: str,
    expected_period: str,
) -> dict[str, str]:
    path = Path(file_path)
    if not path.is_file() or path.stat().st_size <= 0 or path.suffix.lower() != ".csv":
        raise RuntimeError(f"无效的泰国 CSV 账单文件：{path}")
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "utf-16", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                first_line = handle.readline()
                handle.seek(0)
                delimiter = max((",", ";", "\t"), key=first_line.count)
                reader = csv.reader(handle, delimiter=delimiter)
                headers = [str(value or "").strip() for value in next(reader, [])]
                values = next((row for row in reader if any(str(value or "").strip() for value in row)), [])
            normalized = [re.sub(r"\s+", "", value).casefold() for value in headers]
            statement_index = next(
                (
                    index
                    for index, value in enumerate(normalized)
                    if "对账单编号" in value or "statementnumber" in value or "statementno" in value
                ),
                None,
            )
            period_index = next(
                (
                    index
                    for index, value in enumerate(normalized)
                    if "账单周期" in value or "billingperiod" in value or value == "period"
                ),
                None,
            )
            if statement_index is None or period_index is None:
                raise RuntimeError(f"CSV 缺少账单号或账期列：{headers}")
            actual_statement = str(values[statement_index] or "").strip() if statement_index < len(values) else ""
            actual_period = str(values[period_index] or "").strip() if period_index < len(values) else ""
            if actual_statement != expected_statement_number:
                raise RuntimeError(
                    f"CSV 内部账单号不匹配：期望 {expected_statement_number}，实际 {actual_statement or '空'}"
                )
            if re.sub(r"\s+", " ", actual_period).strip() != re.sub(r"\s+", " ", expected_period).strip():
                raise RuntimeError(
                    f"CSV 内部账期不匹配：期望 {expected_period}，实际 {actual_period or '空'}"
                )
            return {"statement_number": actual_statement, "period": actual_period, "sheet": "CSV"}
        except UnicodeError as exc:
            last_error = exc
            continue
    if last_error is not None:
        raise RuntimeError(f"无法识别 CSV 编码：{last_error}") from last_error
    raise RuntimeError("CSV 账单校验失败")


def validate_statement_file(
    file_path: str | Path,
    expected_statement_number: str,
    expected_period: str,
) -> dict[str, str]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".xlsx":
        return validate_statement_workbook(file_path, expected_statement_number, expected_period)
    if suffix == ".csv":
        return validate_statement_csv(file_path, expected_statement_number, expected_period)
    raise RuntimeError(f"不支持的泰国账单格式：{suffix or '无扩展名'}")


def _process_thailand(
    driver: Any,
    job: LazadaWithdrawalStatisticsQuery,
    store_name: str,
    browser_download_dir: str,
    store_dir: Path,
    logger: logging.Logger,
) -> StoreResult:
    _, withdrawal, _, WebdriverUtil = _reference_modules()
    result = StoreResult(job.country, store_name, store_name, job.start_date, job.end_date)
    try:
        if not browser_download_dir or not Path(browser_download_dir).is_dir():
            raise RuntimeError("紫鸟未返回有效下载目录")
        driver._lazada_target_country = "TH"
        util = WebdriverUtil(driver)
        withdrawal._ensure_lazada_login(util)
        driver.get(f"https://{COUNTRIES['TH']['domain']}/")
        withdrawal._wait_page_ready(driver, timeout=30)
        time.sleep(3)
        if not withdrawal.open_lazada_income_statement_page(util):
            raise RuntimeError("无法进入 My Income 页面")
        withdrawal.close_lazada_popups(driver)
        if not withdrawal.open_lazada_income_statement_tab(driver):
            raise RuntimeError("无法切换到 Income Statement")

        target_start = withdrawal._parse_iso_date(job.start_date)
        target_end = withdrawal._parse_iso_date(job.end_date)
        matched_count = 0
        for page_index in range(1, 25):
            rows = withdrawal.wait_for_thailand_income_statement_rows(driver, timeout=35)
            if not rows:
                break
            matched_rows = [
                row
                for row in rows
                if row.get("period_range")
                and withdrawal._date_ranges_intersect(
                    row["period_range"][0], row["period_range"][1], target_start, target_end
                )
            ]
            for original_row in matched_rows:
                matched_count += 1
                statement_number = str(original_row.get("statement_number") or "")
                period = str(original_row.get("period") or "")
                last_error: Exception | None = None
                for attempt in range(1, 3):
                    try:
                        fresh_rows = withdrawal.get_thailand_income_statement_rows(driver)
                        row = next(
                            (item for item in fresh_rows if item.get("statement_number") == statement_number),
                            original_row,
                        )
                        local_file, export_id = _download_statement_exact(
                            driver,
                            row,
                            browser_download_dir,
                            store_dir,
                            withdrawal,
                        )
                        validate_statement_file(local_file, statement_number, period)
                        result.local_files.append(local_file)
                        result.downloaded_count += 1
                        result.drive_paths.append(_upload_file(job, store_name, local_file, logger))
                        result.uploaded_count += 1
                        logger.info(
                            "%s：账单 %s 导出、校验、上传完成（exportId=%s）",
                            store_name,
                            statement_number,
                            export_id,
                        )
                        last_error = None
                        break
                    except Exception as exc:
                        last_error = exc
                        withdrawal._close_thailand_export_modal(driver)
                        if attempt < 2:
                            logger.warning(
                                "%s：账单 %s 第一次处理失败，刷新当前行后重试：%s",
                                store_name,
                                statement_number,
                                exc,
                            )
                            time.sleep(2)
                if last_error is not None:
                    message = f"{statement_number or period}: {last_error}"
                    result.failed_items.append(message)
                    logger.error("%s：账单处理失败：%s", store_name, message)

            oldest_period = rows[-1].get("period_range")
            if oldest_period and oldest_period[1] < target_start:
                break
            if not withdrawal.click_thailand_income_statement_next_page(driver):
                break
            logger.info("%s：进入收入账单第 %s 页", store_name, page_index + 1)

        if matched_count == 0:
            result.status = "no_matching_statement"
            result.message = "统计日期范围内没有匹配账单"
        elif result.uploaded_count == matched_count and not result.failed_items:
            result.status = "success"
            result.message = f"{result.uploaded_count} 个账单已下载、校验并上传"
        elif result.uploaded_count:
            result.status = "partial_success"
            result.message = f"成功 {result.uploaded_count} 个，失败 {len(result.failed_items)} 个"
        else:
            result.status = "failed"
            result.message = f"{matched_count} 个匹配账单均处理失败"
        return result
    except Exception as exc:
        result.status = "failed"
        result.message = str(exc)
        logger.exception("%s 处理失败", store_name)
        return result
    finally:
        result.finished_at = _now_text()


def _run_one_browser(
    loader: Any,
    browser: dict[str, Any],
    job: LazadaWithdrawalStatisticsQuery,
    run_dir: Path,
    logger: logging.Logger,
) -> StoreResult:
    actual_name = str(browser.get("browserName") or "").strip()
    requested_name = _requested_store_for_browser(actual_name, job.store_names)
    store_id = str(browser.get("browserOauth") or browser.get("browserId") or "")
    driver = None
    running_store_id = store_id
    try:
        opened = None
        for attempt in range(1, 3):
            try:
                logger.info("打开店铺：%s（第 %s/2 次）", actual_name, attempt)
                opened = loader.open_store(store_id)
                break
            except Exception:
                if attempt >= 2:
                    raise
                logger.warning("%s：首次启动超时或失败，关闭残留会话后重试", actual_name, exc_info=True)
                try:
                    loader.close_store(store_id)
                except Exception:
                    pass
                time.sleep(2)
        if opened is None:
            raise RuntimeError("紫鸟未返回店铺启动结果")
        running_store_id = str(opened.get("browserOauth") or opened.get("browserId") or store_id)
        driver = loader.get_driver(opened)
        if driver is None:
            raise RuntimeError("紫鸟店铺已打开，但 Selenium 连接失败")
        driver._ziniao_loader = loader
        driver._ziniao_store_name = actual_name
        driver._lazada_target_country = job.country
        driver.set_page_load_timeout(90)
        driver.set_script_timeout(90)
        driver.implicitly_wait(2)
        if job.browser_window_mode == "normal":
            try:
                driver.set_window_size(1500, 950)
            except Exception:
                pass
        else:
            try:
                driver.set_window_position(-32000, -32000)
            except Exception:
                pass
        ip_page = opened.get("ipDetectionPage")
        if ip_page and not loader.open_ip_check(driver, ip_page):
            logger.warning("%s：IP 检测页未识别为成功，继续页面校验", actual_name)
        loader.open_launcher_page(driver, opened.get("launcherPage"))
        store_dir = run_dir / job.country / _safe_name(actual_name)
        store_dir.mkdir(parents=True, exist_ok=True)
        if job.country in {"PH", "MY"}:
            result = _process_ph_my(driver, job, actual_name, store_dir, logger)
        else:
            result = _process_thailand(
                driver,
                job,
                actual_name,
                str(opened.get("downloadPath") or "").strip(),
                store_dir,
                logger,
            )
        result.requested_store_name = requested_name
        return result
    except Exception as exc:
        logger.exception("%s：店铺执行异常", actual_name or requested_name)
        return StoreResult(
            country=job.country,
            requested_store_name=requested_name,
            store_name=actual_name or requested_name,
            start_date=job.start_date,
            end_date=job.end_date,
            status="failed",
            message=str(exc),
            finished_at=_now_text(),
        )
    finally:
        if driver is not None and job.browser_window_mode == "background":
            try:
                driver.set_window_position(50, 50)
            except Exception:
                pass
        if running_store_id:
            try:
                loader.close_store(running_store_id)
            except Exception:
                logger.warning("%s：关闭紫鸟店铺失败", actual_name, exc_info=True)


def run_lazada_withdrawal_statistics(
    job: LazadaWithdrawalStatisticsQuery,
    progress: ProgressCallback | None = None,
) -> dict[str, Any]:
    operator_service, withdrawal, LoadSuperBrowser, _ = _reference_modules()
    root = Path(job.output_root or (Path.cwd() / "outputs"))
    run_dir = root / "lazada_withdrawal_statistics" / datetime.now().strftime("%Y%m%d_%H%M%S")
    run_dir.mkdir(parents=True, exist_ok=True)
    logger, log_file = _make_logger(run_dir, progress)
    results: list[StoreResult] = []
    loader = None
    try:
        logger.info(
            "开始 Lazada %s提现统计：%s 至 %s，共 %s 个店铺",
            COUNTRIES[job.country]["name"],
            job.start_date,
            job.end_date,
            len(job.store_names),
        )
        process_config, global_config = operator_service.load_runtime_configs()
        browser_config = process_config.get("ziniao", {}).get("browser", {}) or {}
        settings = operator_service.LazadaWithdrawalSettings(
            company=job.company,
            username=job.username,
            password=job.password,
            country_mode=job.country,
            start_date=job.start_date,
            end_date=job.end_date,
            store_names=list(job.store_names),
            client_path=job.client_path or str(browser_config.get("client_path") or ""),
            webdriver_path=job.webdriver_path or str(browser_config.get("webdriver_path") or ""),
            browser_window_mode=job.browser_window_mode,
            max_concurrent_stores=job.max_concurrent_stores,
            output_root=str(run_dir),
        )
        process_config = operator_service._merge_settings_into_config(process_config, settings)
        global_config = operator_service._merge_settings_into_config(global_config, settings)
        withdrawal.configure_runtime(output_dir=str(run_dir), logger=logger)
        loader = LoadSuperBrowser(func_dict={}, pcfg=process_config, cfg=global_config, log=logger)
        browser_list = loader.load_super_browser() or []
        if not browser_list:
            raise RuntimeError("未读取到紫鸟浏览器列表")
        matched, unmatched = operator_service._filter_lazada_stores(
            browser_list, list(job.store_names), logger
        )
        for name in unmatched:
            results.append(
                StoreResult(
                    country=job.country,
                    requested_store_name=name,
                    store_name="",
                    start_date=job.start_date,
                    end_date=job.end_date,
                    status="unmatched_store",
                    message="未唯一匹配到 Lazada 紫鸟浏览器",
                    finished_at=_now_text(),
                )
            )
        if matched:
            with ThreadPoolExecutor(max_workers=job.max_concurrent_stores) as executor:
                futures = {
                    executor.submit(_run_one_browser, loader, browser, job, run_dir, logger): browser
                    for browser in matched
                }
                for future in as_completed(futures):
                    results.append(future.result())
        if not matched and not results:
            raise RuntimeError("输入店铺均未匹配到 Lazada 紫鸟浏览器")
    except Exception as exc:
        logger.exception("Lazada 提现统计任务失败")
        if not results:
            results.append(
                StoreResult(
                    country=job.country,
                    requested_store_name="",
                    store_name="",
                    start_date=job.start_date,
                    end_date=job.end_date,
                    status="failed",
                    message=str(exc),
                    finished_at=_now_text(),
                )
            )
    finally:
        if loader is not None:
            try:
                loader.get_exit()
            except Exception:
                logger.warning("退出紫鸟客户端失败", exc_info=True)

    success_statuses = {"success", "no_data"}
    successful = sum(item.status in success_statuses for item in results)
    partial = sum(item.status == "partial_success" for item in results)
    failed = len(results) - successful - partial
    matched_count = sum(item.status != "unmatched_store" for item in results)
    result_file = run_dir / "result.json"
    payload = {
        "success": failed == 0 and partial == 0,
        "message": f"处理完成：成功 {successful}，部分成功 {partial}，失败/未匹配 {failed}",
        "country": job.country,
        "country_name": COUNTRIES[job.country]["name"],
        "start_date": job.start_date,
        "end_date": job.end_date,
        "input_store_count": len(job.store_names),
        "matched_store_count": matched_count,
        "processed_store_count": len(results),
        "success_store_count": successful,
        "partial_store_count": partial,
        "failed_store_count": failed,
        "output_dir": str(run_dir),
        "output_file": str(result_file),
        "log_file": str(log_file),
        "stores": [asdict(item) for item in results],
    }
    result_file.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    logger.info(payload["message"])
    return payload
