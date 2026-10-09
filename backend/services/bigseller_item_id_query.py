from __future__ import annotations

import re
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable
from uuid import uuid4

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

from backend.services.bigseller_sync import (
    BigSellerSyncJob,
    api_headers,
    build_process_config,
    load_bigseller_request_util,
)


ProgressCallback = Callable[[str], None]
MAX_SKU_COUNT = 5000
PREVIEW_ROW_LIMIT = 500
REQUEST_ATTEMPTS = 3
LISTING_PATH = "/v1/product/listing/shopee/pageList.json"
OUTPUT_COLUMNS = ("SKU", "Item ID")


@dataclass(frozen=True)
class BigSellerItemIdQueryJob:
    username: str = field(repr=False)
    password: str = field(repr=False)
    skus: tuple[str, ...]
    output_dir: Path
    captcha_username: str = field(default="", repr=False)
    captcha_password: str = field(default="", repr=False)


class BigSellerQueryError(RuntimeError):
    """A locally generated, credential-free error safe for task logs."""


class BigSellerAuthenticationError(BigSellerQueryError):
    pass


class BigSellerPermissionError(BigSellerQueryError):
    pass


class BigSellerTransientError(BigSellerQueryError):
    pass


class BigSellerRateLimitError(BigSellerQueryError):
    pass


def parse_skus(value: object) -> tuple[str, ...]:
    if not isinstance(value, str):
        raise ValueError("请以文本批量输入 SKU")
    skus = tuple(dict.fromkeys(part.strip() for part in re.split(r"[\r\n,，;；\t]+", value.lstrip("\ufeff")) if part.strip()))
    if not skus:
        raise ValueError("请至少输入一个 SKU")
    if len(skus) > MAX_SKU_COUNT:
        raise ValueError(f"单次最多查询 {MAX_SKU_COUNT} 个不同 SKU，请分批查询")
    if any(ILLEGAL_CHARACTERS_RE.search(sku) or len(sku) > 32767 for sku in skus):
        raise ValueError("SKU 含不支持的控制字符或超出 Excel 单元格长度限制")
    return skus


def validate_bigseller_item_id_query_payload(payload: dict[str, Any]) -> BigSellerItemIdQueryJob:
    username = str(payload.get("username") or "").strip()
    password = str(payload.get("password") or "")
    if not username:
        raise ValueError("请输入 BigSeller 账号")
    if not password.strip():
        raise ValueError("请输入 BigSeller 密码")
    skus = parse_skus(payload.get("sku_text"))
    output_text = str(payload.get("output_dir") or "").strip()
    if not output_text:
        raise ValueError("请选择输出目录")
    output_dir = Path(output_text).expanduser()
    if output_dir.exists() and not output_dir.is_dir():
        raise ValueError("输出目录不能是文件")
    return BigSellerItemIdQueryJob(
        username=username,
        password=password,
        skus=skus,
        output_dir=output_dir,
        captcha_username=str(payload.get("captcha_username") or "").strip(),
        captcha_password=str(payload.get("captcha_password") or ""),
    )


def build_listing_query(sku: str) -> dict[str, Any]:
    # Same read-only endpoint as the Shopee active listings page. No shop filter:
    # the first row is selected across all shops this account can access.
    # Search the SKU field, including variation SKUs exactly as entered.
    # Parent-SKU-only search misses inputs such as hqsZ12027A / hqsZ12027B.
    return {
        "searchType": "sku",
        "searchContent": sku,
        "inquireType": 0,
        "shopeeStatus": "live",
        "status": "active",
        "orderBy": "views",
        "desc": True,
        "pageNo": 1,
        "pageSize": 50,
        "timeType": "create_time",
        "startDateStr": "",
        "endDateStr": "",
    }


def extract_listing_rows(payload: object) -> list[dict[str, Any]]:
    if not isinstance(payload, dict):
        raise BigSellerQueryError("BigSeller 商品列表响应格式异常")
    code = str(payload.get("code"))
    message = str(payload.get("msg") or payload.get("message") or "").lower()
    if code in {"2001", "401", "401006"} or (
        code != "0" and any(word in message for word in ("重新登录", "请登录", "未登录", "login", "session expired"))
    ):
        raise BigSellerAuthenticationError("BigSeller 登录已失效")
    if code == "403":
        raise BigSellerPermissionError("BigSeller 账号无权读取商品列表，请检查账号权限")
    if code == "429":
        raise BigSellerRateLimitError("BigSeller 限制访问频率，请稍后重新查询未完成的 SKU")
    if code != "0":
        # Do not echo server messages: they can contain credentials or request data.
        raise BigSellerQueryError("BigSeller 商品查询接口返回失败，请检查账号权限或稍后重试")
    data = payload.get("data")
    page = data.get("page") if isinstance(data, dict) else None
    rows = page.get("rows") if isinstance(page, dict) else None
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows):
        raise BigSellerQueryError("BigSeller 响应缺少 data.page.rows，无法确认查询结果")
    return rows


def first_item_id(rows: list[dict[str, Any]]) -> str:
    if not rows:
        return ""
    # Never fall back to `id` / `productId` / `listingId` (BigSeller internal IDs),
    # and never skip the first result in favour of another product.
    value = rows[0].get("itemId")
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise BigSellerQueryError("搜索结果第一条缺少有效 Item ID（itemId），未使用其他商品替代")
    item_id = str(value).strip()
    if not item_id or len(item_id) > 32767 or not re.fullmatch(r"[0-9]+", item_id) or not item_id.lstrip("0"):
        raise BigSellerQueryError("搜索结果第一条的 Item ID 格式异常，未使用其他商品替代")
    return item_id


def request_first_item_id(session: Any, util: Any, config: dict[str, Any], sku: str) -> str:
    for attempt in range(REQUEST_ATTEMPTS):
        try:
            response = session.post(
                util.build_api_url(LISTING_PATH),
                json=build_listing_query(sku),
                headers=api_headers(util, config, "shopee", content_type="application/json"),
                timeout=(10, 45),
                allow_redirects=False,
            )
            try:
                if response.status_code == 401 or 300 <= response.status_code < 400:
                    raise BigSellerAuthenticationError("BigSeller 登录已失效")
                if response.status_code == 403:
                    raise BigSellerPermissionError("BigSeller 账号无权读取商品列表，请检查账号权限")
                if response.status_code == 429:
                    raise BigSellerRateLimitError("BigSeller 限制访问频率，请稍后重新查询未完成的 SKU")
                if response.status_code >= 500:
                    raise BigSellerTransientError("BigSeller 暂时繁忙或限制访问，请稍后重试")
                if not 200 <= response.status_code < 300:
                    raise BigSellerQueryError("BigSeller 商品查询请求失败，请检查账号权限或稍后重试")
                try:
                    payload = response.json()
                except ValueError:
                    raise BigSellerQueryError("BigSeller 返回了非 JSON 内容，请检查登录状态或稍后重试") from None
                return first_item_id(extract_listing_rows(payload))
            finally:
                response.close()
        except (requests.Timeout, requests.ConnectionError):
            if attempt == REQUEST_ATTEMPTS - 1:
                raise BigSellerQueryError("BigSeller 商品查询网络超时或连接失败，已重试 3 次") from None
        except BigSellerTransientError:
            if attempt == REQUEST_ATTEMPTS - 1:
                raise BigSellerQueryError("BigSeller 暂时繁忙或限制访问，已重试 3 次") from None
        except requests.RequestException:
            raise BigSellerQueryError("BigSeller 商品查询网络请求失败") from None
        time.sleep(0.5 * (attempt + 1))
    raise BigSellerQueryError("BigSeller 商品查询未完成")


def query_process_config(job: BigSellerItemIdQueryJob) -> dict[str, Any]:
    # Reuse account/captcha resolution only. This module never invokes a sync API.
    config = build_process_config(BigSellerSyncJob(
        username=job.username,
        password=job.password,
        sync_type="product",
        listing_status="live",
        output_dir=job.output_dir,
        captcha_username=job.captcha_username,
        captcha_password=job.captcha_password,
    ))
    config["process_name"] = "BS商品ID查询"
    config["captcha_service"]["max_attempts"] = 3
    return config


def login_for_query(util: Any, config: dict[str, Any], progress: ProgressCallback) -> Any:
    progress("正在登录 BigSeller 并校验登录状态…")
    try:
        # The legacy logger includes CAPTCHA text and raw server errors. Keep it
        # disabled here; expose only locally generated, credential-free messages.
        session = util.login_bigseller_by_request(config, logger=None)
    except Exception as exc:
        # Only recognized reasons are mapped to fixed text. The legacy error may
        # include CAPTCHA answers, account data, or cookies and must not be echoed.
        reason = str(exc)
        if "api.ttshitu.com" in reason and any(status in reason for status in ("500 Server Error", "502 Server Error", "503 Server Error", "504 Server Error")):
            message = "验证码服务返回服务器错误，BigSeller 登录尚未完成，请稍后重试"
        elif "账号或密码错误" in reason:
            message = "BigSeller 登录接口返回账号或密码错误，请核对当前选中账号并重新保存登录凭据"
        elif "账号已停用" in reason:
            message = "BigSeller 登录接口提示账号已停用，请核对当前选中账号及其状态"
        else:
            message = "BigSeller 登录失败，请检查账号密码、验证码服务配置及网络后重试"
        raise BigSellerAuthenticationError(message) from None
    progress("BigSeller 登录成功")
    return session


def export_item_ids(output_dir: Path, rows: list[dict[str, str]]) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = output_dir / f"BS商品ID查询_{datetime.now():%Y%m%d_%H%M%S_%f}_{uuid4().hex[:6]}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SKU-Item ID"
    sheet.append(OUTPUT_COLUMNS)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 30
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
    for index, row in enumerate(rows, 2):
        for column, key in enumerate(("sku", "item_id"), 1):
            cell = sheet.cell(index, column)
            cell.value = row[key]
            # Assigning a leading '=' string defaults to an Excel formula. Force
            # string type explicitly; number_format alone does not prevent that.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.auto_filter.ref = sheet.dimensions
    try:
        workbook.save(output_file)
    finally:
        workbook.close()
    return output_file


def run_bigseller_item_id_query(job: BigSellerItemIdQueryJob, progress: ProgressCallback) -> dict[str, Any]:
    job.output_dir.mkdir(parents=True, exist_ok=True)
    config = query_process_config(job)
    util = load_bigseller_request_util()
    progress(f"准备查询 {len(job.skus)} 个 SKU：Shopee 在售商品，全部可访问店铺")
    progress("规则：SKU（含子SKU） / Fuzzy Search / Views 降序，只取第一条 Item ID")
    session = login_for_query(util, config, progress)
    rows: list[dict[str, str]] = []
    refreshed = False
    stop_reason = ""
    try:
        for index, sku in enumerate(job.skus, 1):
            row = {"sku": sku, "item_id": "", "status": "failed", "message": ""}
            if stop_reason:
                row["message"] = f"未执行：{stop_reason}"
            else:
                try:
                    try:
                        item_id = request_first_item_id(session, util, config, sku)
                    except BigSellerAuthenticationError:
                        if refreshed:
                            raise
                        refreshed = True
                        progress("BigSeller 登录失效，重新登录一次后重试当前 SKU")
                        session.close()
                        session = None
                        session = login_for_query(util, config, progress)
                        item_id = request_first_item_id(session, util, config, sku)
                    row.update(
                        item_id=item_id,
                        status="matched" if item_id else "not_found",
                        message="" if item_id else "未找到匹配商品",
                    )
                except (BigSellerAuthenticationError, BigSellerPermissionError, BigSellerRateLimitError) as exc:
                    stop_reason = str(exc)
                    row["message"] = stop_reason
                    progress("登录、权限或限流异常，停止后续请求；所有未完成 SKU 将保留为空白 Item ID")
                except BigSellerQueryError as exc:
                    row["message"] = str(exc)
                except Exception:
                    # Neither a raw exception nor its traceback may reach logs:
                    # request wrappers can include auth headers and cookie data.
                    row["message"] = "商品查询异常，请检查网络或联系维护人员"
            rows.append(row)
            outcome = f"Item ID：{row['item_id']}" if row["status"] == "matched" else row["message"]
            progress(f"[BS进度 {index}/{len(job.skus)}] {sku} — {outcome}")
    finally:
        if session is not None:
            session.close()

    matched = sum(row["status"] == "matched" for row in rows)
    not_found = sum(row["status"] == "not_found" for row in rows)
    failed = sum(row["status"] == "failed" for row in rows)
    output_file = export_item_ids(job.output_dir, rows)
    progress(f"查询结束：命中 {matched}，未命中 {not_found}，失败 {failed}；Excel 已保存：{output_file}")
    return {
        "sku_count": len(rows),
        "matched_count": matched,
        "not_found_count": not_found,
        "failed_count": failed,
        "rows": rows[:PREVIEW_ROW_LIMIT],
        "preview_limited": len(rows) > PREVIEW_ROW_LIMIT,
        "output_file": str(output_file),
        "output_dir": str(job.output_dir),
    }
