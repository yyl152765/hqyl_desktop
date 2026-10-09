"""Read claim SKUs and append BigSeller dates while preserving the source book."""

from __future__ import annotations

import hashlib
import math
import re
import tempfile
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_SKUS = 10_000
HEADER_SEARCH_ROWS = 20
SKU_HEADERS = ("主SKU", "PARENTSKU", "SKU", "库存SKU", "库存SKU编号")
OUTPUT_HEADERS = {
    "shop_name": "BS最早创建店铺",
    "created_time": "BS创建时间",
    "listed_time": "BS上架时间",
    "status": "BS查询状态",
    "message": "BS查询说明",
}
STATUS_LABELS = {"matched": "已匹配", "not_found": "未找到商品", "failed": "查询失败"}
SITE_LABELS = {"ALL": "全部站点", "ID": "印度尼西亚站点", "PH": "菲律宾站点", "MY": "马来西亚站点",
               "VN": "越南站点", "TH": "泰国站点", "SG": "新加坡站点", "TW": "台湾站点"}
DATE_FORMAT = "yyyy-mm-dd hh:mm:ss"
GENERATED_AUTHOR = "HQYL BigSeller claim"
GENERATED_MARKER = "HQYL_BIGSELLER_CLAIM_V1"
DETAIL_HEADERS = ("来源工作表", "来源行号", "SKU", "最早创建店铺", "Item ID", "创建时间", "上架时间", "查询状态", "说明")
_DATE_IN_TITLE = re.compile(r"\d{1,4}[年./-]\d{1,2}(?:[月./-]\d{1,2})?|\d{1,2}月\d{1,2}日?")
_GROUP_INSTRUCTION = re.compile(r"^(?:以下|以上|下列|下面|本批|本次).*(?:柜子|新品|产品|商品).*(?:按|分组|分配|切开|认领|说明)")
_EMPTY_INLINE_STRING = re.compile(rb'(<c\b(?=[^>]*\bt="inlineStr")[^>]*?)(?:\s*/>|>\s*</c>)')


@dataclass(frozen=True)
class ClaimWorkbookInput:
    source_path: Path
    sheet_name: str
    header_row: int
    sku_column: int
    skus: tuple[str, ...]
    row_numbers: dict[str, tuple[int, ...]]
    warnings: tuple[str, ...] = ()
    source_sha256: str = ""
    skipped_rows: tuple[dict[str, Any], ...] = ()
    row_skus: dict[int, tuple[str, ...]] = field(default_factory=dict)

    @property
    def total_rows(self) -> int:
        """Count source rows once even when a cell contains several SKUs."""
        return len({number for numbers in self.row_numbers.values() for number in numbers})

    @property
    def sku_occurrence_count(self) -> int:
        return sum(len(numbers) for numbers in self.row_numbers.values())

    @property
    def fingerprint(self) -> str:
        return self.source_sha256


def _source_path(path: str | Path) -> Path:
    source = Path(path).expanduser().resolve()
    if source.suffix.lower() != ".xlsx":
        raise ValueError("新品认领查询仅支持 .xlsx，请用 Excel 或 WPS 另存为 .xlsx 后导入")
    if not source.is_file():
        raise ValueError("请选择存在的 Excel 文件")
    return source


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _open_book(source: Path, *, data_only: bool = False):
    # Normal mode derives dimensions from actual cells. Some supplier workbooks
    # incorrectly declare A1 as the entire worksheet dimension.
    try:
        return load_workbook(source, data_only=data_only, keep_links=True)
    except Exception as exc:
        raise ValueError("无法读取 Excel，请确认文件未损坏、未加密且为 .xlsx 格式") from exc


def _header_text(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").lstrip("\ufeff")).upper()


def _find_header(sheet: Any) -> tuple[int, int]:
    candidates = [
        (SKU_HEADERS.index(_header_text(cell.value)), cell.row, cell.column)
        for row in sheet.iter_rows(max_row=min(sheet.max_row, HEADER_SEARCH_ROWS))
        for cell in row if _header_text(cell.value) in SKU_HEADERS
    ]
    if not candidates:
        raise ValueError(f"前 {HEADER_SEARCH_ROWS} 行未找到 SKU、主SKU、Parent SKU 或库存SKU 表头")
    _, row, column = min(candidates)
    return row, column


def _cell_text(cell: Any, cached: Any) -> str:
    value = cell.value
    if cell.data_type == "f":
        value = cached.value
        if value is None:
            raise ValueError("SKU 公式没有缓存值，请用 Excel 计算并保存后重试")
    if cell.data_type == "e" or (cell.data_type == "f" and cached.data_type == "e"):
        raise ValueError("SKU 单元格含 Excel 错误，请修正后重试")
    if value is None:
        return ""
    if isinstance(value, bool):
        raise TypeError("SKU 是布尔值，请改为文本")
    if isinstance(value, (date, datetime)):
        raise TypeError("SKU 是日期值，请核对后改为文本")
    if isinstance(value, (int, float)):
        if not math.isfinite(value):
            raise ValueError("SKU 不是有效数值")
        if int(value) == value:
            text = str(int(value))
            mask = str(cell.number_format or "").split(";", 1)[0].strip()
            if value >= 0 and re.fullmatch(r"0+", mask):
                text = text.zfill(len(mask))
        else:
            text = str(value)
    else:
        text = str(value).strip().lstrip("\ufeff").strip()
    if ILLEGAL_CHARACTERS_RE.search(text):
        raise ValueError("SKU 含无法识别的控制字符")
    return text


def _skip_reason(text: str) -> tuple[str, str] | None:
    # Require a date-like component plus an explicit logistics unit. Do not
    # discard general Chinese identifiers such as 红色杯子-01.
    if _DATE_IN_TITLE.search(text) and re.search(r"柜|\d+\s*件", text):
        return "batch", "批次或物流标题，未作为 SKU 查询"
    if _GROUP_INSTRUCTION.search(text):
        return "invalid", "疑似分组说明，未作为 SKU 查询，请人工核对"
    if _header_text(text) in SKU_HEADERS:
        return "header", "重复 SKU 表头，未作为 SKU 查询"
    return None


def _sheet_input(source: Path, sheet: Any, cached_sheet: Any, fingerprint: str = "") -> ClaimWorkbookInput:
    header_row, sku_column = _find_header(sheet)
    row_numbers: dict[str, list[int]] = {}
    row_skus: dict[int, tuple[str, ...]] = {}
    skipped: list[dict[str, Any]] = []
    spaces: set[str] = set()
    for cells in sheet.iter_rows(min_row=header_row + 1, min_col=sku_column, max_col=sku_column):
        cell = cells[0]
        try:
            text = _cell_text(cell, cached_sheet.cell(cell.row, sku_column))
        except (TypeError, ValueError) as exc:
            skipped.append({"row_number": cell.row, "value": str(cell.value or ""), "kind": "invalid", "reason": str(exc)})
            continue
        tokens = list(dict.fromkeys(part.strip() for part in text.splitlines() if part.strip()))
        valid = []
        for token in tokens:
            reason = _skip_reason(token)
            if reason:
                skipped.append({"row_number": cell.row, "value": token, "kind": reason[0], "reason": reason[1]})
                continue
            valid.append(token)
            row_numbers.setdefault(token, []).append(cell.row)
            if len(row_numbers) > MAX_SKUS:
                raise ValueError(f"不同 SKU 超过 {MAX_SKUS:,} 个，请分批导入")
            if re.search(r"\s", token):
                spaces.add(token)
        if valid:
            row_skus[cell.row] = tuple(valid)
    if not row_numbers:
        raise ValueError("所选工作表没有可查询的 SKU，请检查表头、批次标题或异常单元格")
    warnings = []
    batch_count = sum(row["kind"] == "batch" for row in skipped)
    invalid_count = sum(row["kind"] == "invalid" for row in skipped)
    repeated_count = sum(row["kind"] == "header" for row in skipped)
    if batch_count:
        warnings.append(f"已识别 {batch_count} 条批次或物流标题，不参与查询")
    if invalid_count:
        warnings.append(f"有 {invalid_count} 条异常内容未查询，导出后请查看 BS查询说明 和明细")
    if repeated_count:
        warnings.append(f"已跳过 {repeated_count} 条重复 SKU 表头")
    multi_count = sum(len(skus) > 1 for skus in row_skus.values())
    if multi_count:
        warnings.append(f"有 {multi_count} 个单元格包含多个换行 SKU，将分别查询并按 SKU 标注结果")
    if spaces:
        warnings.append(f"有 {len(spaces)} 个 SKU 含空格，将按完整文本查询，请核对原表")
    return ClaimWorkbookInput(
        source_path=source, sheet_name=sheet.title, header_row=header_row, sku_column=sku_column,
        skus=tuple(row_numbers), row_numbers={sku: tuple(numbers) for sku, numbers in row_numbers.items()},
        warnings=tuple(warnings), source_sha256=fingerprint, skipped_rows=tuple(skipped), row_skus=row_skus,
    )


def _owned_comment(comment: Any, *, kind: str, sheet_name: str, key: str = "") -> bool:
    return bool(comment and comment.author == GENERATED_AUTHOR and comment.text == _marker(kind, sheet_name, key))


def _marker(kind: str, sheet_name: str, key: str = "") -> str:
    return f"{GENERATED_MARKER}\n{kind}\n{sheet_name}\n{key}"


def _is_generated_sheet(sheet: Any) -> bool:
    comment = sheet["A1"].comment
    return bool(comment and comment.author == GENERATED_AUTHOR and comment.text.startswith(f"{GENERATED_MARKER}\ndetail\n"))


def inspect_claim_workbook(path: str | Path) -> dict[str, Any]:
    source = _source_path(path)
    book = _open_book(source)
    cached = _open_book(source, data_only=True)
    try:
        details = []
        for sheet in book.worksheets:
            if _is_generated_sheet(sheet):
                continue
            item = {"name": sheet.title, "sku_count": 0, "row_count": 0, "error": "", "warnings": []}
            try:
                data = _sheet_input(source, sheet, cached[sheet.title])
                item.update(sku_count=len(data.skus), row_count=data.total_rows,
                            skipped_count=len(data.skipped_rows), header_row=data.header_row,
                            sku_column=data.sku_column, warnings=list(data.warnings))
            except ValueError as exc:
                item["error"] = str(exc)
            details.append(item)
        valid = [item for item in details if not item["error"]]
        selected = next((item for item in valid if "新sku上架" in item["name"].lower()), valid[0] if valid else None)
        return {"path": str(source), "source_path": str(source), "sheets": [item["name"] for item in details],
                "sheet_name": selected["name"] if selected else "", "sheet_details": details,
                "warnings": selected["warnings"] if selected else []}
    finally:
        book.close()
        cached.close()


def load_claim_input(source_file: str | Path, sheet_name: str | None = None) -> ClaimWorkbookInput:
    source = _source_path(source_file)
    if not sheet_name:
        sheet_name = inspect_claim_workbook(source)["sheet_name"]
        if not sheet_name:
            raise ValueError("工作簿中没有包含有效 SKU 的工作表")
    fingerprint = _sha256(source)
    book = _open_book(source)
    cached = _open_book(source, data_only=True)
    try:
        if sheet_name not in book.sheetnames or _is_generated_sheet(book[sheet_name]):
            raise ValueError("请选择存在的来源工作表")
        data = _sheet_input(source, book[sheet_name], cached[sheet_name], fingerprint)
        if _sha256(source) != fingerprint:
            raise ValueError("源工作簿在读取过程中已变化，请重新选择文件")
        return data
    finally:
        book.close()
        cached.close()


def _write_value(cell: Any, value: Any) -> None:
    if isinstance(value, str):
        cell.value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
        cell.data_type = "s"
    else:
        cell.value = value


def _timestamp(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            raise ValueError("时间含不明确时区")
        return value
    if isinstance(value, str):
        # Preserve the wall-clock value supplied by BS; Excel cannot store tzinfo.
        return datetime.strptime(value.strip(), "%Y-%m-%d %H:%M:%S")  # noqa: DTZ007
    raise ValueError("时间格式无效")


def _normalize_result(sku: str, result: Mapping[str, Any] | None) -> dict[str, Any]:
    data = dict(result or {})
    data["sku"] = sku
    data["status"] = str(data.get("status") or "failed")
    data["shop_name"] = str(data.get("shop_name") or "").strip()
    data["item_id"] = str(data.get("item_id") or "").strip()
    data["message"] = str(data.get("message") or "")
    if not result:
        data["message"] = "本次未取得查询结果，请重新查询"
    if data["status"] not in STATUS_LABELS:
        data.update(status="failed", message="查询结果状态异常，请重新查询")
    if data["status"] == "matched":
        try:
            data["created_time"] = _timestamp(data.get("created_time"))
            data["listed_time"] = _timestamp(data.get("listed_time"))
            if not data["created_time"] or not data["shop_name"]:
                raise ValueError("缺少店铺或创建时间")
            if data["listed_time"] is None:
                data["message"] = "；".join(filter(None, (data["message"], "BS 未提供该记录的上架时间")))
        except (TypeError, ValueError):
            data.update(status="failed", message="结果缺少有效店铺或日期，请重新查询")
    if data["status"] != "matched":
        data.update(shop_name="", item_id="", created_time=None, listed_time=None)
    return data


def _output_columns(sheet: Any, input_data: ClaimWorkbookInput) -> dict[str, int]:
    columns = {}
    for key, title in OUTPUT_HEADERS.items():
        owned = [cell for cell in sheet[input_data.header_row]
                 if cell.value == title and _owned_comment(cell.comment, kind="column", sheet_name=sheet.title, key=key)]
        if len(owned) > 1:
            raise ValueError(f"检测到多列重复的模块标记：{title}，请保留一列后重试")
        cell = owned[0] if owned else sheet.cell(input_data.header_row, sheet.max_column + 1)
        _write_value(cell, title)
        cell.comment = Comment(_marker("column", sheet.title, key), GENERATED_AUTHOR)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D87")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(cell.column)].width = {
            "shop_name": 28, "created_time": 26, "listed_time": 26, "status": 20, "message": 55,
        }[key]
        columns[key] = cell.column
    return columns


def _detail_sheet(book: Any, source_sheet: str):
    matches = [sheet for sheet in book.worksheets
               if _owned_comment(sheet["A1"].comment, kind="detail", sheet_name=source_sheet)]
    if len(matches) > 1:
        raise ValueError("检测到重复的新品认领查询明细，请保留一份后重试")
    if matches:
        detail = matches[0]
        detail.delete_rows(1, detail.max_row)
    else:
        title = "BS新品认领查询明细"
        suffix = 1
        while title in book.sheetnames:
            suffix += 1
            title = f"BS新品认领查询明细_{suffix}"
        detail = book.create_sheet(title)
    for index, title in enumerate(DETAIL_HEADERS, 1):
        cell = detail.cell(1, index, title)
        cell.font = Font(name="微软雅黑", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="275D87")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    detail["A1"].comment = Comment(_marker("detail", source_sheet), GENERATED_AUTHOR)
    for index, width in enumerate((24, 12, 34, 30, 26, 23, 23, 20, 65), 1):
        detail.column_dimensions[get_column_letter(index)].width = width
    detail.freeze_panes = "D2"
    return detail


def _display(value: Any) -> str:
    return value.strftime("%Y-%m-%d %H:%M:%S") if isinstance(value, datetime) else str(value or "—")


def export_claim_workbook(
    input_data: ClaimWorkbookInput,
    rows: Sequence[Mapping[str, Any]] | Mapping[str, Mapping[str, Any]],
    output_dir: str | Path,
    metadata: Mapping[str, Any] | None = None,
) -> Path:
    source = _source_path(input_data.source_path)
    if input_data.source_sha256 and _sha256(source) != input_data.source_sha256:
        raise ValueError("源工作簿在查询期间已变化，请重新选择文件，避免结果写入错误行")
    destination = Path(output_dir).expanduser().resolve()
    if destination.exists() and not destination.is_dir():
        raise ValueError("输出目录不能是文件")
    lookup = dict(rows) if isinstance(rows, Mapping) else {str(row.get("sku", "")): row for row in rows}
    results = {sku: _normalize_result(sku, lookup.get(sku)) for sku in input_data.skus}
    book = _open_book(source)
    temporary: Path | None = None
    output_path: Path | None = None
    try:
        if input_data.sheet_name not in book.sheetnames:
            raise ValueError("来源工作表已变化，请重新选择文件")
        sheet = book[input_data.sheet_name]
        if _find_header(sheet) != (input_data.header_row, input_data.sku_column):
            raise ValueError("来源表头已变化，请重新选择文件")
        columns = _output_columns(sheet, input_data)
        detail = _detail_sheet(book, sheet.title)
        next_detail_row = 2
        source_rows = dict(input_data.row_skus)
        if not source_rows:
            for sku, numbers in input_data.row_numbers.items():
                for number in numbers:
                    source_rows[number] = (*source_rows.get(number, ()), sku)
        skipped_by_row: dict[int, list[dict[str, Any]]] = {}
        for skipped in input_data.skipped_rows:
            skipped_by_row.setdefault(skipped["row_number"], []).append(skipped)
        # Clear only previously marked result columns, including rows whose SKU
        # was removed since the previous exported workbook became the source.
        for number in range(input_data.header_row + 1, sheet.max_row + 1):
            for column in columns.values():
                sheet.cell(number, column).value = None
        for number in sorted(set(source_rows) | set(skipped_by_row)):
            skus = source_rows.get(number, ())
            skipped = skipped_by_row.get(number, [])
            multiple = len(skus) + len(skipped) > 1
            for key, column in columns.items():
                cell = sheet.cell(number, column)
                if multiple:
                    lines = [f"{sku}：{_display(STATUS_LABELS[results[sku]['status']] if key == 'status' else results[sku].get(key))}"
                             for sku in skus]
                    if key in {"status", "message"}:
                        lines.extend(f"{item['value']}：{'未查询' if key == 'status' else item['reason']}" for item in skipped)
                    _write_value(cell, "\n".join(lines))
                    cell.number_format = "@"
                elif skus:
                    value = STATUS_LABELS[results[skus[0]]["status"]] if key == "status" else results[skus[0]].get(key)
                    _write_value(cell, value)
                    cell.number_format = DATE_FORMAT if key in {"created_time", "listed_time"} else "General"
                elif skipped:
                    _write_value(cell, "未查询" if key == "status" else skipped[0]["reason"] if key == "message" else None)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            for sku in skus:
                result = results[sku]
                values = (sheet.title, number, sku, result["shop_name"], result["item_id"], result.get("created_time"),
                          result.get("listed_time"), STATUS_LABELS[result["status"]], result["message"])
                _append_detail(detail, next_detail_row, values)
                next_detail_row += 1
            for item in skipped:
                _append_detail(detail, next_detail_row, (sheet.title, number, item["value"], "", "", None, None, "未查询", item["reason"]))
                next_detail_row += 1
        detail.auto_filter.ref = f"A1:I{detail.max_row}"
        metadata = dict(metadata or {})
        site = str(metadata.get("site") or "all")
        scope = str(metadata.get("listing_scope") or "live")
        site_label = SITE_LABELS.get(site.upper(), site)
        scope_label = "Shopee在售商品" if scope == "live" else scope
        time_semantics = str(metadata.get("time_semantics") or "采用 BS 平台创建时间").rstrip("；。 ")
        # Metadata lives alongside the audit detail, avoiding changes to the
        # customer's original headings, filters, merged cells or manual values.
        summary = (
            ("查询范围", f"{site_label} / {scope_label} / 当前账号可见店铺"),
            ("上架时间口径", f"{time_semantics}；取同一最早创建商品记录，缺失留空"),
            ("不同 SKU 数", len(results)), ("来源 SKU 行数", input_data.total_rows),
            ("已匹配", sum(row["status"] == "matched" for row in results.values())),
            ("未找到", sum(row["status"] == "not_found" for row in results.values())),
            ("查询失败", sum(row["status"] == "failed" for row in results.values())),
            ("未查询内容数", len(input_data.skipped_rows)),
            ("说明", "同一单元格的多个 SKU 分别查询；店铺、创建时间、上架时间来自同一条最早创建的商品记录。"),
        )
        for index, (label, value) in enumerate(summary, 1):
            _write_value(detail.cell(index, 11), label)
            _write_value(detail.cell(index, 12), value)
            detail.cell(index, 12).alignment = Alignment(vertical="top", wrap_text=True)
        detail.column_dimensions["K"].width = 20
        detail.column_dimensions["L"].width = 72
        if input_data.source_sha256 and _sha256(source) != input_data.source_sha256:
            raise ValueError("源工作簿在导出期间已变化，请重新选择文件")
        destination.mkdir(parents=True, exist_ok=True)
        incomplete = metadata.get("is_complete") is False or any(row["status"] == "failed" for row in results.values())
        stem = re.sub(r'[\\/:*?"<>|]', "_", source.stem)[:70]
        filename = f"{stem}_BS新品认领_{datetime.now().astimezone():%Y%m%d_%H%M%S}_{uuid4().hex[:8]}{'_结果不完整' if incomplete else ''}.xlsx"
        output_path = destination / filename
        with tempfile.NamedTemporaryFile(prefix=".bs-claim-", suffix=".xlsx", dir=destination, delete=False) as stream:
            temporary = Path(stream.name)
        book.save(temporary)
        _preserve_empty_strings(temporary)
        # Exclusive creation guarantees that a collision cannot replace an
        # existing export. Partial copies are removed if disk I/O fails.
        created = False
        try:
            with output_path.open("xb") as target, temporary.open("rb") as source_stream:
                created = True
                for block in iter(lambda: source_stream.read(1024 * 1024), b""):
                    target.write(block)
        except Exception:
            if created:
                output_path.unlink(missing_ok=True)
            raise
        return output_path
    finally:
        book.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


def _append_detail(sheet: Any, number: int, values: Sequence[Any]) -> None:
    for column, value in enumerate(values, 1):
        cell = sheet.cell(number, column)
        _write_value(cell, value)
        cell.number_format = DATE_FORMAT if column in {6, 7} else "@" if column in {3, 5} else "General"
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def _preserve_empty_strings(workbook_file: Path) -> None:
    """Retain explicit empty strings instead of silently turning them into None.

    openpyxl emits an empty inline-string cell without its required text child.
    Adding the empty child keeps source values (including supplier template cells)
    intact when the workbook is reopened. Media and all other ZIP entries are
    copied unchanged.
    """
    with tempfile.NamedTemporaryFile(prefix=".bs-claim-strings-", suffix=".xlsx", dir=workbook_file.parent, delete=False) as stream:
        corrected = Path(stream.name)
    try:
        with ZipFile(workbook_file) as original, ZipFile(corrected, "w") as target:
            for entry in original.infolist():
                content = original.read(entry.filename)
                if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml"):
                    content = _EMPTY_INLINE_STRING.sub(rb"\1><is><t></t></is></c>", content)
                target.writestr(entry, content)
        corrected.replace(workbook_file)
    finally:
        corrected.unlink(missing_ok=True)
