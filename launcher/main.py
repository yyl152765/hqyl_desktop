from __future__ import annotations

import sys
import ctypes
from pathlib import Path


def project_root() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parents[1]


ROOT = project_root()
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from backend.app_bridge import AppBridge


PAGE_ENTRIES = {
    "temu-on-sale-export": "temu-on-sale-export.html",
    "temu-balance-statistics": "temu-balance-statistics.html",
    "lazada-balance-statistics": "lazada-balance-statistics.html",
    "dashboard": "dashboard.html",
    "mabang-arrival-query": "mabang-arrival-query.html",
    "mabang-income-expense-report": "mabang-income-expense-report.html",
    "temu-shipping-channel": "temu-shipping-channel.html",
    "mabang-warehouse-permission": "mabang-warehouse-permission.html",
    "mabang-developer-permission": "mabang-developer-permission.html",
    "sku-inventory-query": "sku-inventory-query.html",
    "developer-sales-income-summary": "developer-sales-income-summary.html",
    "bigseller-item-id-query": "bigseller-item-id-query.html",
    "bigseller-sku-benchmark": "bigseller-sku-benchmark.html",
    "bigseller-claim-query": "bigseller-claim-query.html",
    "lazada-monthly-report": "lazada-monthly-report.html",
    "lazada-ads-data": "lazada-ads-data.html",
    "lazada-bill-detail": "lazada-bill-detail.html",
    "vietnam-income-reconciliation": "vietnam-income-reconciliation.html",
    "kec-reconciliation": "kec-reconciliation.html",
}


def resolve_entry_html(argv: list[str] | None = None) -> Path:
    """Resolve a known desktop page without exposing arbitrary file paths."""
    args = sys.argv[1:] if argv is None else argv
    page_name = "dashboard"
    for argument in args:
        if argument.startswith("--page="):
            page_name = argument.partition("=")[2].strip()
            break

    filename = PAGE_ENTRIES.get(page_name)
    if filename is None:
        choices = ", ".join(sorted(PAGE_ENTRIES))
        raise ValueError(f"unknown desktop page: {page_name}; available: {choices}")

    entry_html = ROOT / "frontend" / "pages" / filename
    if not entry_html.exists():
        raise RuntimeError(f"frontend entry not found: {entry_html}")
    return entry_html


def self_check() -> int:
    from tempfile import TemporaryDirectory
    from types import SimpleNamespace

    from openpyxl import Workbook, load_workbook

    from backend import app_bridge as bridge_module
    from backend.config_store import AppSettings

    for page_name in PAGE_ENTRIES:
        resolve_entry_html([f"--page={page_name}"])

    sample_html = ROOT / "frontend" / "pages" / "sample-registration.html"
    if not sample_html.exists():
        raise RuntimeError(f"sample registration page not found: {sample_html}")

    # Package checks must never read, migrate, or write the user's settings.
    # AppBridge.__init__ only creates store/task objects; no config is loaded.
    bridge = AppBridge()
    bridge.config_store = SimpleNamespace(load=AppSettings)
    original_group_options = bridge_module.get_group_options
    try:
        # The regular group picker can read a source YAML containing credentials.
        bridge_module.get_group_options = lambda: []
        info = bridge.get_app_info()
    finally:
        bridge_module.get_group_options = original_group_options
    if not info.get("ok"):
        raise RuntimeError(f"bridge self-check failed: {info}")

    # Exercise the dynamic imports that PyInstaller cannot discover on its own.
    from backend.services.shopee_ads_recharge import SITE_OPTIONS, get_site_module

    for site in SITE_OPTIONS:
        module = get_site_module(site.code)
        if not hasattr(module, "run_ads_recharge"):
            raise RuntimeError(f"Shopee Ads module is incomplete: {site.module_name}")

    # Check the bundled BigSeller loader without creating a session or logging in.
    from backend.services.bigseller_sync import load_bigseller_request_util
    from backend.services.bigseller_item_id_query import build_listing_query, export_item_ids

    bigseller_util = load_bigseller_request_util()
    for name in (
        "login_bigseller_by_request", "build_api_url", "get_request_user_agent",
        "get_listing_page_url", "build_request_headers",
    ):
        if not callable(getattr(bigseller_util, name, None)):
            raise RuntimeError(f"BigSeller request module is incomplete: {name}")

    query = build_listing_query("self-check SKU")
    expected_query = {
        "searchType": "sku", "searchContent": "self-check SKU", "inquireType": 0,
        "orderBy": "views", "desc": True, "pageNo": 1,
        "shopeeStatus": "live", "status": "active",
    }
    if (
        any(query.get(key) != value for key, value in expected_query.items())
        or query.get("desc") is not True
        or type(query.get("inquireType")) is not int
    ):
        raise RuntimeError("BigSeller Item ID query contract is invalid")

    # Only synthetic values are exported; the temporary directory is removed.
    from backend.services.sample_registration_config import load_config, validate_config

    with TemporaryDirectory(prefix="hqyl-self-check-") as temp_dir:
        check_dir = Path(temp_dir)
        # Verify the new SKU coverage service is present and functional inside
        # the frozen executable without logging in or reading a real account.
        from backend.services.sku_inventory_query import DeveloperOption
        from backend.services.sku_warehouse_coverage import (
            coverage_record, export_coverage_records, parse_scope_warehouses,
            unclassified_warehouses,
        )

        coverage_html = '''<select name="defaultWarehouseId">
            <option value="TH01">泰国TH01仓</option>
            <option value="PH01">菲律宾PH01仓</option>
            <option value="TRANSIT">泰国中转仓</option>
            <option value="1100808">Ola-Ola海外仓</option>
            <option value="1100602">T&amp;C 海外 2 号仓-T&amp;C海外仓</option>
        </select>'''
        coverage_scope = parse_scope_warehouses(coverage_html, {})
        if len(coverage_scope) != 2 or unclassified_warehouses(coverage_html, {}):
            raise RuntimeError("SKU warehouse scope exclusions failed")
        coverage = coverage_record(
            {"stockSku": "=SKU-001", "nameCN": "自检商品", "livenessType": "1",
             "stockWarehouseData": [{"warehouseId": "TH01", "stockWarehouseAvailableInventory": 0}]},
            DeveloperOption("self-check", "自检-开发"), coverage_scope,
        )
        if coverage["missing_warehouse_ids"] != "PH01" or coverage["covered_warehouse_count"] != 1:
            raise RuntimeError("SKU zero-stock warehouse association check failed")
        coverage_output = export_coverage_records([coverage], coverage_scope, check_dir)
        coverage_book = load_workbook(coverage_output, read_only=True, data_only=False)
        try:
            if (coverage_book.sheetnames != ["SKU缺失仓库", "东南亚仓库范围", "查询口径"]
                    or coverage_book.worksheets[0]["B2"].value != "=SKU-001"
                    or coverage_book.worksheets[0]["B2"].data_type != "s"
                    or coverage_book.worksheets[0]["I2"].value != "PH01"
                    or coverage_book.worksheets[1].max_row != 3):
                raise RuntimeError("SKU coverage frozen workbook check failed")
        finally:
            coverage_book.close()
        cfg = load_config(check_dir / "sample_registration.json")
        _ = validate_config(cfg)
        sample_rows = [
            {"sku": "=1+1", "item_id": "001234567890123456789"},
            {"sku": "00001", "item_id": ""},
        ]
        output_file = export_item_ids(check_dir, sample_rows)
        workbook = load_workbook(output_file, read_only=True, data_only=False)
        try:
            sheet = workbook.active
            if (
                len(workbook.worksheets) != 1
                or sheet.max_column != 2
                or sheet.max_row != len(sample_rows) + 1
                or tuple(cell.value for cell in sheet[1]) != ("SKU", "Item ID")
            ):
                raise RuntimeError("BigSeller Excel must contain only SKU / Item ID columns")
            for index, row in enumerate(sample_rows, 2):
                for column, key in enumerate(("sku", "item_id"), 1):
                    cell = sheet.cell(index, column)
                    if (
                        (cell.value or "") != row[key]
                        or cell.number_format != "@"
                        or (row[key] and cell.data_type != "s")
                    ):
                        raise RuntimeError("BigSeller Excel text preservation check failed")
        finally:
            workbook.close()

        # Verify the TEMU workbook pipeline in the frozen runtime with synthetic
        # data. No CLI, browser session, real store, or user settings are accessed.
        from backend.services.temu_on_sale_workbook import COLUMNS, SHEET_NAME, inspect_source, merge_sources

        temu_source = check_dir / "temu-input.xlsx"
        temu_book = Workbook()
        temu_sheet = temu_book.active
        temu_sheet.title = SHEET_NAME
        temu_sheet.append(list(COLUMNS[1:]))
        temu_row = {"商品标题": "自检商品", "SPU ID": "0001", "SKC ID": "0002",
                    "SKU ID": "001234567890123456789", "商品状态": "在售中",
                    "申报价格状态": "已作废", "库存": 0}
        temu_sheet.append([temu_row.get(name) for name in COLUMNS[1:]])
        temu_sheet.row_dimensions[2].hidden = True
        temu_book.save(temu_source)
        temu_book.close()
        temu_output = check_dir / "temu-summary.xlsx"
        temu_info = inspect_source(temu_source)
        merged = merge_sources([("自检店铺", temu_info)], temu_output)
        temu_book = load_workbook(temu_output, read_only=True)
        try:
            values = list(temu_book[SHEET_NAME].iter_rows(values_only=True))
            if (merged["row_count"] != 1 or values[0] != COLUMNS
                    or values[1][0] != "自检店铺" or values[1][4] != temu_row["SKU ID"]
                    or values[1][14] != "已作废" or values[1][15] != 0
                    or inspect_source(temu_source).sha256 != temu_info.sha256):
                raise RuntimeError("TEMU frozen workbook preservation check failed")
        finally:
            temu_book.close()

        # Exercise SQLite's bundled native library and the benchmark exports
        # using synthetic data only; no account lookup or API query is needed.
        from backend.services.bigseller_benchmark_checkpoint import create_checkpoint, read_checkpoint
        from backend.services.bigseller_benchmark_workbook import export_benchmark_workbook, load_workbook_input
        from backend.services.bigseller_sku_benchmark import build_benchmark_query

        source_file = check_dir / "benchmark-input.xlsx"
        source_book = Workbook()
        source_book.active.append(["SKU"])
        source_book.active.append(["00001"])
        source_book.active.append(["00002"])
        source_book.save(source_file)
        source_book.close()
        input_data = load_workbook_input(source_file)
        benchmark_rows = [
            {"sku": "00001", "status": "matched", "shop_name": "Self-check shop",
             "metric_value": 0, "shop_count": 1, "attempts": 1},
            {"sku": "00002", "status": "failed", "message": "Self-check retry",
             "attempts": 2, "retry_not_before": 2000000000.0},
        ]
        summary = {"is_complete": False, "total_rows": 2, "sku_count": 2,
                   "matched_count": 1, "not_found_count": 0, "failed_count": 1,
                   "confirmed_count": 1, "retry_rounds_used": 1,
                   "recovered_count": 0, "resumed_count": 0}
        for metric, label in (("views", "浏览量"), ("sales", "销量")):
            query = build_benchmark_query("00001", metric)
            if query.get("orderBy") != metric or query.get("searchContent") != "00001":
                raise RuntimeError("BigSeller benchmark metric query contract is invalid")
            identity = {"source_file": str(source_file), "source_sha256": input_data.source_sha256,
                        "sheet_name": input_data.sheet_name, "metric": metric,
                        "account_key": "0" * 64, "skus": list(input_data.skus)}
            with create_checkpoint(check_dir, identity) as checkpoint:
                checkpoint.save_results(benchmark_rows)
                checkpoint_file = checkpoint.path
            restored = read_checkpoint(checkpoint_file, identity)
            if (set(restored) != {"00001", "00002"}
                    or restored["00001"]["metric_value"] != 0
                    or restored["00002"]["status"] != "failed"
                    or restored["00002"]["attempts"] != 2
                    or restored["00002"]["retry_not_before"] != 2000000000.0):
                raise RuntimeError("BigSeller benchmark checkpoint recovery check failed")
            output_file = export_benchmark_workbook(input_data, restored, metric, check_dir, summary=summary)
            workbook = load_workbook(output_file, read_only=True, data_only=False)
            try:
                summary_sheet = workbook[f"BigSeller{label}对标汇总"]
                pending_sheet = workbook[f"BigSeller{label}待补查"]
                if ("结果不完整" not in output_file.name
                        or summary_sheet["B2"].value != "结果不完整"
                        or pending_sheet["A2"].value != "00002"
                        or pending_sheet["C2"].value != 2
                        or pending_sheet.max_row != 2
                        or workbook[input_data.sheet_name]["A2"].value != "00001"):
                    raise RuntimeError("BigSeller benchmark Excel integrity check failed")
            finally:
                workbook.close()

        # Keep claim parsing, its SQLite schema, and original-workbook export
        # executable in the packaged runtime. These are synthetic local files.
        from backend.services.bigseller_claim_checkpoint import (
            create_claim_checkpoint,
            read_claim_checkpoint,
        )
        from backend.services.bigseller_claim_workbook import (
            export_claim_workbook,
            load_claim_input,
        )

        claim_source = check_dir / "claim-input.xlsx"
        claim_book = Workbook()
        claim_book.active.title = "Claim source"
        claim_book.active.append(["SKU", "上架时间"])
        claim_book.active.append(["00001", "Original manual time"])
        claim_book.active.append(["00001\n00002", "Keep original content"])
        claim_book.save(claim_source)
        claim_book.close()
        claim_source_bytes = claim_source.read_bytes()
        claim_input = load_claim_input(claim_source)
        if (claim_input.skus != ("00001", "00002") or claim_input.total_rows != 2
                or claim_input.row_numbers != {"00001": (2, 3), "00002": (3,)}):
            raise RuntimeError("BigSeller claim duplicate and multi-SKU parsing check failed")
        claim_rows = [
            {"sku": "00001", "shop_name": "Self-check shop", "item_id": "00123",
             "created_time": "2026-09-01 12:30:00", "listed_time": "2026-08-01 08:00:00",
             "status": "matched", "message": ""},
            {"sku": "00002", "shop_name": "", "item_id": "", "created_time": "", "listed_time": "",
             "status": "not_found", "message": "Self-check no match"},
        ]
        claim_identity = {"source_file": str(claim_source), "source_sha256": claim_input.source_sha256,
                          "sheet_name": claim_input.sheet_name, "account_key": "0" * 64,
                          "skus": list(claim_input.skus), "site": "all", "listing_scope": "live",
                          "query_contract": "parentSku-exact-earliest-create_time-v2"}
        claim_checkpoint = create_claim_checkpoint(check_dir, claim_identity)
        try:
            claim_checkpoint.save_results(claim_rows)
            claim_checkpoint.save_deadline(2000000000.0)
            claim_checkpoint_file = claim_checkpoint.path
        finally:
            claim_checkpoint.close()
        claim_restored, claim_deadline = read_claim_checkpoint(claim_checkpoint_file, claim_identity)
        if (claim_restored != {row["sku"]: row for row in claim_rows}
                or claim_deadline != 2000000000.0):
            raise RuntimeError("BigSeller claim checkpoint recovery check failed")
        claim_output = export_claim_workbook(claim_input, claim_restored, check_dir,
                                             metadata={"is_complete": True, "site": "all", "listing_scope": "live"})
        claim_book = load_workbook(claim_output, read_only=True, data_only=False)
        try:
            claim_sheet = claim_book[claim_input.sheet_name]
            claim_headers = {cell.value: cell.column for cell in claim_sheet[1]}
            shop_column = claim_headers.get("BS最早创建店铺")
            created_column = claim_headers.get("BS创建时间")
            if (not shop_column or not created_column
                    or claim_sheet["B2"].value != "Original manual time"
                    or claim_sheet["B3"].value != "Keep original content"
                    or claim_sheet["A3"].value != "00001\n00002"
                    or claim_sheet.cell(2, shop_column).value != "Self-check shop"
                    or "00001：Self-check shop" not in str(claim_sheet.cell(3, shop_column).value)
                    or "00002：" not in str(claim_sheet.cell(3, shop_column).value)
                    or str(claim_sheet.cell(2, created_column).value) != "2026-09-01 12:30:00"
                    or claim_source.read_bytes() != claim_source_bytes):
                raise RuntimeError("BigSeller claim frozen workbook preservation check failed")
        finally:
            claim_book.close()

    return 0


def main() -> None:
    import webview

    index_html = resolve_entry_html()
    icon_path = (
        ROOT / "HQYLAutomation.ico"
        if getattr(sys, "frozen", False)
        else ROOT / "packaging" / "HQYLAutomation.ico"
    )
    if sys.platform == "win32":
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID("HQYL.Automation")
    api = AppBridge()
    webview.create_window(
        "寰球云联自动化平台",
        index_html.as_uri(),
        js_api=api,
        width=1180,
        height=760,
        min_size=(980, 640),
    )
    webview.start(debug="--debug" in sys.argv, icon=str(icon_path))


if __name__ == "__main__":
    if "--self-check" in sys.argv:
        raise SystemExit(self_check())
    main()
